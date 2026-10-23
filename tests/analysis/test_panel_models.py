"""Models over the waves of a panel on the results page (app/panel_models.py, longitudinal and multilevel concept
2026-10-17, package 2): the choice in the panel "Wellen", the block with its button, the linked persons with FIML, the
invariance over time and the growth curve, the access log, and read access to every wave."""
import math
import random

import numpy as np
import pytest
from werkzeug.datastructures import MultiDict

from app import panel_models
from app.extensions import db
from app.items.personal_code import DEFAULT_PARTS
from app.models import AccessLog, Response, Survey, User
from tests.conftest import add

E = "/surveys/1/editor"


def _code(rng, taken):
    while True:
        c = (rng.choice("ABCDEFGHKLMNPRSTW") + f"{rng.randint(1, 28):02d}" + rng.choice("ABCDEFGHKLMNPRSTW")
             + str(rng.randint(0, 4)) + rng.choice("ABCDEFGHKLMNPRSTW"))
        if c not in taken:
            taken.add(c)
            return c


def _parts(code):
    return dict(zip(DEFAULT_PARTS, [code[0], code[1:3], code[3], code[4], code[5]]))


def _waves(app, client, n=160):
    """Three waves (survey 1 and two duplicates) with a personal code and a scale of three statements; the score
    grows by about 0.3 points a wave; 10 % leave after wave 1, 10 % more after wave 2."""
    add(client, "personal_code")
    client.post(f"{E}/items/2", data=MultiDict([("code", "pcode"), ("question", "Code"), ("parts_shown", "1")]
                                               + [("part", p) for p in DEFAULT_PARTS]))
    add(client, "likert")
    client.post(f"{E}/items/3", data=MultiDict(
        [("code", "zuf"), ("question", "Zufriedenheit"), ("points", "7"), ("labeling", "full"), ("level", "metric")]
        + [x for k in range(1, 4) for x in (("stmt_id", f"s{k}"), ("stmt_code", ""), ("stmt_text", f"Aussage {k}"))]))
    with app.app_context():
        from app import duplicate

        s1 = db.session.get(Survey, 1)
        user = db.session.scalar(db.select(User).filter_by(orcid="0000-0000-0000-001X"))
        s2, _ = duplicate.create(s1, user, "Welle 2")
        s3, _ = duplicate.create(s1, user, "Welle 3")
        rng, taken = random.Random(17), set()
        people = [(_code(rng, taken), rng.gauss(0, 1), rng.gauss(0, 0.3)) for _ in range(n)]
        for t, s in enumerate((s1, s2, s3)):
            items = {i.code: i.id for b in s.blocks for i in b.items}
            for k, (code, level, slope) in enumerate(people):
                if (t == 1 and k % 10 == 0) or (t == 2 and k % 5 == 0):
                    continue
                f = level + (0.3 + slope) * t
                answers = {str(items["pcode"]): _parts(code),
                           str(items["zuf"]): {f"s{j}": max(1, min(7, int(round(3.5 + 0.9 * f + 0.5 * rng.gauss(0, 1)))))
                                               for j in (1, 2, 3)}}
                db.session.add(Response(survey_id=s.id, source="real", is_preview=False, status=Response.COMPLETED,
                                        pid=f"p{t}{k}", qversion=1, answers=answers, layout={}))
        db.session.commit()
        ids = (s1.id, s2.id, s3.id)
    data = MultiDict()
    for sid, label in zip(ids, ("t1", "t2", "t3")):
        data.add("wave_survey", str(sid))
        data.add("wave_label", label)
    assert client.post("/surveys/1/waves", data=data).status_code in (302, 303)
    return ids, data


def test_choice_block_and_models(app, client, survey):
    ids, data = _waves(app, client)
    page = client.get("/surveys/1/researchdesign").data.decode()
    assert 'id="wellen-modell"' in page and 'name="panel_var" value="zuf"' in page
    data.add("panel_form", "1")
    data.add("panel_var", "zuf")
    assert client.post("/surveys/1/waves", data=data).status_code in (302, 303)
    with app.app_context():
        for sid in ids:  # the choice stands in every wave, beside the list
            s = db.session.get(Survey, sid)
            assert s.waves["model"] == {"vars": ["zuf"]} and len(s.waves["waves"]) == 3
    page = client.get("/surveys/1/results?go=1").data.decode()
    assert 'id="panel"' in page and "Längsschnitt rechnen" in page
    app.config["HB_SYNC"] = True  # the job in the request; the page takes its result over
    assert client.post("/surveys/1/results/panel").status_code in (302, 303)
    page = client.get("/surveys/1/results?go=1").data.decode()
    assert "Messinvarianz über die Zeit" in page and "Wachstumskurve" in page and "Längsschnitt rechnen" not in page
    # the same block on the results page of the third wave, from its own computation
    assert "Längsschnitt rechnen" in client.get(f"/surveys/{ids[2]}/results?go=1").data.decode()
    with app.app_context():
        s = db.session.get(Survey, 1)
        user = db.session.scalar(db.select(User).filter_by(orcid="0000-0000-0000-001X"))
        found = panel_models.compute(user, s)
        assert found["report"]["persons"] == 160 and found["report"]["complete"] == 128
        var = found["vars"][0]
        assert [d["n"] for d in var["describe"]] == [160, 144, 128]
        g = var["growth"]
        assert g["fit"]["converged"] and g["fit"]["n"] == 160  # everyone counts, also who left (FIML)
        assert 0.15 < g["means"]["S"]["est"] < 0.45 and g["means"]["S"]["p"] < 0.001
        inv = var["invariance"]
        assert [r["level"] for r in inv["levels"]] == ["configural", "metric", "scalar"]
        assert inv["levels"][0]["fit"]["df"] > 0 and inv["highest"] == "scalar"
        assert inv["means"][1]["est"] > inv["means"][0]["est"] > 0
        log = db.session.scalars(db.select(AccessLog)).all()
        assert any(r.action == "Wellen verknüpft ausgewertet" and "waves=" in r.detail for r in log)


def test_read_access_to_every_wave(app, client, survey):
    ids, data = _waves(app, client, n=40)
    data.add("panel_form", "1")
    data.add("panel_var", "zuf")
    client.post("/surveys/1/waves", data=data)
    with app.app_context():
        from app.models import SurveyMember

        s3 = db.session.get(Survey, ids[2])
        db.session.execute(db.delete(SurveyMember).where(SurveyMember.survey_id == s3.id))  # the third wave is gone
        db.session.commit()
        user = db.session.scalar(db.select(User).filter_by(orcid="0000-0000-0000-001X"))
        assert not s3.can(user, "viewer")
        found = panel_models.page(user, db.session.get(Survey, 1))
        assert found["missing"] == ["t3"]
    page = client.get("/surveys/1/results?go=1").data.decode()
    assert "Leserecht in jeder Welle" in page and "Längsschnitt rechnen" not in page
    assert client.post("/surveys/1/results/panel").status_code in (302, 303)


def test_tables_have_the_expected_degrees_of_freedom():
    cols, tab = panel_models.invariance_table(["a", "b", "c"], 3, "scalar")
    assert len(cols) == 9
    q = len({r[4] or id(r) for r in tab if r[3] is None})
    # p(p + 3)/2 moments minus the free parameters: the scalar model of three statements over three waves
    assert 9 * 10 // 2 + 9 - q == 23
    tab = panel_models.growth_table(["x1", "x2", "x3", "x4"], [0, 1, 2, 3], quadratic=True)
    assert sum(1 for r in tab if r[0] == "=~") == 12 and math.isclose([r for r in tab if r[:2] == ("=~", "Q")][-1][3], 9.0)


# --- the cross-lagged panel models (package 3) --------------------------------------------------------------------

def _panel_data():
    """The seeded panel of tests/analysis/test_sem_table.py with the columns of panel_models.cross_table: the
    statements v1_<k>_t<t> (x) and v2_<k>_t<t> (y), the scores v1_t<t> and v2_t<t>."""
    import numpy as np

    from tests.analysis.test_sem_table import _items, _scores

    cols, x = _items()
    names, s = _scores()
    ren = {f"{a}{k}_t{t}": f"{v}_{k}_t{t}" for t in (1, 2, 3) for k in (1, 2, 3) for a, v in (("x", "v1"), ("y", "v2"))}
    out = [ren[c] for c in cols] + [n.replace("x_", "v1_").replace("y_", "v2_") for n in names]
    return out, np.column_stack([x, s])


@pytest.mark.slow
def test_cross_lagged_models_reference_values():
    """Independently computed reference values (FIML with robust errors, MLR) for the same tables."""
    from app import sem_table

    cols_all, x = _panel_data()
    cases = {  # (variables, kind, equal): chi2, scaled, srmr, (path y_t2 on x_t1: est, se, std)
        "latent RI-CLPM, equal": (([{"id": "v1", "items": 3}, {"id": "v2", "items": 3}], "riclpm", True),
                                  117.240872, 110.498116, 0.027491, ("w_v2_t2", "w_v1_t1", 0.269437, 0.206344, 0.220292)),
        "latent CLPM": (([{"id": "v1", "items": 3}, {"id": "v2", "items": 3}], "clpm", False),
                        121.911416, 114.675615, 0.026130, ("F_v2_t3", "F_v1_t2", 0.049557, 0.047957, 0.051470)),
        "mixed RI-CLPM": (([{"id": "v1", "items": 3}, {"id": "v2", "items": 0}], "riclpm", False),
                          32.393605, 30.209377, 0.019184, ("w_v2_t2", "w_v1_t1", 0.486065, 0.178500, 0.402009)),
        "observed CLPM, equal": (([{"id": "v1", "items": 0}, {"id": "v2", "items": 0}], "clpm", True),
                                 63.941806, 59.611162, 0.038348, ("v2_t2", "v1_t1", 0.174584, 0.029163, 0.177101)),
    }
    for name, ((vs, kind, equal), chi2, scaled, srmr, (lhs, rhs, est, se, std)) in cases.items():
        cols, tab, same = panel_models.cross_table(vs, 3, kind, equal)
        fit = sem_table.fit(x[:, [cols_all.index(c) for c in cols]], cols, tab, same=same)
        assert fit["converged"], name
        assert (fit["chi2"], fit["scaled"]["chi2"], fit["srmr"]) == pytest.approx((chi2, scaled, srmr), abs=2e-5), name
        row = sem_table.find(fit, "~", lhs, rhs)
        assert (row["est"], row["se"], row["std"]) == pytest.approx((est, se, std), abs=1e-5), name


def _two_scales(app, client, n=240):
    """Three waves with two scales x and y (three statements each) from a RI-CLPM: x in one wave raises y in the
    next (0.35 within persons), stable differences between persons besides."""
    add(client, "personal_code")
    client.post(f"{E}/items/2", data=MultiDict([("code", "pcode"), ("question", "Code"), ("parts_shown", "1")]
                                               + [("part", p) for p in DEFAULT_PARTS]))
    for number, code in ((3, "x"), (4, "y")):
        add(client, "likert")
        client.post(f"{E}/items/{number}", data=MultiDict(
            [("code", code), ("question", f"Skala {code}"), ("points", "7"), ("labeling", "full"), ("level", "metric")]
            + [v for k in range(1, 4) for v in (("stmt_id", f"s{k}"), ("stmt_code", ""), ("stmt_text", f"{code} {k}"))]))
    client.post("/surveys/1/researchdesign", data=MultiDict({
        "research_question": "Wirkt x auf y?", "goal": "correlation", "design": "panel",
        "h_id": [""], "h_text": ["x wirkt in der nächsten Welle auf y"], "h_kind": ["lagged"], "h_dir": ["pos"],
        "h_iv": ["x"], "h_dv": ["y"], "h_analysis": [""]}))
    with app.app_context():
        from app import duplicate

        s1 = db.session.get(Survey, 1)
        user = db.session.scalar(db.select(User).filter_by(orcid="0000-0000-0000-001X"))
        s2, _ = duplicate.create(s1, user, "Welle 2")
        s3, _ = duplicate.create(s1, user, "Welle 3")
        rng, taken = random.Random(5), set()
        people = []
        for _ in range(n):
            rx = rng.gauss(0, 0.8)
            ry = 0.4 * rx + rng.gauss(0, 0.7)
            wx, wy = [rng.gauss(0, 1)], [rng.gauss(0, 1)]
            for _t in (1, 2):
                wx.append(0.4 * wx[-1] + rng.gauss(0, 0.9))
                wy.append(0.3 * wy[-1] + 0.35 * wx[-2] + rng.gauss(0, 0.85))
            people.append((_code(rng, taken), [rx + w for w in wx], [ry + w for w in wy]))
        for t, s in enumerate((s1, s2, s3)):
            items = {i.code: i.id for b in s.blocks for i in b.items}
            for k, (code, xs, ys) in enumerate(people):
                if t and k % (12 - 4 * t) == 0:
                    continue

                def stmts(f):
                    return {f"s{j}": max(1, min(7, int(round(4 + 0.9 * f + 0.45 * rng.gauss(0, 1))))) for j in (1, 2, 3)}
                answers = {str(items["pcode"]): _parts(code), str(items["x"]): stmts(xs[t]), str(items["y"]): stmts(ys[t])}
                db.session.add(Response(survey_id=s.id, source="real", is_preview=False, status=Response.COMPLETED,
                                        pid=f"p{t}{k}", qversion=1, answers=answers, layout={}))
        db.session.commit()
        ids = (s1.id, s2.id, s3.id)
    data = MultiDict()
    for sid, label in zip(ids, ("t1", "t2", "t3")):
        data.add("wave_survey", str(sid))
        data.add("wave_label", label)
    client.post("/surveys/1/waves", data=data)
    return ids, data


def test_hypothesis_over_time_judged_in_the_riclpm(app, client, survey):
    ids, data = _two_scales(app, client)
    page = client.get("/surveys/1/researchdesign").data.decode()
    assert 'value="lagged" selected' in page and "Wirkung über die Zeit" in page
    page = client.get("/surveys/1/results?go=1").data.decode()
    assert "Längsschnitt rechnen" in page and "Geprüft im Block „Längsschnitt über die Wellen“" in page
    for flag in ("panel_equal", "panel_latent"):
        data.add(flag, "1")
    data.add("panel_form", "1")
    data.add("panel_var", "x")
    data.add("panel_var", "y")
    client.post("/surveys/1/waves", data=data)
    app.config["HB_SYNC"] = True
    assert client.post("/surveys/1/results/panel").status_code in (302, 303)
    page = client.get("/surveys/1/results?go=1").data.decode()
    assert "RI-CLPM" in page and "kreuzverzögerter Pfad β" in page and "unterstützt" in page
    with app.app_context():
        s = db.session.get(Survey, 1)
        user = db.session.scalar(db.select(User).filter_by(orcid="0000-0000-0000-001X"))
        found = panel_models.compute(user, s)
        c = found["cross"]
        assert c["kind"] == "riclpm" and c["equal"] and c["fit"]["converged"] and c["n"] == 240
        assert set(c["latent"]) == {v["label"] for v in found["vars"]}  # metric invariance holds for both scales
        path = next(p for p in c["paths"] if p["from"] == "x" and p["to"] == "y")
        back = next(p for p in c["paths"] if p["from"] == "y" and p["to"] == "x")
        assert path["std"] > 0.15 and path["p"] < 0.01 and back["p"] > 0.01
        assert c["hyps"]["x>y"]["test"] == "z"
        # the model as R code on the merged wide file (E14): scores from the statements, the lavaan call
        assert 'daten$t1_x <- rowMeans(daten[, c("t1_x_1", "t1_x_2", "t1_x_3")])' not in c["r"]  # latent: no scores
        assert "F_x_t1 =~ 1*t1_x_1" in c["r"] and "w_y_t2 ~ b_v1_v2*w_x_t1" in c["r"] and 'missing = "ml"' in c["r"]
        g = found["vars"][0]["growth"]["r"]
        assert 'daten$t1_x <- rowMeans(daten[, c("t1_x_1", "t1_x_2", "t1_x_3")])' in g and "S =~ 2*t3_x" in g
        from app import methods_text

        with app.test_request_context():
            text = " ".join(" ".join(ps) for _, ps in methods_text.sections(s, None, None, None, "de"))
        assert "Die 3 Wellen wurden über den persönlichen Code verknüpft" in text
        assert "Random-Intercept-Cross-Lagged-Panel-Modell" in text
    assert "Als R-Code (lavaan)" in page and 'class="sem-diagram"' in page
    # the tables for a manuscript
    reply = client.get("/surveys/1/results/panel.docx?lang=de")
    assert reply.status_code == 200 and reply.data[:2] == b"PK"
    import io

    from docx import Document

    doc = Document(io.BytesIO(reply.data))
    words = " ".join(p.text for p in doc.paragraphs)
    assert "Kreuzverzögerte Pfade über die Wellen" in words and "Latente Wachstumskurven" in words
    cells = [c.text for t in doc.tables for r in t.rows for c in r.cells]
    assert any("x → y" in c for c in cells)


def test_power_of_the_hypotheses_over_time(app, client, survey):
    """E11: the simulation from the plan (job kind sem_power with "panel"), as the page shows it."""
    from app import panel_power

    sigma, kind = panel_power.population(["x", "y"], 3, [("x", "y", 0.3)])
    assert kind == "riclpm" and np.allclose(np.diag(sigma), 1.0)
    out = panel_power.run({"key": "k", "n": 250, "reps": 30, "seed": 2, "panel": {
        "vars": ["x", "y"], "waves": 3, "betas": [["x", "y", 0.3]],
        "hyps": [{"number": 1, "from": "x", "to": "y", "direction": "pos"},
                 {"number": 2, "from": "y", "to": "x", "direction": "pos"}]}})
    power = dict(out["power"])
    assert power[1] > 0.7 and power[2] < 0.2 and out["failed"] == 0
    for code in ("x", "y"):
        add(client, "likert")
    client.post("/surveys/1/researchdesign", data=MultiDict({
        "research_question": "Wirkt x auf y?", "goal": "correlation", "design": "panel", "sample_target": "250",
        "h_id": [""], "h_text": ["x wirkt auf y"], "h_kind": ["lagged"], "h_dir": ["pos"], "h_iv": ["Q1"],
        "h_dv": ["Q2"], "h_analysis": [""], "h_effect": ["0.3"]}))
    with app.app_context():
        s = db.session.get(Survey, 1)
        inp = panel_models.power_input(s)
        assert inp and inp["panel"]["waves"] == 3 and inp["n"] == 250
    page = client.get("/surveys/1/researchdesign").data.decode()
    assert "Power der Hypothesen über die Zeit prüfen" in page
    app.config["HB_SYNC"] = True
    assert client.post("/surveys/1/researchdesign/panel-power").status_code in (302, 303)
    page = client.get("/surveys/1/researchdesign").data.decode()
    assert "Panels mit N = 250, 3 Wellen, RI-CLPM" in page


def test_models_run_as_a_job(app, client, survey, monkeypatch):
    """"Längsschnitt rechnen" reads the waves in the request and leaves the models to the job "panel" (estimation
    process or Rechenknecht; thread "Rechenknecht", 2026-10-17): the page shows the bar while it runs, takes the
    result over when it is there, and the job's numbers are those of compute() also after JSON."""
    import json

    from app import jobs, panel_job

    ids, data = _waves(app, client)
    data.add("panel_form", "1")
    data.add("panel_var", "zuf")
    client.post("/surveys/1/waves", data=data)
    started = []

    def start(app_, kind, paths, inp, pool, background=True, name="hb"):
        started.append((kind, paths, inp))
        return True

    monkeypatch.setattr(jobs, "start", start)
    assert client.post("/surveys/1/results/panel").status_code in (302, 303)
    (kind, paths, inp), = started
    assert kind == "panel" and paths[1].endswith("panel.run.json")
    sent = json.loads(json.dumps(inp))  # plain JSON: no participation numbers, no codes
    assert "pid" not in json.dumps(sent) and set(sent["scores"]) == {"zuf"}
    with app.app_context():
        from app import maxdiff_hb

        os_path = paths[1]
        maxdiff_hb._write(os_path, {"state": "running", "started": maxdiff_hb._now().isoformat(),
                                    "beat": maxdiff_hb._now().isoformat(), "progress": 0.5, "key": inp["key"]})
    page = client.get("/surveys/1/results?go=1").data.decode()
    assert "Längsschnitt wird gerechnet" in page and "results/panel/status" in page
    assert client.get("/surveys/1/results/panel/status").get_json()["state"] == "running"
    result = panel_job.run(sent)
    with app.app_context():
        maxdiff_hb._remove(paths[1])
        maxdiff_hb._write(paths[0], json.loads(json.dumps(result)))
        s = db.session.get(Survey, 1)
        user = db.session.scalar(db.select(User).filter_by(orcid="0000-0000-0000-001X"))
        direct = panel_models.compute(user, s)
    assert result["result"]["report"] == direct["report"]
    assert result["result"]["vars"][0]["growth"]["means"]["S"]["est"] == pytest.approx(direct["vars"][0]["growth"]["means"]["S"]["est"])
    page = client.get("/surveys/1/results?go=1").data.decode()
    assert "Wachstumskurve" in page and "Längsschnitt wird gerechnet" not in page
    import os

    assert not os.path.exists(paths[0])  # taken over into the kept results


def test_power_refuses_effects_too_large_for_variance_one(app, client, survey):
    """Audit 2026-10-21, E5: with RI share 0.5 a cross-lagged path of 0.9 explains more than the within variance;
    the floor raised the variance above 1 and simulated another effect. Now it is said, not simulated."""
    from app import panel_power

    assert panel_power.too_large(["a", "b"], 3, [("a", "b", 0.3)]) == []
    assert panel_power.too_large(["a", "b"], 3, [("a", "b", 0.6)]) == []
    assert panel_power.too_large(["a", "b"], 3, [("a", "b", 0.9)]) == ["b"]
    assert np.allclose(np.diag(panel_power.population(["a", "b"], 3, [("a", "b", 0.6)])[0]), 1.0)
    assert not panel_power.check_population(["a", "b"], 3, [("a", "b", 0.9)])
    assert panel_power.check_population(["a", "b"], 3, [("a", "b", 0.3)])
    for _ in ("x", "y"):
        add(client, "likert")
    client.post("/surveys/1/researchdesign", data=MultiDict({
        "research_question": "Wirkt x auf y?", "goal": "correlation", "design": "panel", "sample_target": "250",
        "h_id": [""], "h_text": ["x wirkt auf y"], "h_kind": ["lagged"], "h_dir": ["pos"], "h_iv": ["Q1"],
        "h_dv": ["Q2"], "h_analysis": [""], "h_effect": ["0.9"]}))
    with app.app_context():
        s = db.session.get(Survey, 1)
        assert panel_models.power_input(s)["panel"]["betas"] == [["Q1", "Q2", 0.9]]
        assert panel_models.power_state(s)["state"] == "too_large" and panel_models.start_power(s) == "too_large"
    page = client.get("/surveys/1/researchdesign").data.decode()
    assert "Effekte zu groß für die Planungsannahmen" in page and "Power der Hypothesen über die Zeit prüfen" not in page
    flashed = client.post("/surveys/1/researchdesign/panel-power", follow_redirects=True).data.decode()
    assert "Effekte zu groß für die Planungsannahmen" in flashed
