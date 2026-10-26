"""Structural equation models of a survey (experimental feature "sem", app/sem_model.py): the model from the study
plan, the block on the results page, the verdicts of the hypotheses, the bootstrap, the measurement model without
hypotheses, and the gate of the experimental feature."""
import math
import random

from werkzeug.datastructures import MultiDict

from app import experimental, sem_model, study_plan
from app.extensions import db
from app.models import Response, Survey
from tests.conftest import add, experiments_on

E = "/surveys/1/editor"


def _scale(client, number, code, k):
    add(client, "likert")
    client.post(f"{E}/items/{number}", data=MultiDict(
        [("code", code), ("question", f"Skala {code}"), ("points", "7"), ("labeling", "full"), ("level", "metric")]
        + [x for n in range(1, k + 1) for x in (("stmt_id", f"s{n}"), ("stmt_code", ""), ("stmt_text", f"{code} Aussage {n}"))]))


def _answers(rng, n):
    """tb -> einst (.6) -> absicht (.5), plus tb -> absicht (.2); loadings about .75, seven points."""
    def point(v):
        return max(1, min(7, int(round(4 + 1.4 * v))))

    out = []
    for _ in range(n):
        tb = rng.gauss(0, 1)
        einst = 0.6 * tb + math.sqrt(0.64) * rng.gauss(0, 1)
        absicht = 0.5 * einst + 0.2 * tb + math.sqrt(1 - 0.25 - 0.04 - 2 * 0.5 * 0.2 * 0.6) * rng.gauss(0, 1)

        def items(f, k):
            return {f"s{i}": point(0.75 * f + math.sqrt(1 - 0.5625) * rng.gauss(0, 1)) for i in range(1, k + 1)}

        out.append({"2": items(tb, 4), "3": items(einst, 3), "4": items(absicht, 3)})
    return out


def _survey(app, client, n=300, plan=True):
    for number, code, k in ((2, "tb", 4), (3, "einst", 3), (4, "absicht", 3)):
        _scale(client, number, code, k)
    with app.app_context():
        s = db.session.get(Survey, 1)
        rng = random.Random(4)
        for answers in _answers(rng, n):
            s.responses.append(Response(source="real", status=Response.COMPLETED, answers=answers))
        db.session.commit()
    if plan:
        client.post("/surveys/1/researchdesign", data=MultiDict({
            "research_question": "Technikbereitschaft und Nutzung", "goal": "correlation", "design": "cross",
            "h_id": ["", "", ""], "h_text": ["TB wirkt auf Einstellung", "Einstellung wirkt auf Absicht", "TB über Einstellung"],
            "h_kind": ["effect", "effect", "mediation"], "h_dir": ["pos", "pos", "pos"],
            "h_iv": ["tb", "einst", "tb"], "h_dv": ["einst", "absicht", "absicht"], "h_mv": ["", "", "einst"],
            "h_analysis": ["sem", "sem", ""], "sem_method": "cb"}))


def test_everyone_has_the_plan_fields_and_the_block(app, client, survey):
    """For everyone since 2026-10-11 (Ralf; before the experimental feature "sem"): no switch needed."""
    _survey(app, client, n=150, plan=False)
    page = client.get("/surveys/1/researchdesign").data.decode()
    assert "experimentell" not in page[page.index("plan-hyps"):]
    assert "Vermittlung (indirekter Effekt)" in page and "Strukturgleichungsmodell schätzen" in page
    # without hypotheses: the measurement model of all scales
    page = client.get("/surveys/1/results?go=1").data.decode()
    assert 'id="sem"' in page and "Messmodell aller Skalen" in page and "CFI" in page
    assert client.post("/surveys/1/results/sem-bootstrap").status_code in (302, 303)


def test_model_from_the_plan_verdicts_and_bootstrap(app, client, survey):
    experiments_on(client, "sem")
    _survey(app, client)
    with app.app_context():
        s = db.session.get(Survey, 1)
        plan = study_plan.plan_of(s)
        assert plan["hypotheses"][2]["mv"] == "einst" and plan["sem_method"] == "cb"
        assert study_plan.suggest(s, plan["hypotheses"][2]) == "sem" and sem_model.in_use(s)
        model = sem_model.build(s)
        assert model["mode"] == "structural"
        assert set(model["spec"]["paths"]) == {("tb", "einst"), ("einst", "absicht"), ("tb", "absicht")}
        assert model["spec"]["indirect"] == [("tb", "einst", "absicht")]
        assert "Strukturgleichungsmodell" in study_plan.markdown(s)
    page = client.get("/surveys/1/results?go=1").data.decode()
    assert 'id="sem"' in page and "Strukturgleichungsmodell" in page and "RMSEA" in page and "<svg" in page
    assert "Bootstrap rechnen" in page and "Modellfit für den Bericht" in page
    # W68: the fit verdict as a sign and a word too; W69: sources linked, intervals "[a; b]", "Aussagen" in the nodes
    part = page[page.index('id="sem"'):]
    assert 'class="sem-mark"' in part and "(Hu &amp; Bentler, 1999)" not in part and 'href="#ref-' in part
    assert "Items" not in part and "95% CI" not in part
    # H1, H2 and the mediation (joint significance of both paths, SEM-1) decided from the model; the bootstrap
    # interval of the indirect effect follows
    hyp = page[page.index('id="hypothesen"'):page.index('id="sem"')]
    assert hyp.count("Unterstützt") == 3 and "β = 0," in hyp and "gemeinsame Signifikanz beider Pfade" in hyp
    assert "Bootstrap-Intervall des indirekten Effekts fehlt noch" in hyp
    assert "normaltheoretisch" in page
    # SEM extensions 2026-10-17: robust fit indices beside them (1), the largest residuals as a diagnosis (12)
    assert "<sub>robust</sub>" in part and "robuster CFI = " in part and 'id="sem-residuals"' in part
    assert "Modifikationsindizes" in part
    app.config["HB_SYNC"] = True
    client.post("/surveys/1/results/sem-bootstrap")
    page = client.get("/surveys/1/results").data.decode()
    hyp = page[page.index('id="hypothesen"'):page.index('id="sem"')]
    assert hyp.count("Unterstützt") == 3 and "indirekter Effekt β = 0," in hyp and "(Bootstrap)" in hyp
    assert "Bootstrap mit 2000 Ziehungen" in page and "Intervalle unsicher" not in page
    status = client.get("/surveys/1/results/sem-bootstrap/status").get_json()
    assert status["state"] == "done"


def test_pls_and_too_few_cases(app, client, survey):
    experiments_on(client, "sem")
    _survey(app, client, n=60)
    page = client.get("/surveys/1/results?go=1").data.decode()
    assert "Zu wenige vollständige Teilnahmen: 60 von mindestens" in page
    client.post("/surveys/1/researchdesign", data=MultiDict({
        "h_id": ["", ""], "h_text": ["TB wirkt auf Einstellung", "Einstellung wirkt auf Absicht"],
        "h_kind": ["effect", "effect"], "h_dir": ["", ""], "h_iv": ["tb", "einst"], "h_dv": ["einst", "absicht"],
        "h_mv": ["", ""], "h_analysis": ["sem", "sem"], "sem_method": "pls"}))
    with app.app_context():
        s = db.session.get(Survey, 1)
        rng = random.Random(9)
        for answers in _answers(rng, 120):
            s.responses.append(Response(source="real", status=Response.COMPLETED, answers=answers))
        db.session.commit()
    page = client.get("/surveys/1/results?go=1").data.decode()
    assert "varianzbasiert (PLS" in page and "PLS prüft Pfade nur mit dem Bootstrap" in page
    # W70: without the bootstrap the paths say so, and its button is the main one, above the paths
    part = page[page.index('id="sem"'):]
    assert "Bootstrap fehlt" in part and 'class="button button-small">Bootstrap rechnen' in part
    assert part.index("Bootstrap rechnen") < part.index("sem-paths")
    app.config["HB_SYNC"] = True
    client.post("/surveys/1/results/sem-bootstrap")
    page = client.get("/surveys/1/results").data.decode()
    hyp = page[page.index('id="hypothesen"'):page.index('id="sem"')]
    assert hyp.count("Unterstützt") == 2


def test_variables_that_cannot_enter(app, client, survey):
    experiments_on(client, "sem")
    _scale(client, 2, "tb", 4)
    add(client, "single_choice")
    client.post(f"{E}/items/3", data=MultiDict({"code": "gruppe", "question": "Gruppe?", "option_id": ["a", "b", "c"],
                                                "option_code": ["1", "2", "3"], "option_label": ["A", "B", "C"]}))
    client.post("/surveys/1/researchdesign", data=MultiDict({
        "h_id": [""], "h_text": ["Gruppe wirkt"], "h_kind": ["effect"], "h_dir": [""], "h_iv": ["gruppe"],
        "h_dv": ["tb"], "h_mv": [""], "h_analysis": ["sem"]}))
    with app.app_context():
        s = db.session.get(Survey, 1)
        model = sem_model.build(s)
        # stage 3 of the drawing: more than two groups explain with one path per group against the first
        assert model["hyps"][1]["errors"] == [] and model["hyps"][1]["paths"] == [("gruppe:2", "tb"), ("gruppe:3", "tb")]


def test_power_of_the_paths_and_the_model_test_in_the_plan(app, client, survey):
    experiments_on(client, "sem")
    _survey(app, client, n=10, plan=False)
    client.post("/surveys/1/researchdesign", data=MultiDict({
        "h_id": ["", ""], "h_text": ["TB wirkt auf Einstellung", "TB über Einstellung"],
        "h_kind": ["effect", "mediation"], "h_dir": ["pos", "pos"], "h_iv": ["tb", "tb"],
        "h_dv": ["einst", "absicht"], "h_mv": ["", "einst"], "h_analysis": ["sem", ""], "h_effect": ["0,3", "0,3"],
        "h_basis": ["literature", "literature"], "h_basis_note": ["", ""], "h_power": ["0.8", "0.8"], "sem_method": "cb"}))
    with app.app_context():
        s = db.session.get(Survey, 1)
        h1, h2 = study_plan.plan_of(s)["hypotheses"]
        p1, p2 = study_plan.power_of(s, h1), study_plan.power_of(s, h2)
        assert p1["measure"] == "β" and p1["n"] and p2["n"] > p1["n"]
        assert study_plan.power_summary(s)["n"] == p2["n"]
        plan = sem_model.rmsea_plan(s)
        assert plan["df"] > 0 and plan["n"]
        assert "β = 0,30" in study_plan.markdown(s)
    page = client.get("/surveys/1/researchdesign").data.decode()
    assert "Pfad im geplanten Modell" in page and "beide Pfade signifikant" in page and "Modelltest (RMSEA" in page
    page = client.get(f"/sample-size/?survey=1&design=sem&df={plan['df']}").data.decode()
    assert "Modelltest (RMSEA)" in page and f'value="{plan["df"]}" data-ss="semdf"' in page


def test_r_script_methods_draft_and_glossary(app, client, survey):
    assert "Strukturgleichungsmodelle" in client.get("/learn/glossary").data.decode()  # for everyone
    experiments_on(client, "sem")
    _survey(app, client)
    script = client.get("/surveys/1/export/model.R")
    assert script.status_code == 200
    text = script.data.decode()
    assert "library(lavaan)" in text and "F_tb =~ tb_1 + tb_2 + tb_3 + tb_4" in text
    assert "F_absicht ~ p2*F_einst + p3*F_tb" in text or "F_absicht ~" in text
    assert "ind1 := " in text and 'estimator = "MLM", fixed.x = FALSE' in text and "readRDS(" in text
    # SEM-5: the bootstrap of the standardised indirect effect with a seed; SEM-6: no exogenous covariance here
    # (one exogenous construct), see test_r_script_writes_the_exogenous_covariances
    assert "set.seed(1)" in text and "bootstrapLavaan(fit_ml, R = 2000" in text and "$est.std" in text
    assert " ~~ " not in text
    assert '"cfi.robust", "tli.robust", "rmsea.robust"' in text and "# modindices(fit" in text
    assert "Strukturgleichungsmodell für R" in client.get("/surveys/1/data").data.decode()
    md = client.get("/surveys/1/export/methods-section.md").data.decode()
    assert "kovarianzbasierten Strukturgleichungsmodell" in md and "Perzentil-Bootstrap" in md
    assert "gemeinsame Signifikanz beider Pfade" in md and "normaltheoretisch" in md
    assert "Das Strukturgleichungsmodell (N = 300) passte so zu den Daten: χ²(" in md and "β = 0," in md
    en = client.get("/surveys/1/export/methods-section.md?lang=en").data.decode()
    assert "covariance-based structural equation model" in en and "95% CI [" in en
    page = client.get("/learn/glossary").data.decode()
    assert "Strukturgleichungsmodelle" in page and "Heterotrait-Monotrait-Verhältnis" in page


def test_holm_counts_the_paths_and_the_mediation(app):
    """SEM-1 (audit 2026-10-06, t5_holm.py): two paths with p = .030 and .040 and a mediation over both; with Holm
    for k = 3 none stands (p_Holm .09). The mediation's p is the larger p of its paths (decision 2)."""
    from app import hypotheses

    plan = [{"id": 1, "kind": "relation", "analysis": "sem", "iv": "x", "dv": "A", "direction": "pos"},
            {"id": 2, "kind": "relation", "analysis": "sem", "iv": "A", "dv": "y", "direction": "pos"},
            {"id": 3, "kind": "mediation", "analysis": "sem", "iv": "x", "mv": "A", "dv": "y", "direction": "pos"}]
    fit = {"admissible": True, "paths": [
        {"from": "x", "to": "A", "est": .2, "se": .092, "z": 2.17, "p": 0.030, "std": .15, "std_se": .07, "ci": (.01, .29)},
        {"from": "A", "to": "y", "est": .2, "se": .097, "z": 2.05, "p": 0.040, "std": .14, "std_se": .07, "ci": (.005, .28)},
        {"from": "x", "to": "y", "est": .0, "se": .1, "z": 0.0, "p": 1.0, "std": .0, "std_se": .07, "ci": (-.1, .1)}],
        "indirect": [{"x": "x", "m": "A", "y": "y", "std": .021, "std_se": .012, "p": .08, "ci": (-.002, .044)}]}
    result = {"fit": fit, "method": "cb", "n": 300, "nodes": {"x": ("x", "x"), "A": ("A", "A"), "y": ("y", "y")},
              "hyps": {1: {"paths": [("x", "A")]}, 2: {"paths": [("A", "y")]},
                       3: {"paths": [("x", "A"), ("A", "y")], "indirect": [("x", "A", "y")], "direct": [("x", "y")]}}}
    boot = {"paths": {}, "indirect": {("x", "A", "y"): {"ci": (0.001, 0.05), "se": .012}}}
    with app.test_request_context():
        found = [{"text": h["kind"], "kind": h["kind"], "verdict": "offen", "sentence": "", "notes": [], "groups": []}
                 for h in plan]
        sem_model.apply(found, result, plan, boot, "de")
        assert [(r["verdict"], r.get("p")) for r in found] == [("unterstützt", 0.03), ("unterstützt", 0.04),
                                                               ("unterstützt", 0.04)]
        assert "gemeinsame Signifikanz beider Pfade (größeres p): p = ,040" in found[2]["sentence"]
        hypotheses.holm(found, 3, "de")
        assert [r["verdict"] for r in found] == ["nicht unterstützt"] * 3
        assert [round(r["p_holm"], 3) for r in found] == [0.09, 0.09, 0.09]
        # a mediation with one path not significant is not supported, whatever the bootstrap interval says
        fit["paths"][1]["p"] = 0.2
        found = [{"text": h["kind"], "kind": h["kind"], "verdict": "offen", "sentence": "", "notes": [], "groups": []}
                 for h in plan]
        sem_model.apply(found, result, plan, boot, "de")
        assert found[2]["verdict"] == "nicht unterstützt" and found[2]["p"] == 0.2
        # not admissible: no verdict and no p, the hypothesis counts in k with p = 1
        fit["admissible"] = False
        found = [{"text": h["kind"], "kind": h["kind"], "verdict": "offen", "sentence": "", "notes": [], "groups": []}
                 for h in plan]
        sem_model.apply(found, result, plan, boot, "de")
        assert all(r["verdict"] == "offen" and r.get("p") is None for r in found)


def test_r_script_writes_the_exogenous_covariances(app, client, survey):
    """SEM-6: with fixed.x = FALSE and the covariances of all exogenous variables written out, lavaan estimates the
    model the app estimates (here two exogenous constructs)."""
    experiments_on(client, "sem")
    _survey(app, client, plan=False)
    client.post("/surveys/1/researchdesign", data=MultiDict({
        "research_question": "Technikbereitschaft und Nutzung", "goal": "correlation", "design": "cross",
        "h_id": ["", ""], "h_text": ["TB wirkt auf Absicht", "Einstellung wirkt auf Absicht"],
        "h_kind": ["effect", "effect"], "h_dir": ["pos", "pos"], "h_iv": ["tb", "einst"],
        "h_dv": ["absicht", "absicht"], "h_mv": ["", ""], "h_analysis": ["sem", "sem"], "sem_method": "cb"}))
    text = client.get("/surveys/1/export/model.R").data.decode()
    assert "F_tb ~~ F_einst" in text and "fixed.x = FALSE" in text and "set.seed" not in text



def test_the_diagram_of_an_inadmissible_solution_shows_no_values():
    from app import sem_model

    result = {"nodes": {"a": ("a", "a"), "b": ("b", "b")}, "spec": {"paths": [("a", "b")], "latent": {}},
              "fit": {"admissible": False, "problems": ["extreme"], "r2": {"b": -0.01},
                      "paths": [{"from": "a", "to": "b", "std": 5.6e40, "p": 0.0}]}, "labels": {}}
    d = sem_model.diagram(result)
    assert d["edges"][0]["value"] is None and not d["edges"][0]["sig"] and "R²" not in d["nodes"][1]["sub"]


def test_the_diagram_avoids_crossing_paths():
    """A source pointing only at the third column stands in the second (a straight path, no arc); the rows are
    ordered so that no paths cross where that is possible (Ralf, 2026-10-07)."""
    from app import sem_model

    cols = sem_model._layers(["hinweis", "bild", "vertrauen", "absicht", "nutzung"],
                             [("hinweis", "vertrauen"), ("hinweis", "absicht"), ("bild", "absicht"),
                              ("vertrauen", "absicht"), ("absicht", "nutzung")])
    assert cols == {0: ["hinweis"], 1: ["vertrauen", "bild"], 2: ["absicht"], 3: ["nutzung"]}
    cols = sem_model._layers(list("abcdef"), [("a", "e"), ("b", "d"), ("c", "f"), ("a", "f"), ("b", "e")])
    assert cols == {0: ["b", "a", "c"], 1: ["d", "e", "f"]}


def test_a_randomisation_is_drawn_as_an_assignment():
    from app import sem_model

    result = {"nodes": {"hinweis": ("hinweis", "Datenschutzhinweis"), "b": ("b", "b")}, "conditions": ["hinweis"],
              "spec": {"paths": [("hinweis", "b")], "latent": {}},
              "fit": {"admissible": True, "r2": {"b": 0.1}, "paths": [{"from": "hinweis", "to": "b", "std": 0.2, "p": 0.01}]},
              "labels": {}}
    node = sem_model.diagram(result)["nodes"][0]
    assert node["condition"] and node["label"] == "Zuteilung: Datenschutzhinweis"


def test_a_collapsed_estimate_shows_only_why(app, client, survey, monkeypatch):
    """Impossible values (paths of 10^40, a negative χ²) are no result: the page says why and shows no paths or fit
    line (Ralf, 2026-10-07)."""
    from app import sem

    experiments_on(client, "sem")
    _survey(app, client)
    real = sem.fit_ml

    def collapsed(*args, **kw):
        out = real(*args, **kw)
        return {**out, "chi2": -3962445288873.65, "problems": ["extreme"], "admissible": False}

    monkeypatch.setattr(sem, "fit_ml", collapsed)
    page = client.get("/surveys/1/results?go=1").data.decode()
    part = page[page.index('id="sem"'):]
    assert "Die Schätzung ist gekippt" in part and "sem-paths" not in part and "Modellfit für den Bericht" not in part
    assert "Höchste Korrelationen zwischen Einflussgrößen" in part and "Aussagen, α = " in part   # the diagnosis


def test_the_diagnosis_marks_predictors_that_are_the_same():
    import random as rnd

    from app import sem_model

    r = rnd.Random(4)
    rows = []
    for _ in range(200):
        f = r.gauss(0, 1)
        rows.append([f + r.gauss(0, .3), f + r.gauss(0, .3), f + r.gauss(0, .01), r.gauss(0, 1)])
    spec = {"latent": {"A": ["a1", "a2"]}, "observed": ["b", "y"], "paths": [("A", "y"), ("b", "y")]}
    d = sem_model.diagnosis(rows, ["a1", "a2", "b", "y"], spec, {"A": ("a", "Konstrukt A"), "b": ("b", "Frage B")})
    (a, b, rr), = d["pairs"]
    assert {a, b} == {"Konstrukt A", "Frage B"} and rr > sem_model.HIGH_R
    assert d["scales"][0][:2] == ("Konstrukt A", 2) and d["scales"][0][2] > 0.8


def test_methods_section_takes_the_model_the_results_page_keeps(app, client, survey, monkeypatch):
    """Code review 2026-10-15, A10: every download of the methods section estimated the model anew; now it takes the
    one the results page keeps for the same pool, and keeps it when the page has not estimated it yet."""
    from app import sem_model

    experiments_on(client, "sem")
    _survey(app, client)
    counted = []
    real = sem_model.estimate
    monkeypatch.setattr(sem_model, "estimate", lambda *a, **k: counted.append(1) or real(*a, **k))
    first = client.get("/surveys/1/export/methods-section.md").data.decode()
    assert "Das Strukturgleichungsmodell (N = 300)" in first and counted == [1]   # estimated and kept
    assert client.get("/surveys/1/export/methods-section.md").data.decode() == first and counted == [1]
    client.get("/surveys/1/results?go=1")
    assert counted == [1]                                                          # the page takes it too


def test_control_variables_and_a_single_question_with_its_reliability(app, client, survey):
    """SEM extensions 2026-10-17, points 6 and 7: a control variable gets a path onto every explained variable, without
    a hypothesis and outside Holm; a single question enters as a construct of one indicator whose error variance is
    1 − reliability, so its path is not shrunk. Both stored in the plan, in the pre-registration text, the methods
    draft and the R script."""
    for number, code, k in ((2, "tb", 4), (3, "einst", 3), (4, "absicht", 3)):
        _scale(client, number, code, k)
    for number, code in ((5, "alter"), (6, "zufr")):
        add(client, "rating")
        client.post(f"{E}/items/{number}", data={"code": code, "question": f"{code}?", "style": "slider", "min": "0",
                                                  "max": "100", "step": "1", "level": "metric"})
    rng = random.Random(9)
    with app.app_context():
        s = db.session.get(Survey, 1)
        for answers in _answers(rng, 400):
            tb = sum(answers["2"].values()) / 4 - 4
            answers["5"] = round(40 + 12 * rng.gauss(0, 1))
            true = 0.5 * tb / 1.1 + math.sqrt(0.75) * rng.gauss(0, 1)
            answers["6"] = round(50 + 15 * (math.sqrt(0.6) * true + math.sqrt(0.4) * rng.gauss(0, 1)))
            s.responses.append(Response(source="real", status=Response.COMPLETED, answers=answers))
        db.session.commit()
    form = {"research_question": "Technikbereitschaft", "goal": "correlation", "design": "cross",
            "h_id": ["", "", ""], "h_text": ["TB wirkt auf Einstellung", "Einstellung wirkt auf Absicht", "TB auf Zufriedenheit"],
            "h_kind": ["effect", "effect", "effect"], "h_dir": ["pos", "pos", "pos"],
            "h_iv": ["tb", "einst", "tb"], "h_dv": ["einst", "absicht", "zufr"], "h_mv": ["", "", ""],
            "h_analysis": ["sem", "sem", "sem"], "sem_method": "cb", "sem_shown": "1", "sem_control": ["alter"]}
    client.post("/surveys/1/researchdesign", data=MultiDict(form))
    page = client.get("/surveys/1/researchdesign").data.decode()
    assert 'name="sem_control" value="alter" checked' in page and 'name="sem_single" value="zufr"' in page
    with app.app_context():
        s = db.session.get(Survey, 1)
        model = sem_model.build(s)
        assert model["controls"] == {"alter": "alter"}
        assert {("alter", "einst"), ("alter", "absicht"), ("alter", "zufr")} <= set(model["spec"]["paths"])
        assert all(("alter", b) not in e["paths"] for e in model["hyps"].values() for b in ("einst", "absicht", "zufr"))
    page = client.get("/surveys/1/results?go=1").data.decode()
    part = page[page.index('id="sem"'):]
    assert "Kontrolliert für alter" in part and "Kontrollvariable" in part
    hyp = page[page.index('id="hypothesen"'):page.index('id="sem"')]
    assert hyp.count("Unterstützt") == 3
    # the single question with its reliability
    form.update(sem_single=["zufr"], sem_rel_zufr="0,6")
    client.post("/surveys/1/researchdesign", data=MultiDict(form))
    with app.app_context():
        s = db.session.get(Survey, 1)
        assert study_plan.plan_of(s)["sem"] == {"controls": ["alter"], "single": {"zufr": 0.6}, "formative": [], "global": {}, "invariance": "", "missing": "", "ordinal": "", "second": [], "alternative": [], "hoc": [], "higher_formative": []}
        model = sem_model.build(s)
        assert model["spec"]["latent"]["zufr"] == ["zufr°"] and model["spec"]["single"] == {"zufr": 0.6}
        assert "zufr" not in model["spec"]["observed"]
        text = sem_model.r_script(s, "daten.rds")
        assert "daten$zufr_ <- daten$zufr" in text and 'modell <- paste0(modell, "\\n  zufr_ ~~ ", (1 - 0.6)' in text
        assert "Kontrollvariablen (Pfade auf alle erklärten Größen, ohne Hypothese): alter" in study_plan.markdown(s, "de")
        assert "zufr (0,60)" in study_plan.markdown(s, "de")
    page = client.get("/surveys/1/results?go=1").data.decode()
    part = page[page.index('id="sem"'):]
    assert "Einzelfragen mit angenommener Reliabilität: zufr (0,60)" in part
    md = client.get("/surveys/1/export/methods-section.md").data.decode()
    assert "Kontrolliert wurde für alter" in md and "zufr: 0,60" in md
    # renaming a variable in the editor: the settings follow
    with app.app_context():
        s = db.session.get(Survey, 1)
        study_plan.rename_codes(s, {"alter": "age", "zufr": "sat"})
        assert study_plan.plan_of(s)["sem"] == {"controls": ["age"], "single": {"sat": 0.6}, "formative": [], "global": {}, "invariance": "", "missing": "", "ordinal": "", "second": [], "alternative": [], "hoc": [], "higher_formative": []}


def test_a_single_question_with_a_reliability_recovers_the_path():
    """Point 7: with the reliability of the single question fixed, the path to it is the path to its true score."""
    import numpy as np

    from app import sem

    rng = np.random.default_rng(3)
    n = 5000
    t = rng.normal(size=n)
    x = math.sqrt(0.7) * t + math.sqrt(0.3) * rng.normal(size=n)
    f = 0.5 * t + math.sqrt(0.75) * rng.normal(size=n)
    ys = np.column_stack([0.7 * f + math.sqrt(0.51) * rng.normal(size=n) for _ in range(3)])
    data = np.column_stack([ys, x])
    single = sem.fit_ml(data, ["y1", "y2", "y3", "x°"],
                        {"latent": {"Y": ["y1", "y2", "y3"], "x": ["x°"]}, "paths": [("x", "Y")], "single": {"x": 0.7}})
    plain = sem.fit_ml(data, ["y1", "y2", "y3", "x"], {"latent": {"Y": ["y1", "y2", "y3"]}, "observed": ["x"],
                                                        "paths": [("x", "Y")]})
    assert abs(single["paths"][0]["std"] - 0.5) < 0.03 and plain["paths"][0]["std"] < 0.45
    assert single["admissible"] and single["df"] == plain["df"] and "x" not in single["quality"]["ave"]
    assert single["loadings"]["x"][0]["std"] == pytest_approx(math.sqrt(0.7))


def pytest_approx(v):
    import pytest

    return pytest.approx(v, abs=1e-6)


def test_pls_figures_formative_constructs_and_scores(app, client, survey):
    """SEM extensions 2026-10-17, points 9, 13 and 17: with PLS the paths carry f² and VIF, the explained variables R²
    and adjusted R², PLSpredict judges the predictive power, a scale can be formative (weights with bootstrap p and
    VIF), HTMT gets its upper bound from the bootstrap, and the construct scores come as a CSV per resp_id for
    "Bearbeiten" and up."""
    _survey(app, client, n=300, plan=False)
    form = {"research_question": "TB", "goal": "correlation", "design": "cross",
            "h_id": ["", ""], "h_text": ["TB wirkt auf Einstellung", "Einstellung wirkt auf Absicht"],
            "h_kind": ["effect", "effect"], "h_dir": ["pos", "pos"], "h_iv": ["tb", "einst"], "h_dv": ["einst", "absicht"],
            "h_mv": ["", ""], "h_analysis": ["sem", "sem"], "sem_method": "pls", "sem_shown": "1"}
    client.post("/surveys/1/researchdesign", data=MultiDict(form))
    page = client.get("/surveys/1/researchdesign").data.decode()
    assert 'name="sem_mode_tb" value="formative"' in page and "Warum? Reflektiv oder formativ?" in page
    form["sem_mode_tb"] = "formative"
    client.post("/surveys/1/researchdesign", data=MultiDict(form))
    with app.app_context():
        s = db.session.get(Survey, 1)
        assert study_plan.plan_of(s)["sem"]["formative"] == ["tb"]
        assert sem_model.build(s)["spec"]["formative"] == ["tb"]
        assert "Formative Konstrukte (PLS, Modus B): tb" in study_plan.markdown(s, "de")
        assert 'composite("F_tb", c("tb_1", "tb_2", "tb_3", "tb_4"), weights = mode_B)' in sem_model.r_script(s, "d.rds")
    app.config["HB_SYNC"] = True
    client.get("/surveys/1/results?go=1")
    client.post("/surveys/1/results/sem-bootstrap")
    page = client.get("/surveys/1/results").data.decode()
    part = page[page.index('id="sem"'):]
    assert "R² korrigiert" in part
    assert 'id="sem-predict"' in part and "Q²<sub>predict</sub>" in part
    assert "Formative Konstrukte" in part
    assert "obere Grenze (95 %, einseitig)" in part  # einst and absicht: the one reflective pair
    assert "Konstruktwerte herunterladen (CSV)" in part
    csv_text = client.get("/surveys/1/export/sem-scores.csv").data.decode()
    lines = csv_text.strip().splitlines()
    assert lines[0] == "resp_id,tb,einst,absicht" and len(lines) == 301
    md = client.get("/surveys/1/export/methods-section.md").data.decode()
    assert "PLSpredict" in md and "Formativ gemessen (Modus B) wurden tb" in md
    # a viewer gets no response data
    from app.models import SurveyMember, User

    with app.app_context():
        u = User(orcid="0000-0002-1825-0097", name="Gast")
        db.session.add(u)
        db.session.flush()
        db.session.add(SurveyMember(survey_id=1, user_id=u.id, orcid=u.orcid, role="viewer"))
        db.session.commit()
    other = app.test_client()
    other.get("/auth/dev-login?orcid=0000-0002-1825-0097&name=Gast")
    other.post("/terms/", data={"accept": "1"})
    assert other.get("/surveys/1/export/sem-scores.csv").status_code in (302, 303)


def test_word_tables_and_lavaan_syntax(app, client, survey):
    """SEM extensions 2026-10-17, points 15 and 18: the model as APA tables in Word (German and English) and its lines
    in lavaan syntax to copy for JASP and jamovi."""
    import io

    from docx import Document

    _survey(app, client)
    page = client.get("/surveys/1/results?go=1").data.decode()
    part = page[page.index('id="sem"'):]
    assert 'id="sem-syntax"' in part and "F_tb =~ tb_1 + tb_2 + tb_3 + tb_4" in part
    assert 'role="button"' in part[part.index('id="sem-syntax"'):] and "/results/sem.docx?lang=en" in part
    for lang, title, word in (("de", "Pfadkoeffizienten des Strukturgleichungsmodells", "Indirekte Effekte"),
                              ("en", "Path Coefficients of the Structural Equation Model", "Indirect Effects")):
        r = client.get(f"/surveys/1/results/sem.docx?lang={lang}")
        assert r.status_code == 200 and r.headers["Content-Type"].startswith("application/vnd.openxml")
        doc = Document(io.BytesIO(r.data))
        text = "\n".join(p.text for p in doc.paragraphs)
        assert title in text and word in text and len(doc.tables) == 3
        cells = [c.text for c in doc.tables[0].rows[1].cells]
        assert "(H1, H3)" in cells[0] and len(cells) == 6  # the path of H1 and the first leg of the mediation H3
        assert ("Modellfit: χ²(" if lang == "de" else "Model fit: χ²(") in text
    with app.app_context():
        s = db.session.get(Survey, 1)
        syntax = sem_model.lavaan_syntax(s)
        assert syntax.splitlines()[0] == "F_tb =~ tb_1 + tb_2 + tb_3 + tb_4" and "ind1 := " in syntax


def test_measurement_invariance_across_the_languages(app, client, survey):
    """SEM extensions 2026-10-17, point 4: the plan names the language of the survey as groups; the results page
    tests the measurement model level by level (configural to strict), with Δχ² and ΔCFI, and with scalar invariance
    compares the latent means. The same in the R script, the methods draft and the pre-registration text."""
    _survey(app, client, n=300)
    with app.app_context():
        s = db.session.get(Survey, 1)
        rng = random.Random(21)
        for answers in _answers(rng, 300):  # a second language, higher on tb by about half an SD
            answers["2"] = {k: min(7, v + (1 if rng.random() < 0.5 else 0)) for k, v in answers["2"].items()}
            s.responses.append(Response(source="real", status=Response.COMPLETED, answers=answers, language="en"))
        db.session.commit()
        assert ("_lang", "Sprache der Befragung") in sem_model.invariance_candidates(s)
    form = {"research_question": "TB", "goal": "correlation", "design": "cross",
            "h_id": ["", "", ""], "h_text": ["TB wirkt auf Einstellung", "Einstellung wirkt auf Absicht", "TB über Einstellung"],
            "h_kind": ["effect", "effect", "mediation"], "h_dir": ["pos", "pos", "pos"],
            "h_iv": ["tb", "einst", "tb"], "h_dv": ["einst", "absicht", "absicht"], "h_mv": ["", "", "einst"],
            "h_analysis": ["sem", "sem", ""], "sem_method": "cb", "sem_shown": "1", "sem_invariance": "_lang"}
    client.post("/surveys/1/researchdesign", data=MultiDict(form))
    with app.app_context():
        s = db.session.get(Survey, 1)
        assert study_plan.plan_of(s)["sem"]["invariance"] == "_lang"
        assert "Messinvarianz (konfigural, metrisch, skalar, strikt) über: Sprache der Befragung" in study_plan.markdown(s, "de")
        text = sem_model.r_script(s, "d.rds")
        assert 'group = "lang"' in text and 'group.equal = c("loadings", "intercepts")' in text and "lavTestLRT(" in text
    page = client.get("/surveys/1/results?go=1").data.decode()
    part = page[page.index('id="sem-invariance"'):]
    assert "über Sprache der Befragung" in part and "skalar (+ gleiche Achsenabschnitte)" in part
    assert part.count('<span class="sem-mark">') >= 2 and "Höchste Stufe, die in Folge hält" in part
    md = client.get("/surveys/1/export/methods-section.md").data.decode()
    assert "Die Messinvarianz über die Sprache der Befragung wurde" in md


def test_invariance_levels_against_simulated_groups():
    """Point 4: the configural model is the sum of the single groups; an intercept that differs breaks scalar
    invariance; a mean that differs is found as the latent mean of the second group (d near the true 0.5)."""
    import numpy as np

    from app import sem, sem_means

    rng = np.random.default_rng(8)

    def group(n, shift=0.0, intercept=0.0):
        f = rng.normal(size=n) + shift
        x = np.column_stack([0.7 * f + math.sqrt(0.51) * rng.normal(size=n) for _ in range(4)]) + 3
        x[:, 3] += intercept
        return x

    cols = [f"a{i}" for i in range(4)]
    spec = {"latent": {"A": cols}}
    xs = [group(3000), group(3000, shift=0.5)]
    found = sem_means.invariance(xs, cols, spec)
    single = sum(sem.fit_ml(x, cols, spec, robust=False)["chi2"] for x in xs)
    assert found["levels"][0]["fit"]["chi2"] == pytest_approx_rel(single)
    assert found["highest"] == "strict"
    [mean] = found["levels"][2]["fit"]["latent_means"]["A"]
    assert abs(mean["d"] - 0.5) < 0.06 and mean["p"] < 0.001
    broken = sem_means.invariance([group(3000), group(3000, intercept=0.6)], cols, spec)
    assert broken["highest"] == "metric" and broken["levels"][2]["holds"] is False
    diff = broken["levels"][2]["diff"]
    assert diff["df"] == 3 and diff["p"] < 0.001 and diff["scaled"]


def pytest_approx_rel(v):
    import pytest

    return pytest.approx(v, rel=1e-6)


def test_fiml_and_ordered_categories_from_the_plan(app, client, survey, monkeypatch):
    """SEM extensions 2026-10-17, points 2 and 3: "alle Teilnahmen (Full Information ML)" keeps the participations
    with gaps; "als geordnete Kategorien (WLSMV)" estimates on polychoric correlations, also in the bootstrap; both in
    the R script, the methods draft and the pre-registration text."""
    _survey(app, client, n=300)
    with app.app_context():
        s = db.session.get(Survey, 1)
        rng = random.Random(5)
        for r in s.responses[:60]:  # 60 participations skip one statement each
            answers = {k: dict(v) for k, v in r.answers.items()}
            block = rng.choice(["2", "3", "4"])
            answers[block].pop(rng.choice(sorted(answers[block])))
            r.answers = answers
        db.session.commit()
    form = {"research_question": "TB", "goal": "correlation", "design": "cross",
            "h_id": ["", "", ""], "h_text": ["TB wirkt auf Einstellung", "Einstellung wirkt auf Absicht", "TB über Einstellung"],
            "h_kind": ["effect", "effect", "mediation"], "h_dir": ["pos", "pos", "pos"],
            "h_iv": ["tb", "einst", "tb"], "h_dv": ["einst", "absicht", "absicht"], "h_mv": ["", "", "einst"],
            "h_analysis": ["sem", "sem", ""], "sem_method": "cb", "sem_shown": "1"}
    page = client.get("/surveys/1/results?go=1").data.decode()
    assert "N = 240 vollständige Teilnahmen (60 mit Lücken" in page
    client.post("/surveys/1/researchdesign", data=MultiDict({**form, "sem_missing": "fiml"}))
    page = client.get("/surveys/1/results?go=1").data.decode()
    part = page[page.index('id="sem"'):]
    assert "N = 300 Teilnahmen mit mindestens einer Antwort im Modell, davon 240 vollständig" in part
    hyp = page[page.index('id="hypothesen"'):page.index('id="sem"')]
    assert hyp.count("Unterstützt") == 3
    with app.app_context():
        s = db.session.get(Survey, 1)
        text = sem_model.r_script(s, "d.rds")
        assert 'estimator = "MLR", missing = "ml"' in text and "rowSums(!is.na(" in text
        assert "Full Information Maximum Likelihood (robust, MLR)" in study_plan.markdown(s, "de")
    assert "Full Information Maximum Likelihood behandelt" in client.get("/surveys/1/export/methods-section.md").data.decode()
    app.config["HB_SYNC"] = True
    monkeypatch.setattr(sem_model, "BOOT_DRAWS", 60)  # the draws keep their gaps; few of them for the test
    client.post("/surveys/1/results/sem-bootstrap")
    page = client.get("/surveys/1/results").data.decode()
    assert "Bootstrap mit 60 Ziehungen" in page
    client.post("/surveys/1/researchdesign", data=MultiDict({**form, "sem_ordinal": "ordinal"}))
    page = client.get("/surveys/1/results?go=1").data.decode()
    part = page[page.index('id="sem"'):]
    assert "χ²<sub>WLSMV</sub>" in part and "polychorische Korrelationen" in part and "N = 240 vollständige" in part
    hyp = page[page.index('id="hypothesen"'):page.index('id="sem"')]
    assert hyp.count("Unterstützt") == 3
    with app.app_context():
        s = db.session.get(Survey, 1)
        assert 'ordered = TRUE, estimator = "WLSMV"' in sem_model.r_script(s, "d.rds")
    client.post("/surveys/1/results/sem-bootstrap")
    assert "Bootstrap mit 60 Ziehungen" in client.get("/surveys/1/results").data.decode()


def test_weighted_model_on_the_weighted_page(app, client, survey, monkeypatch):
    """SEM extensions 2026-10-17, point 5: with weights the covariance-based model takes them (only the weighted
    participations), names n_eff, its bootstrap key differs from the unweighted one and the bootstrap job draws with
    the weights."""
    from app import compare, compute, weights

    _survey(app, client, n=300)
    monkeypatch.setattr(weights, "compute", lambda s, table: {"weights": {r.get("resp_id"): (2.0 if k % 3 else 0.5)
                                                                          for k, r in enumerate(table) if k % 10}})
    with app.app_context():
        s = db.session.get(Survey, 1)
        f = {"all": False, "sources": ["real"], "speeder": None, "drop": False, "speed_real": True}
        table = compare.table_of(s, f)
        plain = sem_model.estimate(s, table)
        weighted = sem_model.estimate(s, table, weighted=True)
        assert weighted["weighted"] and weighted["n"] == 270 and plain["n"] == 300 and weighted["n_eff"] < 270
        assert weighted["fit"]["weighted"] and weighted["boot_key"] != plain["boot_key"]
        assert weighted["fit"]["chi2"] != plain["fit"]["chi2"]
        inp = {"key": "k", "rows": weighted["_rows"], "cols": weighted["cols"], "method": "cb", "draws": 30, "seed": 1,
               "weights": weighted["_weights"], "spec": {**weighted["spec"], "paths": [list(p) for p in weighted["spec"]["paths"]],
                                                          "indirect": [list(t) for t in weighted["spec"]["indirect"]]}}
        out = compute.runner("sem")(inp, lambda *_: None)
        assert out["failed"] == 0 and out["indirect"][0][4] > out["indirect"][0][3]
        page = client.get("/surveys/1/results?go=1").data.decode()
        assert "ungewichtet" in page[page.index('id="sem"'):]


def test_second_order_factor_and_a_competing_model(app, client, survey):
    """SEM extensions 2026-10-17, points 8 and 11: an instrument with three subscales ticked as a factor of the second
    order enters with one path (its subscales load on it); a competing model without the direct effect of the
    mediation is compared by the scaled difference test, AIC and BIC."""
    add(client, "likert")                                                                      # 2: three subscales
    form = [("code", "tech"), ("question", "Technik"), ("points", "7"), ("labeling", "full"), ("level", "metric"),
            ("subscales_shown", "1"), ("checks_shown", "1"), ("subscales", "1")]
    for k in range(9):
        form += [("stmt_id", f"s{k + 1}"), ("stmt_code", ""), ("stmt_text", f"Aussage {k + 1}"),
                 ("stmt_subscale", "ABC"[k // 3])]
    client.post(f"{E}/items/2", data=MultiDict(form))
    _scale(client, 3, "einst", 3)
    _scale(client, 4, "absicht", 3)
    rng = random.Random(12)

    def point(v):
        return max(1, min(7, int(round(4 + 1.4 * v))))

    with app.app_context():
        s = db.session.get(Survey, 1)
        for _ in range(500):
            g = rng.gauss(0, 1)
            facets = [0.7 * g + math.sqrt(0.51) * rng.gauss(0, 1) for _ in range(3)]
            einst = 0.5 * g + math.sqrt(0.75) * rng.gauss(0, 1)
            absicht = 0.6 * einst + math.sqrt(0.64) * rng.gauss(0, 1)

            def items(f, k):
                return [point(0.75 * f + math.sqrt(1 - 0.5625) * rng.gauss(0, 1)) for _ in range(k)]

            tech = [v for f in facets for v in items(f, 3)]
            s.responses.append(Response(source="real", status=Response.COMPLETED, answers={
                "2": {f"s{k + 1}": v for k, v in enumerate(tech)},
                "3": {f"s{k + 1}": v for k, v in enumerate(items(einst, 3), )},
                "4": {f"s{k + 1}": v for k, v in enumerate(items(absicht, 3))}}))
        db.session.commit()
    plan = {"research_question": "Technik", "goal": "correlation", "design": "cross",
            "h_id": ["", ""], "h_text": ["Technik über Einstellung auf Absicht", "Technik wirkt auf Einstellung"],
            "h_kind": ["mediation", "effect"], "h_dir": ["pos", "pos"], "h_iv": ["tech", "tech"],
            "h_dv": ["absicht", "einst"], "h_mv": ["einst", ""], "h_analysis": ["", "sem"], "sem_method": "cb",
            "sem_shown": "1"}
    client.post("/surveys/1/researchdesign", data=MultiDict(plan))
    page = client.get("/surveys/1/researchdesign").data.decode()
    assert 'name="sem_second" value="tech"' in page and 'name="sem_alternative"' in page
    with app.app_context():
        s = db.session.get(Survey, 1)
        ids = [h["id"] for h in study_plan.plan_of(s)["hypotheses"]]
        hid = ids[0]
        before = sem_model.build(s)
        assert len([p for p in before["spec"]["paths"] if p[1] == "einst"]) == 3  # one path per subscale
    client.post("/surveys/1/researchdesign", data=MultiDict({**plan, "h_id": ids, "sem_second": ["tech"],
                                                            "sem_alternative": [f"d:{hid}"]}))
    with app.app_context():
        s = db.session.get(Survey, 1)
        model = sem_model.build(s)
        assert model["spec"]["higher"] == {"tech": ["tech:A", "tech:B", "tech:C"]}
        assert ("tech", "einst") in model["spec"]["paths"] and ("tech", "absicht") in model["spec"]["paths"]
        text = sem_model.r_script(s, "d.rds")
        assert "F_tech =~ F_tech_A + F_tech_B + F_tech_C" in text
        assert "Konkurrenzmodell ohne: direkter Effekt in H1" in study_plan.markdown(s, "de")
        assert "Faktoren zweiter Ordnung über den Subskalen: tech" in study_plan.markdown(s, "de")
    page = client.get("/surveys/1/results?go=1").data.decode()
    part = page[page.index('id="sem"'):]
    assert "tech (zweite Ordnung)" in part and "3 Subskalen" in part
    hyp = page[page.index('id="hypothesen"'):page.index('id="sem"')]
    assert hyp.count("Unterstützt") == 2
    comp = part[part.index('id="sem-compare"'):]
    assert "ohne direkter Effekt in H1" in comp and "ΔAIC = " in comp and "Δχ²<sub>skaliert</sub>(1)" in comp
    md = client.get("/surveys/1/export/methods-section.md").data.decode()
    assert "Konkurrenzmodell ohne den direkten Effekt in H1" in md


def test_power_by_simulation_from_the_plan(app, client, survey, monkeypatch):
    """Point 16: with expected effects and a target the plan offers „Power per Simulation prüfen“; the simulation
    runs as a job and its shares stand at the plan; a changed plan needs a new one."""
    _survey(app, client, n=10)
    form = {"research_question": "TB", "goal": "correlation", "design": "cross", "sample_target": "250",
            "h_id": ["", "", ""], "h_text": ["TB wirkt auf Einstellung", "Einstellung wirkt auf Absicht", "TB über Einstellung"],
            "h_kind": ["effect", "effect", "mediation"], "h_dir": ["pos", "pos", "pos"],
            "h_iv": ["tb", "einst", "tb"], "h_dv": ["einst", "absicht", "absicht"], "h_mv": ["", "", "einst"],
            "h_analysis": ["sem", "sem", ""], "h_effect": ["0,3", "0,3", "0,3"], "sem_method": "cb", "sem_shown": "1"}
    client.post("/surveys/1/researchdesign", data=MultiDict(form))
    page = client.get("/surveys/1/researchdesign").data.decode()
    assert "Power per Simulation prüfen" in page and "N = 250" in page
    monkeypatch.setattr(sem_model, "POWER_REPS", 40)
    app.config["HB_SYNC"] = True
    client.post("/surveys/1/researchdesign/sem-power")
    page = client.get("/surveys/1/researchdesign").data.decode()
    part = page[page.index('id="plan-sem-power"'):]
    assert "40 Datensätze mit N = 250" in part and "H1 <strong>" in part and "H3 <strong>" in part
    client.post("/surveys/1/researchdesign", data=MultiDict({**form, "sample_target": "300"}))
    assert "Power per Simulation prüfen" in client.get("/surveys/1/researchdesign").data.decode()


def test_r_script_of_a_diary_study_adds_the_two_level_model(app, client, survey):
    """SEM extensions 2026-10-17, point 10 (only the R script, Ralf: as recommended): entries nested in persons get a
    two-level lavaan model with the same paths on both levels, without labels that would make them equal."""
    _survey(app, client, n=10)
    with app.app_context():
        s = db.session.get(Survey, 1)
        plain = sem_model.r_script(s, "d.rds")
        assert "level: 1" not in plain
        s.sampling = {"kind": "diary"}
        text = sem_model.r_script(s, "d.rds")
        assert "level: 1" in text and "level: 2" in text and 'cluster = "person"' in text
        two = text[text.index("modell_2e <- '"):]
        assert "p1*" not in two[:two.index("fit_2e")] and "F_einst ~ F_tb" in two


def test_plan_selects_keep_their_name_with_learning_aids(app, client, survey):
    """UI/UX audit 2026-10-17, W-A10: with the learning aids on, the glossary "?" of "Fehlende Werte", "Antworten der
    Aussagen" and "Power" stands beside the label, not in it: a label binds its first control, which was the "?"
    button, and left the select without a name."""
    import re

    _survey(app, client, n=10)
    page = client.get("/surveys/1/researchdesign").data.decode()
    assert 'data-term-name' in page                                                  # the aids are on
    for name, field_id, text in (("sem_missing", "sem-missing", "Fehlende Werte"),
                                 ("sem_ordinal", "sem-ordinal", "Antworten der Aussagen")):
        assert f'<select name="{name}" id="{field_id}">' in page
        assert f'<label for="{field_id}">{text}</label>' in page
    powers = re.findall(r'<label for="(h-power-[^"]+)">Power</label>', page)
    assert powers and all(f'name="h_power" id="{p}"' in page for p in powers)
    # no label holds both a button and a field any more
    for label in re.findall(r"<label\b[^>]*>(.*?)</label>", page, re.S):
        assert not ("<button" in label and re.search(r"<(select|input|textarea)\b", label)), label[:200]
