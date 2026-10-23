"""Market simulator (CBC) and "Auswahl aus Aussagen" (MaxDiff): app/choice_predict.py, app/simulator.py."""
import io
import math
import random
import re

import numpy as np
import pytest

from app.extensions import db
from app.models import Item, Response, Survey, User
from tests.analysis.test_cbc_results import setup as cbc_setup
from tests.analysis.test_conjoint_hb import HB_FAST, simulate_segments

E = "/surveys/1/editor"


# --- the arithmetic ----------------------------------------------------------------------------------------------

def _two_persons():
    """One attribute with two levels and a linear price (centre 10): person 1 likes level 1 a lot, person 2 less and
    is more price sensitive."""
    return [{"pw": {0: [1.0, -1.0]}, "price": -0.1, "center": 10.0, "asc": 0.2},
            {"pw": {0: [0.5, -0.5]}, "price": -0.2, "center": 10.0, "asc": -3.0}]


def test_shares_by_hand():
    from app import choice_predict as cp

    people = _two_persons()
    U = np.array([cp.task_utilities(m, [[0, None], [1, None]], [10, 5]) for m in people])
    # person 1: A = 1, B = -1 + (-0.1)(5 - 10) = -0.5; person 2: A = 0.5, B = -0.5 + (-0.2)(-5) = 0.5 (a tie)
    assert np.allclose(U, [[1.0, -0.5], [0.5, 0.5]])
    P = np.array([cp.shares(list(u)) for u in U])
    p1 = math.exp(1) / (math.exp(1) + math.exp(-0.5))
    assert np.allclose(P, [[p1, 1 - p1], [0.5, 0.5]])
    mean, lo, hi = cp.mean_shares(P)
    assert np.allclose(mean, [(p1 + 0.5) / 2, (1 - p1 + 0.5) / 2])
    assert lo is not None and np.all(lo <= mean + 1e-12) and np.all(mean <= hi + 1e-12)
    assert np.allclose(cp.mean_shares(P)[1], lo)                              # fixed seed: the same interval
    # "none" is the constant alone, as the last column
    U3 = np.array([cp.task_utilities(m, [[0, None], [1, None]], [10, 5], with_none=True) for m in people])
    assert U3.shape == (2, 3) and np.allclose(U3[:, 2], [0.2, -3.0])


def test_first_choice_splits_ties():
    from app import choice_predict as cp

    F = np.array([cp.first_choice(u) for u in ([1.0, -0.5], [0.5, 0.5], [0.0, 2.0])])
    assert np.allclose(F, [[1, 0], [0.5, 0.5], [0, 1]])
    assert np.allclose(cp.mean_shares(F, draws=0)[0], [1.5 / 3, 1.5 / 3])


def test_weighted_mean_and_single_row():
    from app import choice_predict as cp

    P = np.array([[0.8, 0.2], [0.4, 0.6]])
    mean, _, _ = cp.mean_shares(P, weights=[3, 1])
    assert np.allclose(mean, [(3 * 0.8 + 0.4) / 4, (3 * 0.2 + 0.6) / 4])
    assert np.allclose(cp.mean_shares(P, weights=[1, 0])[0], [0.8, 0.2])    # weight 0: left out
    mean, lo, hi = cp.mean_shares(P[:1])                                    # the pooled model: one row, no interval
    assert np.allclose(mean, [0.8, 0.2]) and lo is None and hi is None


def test_price_between_levels_is_linear():
    from app import choice_predict as cp

    model = _two_persons()[0]
    u = lambda p: cp.concept_utility(model, [0, None], p)
    assert np.allclose(u(7.5), (u(5) + u(10)) / 2) and np.allclose(u(12.5), (u(10) + u(15)) / 2)


def _hb_shares(result, profiles, prices, with_none=False, rows=None):
    """Mean logit shares (in percent) by hand from the kept draws of a stored HB result: per person and draw the
    logit rule, averaged over the draws, then over the persons (audit 2026-10-05, W-1)."""
    from app import conjoint_hb

    par = conjoint_hb.draw_parameters(result)
    spans = {a: start for a, start, _ in result["spans"]}
    center = result["coding"]["price_center"]
    U = []
    for profile, x in zip(profiles, prices):
        u = sum(par["util"][:, :, spans[a] + lv] for a, lv in enumerate(profile) if lv is not None and a in spans)
        U.append(u + (par["price"] * (x - center) if par["price"] is not None else 0))
    if with_none:
        U.append(par["asc"])
    U = np.stack(U, axis=2)
    E = np.exp(U - U.max(axis=2, keepdims=True))
    P = (E / E.sum(axis=2, keepdims=True)).mean(axis=0)
    return (P if rows is None else P[rows]).mean(axis=0) * 100


# --- the page ------------------------------------------------------------------------------------------------------

def _hb(app, client):
    cbc_setup(client, price_attr=True, none=True, tasks=12, versions=6)
    app.config.update(HB_FAST)
    simulate_segments(app, 40, none_asc=-0.5)
    assert client.post("/surveys/1/results/hb/2", data={"source": "real", "speeder": "50"}).status_code == 302


def _shares(html):
    bars = re.search(r'<ul class="plain res-bars sim-bars">(.*?)</ul>', html, re.S).group(1)
    return [(name, float(x.replace(",", "."))) for name, x in
            re.findall(r'<span class="res-bar-label">([^<]+?)\s*(?:<span class="hint">[^<]*</span>)?</span>.*?'
                       r'<span class="res-bar-n">([\d,]+) %', bars, re.S)]


def test_pooled_without_hb_and_hb_page(app, client, survey):
    cbc_setup(client, price_attr=True, none=True, tasks=12, versions=6)
    simulate_segments(app, 40, none_asc=-0.5)
    html = client.get("/surveys/1/results").data.decode()
    assert 'id="sim-2"' in html and "Marktsimulator" in html and "aggregierten Logit-Modell" in html
    assert "sim-result" not in html and "First Choice" not in html          # no persons: no first choice
    app.config.update(HB_FAST)
    client.post("/surveys/1/results/hb/2", data={"source": "real", "speeder": "50"})
    query = ("sim=2&sim_n=3&sim_0_0=0&sim_0_2=0&sim_0_price=75&sim_1_0=1&sim_1_2=0&sim_1_price=50"
             "&sim_2_0=2&sim_2_2=1&sim_2_price=150&sim_none=1")
    html = client.get("/surveys/1/results?" + query).data.decode()
    assert "aggregierten Logit-Modell. Es behandelt" not in html and "individuellen HB-Nutzen von 40 Personen" in html
    assert re.search(r'<section class="sim-fold sim-section" id="sim-2"[^>]*>', html)
    shares = _shares(html)
    assert len(shares) == 4 and shares[-1][0] == "Keine davon"               # "none" last
    assert [x for _, x in shares[:-1]] == sorted([x for _, x in shares[:-1]], reverse=True)
    assert abs(sum(x for _, x in shares) - 100) < 0.3
    assert "First Choice" in html and "Präferenzanteile sind keine Marktanteile" in html
    assert re.search(r"\d+,\d %", html)                                       # one decimal
    assert "Orme, B. K. (2019). Market simulators for conjoint analysis" in html and "Efron" in html
    assert 'name="sim_0_price" value="75"' in html and re.search(r'min="50(\.0)?" max="150(\.0)?"', html)
    # the same numbers as the computation by hand from the stored utilities
    with app.app_context():
        from app import conjoint_hb

        mean = _hb_shares(conjoint_hb.load(db.session.get(Item, 2)), [[0, None, 0], [1, None, 0], [2, None, 1]],
                          [75, 50, 150], with_none=True)
    product_a = next(x for name, x in shares if name.startswith("Produkt A"))
    assert abs(product_a - mean[0]) < 0.051
    # outside the price range: locked, nothing computed
    html = client.get("/surveys/1/results?" + query.replace("sim_0_price=75", "sim_0_price=160")).data.decode()
    assert "außerhalb der Preisspanne der Befragung (50 bis 150)" in html and "sim-bars" not in html
    html = client.get("/surveys/1/results?" + query.replace("sim_0_price=75", "sim_0_price=49,5")).data.decode()
    assert "außerhalb der Preisspanne" in html
    # add and drop products: 1 to 6
    html = client.get("/surveys/1/results?" + query + "&sim_act=add").data.decode()
    assert "Produkt D" in html
    html = client.get("/surveys/1/results?" + query + "&sim_act=drop0").data.decode()
    assert "Produkt C" not in html and len(_shares(html)) == 3


FETCH = {"X-Requested-With": "fetch"}


def test_fold_alone_without_the_page(app, client, survey):
    """static/simulator.js: the fold alone from the kept results (no page, nothing computed again); one product
    left; without kept results 409 (the script loads the page); the scenario forms answer with JSON."""
    _hb(app, client)
    client.get("/surveys/1/results")                                          # the kept results
    q = "sim=2&sim_n=2&sim_0_0=0&sim_0_2=0&sim_0_price=75&sim_1_0=1&sim_1_2=0&sim_1_price=50"
    r = client.get("/surveys/1/results?" + q + "&sim_part=2", headers=FETCH)
    html = r.data.decode()
    assert r.status_code == 200 and html.strip().startswith('<section class="sim-fold sim-section" id="sim-2"')
    assert "<html" not in html and "Literatur zu dieser Seite" not in html and len(_shares(html)) == 2
    assert 'data-sim-drop="Produkt B"' in html
    # down to one product (and not below)
    html = client.get("/surveys/1/results?" + q + "&sim_act=drop1&sim_part=2", headers=FETCH).data.decode()
    assert "Produkt B" not in html and len(_shares(html)) == 1 and "immer bei 100 %" in html
    assert "data-sim-drop" not in html
    one = "sim=2&sim_n=1&sim_0_0=0&sim_0_2=0&sim_0_price=75"
    html = client.get(f"/surveys/1/results?{one}&sim_act=drop0&sim_part=2", headers=FETCH).data.decode()
    assert "Produkt A" in html and len(_shares(html)) == 1
    # another pool, nothing kept: the question alone is computed, never the page (Ralf, 2026-10-05)
    r = client.get("/surveys/1/results?all=1&" + q + "&sim_act=opt&sim_part=2", headers=FETCH)
    assert r.status_code == 200 and r.data.decode().strip().startswith("<section") and b"sim-opt-table" in r.data
    # the page that computes first keeps the scenario, so the simulator opens again
    app.config["RESULTS_WAIT_PAGE"] = True
    html = client.get("/surveys/1/results?all=1&" + q).data.decode()
    assert 'action="#sim-2"' in html and 'name="sim_0_price" value="75"' in html and 'name="sim_part"' not in html
    # saving in the background: JSON, the script shows the saved scenario
    r = client.post("/surveys/1/results/sim/2/save", data={"sim": "2", "sim_n": "2", "sim_name": "Basis"}, headers=FETCH)
    assert r.get_json() == {"message": "Szenario „Basis“ gespeichert.", "error": False, "load": 0}
    r = client.post("/surveys/1/results/sim/2/delete", data={"sim_delete": "0"}, headers=FETCH)
    assert r.get_json()["message"] == "Szenario „Basis“ gelöscht."


def test_optimise_one_product(app, client, survey):
    """"Produkt optimieren": the price held (default), every combination of the other attributes against product B
    and "none"; the best by hand from the kept draws."""
    _hb(app, client)
    q = "sim=2&sim_n=2&sim_0_0=0&sim_0_2=0&sim_0_price=75&sim_1_0=1&sim_1_2=0&sim_1_price=50&sim_none=1"
    html = client.get("/surveys/1/results?" + q + "&sim_act=opt").data.decode()
    assert "Varianten von Produkt A durchgerechnet" in html and 'name="sim_opt" value="1"' in html
    assert "Kohli, R., &amp; Krishnamurti, R. (1987)" in html
    rows = re.findall(r"<tr(?: class=\"sim-opt-current\")?><td>(\d+)</td>.*?<td>([\d,]+) %</td>", html, re.S)
    with app.app_context():
        from app import conjoint_hb

        item = db.session.get(Item, 2)
        counts = item.kind.level_counts(item)
        result = conjoint_hb.load(item)
        best = max(_hb_shares(result, [[a, None, b], [1, None, 0]], [75, 50], with_none=True)[0]
                   for a in range(counts[0]) for b in range(counts[2]))
    assert f"{counts[0] * counts[2]} Varianten" in html
    assert rows[0][0] == "1" and abs(float(rows[0][1].replace(",", ".")) - best) < 0.051
    # a variant becomes the product
    html = client.get("/surveys/1/results?" + q + "&sim_opt=1&sim_act=take0:2,2:1").data.decode()
    assert re.search(r'name="sim_0_0"[^>]*>.*?<option value="2" selected>', html, re.S)
    # alone, without "none": by the mean utility; everything held: nothing to compute
    one = "sim=2&sim_n=1&sim_0_0=0&sim_0_2=0&sim_0_price=75&sim_act=opt"
    assert "nach dem mittleren Nutzen" in client.get("/surveys/1/results?" + one).data.decode()
    html = client.get("/surveys/1/results?" + one + "&sim_fixset=1&sim_fix=0&sim_fix=1&sim_fix=2").data.decode()
    assert "Alle Eigenschaften sind festgehalten" in html
    # off again
    assert "Varianten von" not in client.get("/surveys/1/results?" + q + "&sim_opt=1&sim_act=optoff").data.decode()
    # ticked while off: stays ticked (Ralf, 2026-10-05: the tick vanished at once)
    html = client.get("/surveys/1/results?" + q + "&sim_fixset=1&sim_fix=0&sim_fix=2").data.decode()
    assert 'name="sim_opt" value="1"' not in html
    assert re.search(r'name="sim_fix" value="0" checked', html) and re.search(r'name="sim_fix" value="2" checked', html)
    assert not re.search(r'name="sim_fix" value="1" checked', html)


def test_weighting_acts_on_the_mean_over_persons(app, client, survey):
    _hb(app, client)
    with app.app_context():
        from app import conjoint_hb, results, simulator

        item = db.session.get(Item, 2)
        data = results.compute(db.session.get(Survey, 1))
        entry = next(e for e in data["items"] if e["item"].id == 2)["cbc"]
        state, _ = simulator.cbc_state(item, {}, own=False)
        pids = conjoint_hb.load(item)["pids"]
        half = {pid: 2.0 for pid in pids[:20]}                                # the other half has no weight
        weighted = simulator.cbc_compute(item, entry, state, "de", {"pids": half})
        plain = simulator.cbc_compute(item, entry, state, "de")
        assert weighted["persons"] == 20 and plain["persons"] == 40
        expect = _hb_shares(conjoint_hb.load(item), [p["levels"] for p in state["products"]],
                            [p["price"] for p in state["products"]], rows=slice(0, 20))
        got = {r["label"]["name"]: float(r["sop"].replace(",", ".")) for r in weighted["rows"]}
        assert abs(got["Produkt A"] - expect[0]) < 0.051 and abs(got["Produkt B"] - expect[1]) < 0.051
        assert simulator.cbc_compute(item, entry, state, "de", {"pids": {}})["error"]


def test_save_undo_and_word(app, client, survey):
    _hb(app, client)
    with app.app_context():
        from app import results_store

        s = db.session.get(Survey, 1)
        pool = {"sources": ["real"], "all": False, "drop": True, "speeder": 50}
        before = results_store.fingerprint(s, pool)
    form = {"sim": "2", "sim_n": "2", "sim_0_0": "1", "sim_0_2": "0", "sim_0_price": "80", "sim_1_0": "0",
            "sim_1_2": "1", "sim_1_price": "120", "sim_name": "Basis", "back": "gewichtet=1&evil=x"}
    r = client.post("/surveys/1/results/sim/2/save", data=form)
    assert r.status_code == 302 and "sim_load=0" in r.location and "gewichtet=1" in r.location and "evil" not in r.location
    with app.app_context():
        from app import results_store, simulator

        item = db.session.get(Item, 2)
        saved = simulator.saved(item)
        look = item.config["attributes"][0]
        assert saved == [{"name": "Basis", "none": False, "products": [
            {"levels": {look["id"]: look["levels"][1]["id"], item.config["attributes"][2]["id"]:
                        item.config["attributes"][2]["levels"][0]["id"]}, "price": 80.0},
            {"levels": {look["id"]: look["levels"][0]["id"], item.config["attributes"][2]["id"]:
                        item.config["attributes"][2]["levels"][1]["id"]}, "price": 120.0}]}]
        # no change of the questionnaire: results cache and HB key stay
        assert results_store.fingerprint(db.session.get(Survey, 1), pool) == before
    html = client.get("/surveys/1/results").data.decode()
    assert "<strong>fertig</strong>" in html and ">Basis</a>" in html            # HB still current; first scenario shown
    assert 'name="sim_0_price" value="80"' in html
    # a second one, then the simulator's own undo and redo
    client.post("/surveys/1/results/sim/2/save", data={**form, "sim_name": "Teuer", "sim_0_price": "150"})
    with app.app_context():
        from app import simulator

        assert [x["name"] for x in simulator.saved(db.session.get(Item, 2))] == ["Basis", "Teuer"]
    client.post("/surveys/1/results/sim/2/undo", data={})
    with app.app_context():
        assert [x["name"] for x in simulator.saved(db.session.get(Item, 2))] == ["Basis"]
    client.post("/surveys/1/results/sim/2/undo", data={})
    with app.app_context():
        assert [x["name"] for x in simulator.saved(db.session.get(Item, 2))] == ["Basis", "Teuer"]
    client.post("/surveys/1/results/sim/2/undo", data={})
    with app.app_context():
        assert [x["name"] for x in simulator.saved(db.session.get(Item, 2))] == ["Basis"]
        from app import undo

        steps = undo.history(db.session.get(Survey, 1), db.session.get(User, 1))["undo"]
        assert not any("Szenari" in s.label for s in steps)                  # no step for the arrows of the editor
    # Word: the saved scenarios as tables
    from docx import Document

    r = client.get("/surveys/1/results/scales.docx?lang=de")
    doc = Document(io.BytesIO(r.data))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Marktsimulator" in text and "Szenario „Basis“" in text and "keine Marktanteile" in text
    assert doc.tables and doc.tables[0].rows[0].cells[1].text == "Share of Preference"
    assert len(doc.tables[0].rows) == 3 and re.match(r"\d+,\d \[\d+,\d; \d+,\d\]", doc.tables[0].rows[1].cells[1].text)
    # delete
    client.post("/surveys/1/results/sim/2/delete", data={"sim_delete": "0"})
    with app.app_context():
        assert simulator.saved(db.session.get(Item, 2)) == []


def test_scenarios_allowed_after_start_and_travel_with_copies(app, client, survey):
    _hb(app, client)
    client.post("/surveys/1/results/sim/2/save", data={"sim": "2", "sim_n": "2", "sim_name": "Feld"})
    with app.app_context():
        from app import simulator

        s = db.session.get(Survey, 1)
        s.status = Survey.RUNNING
        db.session.commit()
    r = client.post("/surveys/1/results/sim/2/save", data={"sim": "2", "sim_n": "2", "sim_name": "Nach Start"})
    assert r.status_code == 302
    with app.app_context():
        item = db.session.get(Item, 2)
        assert [x["name"] for x in simulator.saved(item)] == ["Feld", "Nach Start"]
        # survey files: the rule keeps the known shape and drops the rest
        config, unfit, unknown = item.kind.import_config({**item.config})
        assert "sim" not in unfit and "sim" not in unknown
        assert [x["name"] for x in config["sim"]["scenarios"]] == ["Feld", "Nach Start"] and config["sim"]["prev"] is None
        assert simulator.clean({"scenarios": [{"name": 3}, {"name": "x", "products": [{"levels": {"a": 1}}]}]}) == \
            {"scenarios": [{"name": "x", "products": [{"levels": {}, "price": None}], "none": False}], "prev": None}
        assert simulator.clean("kaputt") == {"scenarios": [], "prev": None}


# --- MaxDiff ---------------------------------------------------------------------------------------------------

def _maxdiff(app, client, persons=40):
    from tests.analysis.test_maxdiff import setup

    setup(client)
    rng = random.Random(5)
    with app.app_context():
        s = db.session.get(Survey, 1)
        item = db.session.get(Item, 2)
        sets = item.config["design"]["versions"][0]
        true = [1.5 - 0.4 * i for i in range(len(item.config["statements"]))]
        for _ in range(persons):
            u = [x + rng.gauss(0, 0.5) for x in true]
            answers = {"v": 0}
            for t, st in enumerate(sets):
                g = {i: u[i] - math.log(-math.log(rng.random())) for i in st}
                b = max(st, key=g.get)
                w = min((i for i in st if i != b), key=g.get)
                answers[f"t{t}"] = {"b": b, "w": w}
            s.responses.append(Response(answers={"2": answers}, layout={"cj": {"2": {"version": 0}}},
                                        status=Response.COMPLETED))
        db.session.commit()


def test_maxdiff_selection_pooled_and_hb(app, client, survey):
    _maxdiff(app, client)
    html = client.get("/surveys/1/results?sim=2&sim_s=0&sim_s=1&sim_s=3").data.decode()
    assert "Auswahl aus Aussagen" in html and "aggregierten Logit-Modell" in html
    shares = _shares(html)
    assert len(shares) == 3 and abs(sum(x for _, x in shares) - 100) < 0.2
    assert [x for _, x in shares] == sorted((x for _, x in shares), reverse=True)  # highest share first
    with app.app_context():
        from app import results

        entry = next(e for e in results.compute(db.session.get(Survey, 1))["items"] if e["item"].id == 2)["maxdiff"]
        u = {r["n"] - 1: r["u"] for r in entry["mnl"]["rows"]}
        e = np.exp([u[0], u[1], u[3]])
    got = {name.split(".")[0]: x for name, x in shares}
    assert abs(got["1"] - 100 * e[0] / e.sum()) < 0.051
    html = client.get("/surveys/1/results?sim=2&sim_s=0").data.decode()
    assert "Bitte 2 bis 10 Aussagen ankreuzen" in html
    app.config.update(HB_FAST)
    client.post("/surveys/1/results/hb/2", data={"source": "real", "speeder": "50"})
    html = client.get("/surveys/1/results?sim=2&sim_s=0&sim_s=1&sim_s=3").data.decode()
    assert "individuellen HB-Nutzen von 40 Personen" in html and "First Choice" in html
    with app.app_context():
        from app import maxdiff_hb

        res = maxdiff_hb.load(db.session.get(Item, 2))
        util = maxdiff_hb.draw_utilities(res)                                   # draws x persons x statements
        assert util.shape[0] > 1 and util.shape[1] == 40
        col = {i: c for c, i in enumerate(res["est"])}
        U = util[:, :, [col[0], col[1], col[3]]]
        P = (np.exp(U) / np.exp(U).sum(axis=2, keepdims=True)).mean(axis=0)  # per person over the draws (W-1)
    assert abs({n.split(".")[0]: x for n, x in _shares(html)}["1"] - 100 * P[:, 0].mean()) < 0.051
    client.post("/surveys/1/results/sim/2/save", data={"sim": "2", "sim_s": ["0", "3"], "sim_name": "Zwei"})
    with app.app_context():
        from app import simulator

        item = db.session.get(Item, 2)
        ids = [x["id"] for x in item.config["statements"]]
        assert simulator.saved(item) == [{"name": "Zwei", "statements": [ids[0], ids[3]]}]


@pytest.mark.parametrize("weights,expected", [(None, 0.5), ([1, 3], 0.25)])
def test_first_choice_mean_with_weights(weights, expected):
    from app import choice_predict as cp

    F = np.array([cp.first_choice(u) for u in ([2.0, 1.0], [0.0, 1.0])])
    assert np.isclose(cp.mean_shares(F, weights, draws=0)[0][0], expected)


# --- ACBC ------------------------------------------------------------------------------------------------------------

def test_acbc_on_hb_utilities_with_list_price(app, client, survey):
    """ACBC: the simulator only on the HB utilities, the summed price as its own row (list price when emptied, within
    the prices shown), no "Keine davon"; shares by hand from the kept draws; "Produkt optimieren" gives each variant
    its list price."""
    from tests.analysis import test_acbc_results as acbc

    acbc.setup(client, screens=5, per_screen=4, price=True)
    app.config.update(HB_FAST)
    with app.app_context():
        acbc.add_responses(db.session.get(Survey, 1), acbc.simulate(db.session.get(Item, 2), 60, price_beta=-0.01))
    html = client.get("/surveys/1/results").data.decode()
    assert 'id="sim-2"' not in html                                              # no HB yet: no simulator
    # ... but its place and the script: after the estimation in place (static/hb.js) it is fetched (Ralf, 2026-10-06)
    assert 'data-sim-slot="2"' in html and "simulator.js" in html
    client.post("/surveys/1/results/hb/2", data={"source": "real", "speeder": "50"})
    html = client.get("/surveys/1/results").data.decode()
    assert 'id="sim-2"' in html and "data-sim-summed" in html and 'name="sim_0_price"' in html
    assert "„Keine davon“ anbieten" not in html and "keine Wahl von nichts" in html
    with app.app_context():
        from app import conjoint_hb

        item = db.session.get(Item, 2)
        result = conjoint_hb.load(item)
        lo, hi = result["price_range"]
        listed = [item.kind.byo_price(item, [0, 0, 0, 0]), item.kind.byo_price(item, [1, 1, 1, 1])]
    price = round((lo + hi) / 2, 2)
    q = (f"sim=2&sim_n=2&sim_0_0=0&sim_0_1=0&sim_0_2=0&sim_0_3=0&sim_0_price={price}"
         "&sim_1_0=1&sim_1_1=1&sim_1_2=1&sim_1_3=1&sim_1_price=")
    html = client.get("/surveys/1/results?" + q).data.decode()
    shares = _shares(html)
    assert len(shares) == 2 and "First Choice" in html
    with app.app_context():
        expect = _hb_shares(conjoint_hb.load(db.session.get(Item, 2)), [[0, 0, 0, 0], [1, 1, 1, 1]],
                            [price, min(max(listed[1], lo), hi)])
    got = {name: x for name, x in shares}
    assert abs(got["Produkt A"] - expect[0]) < 0.051                            # B: its list price (field emptied)
    html = client.get("/surveys/1/results?" + q.replace(f"sim_0_price={price}", f"sim_0_price={hi + 50}")).data.decode()
    assert "außerhalb der Preisspanne" in html
    html = client.get("/surveys/1/results?" + q + "&sim_act=opt").data.decode()
    assert "Varianten von Produkt A durchgerechnet" in html and "<th>Preis</th>" in html
    assert re.search(r'value="take0:\d,1:\d,2:\d,3:\d,p:[\d.]+"', html)
    # saving keeps the price
    client.post("/surveys/1/results/sim/2/save", data={"sim": "2", "sim_n": "1", "sim_0_price": str(price),
                                                        "sim_name": "Mitte"})
    with app.app_context():
        from app import simulator

        assert simulator.saved(db.session.get(Item, 2))[0]["products"][0]["price"] == price


# --- math audit 2026-10-05 ---------------------------------------------------------------------------------------------

def test_bootstrap_means_are_paired_and_give_the_interval():
    """M-1: the samples of mean_shares for every column at once: the same interval, and identical columns are the
    best in half of the samples each with a distance of exactly 0."""
    from app import choice_predict as cp

    rng = np.random.default_rng(3)
    P = rng.uniform(size=(50, 3))
    M = cp.bootstrap_means(P)
    mean, lo, hi = cp.mean_shares(P)
    assert M.shape == (cp.DRAWS, 3) and np.allclose(np.percentile(M, [2.5, 97.5], axis=0), [lo, hi])
    twins = cp.bootstrap_means(np.column_stack([P[:, 0], P[:, 0]]))
    assert np.allclose(twins[:, 0] - twins[:, 1], 0)


def test_optimise_shows_how_often_a_variant_is_best_and_its_paired_distance(app, client, survey):
    """M-1: "Produkt optimieren" no longer gives an interval per variant: per variant how often it is the best in the
    bootstrap samples and its distance to rank 1 with a paired interval; the sentence on the optimizer's curse."""
    _hb(app, client)
    q = "sim=2&sim_n=2&sim_0_0=0&sim_0_2=0&sim_0_price=75&sim_1_0=1&sim_1_2=0&sim_1_price=50&sim_none=1"
    with app.app_context():
        from werkzeug.datastructures import MultiDict

        from app import results, simulator

        item = db.session.get(Item, 2)
        data = results.compute(db.session.get(Survey, 1))
        entry = next(e for e in data["items"] if e["item"].id == 2)["cbc"]
        state, _ = simulator.cbc_state(item, MultiDict([tuple(x.split("=")) for x in q.split("&")] + [("sim_act", "opt")]))
        opt = simulator.cbc_optimize(item, entry, state, "de")
        rows = opt["rows"]
        assert opt["paired"] and "gap" not in rows[0] and all("gap" in r for r in rows[1:])
        best = [float(r["best"].replace(",", ".").rstrip(" %")) for r in rows]
        assert sum(best) <= 100.1 and best[0] > 0                         # rank 1 need not be the best most often
        for r in rows[1:]:
            lo, hi = (float(x.replace(",", ".").replace("−", "-").replace("+", "")) for x in (r["gap_lo"], r["gap_hi"]))
            assert lo <= hi and r["tie"] == (hi >= 0)
        assert "sop_lo" not in rows[0]
    html = client.get("/surveys/1/results?" + q + "&sim_act=opt").data.decode()
    assert "Abstand zu Rang 1" in html and "der Stichproben" in html and "fällt deshalb eher zu hoch aus" in html
    assert "Smith, J. E., &amp; Winkler, R. L. (2006)" in html


def test_similar_products_none_and_linear_price_notes(app, client, survey):
    """M-2: two products alike in all but one attribute: the hint that share of preference overstates them together,
    also with HB (identical products: SoP of the pair above the single product's share, first choice not). M-6: the
    note on "Keine davon" and the test of the linear price."""
    _hb(app, client)
    two = "sim=2&sim_n=2&sim_0_0=0&sim_0_2=0&sim_0_price=75&sim_1_0=1&sim_1_2=1&sim_1_price=50"
    three = two.replace("sim_n=2", "sim_n=3") + "&sim_2_0=0&sim_2_2=0&sim_2_price=75"   # C = A
    html = client.get("/surveys/1/results?" + two).data.decode()
    assert "Ähnliche Produkte" not in html and "keine Nichtkaufquote" not in html
    html = client.get("/surveys/1/results?" + three + "&sim_none=1").data.decode()
    assert "Ähnliche Produkte (A und C)" in html and "First Choice über die HB-Ziehungen ist hier die verlässlichere" in html
    assert "keine Nichtkaufquote" in html
    with app.app_context():
        from werkzeug.datastructures import MultiDict

        from app import results, simulator

        item = db.session.get(Item, 2)
        entry = next(e for e in results.compute(db.session.get(Survey, 1))["items"] if e["item"].id == 2)["cbc"]

        def shares(query):
            out = simulator.cbc_compute(item, entry, simulator.cbc_state(
                item, MultiDict([tuple(x.split("=")) for x in query.split("&")]))[0], "de")
            return ({r["label"]["name"]: float(r["sop"].replace(",", ".")) for r in out["rows"]},
                    {r["label"]["name"]: float(r["fc"].replace(",", ".")) for r in out["rows"]})

        sop2, fc2 = shares(two)
        sop3, fc3 = shares(three)
        assert sop3["Produkt A"] + sop3["Produkt C"] > sop2["Produkt A"] + 1          # IIA: the twin takes from B
        assert abs(fc3["Produkt A"] + fc3["Produkt C"] - fc2["Produkt A"]) < 0.2     # first choice: shared
        model = entry["model"]
        lt = model["price_linear"]
        assert lt["df"] == 1 and lt["chi2"] >= 0 and 0 <= lt["p"] <= 1                # three price levels
    html = client.get("/surveys/1/results").data.decode()
    assert "Linear oder in Stufen? Likelihood-Quotienten-Test" in html and "zwischen den Stufen linear angenommen" in html


def test_price_linearity_by_hand(app, client, survey):
    """M-6: the likelihood ratio of the linear price against a part-worth per price level, by hand."""
    from app.choice_model import fit

    cbc_setup(client, price_attr=True, none=True, tasks=12, versions=6)
    simulate_segments(app, 40, none_asc=-0.5)
    with app.app_context():
        from app import cbc_analysis, export, group_tests as gt
        from app.choice_model import responses_for

        item = db.session.get(Item, 2)
        responses = responses_for(item, export.rows(db.session.get(Survey, 1), with_hb=False))
        model = cbc_analysis.estimate(item, responses, "de")
        sets, coding, _, _ = cbc_analysis.choice_sets(item, responses, price_in_levels=True)
        assert coding.price is None and coding.levels[item.kind.price_attribute(item)] == 3
        ll = fit(sets, coding.k)[2]
        chi2 = 2 * (ll - model["loglik"])
        assert abs(model["price_linear"]["chi2"] - max(chi2, 0)) < 1e-6
        assert abs(model["price_linear"]["p"] - gt.chi2_sf(max(chi2, 0), 1)) < 1e-9


def test_acbc_optimise_leaves_out_list_prices_outside_the_shown_prices(app, client, survey, monkeypatch):
    """M-4: a variant whose list price lies above the highest price shown was computed and shown with that highest
    price; now it is left out and counted."""
    from tests.analysis import test_acbc_results as acbc

    acbc.setup(client, screens=5, per_screen=4, price=True)
    app.config.update(HB_FAST)
    with app.app_context():
        acbc.add_responses(db.session.get(Survey, 1), acbc.simulate(db.session.get(Item, 2), 60, price_beta=-0.01))
    client.post("/surveys/1/results/hb/2", data={"source": "real", "speeder": "50"})
    with app.app_context():
        from werkzeug.datastructures import MultiDict

        from app import conjoint_hb, results, simulator

        item = db.session.get(Item, 2)
        lo, hi = conjoint_hb.load(item)["price_range"]
        # level 1 of the first attribute costs more than any price shown
        monkeypatch.setattr(simulator, "_list_price", lambda it, levels: hi + 100 if levels[0] == 1 else lo)
        entry = next(e for e in results.compute(db.session.get(Survey, 1))["items"] if e["item"].id == 2)["acbc"]
        args = MultiDict([("sim", "2"), ("sim_n", "2"), ("sim_0_price", str(lo)), ("sim_1_price", str(lo)),
                          ("sim_act", "opt"), ("sim_fixset", "1"), ("sim_fix", "2"), ("sim_fix", "3")])
        state, _ = simulator.cbc_state(item, args)
        opt = simulator.cbc_optimize(item, entry, state, "de")
        counts = item.kind.level_counts(item)
        assert opt["outside"] == counts[1]                                   # every variant with level 1 of attribute 0
        assert all(r["take"].split(",")[0] != "0:1" for r in opt["rows"])
        assert all(r["levels"][-1] == simulator._money(lo) for r in opt["rows"])
        monkeypatch.setattr(simulator, "_list_price", lambda it, levels: hi + 100)
        assert "außerhalb der gezeigten Preise" in simulator.cbc_optimize(item, entry, state, "de")["error"]


def test_effective_persons_of_a_segment():
    """M-5: the memberships (0.9, 0.8, 0.5, 0.2, 0, 0.3) are 2.7 persons, effectively 4.0; two sure persons among 28
    almost outside: 30 with a weight above 0, effectively 2."""
    from app.simulator import _effective

    assert np.isclose(_effective([0.9, 0.8, 0.5, 0.2, 0.0, 0.3]), 2.7 ** 2 / 1.83)
    assert abs(_effective([1.0, 1.0] + [1e-4] * 28) - 2.0) < 0.01


# --- audit 2026-10-06 (L-1, L-4, L-5) ------------------------------------------------------------------------------

def test_price_levels_of_a_price_for_prohibitions():
    """L-1: the price stands for its level; a price between two levels for both neighbours; outside none."""
    from app.simulator import _price_levels

    assert _price_levels([50.0, 100.0, 150.0], 100.0) == [1]
    assert _price_levels([50.0, 100.0, 150.0], 75.0) == [0, 1]
    assert _price_levels([150.0, 50.0, 100.0], 120.0) == [0, 2]
    assert _price_levels([50.0, 100.0, 150.0], 200.0) == [2]
    assert _price_levels([50.0, 100.0, 150.0], None) == []


def test_optimise_leaves_out_prohibitions_with_a_price_level(app, client, survey, monkeypatch):
    """L-1, L-4 (kombi/k1): a prohibition "level of attribute 0 x price 50" held for no variant, because the price
    stood as None in the checked profile; the best variant shown could be a forbidden one, and "{n} Varianten" counted
    the forbidden ones. Now: free price, held price at the level and held price between two levels."""
    _hb(app, client)
    with app.app_context():
        from werkzeug.datastructures import MultiDict

        from app import results, simulator

        item = db.session.get(Item, 2)
        entry = next(e for e in results.compute(db.session.get(Survey, 1))["items"] if e["item"].id == 2)["cbc"]
        pa, numbers = simulator._price(item)
        assert pa == 1 and sorted(numbers) == [50.0, 100.0, 150.0]
        cheap = numbers.index(50.0)
        # the prohibition as if the design had it (a new design would need a new HB estimate)
        sig = item.kind.signature(item)
        monkeypatch.setattr(type(item.kind), "prohibitions", lambda self, it: [((0, 0), (pa, cheap))])
        monkeypatch.setattr(type(item.kind), "signature", lambda self, it: sig)
        counts = item.kind.level_counts(item)

        def optimise(price, fixed):
            args = MultiDict([("sim", "2"), ("sim_n", "2"), ("sim_0_0", "0"), ("sim_0_2", "0"),
                              ("sim_0_price", price), ("sim_1_0", "1"), ("sim_1_2", "0"), ("sim_1_price", "100"),
                              ("sim_none", "1"), ("sim_act", "opt"), ("sim_fixset", "1")]
                             + [("sim_fix", str(a)) for a in fixed])
            state, _ = simulator.cbc_state(item, args)
            return simulator.cbc_optimize(item, entry, state, "de")

        free = optimise("100", [])
        total = counts[0] * len(set(numbers)) * counts[2]
        assert free["forbidden"] == counts[2] and free["count"] == total - counts[2]
        assert all(not r["take"].startswith("0:0,p:50,") for r in free["rows"])
        # held at 50: every variant with that level is forbidden; at 75 (between 50 and 100) as well; at 100 none
        for price, forbidden, between in (("50", counts[2], False), ("75", counts[2], True), ("100", 0, False)):
            out = optimise(price, [pa])
            assert out["forbidden"] == forbidden and out["count"] == counts[0] * counts[2] - forbidden
            assert bool(out.get("between")) == between
            assert not forbidden or all(not r["take"].startswith("0:0,") for r in out["rows"])
    # the page counts only the variants computed, names the ones left out and the rule between two levels
    html = client.get("/surveys/1/results?sim=2&sim_n=2&sim_0_0=0&sim_0_2=0&sim_0_price=75&sim_1_0=1&sim_1_2=0"
                      "&sim_1_price=100&sim_none=1&sim_act=opt").data.decode()
    assert f"{counts[0] * counts[2] - counts[2]} Varianten von Produkt A durchgerechnet" in html
    assert f"{counts[2]} Varianten enthalten ein verbotenes Paar von Stufen und sind nicht gerechnet" in html
    assert "wenn es eine der beiden Nachbarstufen ist" in html


def test_small_segment_by_sum_or_effective_number(monkeypatch):
    """L-5 (kombi/k5): 200 persons with a membership of 0.1 are 20 persons, effectively 200; "small" hung on the
    effective number alone, which is never smaller than the sum of memberships under 1. Now the smaller of both."""
    from app import segments, simulator

    pids = [f"p{n}" for n in range(200)]
    monkeypatch.setattr(segments, "solution", lambda stored, K: {"tau": [[0.1, 0.9]] * 200})
    size = simulator._segment_size({"pids": pids}, 2, 1, pids, {"pids": {pid: 0.1 for pid in pids}})
    assert size["sum"] == "20,0" and size["effective"] == "200,0" and size["small"]
    size = simulator._segment_size({"pids": pids}, 2, 2, pids, {"pids": {pid: 0.9 for pid in pids}})
    assert not size["small"]
    # 30 sure persons and 270 nearly outside: sum 33.9, effectively 42.3, not small
    assert not simulator._small(30 * 0.95 + 270 * 0.02, simulator._effective([0.95] * 30 + [0.02] * 270))
    assert simulator._small(29.9, 300.0) and simulator._small(300.0, 29.9)


@pytest.mark.parametrize("with_hb", [False, True])
def test_price_not_linear_names_differing_price_sensitivity(app, client, survey, monkeypatch, with_hb):
    """L-8 (l1_preis_linear): with persons who differ in price sensitivity the aggregate logit is curved in the price
    although every person values it linearly; the text named only price thresholds. With HB the simulator computes
    linearly per person."""
    from app import cbc_analysis

    monkeypatch.setattr(cbc_analysis, "price_linearity", lambda *a, **k: {"chi2": 9.0, "df": 1, "p": 0.003})
    if with_hb:
        _hb(app, client)
    else:
        cbc_setup(client, price_attr=True, none=True, tasks=12, versions=6)
        simulate_segments(app, 40, none_asc=-0.5)
    html = client.get("/surveys/1/results").data.decode()
    assert "weichen deutlich von einer Geraden ab" in html and "unterschiedlich preisempfindliche Personen" in html
    assert ("rechnet der Simulator je Person linear" in html) == with_hb



def test_intervals_leave_out_persons_without_a_weight():
    """Code review 2026-10-15, C12: persons with weight 0 were drawn into the bootstrap samples and took the place of
    persons who count: the interval was wider than with the weighted persons alone."""
    from app import choice_predict

    rng = np.random.default_rng(3)
    P = rng.dirichlet([2, 2, 2], size=40)
    w = np.r_[np.ones(20), np.zeros(20)]
    with_zero = choice_predict.mean_shares(P, w)
    alone = choice_predict.mean_shares(P[:20], np.ones(20))
    for a, b in zip(with_zero, alone):
        assert np.allclose(a, b)
    assert np.allclose(choice_predict.bootstrap_means(P, w), choice_predict.bootstrap_means(P[:20], np.ones(20)))
