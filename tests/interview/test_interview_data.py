"""Interviews of an interview guide ("Interviews erheben", app/interview_data.py): consent step, saving
from the interview view, forms, end with deletion code, deletion on request, retention, account deletion."""
import json
import re
from datetime import timedelta

from app.extensions import db
from app.models import Item, Response, Survey, User, utcnow
from tests.interview.test_interview import first, make

WEBM = b"\x1a\x45\xdf\xa3"  # the first bytes of a WebM recording (app/items/voice.py sniff)


def switch_on(app):
    """Nothing to switch on any more ("Interviews erheben" is regular since 2026-10-05); kept so the tests read alike."""


def start(app, client, code, **extra):
    data = {"label": "", "date": "2026-10-12", "mode": "praesenz", "consent_form": "schriftlich", "consent": "1", "audio": "1", **extra}
    r = client.post(f"/guide/{code}/interviews/new", data=data)
    m = re.search(r"interview=(\d+)", r.headers.get("Location", ""))
    return int(m.group(1)) if m else None


def test_regular_without_any_tick(app, client):
    """Since 2026-10-05 no experimental feature: the tab is there without a tick on "Konto"."""
    from app import experimental

    assert "interviews" not in experimental.KEYS
    code, sid = make(app, client)
    assert client.get(f"/guide/{code}/interviews").status_code == 200
    assert f"/guide/{code}/interviews" in client.get(f"/guide/{code}").data.decode()
    assert "erhebt bei Interviewleitfäden keine Daten" not in client.get(f"/guide/{code}/settings").data.decode()


def test_consent_step_and_identifier(app, client):
    code, sid = make(app, client)
    switch_on(app)
    assert f"/guide/{code}/interviews" in client.get(f"/guide/{code}").data.decode()     # the tab
    page = client.get(f"/guide/{code}/interviews").data.decode()
    assert 'value="I01"' in page and "Einwilligung aller Sprechenden liegt vor" in page
    assert start(app, client, code, consent="") is None                                    # no interview without consent
    rid = start(app, client, code)
    assert rid and start(app, client, code, label="I01") is None                          # taken
    with app.app_context():
        row = db.session.get(Response, rid)
        consent = dict(row.interview["consent"])
        assert row.interview["label"] == "I01" and consent.pop("all") and consent == {"form": "schriftlich", "audio": True}
        assert row.source == "real" and row.interview["by_name"] == "Philipp" and row.pid
    client.post(f"/guide/{code}/settings", data={"part": "interview", "order": "free", "short": "after", "post": "1", "prefix": "E"})
    assert 'value="E01"' in client.get(f"/guide/{code}/interviews").data.decode()


def test_saving_from_the_interview_view(app, client):
    code, sid = make(app, client)
    switch_on(app)
    rid = start(app, client, code)
    q = first(app, sid)
    with app.app_context():
        item = db.session.get(Item, q)
        a1, p1 = item.config["aspects"][0], item.config["probe"][0]
    page = client.get(f"/guide/{code}/conversation?interview={rid}").data.decode()
    assert "data-save=" in page and "data-iv-note" in page and "data-iv-record" in page and "I01" in page
    payload = {"states": {str(q): "spontaneous", "999": "asked"}, "ticks": [f"{q}:aspects:{a1}", f"{q}:probe:{p1}", f"{q}:aspects:fake"],
               "notes": {str(q): "Sturz 2023, Bruder weit weg"}, "marks": [{"t": 754, "item": str(q), "text": "Sturz"}, {"t": -1}],
               "elapsed": 1_500_000, "blocks": {"5": 600_000, "x": 1}}
    r = client.post(f"/guide/{code}/interviews/{rid}/save", data=json.dumps(payload), content_type="application/json")
    assert r.json["ok"]
    with app.app_context():
        row = db.session.get(Response, rid)
        assert row.answers[str(q)] == {"state": "spontaneous", "aspects": [a1], "probes": [p1], "note": "Sturz 2023, Bruder weit weg"}
        assert "999" not in row.answers and row.interview["marks"] == [{"t": 754, "item": str(q), "text": "Sturz"}]
        assert row.interview["elapsed"] == 1500 and row.interview["blocks"] == {"5": 600}
    rec = client.get(f"/guide/{code}/conversation?interview={rid}").data.decode()
    assert "spontaneous" in rec and "Sturz 2023" in rec                                     # the view starts from the stored state


def test_forms_end_and_deletion_code(app, client):
    from app import interview

    code, sid = make(app, client)
    switch_on(app)
    rid = start(app, client, code)
    with app.app_context():
        survey = db.session.get(Survey, sid)
        short = interview.shown_form(survey, "short")
        age = next(i for i in short.items if i.type == "text_input")
        gender = next(i for i in short.items if i.type == "single_choice")
        option = gender.config["options"][0]["code"]
    r = client.post(f"/guide/{code}/interviews/{rid}/form", data={f"i{age.id}": "57", f"i{gender.id}": option})
    assert r.json["ok"]
    with app.app_context():
        answers = db.session.get(Response, rid).answers
        assert answers[str(age.id)] == 57 and answers[str(gender.id)] == option
    assert 'value="57"' in client.get(f"/guide/{code}/conversation?interview={rid}").data.decode()
    page = client.post(f"/guide/{code}/interviews/{rid}/end").data.decode()
    assert "Löschcode" in page
    code_ = re.search(r'class="erase-code"><code>([^<]+)</code>', page).group(1)
    with app.app_context():
        assert db.session.get(Response, rid).status == "completed"
    assert "beendet" in client.get(f"/guide/{code}/interviews").data.decode()
    client.post(f"/guide/{code}/interviews/deletion-code", data={"term": code_})
    with app.app_context():
        assert db.session.get(Response, rid) is None


def test_delete_and_retention(app, client):
    from app import interview_data

    code, sid = make(app, client)
    switch_on(app)
    r1, r2 = start(app, client, code), start(app, client, code)
    client.post(f"/guide/{code}/interviews/{r1}/delete")
    with app.app_context():
        assert db.session.get(Response, r1) is None
        row = db.session.get(Response, r2)
        row.status, row.completed_at = "completed", utcnow() - timedelta(days=365 * 11)  # 10 years are over
        db.session.commit()
        assert interview_data.run() == 1 and db.session.get(Response, r2) is None


def test_account_deletion_keeps_the_name(app, client):
    from app import accounts

    code, sid = make(app, client)
    switch_on(app)
    rid = start(app, client, code)
    with app.app_context():
        user = db.session.scalar(db.select(User))
        uid = user.id
        from app.models import SurveyMember

        survey = db.session.get(Survey, sid)
        survey.members.append(SurveyMember(orcid="0000-0002-0000-0009", name="Kollegin", role="admin", position=1))
        db.session.commit()
        accounts.delete(user, by="admin", admin=None, choices={})
        row = db.session.get(Response, rid)
        assert row.interview["by"] is None and row.interview["by_name"] == "Philipp" and uid


def test_evaluation_and_exports(app, client):
    import io
    import zipfile

    code, sid = make(app, client)
    switch_on(app)
    q = first(app, sid)
    with app.app_context():
        item = db.session.get(Item, q)
        a1, item_code = item.config["aspects"][0], item.code
    r1, r2 = start(app, client, code), start(app, client, code)
    client.post(f"/guide/{code}/interviews/{r1}/save", content_type="application/json", data=json.dumps(
        {"states": {str(q): "asked"}, "ticks": [f"{q}:aspects:{a1}"], "notes": {str(q): "Sturz 2023"},
         "marks": [{"t": 3725, "item": str(q), "text": "Krankenhaus"}], "elapsed": 1_800_000, "blocks": {}}))
    client.post(f"/guide/{code}/interviews/{r2}/save", content_type="application/json", data=json.dumps(
        {"states": {str(q): "spontaneous"}, "ticks": [], "notes": {}, "marks": [], "elapsed": 600_000, "blocks": {}}))
    start(app, client, code, trial="1")
    page = client.get(f"/guide/{code}/interviews/analysis").data.decode()
    assert "Abdeckung" in page and "Sturz 2023" in page and "#01:02:05-0#" in page and "50,0 %" in page   # 1 of 2 interviews
    assert "I03" not in page and "I03" in client.get(f"/guide/{code}/interviews/analysis?probe=1").data.decode()
    csv_ = client.get(f"/guide/{code}/interviews/cases.csv").data.decode("utf-8-sig")
    head, first_row = csv_.splitlines()[:2]
    assert f"{item_code}_stand" in head and f"{item_code}_a1" in head and first_row.startswith("I01;0;")
    assert ";gestellt;1;" in first_row
    with zipfile.ZipFile(io.BytesIO(client.get(f"/guide/{code}/interviews/notes.zip").data)) as z:
        assert sorted(z.namelist()) == ["I01.txt", "I02.txt"]
        text = z.read("I01.txt").decode()
        assert "#01:02:05-0# Krankenhaus" in text and "Notiz: Sturz 2023" in text
    assert client.get(f"/guide/{code}/interviews/notes.docx").data[:2] == b"PK"


def test_recording(app, client):
    import os

    from app import interview_audio

    code, sid = make(app, client)
    switch_on(app)
    rid = start(app, client, code)
    no = start(app, client, code, audio="")
    page = client.get(f"/guide/{code}/conversation?interview={rid}").data.decode()
    assert "data-iv-rec-toggle" in page and "data-iv-rec-toggle" not in client.get(f"/guide/{code}/conversation?interview={no}").data.decode()
    url = f"/guide/{code}/interviews/{rid}/audio"
    send = lambda part, seq, data, mime="audio/webm": client.post(f"{url}?part={part}&seq={seq}&start=12", data=data, content_type=mime)
    assert send(1, 0, WEBM + b"AAAA").json["next"] == 1
    assert send(1, 0, WEBM + b"AAAA").json["next"] == 1                                  # sent twice: kept once
    assert send(1, 2, b"CCCC").status_code == 409                                 # a gap
    assert send(1, 1, b"BBBB").json["next"] == 2
    assert send(3, 0, b"x").status_code == 409 and send(2, 0, b"x", "text/plain").status_code == 415
    assert send(2, 0, WEBM + b"ZZ").json["next"] == 1
    assert client.post(f"/guide/{code}/interviews/{no}/audio?part=1&seq=0", data=b"x", content_type="audio/webm").status_code == 403
    r = client.get(f"{url}/1")
    assert r.status_code == 200 and r.data == WEBM + b"AAAABBBB" and r.mimetype == "audio/webm"
    r.close()  # Windows keeps a sent file open until the response is closed
    with app.app_context():
        survey, row = db.session.get(Survey, sid), db.session.get(Response, rid)
        assert [(p["part"], p["bytes"], p["start"]) for p in row.interview["audio"]] == [(1, 12, 12), (2, 6, 12)]
        where = interview_audio.folder(survey, row)
        assert os.path.isfile(os.path.join(where, "teil-1.webm"))
        assert not where.startswith(app.config["UPLOAD_FOLDER"])                  # never among the served uploads
    assert "Aufnahme (2 Teile" in client.get(f"/guide/{code}/interviews").data.decode()
    # the period: with "transcript" only once the transcript is checked; with months after the end
    with app.app_context():
        row = db.session.get(Response, rid)
        row.status, row.completed_at = "completed", utcnow() - timedelta(days=400)
        db.session.commit()
        assert interview_audio.run() == 0
    client.post(f"/guide/{code}/settings", data={"part": "interview", "order": "free", "short": "after", "post": "1", "audio": "1", "audio_keep": "12"})
    with app.app_context():
        assert interview_audio.run() == 1 and not os.path.isdir(where)
        assert db.session.get(Response, rid).interview["audio"] == []
    # deleting the interview takes the recording along
    send(3, 0, WEBM + b"QQ")
    client.post(f"/guide/{code}/interviews/{rid}/delete")
    assert not os.path.isdir(where)


def test_transcription_by_the_own_whisper_server(app, client, monkeypatch):
    import os

    from app import interview_audio, interview_transcribe

    code, sid = make(app, client)
    switch_on(app)
    rid = start(app, client, code)
    url = f"/guide/{code}/interviews/{rid}/audio"
    client.post(f"{url}?part=1&seq=0&start=0", data=WEBM + b"RIFF", content_type="audio/webm")
    client.post(f"{url}?part=2&seq=0&start=600", data=WEBM + b"RIFF", content_type="audio/webm")
    assert "Transkribieren" not in client.get(f"/guide/{code}/interviews").data.decode()       # no server named
    app.config.update(WHISPER_URL="http://whisper.intern:8000", WHISPER_TOKEN="geheim")
    calls = []

    class R:
        status_code = 200

        def __init__(self, data):
            self._d = data

        def json(self):
            return self._d

    def fake_post(url, headers=None, timeout=None, files=None, data=None):
        calls.append((url, headers, data, files["file"][2]))
        return R({"segments": [{"start": 0.0, "end": 4.0, "text": " Erzählen Sie doch einmal."},
                               {"start": 4.5, "end": 9.0, "text": " Also, das war so."},
                               {"start": 20.0, "end": 25.0, "text": " Später dann."}]})

    monkeypatch.setattr(interview_transcribe.requests, "post", fake_post)
    assert "Transkribieren" in client.get(f"/guide/{code}/interviews").data.decode()
    client.post(f"/guide/{code}/interviews/{rid}/transcribe")
    assert calls[0][0] == "http://whisper.intern:8000/v1/audio/transcriptions" and calls[0][1] == {"Authorization": "Bearer geheim"}
    assert calls[0][2]["language"] == "de" and calls[0][3] == "audio/webm" and len(calls) == 2
    with app.app_context():
        tr = db.session.get(Response, rid).interview["transcript"]
    assert tr["state"] == "done"
    assert tr["text"].startswith("Erzählen Sie doch einmal. Also, das war so. #00:00:09-0#\n\nSpäter dann. #00:00:25-0#")
    assert "#00:10:09-0#" in tr["text"]                                                 # part 2 began at 10 minutes
    page = client.get(f"/guide/{code}/interviews/{rid}/transcript").data.decode()
    assert "Erzählen Sie doch einmal" in page and "Transkript geprüft" in page
    client.post(f"/guide/{code}/interviews/{rid}/transcript", data={"text": "I: Erzählen Sie. #00:00:09-0#", "checked_shown": "1", "checked": "1"})
    assert client.get(f"/guide/{code}/interviews/{rid}/transcript.txt").data.decode() == "I: Erzählen Sie. #00:00:09-0#"
    with app.app_context():
        survey, row = db.session.get(Survey, sid), db.session.get(Response, rid)
        assert row.interview["transcript"]["checked"]
        where = interview_audio.folder(survey, row)
        assert interview_audio.run() == 0                                              # a week of grace after the tick
        assert interview_audio.run(utcnow() + timedelta(days=8)) == 1 and not os.path.isdir(where)  # "bis die Transkription geprüft ist"
    # a server that is down: the reason in the list
    client.post(f"{url}?part=1&seq=0&start=0", data=WEBM + b"RIFF", content_type="audio/webm")

    def down(*a, **k):
        raise interview_transcribe.requests.ConnectionError()

    monkeypatch.setattr(interview_transcribe.requests, "post", down)
    client.post(f"/guide/{code}/interviews/{rid}/transcribe")
    assert "nicht erreichbar" in client.get(f"/guide/{code}/interviews").data.decode()


def test_upload_a_transcript(app, client):
    import io

    from docx import Document

    code, sid = make(app, client)
    switch_on(app)
    rid = start(app, client, code)
    assert "Transkript hochladen" in client.get(f"/guide/{code}/interviews").data.decode()
    up = lambda name, data: client.post(f"/guide/{code}/interviews/{rid}/transcript/upload",
                                        data={"file": (io.BytesIO(data), name)}, content_type="multipart/form-data", follow_redirects=True)
    srt = "1\n00:00:01,000 --> 00:00:04,500\nI: Erzählen Sie doch einmal.\n\n2\n00:00:05,000 --> 00:00:09,000\nB: Also, das war so.\n"
    assert "Transkript übernommen" in up("I01.srt", srt.encode()).data.decode()
    tr = lambda: db.session.get(Response, rid).interview["transcript"]
    with app.app_context():
        assert tr()["text"] == "I: Erzählen Sie doch einmal. #00:00:04-0#\n\nB: Also, das war so. #00:00:09-0#" and tr()["source"] == "upload"
    vtt = "WEBVTT\n\n00:01:00.000 --> 00:01:03.000\n<v I>Und dann?\n"
    up("x.vtt", vtt.encode())
    with app.app_context():
        assert tr()["text"] == "I: Und dann? #00:01:03-0#"                           # the WebVTT voice as the speaker
    doc = Document()
    doc.add_paragraph("I: Wie fing es an? #00:00:12-0#")
    doc.add_paragraph("B: Mit einem Sturz.")
    buf = io.BytesIO()
    doc.save(buf)
    up("I01.docx", buf.getvalue())
    with app.app_context():
        assert tr()["text"] == "I: Wie fing es an? #00:00:12-0#\n\nB: Mit einem Sturz." and tr()["checked"] is None
    up("text.txt", "I: Grüß Gott.".encode("cp1252"))
    with app.app_context():
        assert tr()["text"] == "I: Grüß Gott."
    assert "Bitte eine Datei vom Typ" in up("bild.png", b"\x89PNG").data.decode()
    assert "keine Einträge" in up("leer.srt", b"nichts").data.decode()
    assert "I01-transkript.txt" in __import__("zipfile").ZipFile(io.BytesIO(client.get(f"/guide/{code}/interviews/notes.zip").data)).namelist()


def test_project_for_qda_programs(app, client):
    import io
    import zipfile
    from xml.etree import ElementTree as ET

    from app import interview

    code, sid = make(app, client)
    switch_on(app)
    r1, r2 = start(app, client, code), start(app, client, code)
    start(app, client, code, trial="1")
    with app.app_context():
        short = interview.shown_form(db.session.get(Survey, sid), "short")
        age = next(i for i in short.items if i.type == "text_input")
    client.post(f"/guide/{code}/interviews/{r1}/form", data={f"i{age.id}": "57"})
    client.post(f"/guide/{code}/interviews/{r1}/transcript/upload",
                data={"file": (io.BytesIO("I: Erzählen Sie. #00:00:09-0#".encode()), "I01.txt")}, content_type="multipart/form-data")
    assert "project.qdpx" in client.get(f"/guide/{code}/interviews/analysis").data.decode()
    data = client.get(f"/guide/{code}/interviews/project.qdpx").data
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        names = z.namelist()
        root = ET.fromstring(z.read("project.qde"))
        ns = {"p": "urn:QDA-XML:project:1.0"}
        assert root.tag == "{urn:QDA-XML:project:1.0}Project" and root.find("p:CodeBook/p:Codes/p:Code", ns).get("name") == "Einstieg"
        cases = root.findall("p:Cases/p:Case", ns)
        assert [c.get("name") for c in cases] == ["I01", "I02"]                       # no Probeinterview
        values = [v.find("p:TextValue", ns).text for v in cases[0].findall("p:VariableValue", ns)]
        assert "57" in values and "Persönliches Interview" in values
        src = root.find("p:Sources/p:TextSource", ns)
        path = src.get("plainTextPath").replace("internal://", "Sources/")
        assert path in names and z.read(path).decode() == "I: Erzählen Sie. #00:00:09-0#"
        assert cases[0].find("p:SourceRef", ns).get("targetGUID") == src.get("guid") and cases[1].find("p:SourceRef", ns) is None
        note = root.find("p:Notes/p:Note", ns)
        assert note.get("guid") == src.find("p:NoteRef", ns).get("targetGUID") and "Interview I01" in note.find("p:PlainTextContent", ns).text


def test_search_over_transcripts(app, client):
    import io

    code, sid = make(app, client)
    switch_on(app)
    r1, r2 = start(app, client, code), start(app, client, code)
    up = lambda rid, text: client.post(f"/guide/{code}/interviews/{rid}/transcript/upload",
                                       data={"file": (io.BytesIO(text.encode()), "t.txt")}, content_type="multipart/form-data")
    up(r1, "I: Wie kam es zur Pflege? #00:00:05-0#\n\nB: Meine Mutter brauchte Pflege, der Pflegedienst kam zweimal. #00:01:10-0#")
    up(r2, "I: Und heute? #00:00:03-0#\n\nB: Heute pflege ich allein, die Pflege ist schwer. #00:00:40-0#")
    page = client.get(f"/guide/{code}/interviews/search?q=pflege").data.decode()
    assert "<mark>Pflege</mark>" in page and "#00:01:10-0#" in page and "I01: 2" in page and "I02: 2" in page
    wild = client.get(f"/guide/{code}/interviews/search?q=pfleg*").data.decode()
    assert "<mark>Pflegedienst</mark>" in wild and "<mark>pflege</mark>" in wild
    only_b = client.get(f"/guide/{code}/interviews/search?q=pflege&sprecher=B").data.decode()
    assert "Wie kam es zur" not in only_b and "brauchte" in only_b
    assert "pflege <small>" in page                                                       # in the word frequencies
    assert client.get(f"/guide/{code}/interviews/search?q=%3Cscript%3E").status_code == 200


def test_case_summaries(app, client):
    from app import interview

    code, sid = make(app, client)
    switch_on(app)
    r1, r2 = start(app, client, code), start(app, client, code)
    with app.app_context():
        block = str(interview.guide_blocks(db.session.get(Survey, sid))[0].id)
    url = lambda rid: f"/guide/{code}/interviews/{rid}/summary"
    assert client.post(url(r1), data={"key": block, "text": "Pflege als Pflicht <b>"}).status_code == 200
    assert client.post(url(r1), data={"key": "case", "text": "Typ: Verpflichtete"}).status_code == 200
    assert client.post(url(r2), data={"key": block, "text": "Pflege mit Hilfe"}).status_code == 200
    assert client.post(url(r1), data={"key": "irgendwas", "text": "x"}).status_code == 400
    with app.app_context():
        assert db.session.get(Response, r1).interview["summaries"] == {block: "Pflege als Pflicht <b>", "case": "Typ: Verpflichtete"}
    page = client.get(f"/guide/{code}/interviews/summaries").data.decode()
    assert "Pflege als Pflicht &lt;b&gt;" in page and "Kelle, U., &amp; Kluge, S." in page
    csv_ = client.get(f"/guide/{code}/interviews/summaries.csv").data.decode("utf-8-sig")
    assert "Typ: Verpflichtete" in csv_ and "Pflege mit Hilfe" in csv_
    assert client.get(f"/guide/{code}/interviews/summaries.docx").data[:2] == b"PK"
    client.post(url(r2), data={"key": block, "text": "  "})                              # empty text removes the cell
    with app.app_context():
        assert block not in db.session.get(Response, r2).interview.get("summaries", {})
    assert "/interviews/summaries" in client.get(f"/guide/{code}/interviews/analysis").data.decode()


def test_second_device_records(app, client):
    from app import interview_remote

    code, sid = make(app, client)
    switch_on(app)
    rid = start(app, client, code)
    no = start(app, client, code, audio="")
    view = client.get(f"/guide/{code}/conversation?interview={rid}").data.decode()
    assert f"/interviews/{rid}/recorder/qr.svg" in view and "data-iv-remote" in view
    assert "data-iv-remote" not in client.get(f"/guide/{code}/conversation?interview={no}").data.decode()
    base = f"/guide/{code}/interviews/{rid}/recorder"
    phone = client.get(base).data.decode()
    assert "data-recorder" in phone and "data-rec-toggle" in phone and "data-iv-step" not in phone   # no guide there
    assert "data-mic-select" in phone and "mic.js" in phone and "data-mic-select" in view          # the microphone choice
    assert "keine Einwilligung" in client.get(f"/guide/{code}/interviews/{no}/recorder").data.decode()
    assert client.get(base + "/qr.svg").data.startswith(b"<svg")
    # the clock as the tablet sent it: a part from the phone (no start) begins there
    client.post(f"/guide/{code}/interviews/{rid}/save", json={"elapsed": 65000, "running": False})
    assert client.post(f"/guide/{code}/interviews/{rid}/audio?part=1&seq=0", data=WEBM + b"AA", content_type="audio/webm").status_code == 200
    assert client.get(base + "/state").json == {"state": "", "ms": 0, "alive": False, "stop": False, "pending": 0, "problem": "",
                                                "piece": 0}
    report = lambda state, ms: client.post(base + "/state", json={"state": state, "ms": ms}).json
    assert report("recording", 5000) == {"ok": True, "stop": False}
    s = client.get(base + "/state").json
    assert s["state"] == "recording" and s["alive"] and s["ms"] >= 5000
    assert client.post(base + "/stop").json["stop"] is True                     # the tablet ends the interview
    assert report("recording", 8000)["stop"] is True                              # the phone learns it
    assert report("stopped", 9000)["stop"] is True
    assert client.get(base + "/state").json["state"] == "stopped"
    assert report("recording", 0)["stop"] is False                                # a new recording starts afresh
    assert client.post(base + "/state", json={"state": "x"}).status_code == 400
    assert client.post(f"/guide/{code}/interviews/{no}/recorder/state", json={"state": "recording"}).status_code == 403
    with app.app_context():
        row = db.session.get(Response, rid)
        assert row.interview["audio"][0]["start"] == 65
        assert interview_remote.status(row, now=utcnow() + timedelta(seconds=60))["alive"] is False


def test_two_devices_never_share_a_part(app, client):
    code, sid = make(app, client)
    switch_on(app)
    rid = start(app, client, code)
    url = f"/guide/{code}/interviews/{rid}/audio"
    send = lambda part, seq, rec, data: client.post(f"{url}?part={part}&seq={seq}&start=0&rec={rec}", data=data, content_type="audio/webm")
    assert send(1, 0, "tablet", WEBM + b"TT").json["next"] == 1                         # the tablet records first
    r = send(1, 0, "phone", WEBM + b"PP")                                                 # the phone also starts at part 1
    assert r.status_code == 409 and r.json["part"] == 2
    assert send(1, 1, "phone", WEBM + b"PP").status_code == 409                          # nor later pieces into the tablet's part
    assert send(2, 0, "phone", WEBM + b"PP").json["next"] == 1
    assert send(1, 0, "tablet", WEBM + b"TT").json["next"] == 1                          # the tablet's piece sent twice: kept once
    assert send(1, 1, "tablet", b"tt").json["next"] == 2
    with app.app_context():
        audio = db.session.get(Response, rid).interview["audio"]
        assert [(p["part"], p["bytes"], p["rec"]) for p in audio] == [(1, 8, "tablet"), (2, 6, "phone")]


def test_recording_pages_may_use_the_microphone(app, client):
    """Every other page says microphone=() (security audit 2026-10-04, H2); the pages that record must not."""
    code, sid = make(app, client)
    switch_on(app)
    rid = start(app, client, code)
    assert "microphone=(self)" in client.get(f"/guide/{code}/conversation?interview={rid}").headers["Permissions-Policy"]
    assert "microphone=(self)" in client.get(f"/guide/{code}/interviews/{rid}/recorder").headers["Permissions-Policy"]
    assert "microphone=()" in client.get(f"/guide/{code}/conversation").headers["Permissions-Policy"]   # stores nothing, records nothing
