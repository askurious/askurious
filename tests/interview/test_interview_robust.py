"""Interview recordings and transcription made robust (data audit 2026-10-05, package A): pieces written at the length
the database confirms (L-2, I-1, I-2), a file per recording, the upload route trying a refused commit again, the
phone's part start (I-6), the disk check (I-22), checked parts and a week of grace (I-3), periods of interviews never
ended (L-3), folders by pid and the sweep (L-4), the transcription that holds no row while Whisper works (L-6, I-4,
I-25), its sign of life and cancelling (I-7), parts stored one by one (I-21), versions with a file left (L-7),
uploaded subtitles and text files (I-8, I-9), private caching (I-18), the draft kept only while it differs (L-14),
the prompt and the stored model and duration (I-24). Whisper is a fake (no microphone, no server in the tests)."""
import io
import os
import time
from datetime import timedelta

import pytest
from sqlalchemy.orm.exc import StaleDataError

from app.extensions import db
from app.models import Response, Survey, utcnow
from tests.interview.test_interview import make
from tests.interview.test_interview_data import start

WEBM = b"\x1a\x45\xdf\xa3"  # the first bytes of a WebM recording (app/items/voice.py sniff)


def _rows(sid, rid):
    return db.session.get(Survey, sid), db.session.get(Response, rid)


def _file(app, sid, rid, part=1):
    from app import interview_audio

    with app.app_context():
        survey, row = _rows(sid, rid)
        full, found = interview_audio.path(survey, row, part)
        with open(full, "rb") as f:
            return f.read(), found


def _send(client, code, rid, part, seq, data, rec="tab", extra="&start=0"):
    return client.post(f"/guide/{code}/interviews/{rid}/audio?part={part}&seq={seq}&rec={rec}{extra}", data=data,
                       content_type="audio/webm")


# --- the recording -------------------------------------------------------------------------------------------------


def test_piece_of_a_refused_commit_is_not_doubled(app, client):
    """L-2 / I-1: the phone's report commits between loading and committing a piece; the piece is refused, the
    browser sends it again, and the file holds it once."""
    from app import interview_audio, interview_remote

    code, sid = make(app, client)
    rid = start(app, client, code)
    assert _send(client, code, rid, 1, 0, WEBM + b"HEAD").json["next"] == 1
    with app.app_context():
        survey_a, row_a = _rows(sid, rid)
        _ = row_a.interview
        with app.app_context():
            survey_b, row_b = _rows(sid, rid)
            interview_remote.report(row_b, "recording", 10000)
            db.session.commit()
        interview_audio.append(survey_a, row_a, 1, 1, "audio/webm", b"PIECE1", 0, "tab")
        with pytest.raises(StaleDataError):
            db.session.commit()
        db.session.rollback()
    assert _send(client, code, rid, 1, 1, b"PIECE1").json["next"] == 2
    data, found = _file(app, sid, rid)
    assert data == WEBM + b"HEADPIECE1" and found["bytes"] == 14 and found["file"] == "teil-1-tab.webm"


def test_upload_tries_a_refused_commit_once_more(app, client, monkeypatch):
    """L-2: a conflict with another request is not the browser's problem: the route reloads and stores the piece."""
    code, sid = make(app, client)
    rid = start(app, client, code)
    _send(client, code, rid, 1, 0, WEBM + b"HEAD")
    real, calls = db.session.commit, {"n": 0}

    def once():
        calls["n"] += 1
        if calls["n"] == 1:
            raise StaleDataError("UPDATE statement on table 'response' expected to update 1 row(s); 0 were matched.")
        return real()

    monkeypatch.setattr(db.session, "commit", once, raising=False)
    r = _send(client, code, rid, 1, 1, b"PIECE1")
    monkeypatch.setattr(db.session, "commit", real, raising=False)
    assert r.status_code == 200 and r.json["next"] == 2 and calls["n"] == 2
    data, found = _file(app, sid, rid)
    assert data == WEBM + b"HEADPIECE1" and found["pieces"] == 2


def test_failed_commit_and_resend_keep_the_file_whole(app, client, monkeypatch):
    """L-2: a commit that fails for another reason (database locked) and the piece sent again: once in the file."""
    code, sid = make(app, client)
    rid = start(app, client, code)
    _send(client, code, rid, 1, 0, WEBM + b"HEAD")
    real, calls = db.session.commit, {"n": 0}

    def locked():
        calls["n"] += 1
        if calls["n"] == 1:
            raise RuntimeError("database is locked")
        return real()

    monkeypatch.setattr(db.session, "commit", locked, raising=False)
    app.testing = False
    app.config["PROPAGATE_EXCEPTIONS"] = False
    assert _send(client, code, rid, 1, 1, b"PIECE1").status_code == 500
    monkeypatch.setattr(db.session, "commit", real, raising=False)
    assert _send(client, code, rid, 1, 1, b"PIECE1").status_code == 200
    assert _file(app, sid, rid)[0] == WEBM + b"HEADPIECE1"


def test_two_devices_starting_at_once_write_files_of_their_own(app, client):
    """I-2: both first pieces arrive at the same moment; each recording has its own file, the refused one moves on."""
    from app import interview_audio

    code, sid = make(app, client)
    rid = start(app, client, code)
    with app.app_context():
        survey_a, row_a = _rows(sid, rid)
        _ = row_a.interview
        with app.app_context():
            survey_b, row_b = _rows(sid, rid)
            interview_audio.append(survey_b, row_b, 1, 0, "audio/webm", WEBM + b"PHONEHEAD", 0, "phone")
            db.session.commit()
        interview_audio.append(survey_a, row_a, 1, 0, "audio/webm", WEBM + b"TABLETHEAD", 0, "tablet")
        with pytest.raises(StaleDataError):
            db.session.commit()
        db.session.rollback()
    r = _send(client, code, rid, 1, 0, WEBM + b"TABLETHEAD", rec="tablet")
    assert r.status_code == 409 and r.json["part"] == 2
    assert _send(client, code, rid, 2, 0, WEBM + b"TABLETHEAD", rec="tablet").status_code == 200
    assert _send(client, code, rid, 1, 1, b"phone2", rec="phone").status_code == 200
    assert _file(app, sid, rid, 1)[0] == WEBM + b"PHONEHEADphone2" and _file(app, sid, rid, 2)[0] == WEBM + b"TABLETHEAD"


def test_phone_part_begins_when_its_recording_began(app, client, monkeypatch):
    """I-6: the phone sends how long ago its recording began; a first piece that got through late does not move it."""
    from app import interview_audio

    code, sid = make(app, client)
    rid = start(app, client, code)
    client.post(f"/guide/{code}/interviews/{rid}/save", json={"elapsed": 60000, "running": True})
    later = utcnow() + timedelta(seconds=300)  # the first piece got through five minutes after the recording began
    monkeypatch.setattr(interview_audio, "utcnow", lambda: later)
    assert _send(client, code, rid, 1, 0, WEBM + b"AA", rec="phone", extra="&age=300000").status_code == 200
    with app.app_context():
        assert abs(db.session.get(Response, rid).interview["audio"][0]["start"] - 60) <= 2


def test_full_disk_refuses_the_piece_and_writes_nothing(app, client, monkeypatch):
    """I-22: below the reserve the server answers 507; the browser keeps the piece and sends it later."""
    from collections import namedtuple

    from app import interview_audio

    code, sid = make(app, client)
    rid = start(app, client, code)
    _send(client, code, rid, 1, 0, WEBM + b"HEAD")
    usage = namedtuple("usage", "total used free")
    monkeypatch.setattr(interview_audio.shutil, "disk_usage", lambda where: usage(10, 10, 1000))
    r = _send(client, code, rid, 1, 1, b"PIECE1")
    assert r.status_code == 507 and r.json["error"] == "full"
    monkeypatch.undo()
    data, found = _file(app, sid, rid)
    assert data == WEBM + b"HEAD" and found["pieces"] == 1


def test_only_checked_parts_go_after_a_week(app, client):
    """I-3: "Transkript geprüft" notes the parts it holds; a part recorded after "Fortsetzen" stays, the checked part
    goes a week after the tick, and a tick taken back in time keeps everything."""
    from app import interview_audio

    code, sid = make(app, client)
    rid = start(app, client, code)
    _send(client, code, rid, 1, 0, WEBM + b"PART1", rec="a")
    url = f"/guide/{code}/interviews/{rid}/transcript"
    page = client.get(url).data.decode()
    assert "7 Tage nach dem Haken gelöscht" in page
    client.post(url, data={"text": "I: Erste Sitzung. #00:00:05-0#", "checked_shown": "1", "checked": "1"})
    with app.app_context():
        assert interview_audio.run() == 0                                   # a week of grace
        assert db.session.get(Response, rid).interview["transcript"]["checked_parts"] == {"1": 9}
    client.post(f"/guide/{code}/interviews/{rid}/resume")
    assert _send(client, code, rid, 2, 0, WEBM + b"PART2-NEW", rec="b").status_code == 200
    with app.app_context():
        assert interview_audio.run(utcnow() + timedelta(days=8)) == 1
        audio = db.session.get(Response, rid).interview["audio"]
        assert [p["part"] for p in audio] == [2]                            # the new part stays
        survey, row = _rows(sid, rid)
        assert sorted(os.listdir(interview_audio.folder(survey, row))) == ["teil-2-b.webm"]
    assert _send(client, code, rid, 2, 1, b"MORE", rec="b").status_code == 200   # the recording goes on
    # a tick taken back within the week: nothing goes
    rid2 = start(app, client, code)
    _send(client, code, rid2, 1, 0, WEBM + b"X", rec="c")
    url2 = f"/guide/{code}/interviews/{rid2}/transcript"
    client.post(url2, data={"text": "B: Ja.", "checked_shown": "1", "checked": "1"})
    client.post(url2, data={"text": "B: Ja.", "checked_shown": "1"})
    with app.app_context():
        assert interview_audio.run(utcnow() + timedelta(days=30)) == 0
        assert db.session.get(Response, rid2).interview["audio"]


def test_unfinished_interview_counts_from_its_last_activity(app, client):
    """L-3: an interview never ended goes by both periods, counted from its last activity; the list marks it."""
    from app import interview_audio, interview_data

    code, sid = make(app, client)
    rid = start(app, client, code)
    _send(client, code, rid, 1, 0, WEBM + b"RIFF")
    client.post(f"/guide/{code}/settings", data={"part": "interview", "order": "free", "short": "after", "post": "1",
                                                          "audio": "1", "audio_keep": "12", "data_keep": "60"})
    with app.app_context():
        row = db.session.get(Response, rid)
        row.started_at = row.last_active_at = utcnow() - timedelta(days=10)
        db.session.commit()
    assert "nicht beendet" in client.get(f"/guide/{code}/interviews").data.decode()
    with app.app_context():
        row = db.session.get(Response, rid)
        row.started_at = row.last_active_at = utcnow() - timedelta(days=365 * 2)
        db.session.commit()
        assert interview_audio.run() == 1                                   # the recording after 12 months
        assert interview_data.run() == 0
        row = db.session.get(Response, rid)
        row.last_active_at = utcnow() - timedelta(days=365 * 6)
        db.session.commit()
        assert interview_data.run() == 1 and db.session.get(Response, rid) is None   # all data after 60 months


def test_folders_by_pid_old_folders_move_and_left_ones_are_swept(app, client, monkeypatch):
    """L-4: the folder is named after the pid, so a recycled row id never finds the voice of a deleted interview; a
    folder from before (named after the id) moves at the first access; the sweep removes what no interview holds."""
    from app import interview_audio

    code, sid = make(app, client)
    rid = start(app, client, code)
    _send(client, code, rid, 1, 0, WEBM + b"OLDVOICE", rec="x")
    with app.app_context():
        survey, row = _rows(sid, rid)
        where, pid = interview_audio.folder(survey, row), row.pid
        assert os.path.basename(where) == pid
    monkeypatch.setattr(interview_audio.shutil, "rmtree", lambda *a, **k: None)  # the removal fails silently
    client.post(f"/guide/{code}/interviews/{rid}/delete")
    monkeypatch.undo()
    new = start(app, client, code)
    _send(client, code, new, 1, 0, WEBM + b"NEW", rec="y")
    assert _file(app, sid, new)[0] == WEBM + b"NEW"
    with app.app_context():
        os.utime(where, (time.time() - 7200, time.time() - 7200))
        assert interview_audio.sweep() >= 1 and not os.path.isdir(where)
        survey, row = _rows(sid, new)
        # a recording as it lay before 2026-10-05: under the row's id
        legacy = os.path.join(interview_audio.root(), survey.uid, str(row.id))
        os.rename(interview_audio.folder(survey, row), legacy)
        assert interview_audio.folder(survey, row).endswith(row.pid) and not os.path.isdir(legacy)
    assert _file(app, sid, new)[0] == WEBM + b"NEW"


def test_retention_purge_takes_interview_recordings(app, client):
    """The rule of the project: every store of answers in retention._clean_files."""
    from app import interview_audio, retention

    code, sid = make(app, client)
    rid = start(app, client, code)
    _send(client, code, rid, 1, 0, WEBM + b"RIFF")
    with app.app_context():
        survey = db.session.get(Survey, sid)
        base = os.path.join(interview_audio.root(), survey.uid)
        assert os.path.isdir(base)
        retention._clean_files(survey.uid, set())
        assert not os.path.isdir(base)


def test_interview_files_are_never_cached(app, client):
    """I-18: recordings, transcripts and the exports of notes with "private, no-store"."""
    code, sid = make(app, client)
    rid = start(app, client, code)
    _send(client, code, rid, 1, 0, WEBM + b"RIFF")
    client.post(f"/guide/{code}/interviews/{rid}/transcript", data={"text": "abc"})
    for url in (f"/guide/{code}/interviews/{rid}/audio/1", f"/guide/{code}/interviews/{rid}/transcript.txt",
                f"/guide/{code}/interviews/project.qdpx", f"/guide/{code}/interviews/notes.zip",
                f"/guide/{code}/interviews/summaries.csv"):
        r = client.get(url)
        assert r.status_code == 200 and r.headers["Cache-Control"] == "private, no-store", url
        r.close()


def test_device_report_and_stop_survive_a_conflict(app, client, monkeypatch):
    """L-2 / I-20 / I-26: the request to stop is stored although another request wrote the interview meanwhile; the
    view learns the pieces the phone still holds and its problem."""
    code, sid = make(app, client)
    rid = start(app, client, code)
    base = f"/guide/{code}/interviews/{rid}/recorder"
    assert client.post(base + "/state", json={"state": "recording", "ms": 5000, "pending": 3, "problem": "Mikrofon getrennt"}).json["ok"]
    s = client.get(base + "/state").json
    assert s["pending"] == 3 and s["problem"] == "Mikrofon getrennt" and s["piece"] is None
    real, calls = db.session.commit, {"n": 0}

    def once():
        calls["n"] += 1
        if calls["n"] == 1:
            raise StaleDataError("stale")
        return real()

    monkeypatch.setattr(db.session, "commit", once, raising=False)
    assert client.post(base + "/stop").json["stop"] is True
    monkeypatch.setattr(db.session, "commit", real, raising=False)
    assert client.post(base + "/state", json={"state": "recording", "ms": 9000}).json["stop"] is True


def test_recording_pages_buffer_pieces_and_watch_the_microphone(app, client):
    """I-19, I-20, I-23, I-11, I-26 live in the browser: the pages load the store and the scripts hold the rules."""
    code, sid = make(app, client)
    rid = start(app, client, code)
    view = client.get(f"/guide/{code}/conversation?interview={rid}").data.decode()
    phone = client.get(f"/guide/{code}/interviews/{rid}/recorder").data.decode()
    assert "audiostore.js" in view and "audiostore.js" in phone and 'data-audio-store="iv-' in view and 'data-audio-store="rec-' in phone
    static = os.path.join(app.root_path, "static")
    read = lambda name: open(os.path.join(static, name), encoding="utf-8").read()
    for name in ("interview.js", "recorder.js"):
        js = read(name)
        assert "audioBitsPerSecond: 32000" in js and "onerror" in js and "onended" in js and "onmute" in js
        assert "AskuriousAudioStore" in js and "store.drop(piece)" in js and "507" in js
    assert "&age=" in read("recorder.js")
    assert "channelCount: 1" in read("mic.js") and "indexedDB" in read("audiostore.js")
    assert "Noch nicht übertragen" in read("interview.js") and "d.stop" in read("interview.js")


# --- the transcription ---------------------------------------------------------------------------------------------


class _R:
    def __init__(self, data, status=200):
        self._d, self.status_code = data, status

    def json(self):
        return self._d


def _whisper(app, client, monkeypatch, answer=None, parts=1):
    """An interview with recorded parts and a fake Whisper: answer(file name, data) gives the response."""
    from app import interview_transcribe

    code, sid = make(app, client)
    rid = start(app, client, code)
    for n in range(1, parts + 1):
        _send(client, code, rid, n, 0, WEBM + b"RIFF%d" % n, rec=f"r{n}", extra=f"&start={(n - 1) * 600}")
    app.config.update(WHISPER_URL="http://whisper.intern:8000")
    calls = []

    def fake_post(url, headers=None, timeout=None, files=None, data=None):
        calls.append((files["file"][0], dict(data)))
        if answer:
            return answer(files["file"][0], data)
        return _R({"segments": [{"start": 0.0, "end": 4.0, "text": " Whisper-Entwurf."}]})

    monkeypatch.setattr(interview_transcribe.requests, "post", fake_post)
    return code, sid, rid, calls


def _tr(app, rid):
    with app.app_context():
        return db.session.get(Response, rid).interview.get("transcript") or {}


def test_a_write_during_transcription_keeps_both(app, client, monkeypatch):
    """L-6 / I-4 / I-25: a case summary stored while Whisper works stays, and the draft is stored too."""
    def answer(name, data):
        with app.app_context():
            client.post(f"/guide/{code}/interviews/{rid}/summary", data={"key": "case", "text": "Während Whisper lief"})
        return _R({"segments": [{"start": 0.0, "end": 4.0, "text": " Whisper-Entwurf."}]})

    code, sid, rid, calls = _whisper(app, client, monkeypatch, answer)
    client.post(f"/guide/{code}/interviews/{rid}/transcribe")
    tr = _tr(app, rid)
    assert tr["state"] == "done" and tr["text"] == "Whisper-Entwurf. #00:00:04-0#" and "draft" not in tr   # L-14
    assert tr["covers"] == {"1": 9} and "parts" not in tr and "run" not in tr
    with app.app_context():
        assert db.session.get(Response, rid).interview["summaries"] == {"case": "Während Whisper lief"}


def test_a_write_during_transcription_keeps_a_corrected_text(app, client, monkeypatch):
    """L-6: a text corrected meanwhile is not overwritten; Whisper's draft stays beside it while it differs (L-14)."""
    def answer(name, data):
        with app.app_context():
            client.post(f"/guide/{code}/interviews/{rid}/transcript", data={"text": "I: Von Hand. #00:00:04-0#"})
        return _R({"segments": [{"start": 0.0, "end": 4.0, "text": " Whisper-Entwurf."}]})

    code, sid, rid, calls = _whisper(app, client, monkeypatch, answer)
    client.post(f"/guide/{code}/interviews/{rid}/transcribe")
    tr = _tr(app, rid)
    assert tr["state"] == "done" and tr["text"] == "I: Von Hand. #00:00:04-0#" and tr["draft"] == "Whisper-Entwurf. #00:00:04-0#"
    client.post(f"/guide/{code}/interviews/{rid}/transcript", data={"text": "Whisper-Entwurf. #00:00:04-0#"})
    assert "draft" not in _tr(app, rid)                                    # the same as the text: not kept twice


def test_broken_off_transcription_may_start_again_and_be_cancelled(app, client, monkeypatch):
    """I-7: "running" counts only with a fresh sign of life; a run without one (a restart) may start again;
    "Abbrechen" ends the waiting."""
    from app import interview_transcribe

    code, sid, rid, calls = _whisper(app, client, monkeypatch)
    with app.app_context():
        row = db.session.get(Response, rid)
        old = (utcnow() - timedelta(minutes=10)).isoformat(timespec="seconds")
        interview_transcribe._set(row, state="running", started=old, beat=old, run="weg")
        db.session.commit()
        assert not interview_transcribe.running(row)
    page = client.get(f"/guide/{code}/interviews").data.decode()
    assert "Transkription unterbrochen" in page and "Transkribieren" in page
    client.post(f"/guide/{code}/interviews/{rid}/transcribe")
    assert len(calls) == 1 and _tr(app, rid)["state"] == "done"
    with app.app_context():
        row = db.session.get(Response, rid)
        interview_transcribe._set(row, state="running", started=utcnow().isoformat(timespec="seconds"),
                                  beat=utcnow().isoformat(timespec="seconds"), run="läuft")
        db.session.commit()
    page = client.get(f"/guide/{code}/interviews").data.decode()
    assert "Transkription läuft" in page and "/transcribe/cancel" in page
    client.post(f"/guide/{code}/interviews/{rid}/transcribe/cancel")
    tr = _tr(app, rid)
    assert tr["state"] == "error" and tr["error"] == "Abgebrochen."


def test_sign_of_life_while_waiting_and_cancel_drops_the_result(app, client, monkeypatch):
    """I-7: while Whisper works the run leaves a sign of life; a run cancelled meanwhile drops what Whisper sends."""
    from app import interview_transcribe

    monkeypatch.setattr(interview_transcribe, "BEAT", 0.05)
    seen = {}

    def answer(name, data):
        with app.app_context():
            row = db.session.get(Response, rid)
            interview_transcribe._set(row, beat="2000-01-01T00:00:00")
            db.session.commit()
        time.sleep(0.4)
        with app.app_context():
            seen["beat"] = db.session.get(Response, rid).interview["transcript"]["beat"]
            row = db.session.get(Response, rid)
            interview_transcribe._set(row, state="error", error="Abgebrochen.", run="")   # "Abbrechen" meanwhile
            db.session.commit()
        time.sleep(0.3)
        return _R({"segments": [{"start": 0.0, "end": 4.0, "text": " Zu spät."}]})

    code, sid, rid, calls = _whisper(app, client, monkeypatch, answer)
    client.post(f"/guide/{code}/interviews/{rid}/transcribe")
    assert seen["beat"] > "2000-01-01T00:00:00"
    tr = _tr(app, rid)
    assert tr["state"] == "error" and not tr.get("text") and not tr.get("parts")
    # the cancelled request runs out by itself; until then busy(rid) holds, and the next test's interview has the
    # same id: it would find "still running" and not start (the next test failed now and then)
    deadline = time.monotonic() + 5
    while interview_transcribe.busy(rid) and time.monotonic() < deadline:
        time.sleep(0.02)
    assert not interview_transcribe.busy(rid)


def test_job_cancelled_by_abbrechen_is_no_failure_in_the_log(app, client, monkeypatch, caplog):
    """"Abbrechen" while Whisper works: the service then answers with an error (its job was cancelled). The run ends
    quietly, keeps the state "Abgebrochen." and logs no failure (log check after the deploy 2026-10-29); a failure of
    the current run is still logged."""
    from app import interview_transcribe

    def answer(name, data):
        with app.app_context():
            row = db.session.get(Response, rid)
            interview_transcribe._set(row, state="error", error="Abgebrochen.", run="")   # "Abbrechen" meanwhile
            db.session.commit()
        raise RuntimeError("Der Transkriptionsauftrag wurde abgebrochen oder konnte nicht verarbeitet werden.")

    code, sid, rid, calls = _whisper(app, client, monkeypatch, answer)
    with caplog.at_level("INFO"):
        client.post(f"/guide/{code}/interviews/{rid}/transcribe")
    tr = _tr(app, rid)
    assert tr["state"] == "error" and tr["error"] == "Abgebrochen."
    assert not [r for r in caplog.records if r.levelname == "ERROR" and "failed" in r.getMessage()]
    assert any("cancelled or started anew" in r.getMessage() for r in caplog.records)

    def broken(name, data):
        raise RuntimeError("Whisper kaputt.")

    monkeypatch.setattr(interview_transcribe.requests, "post", lambda *a, **k: broken(None, None))
    caplog.clear()
    client.post(f"/guide/{code}/interviews/{rid}/transcribe")
    tr = _tr(app, rid)
    assert tr["state"] == "error" and tr["error"] == "Whisper kaputt."
    assert [r for r in caplog.records if r.levelname == "ERROR" and "failed" in r.getMessage()]


def test_parts_are_kept_one_by_one_and_only_missing_ones_sent_again(app, client, monkeypatch):
    """I-21: part 1 done, part 2 fails: part 1 stays, the next run sends part 2 only; a part whose file is gone is
    noted in the draft."""
    from app import interview_audio

    fail = {"on": True}

    def answer(name, data):
        if name.startswith("teil-2") and fail["on"]:
            return _R({}, status=500)
        return _R({"segments": [{"start": 0.0, "end": 3.0, "text": f" Aus {name.split('-')[1]}."}]})

    code, sid, rid, calls = _whisper(app, client, monkeypatch, answer, parts=3)
    client.post(f"/guide/{code}/interviews/{rid}/transcribe")
    tr = _tr(app, rid)
    assert tr["state"] == "error" and "500" in tr["error"] and sorted(tr["parts"]) == ["1"]
    assert "1 Teil ist schon fertig" in client.get(f"/guide/{code}/interviews").data.decode()
    fail["on"] = False
    with app.app_context():
        survey, row = _rows(sid, rid)
        os.remove(interview_audio.path(survey, row, 3)[0])                  # part 3 got lost on the disk
    calls.clear()
    client.post(f"/guide/{code}/interviews/{rid}/transcribe")
    assert [c[0] for c in calls] == ["teil-2-r2.webm"]
    tr = _tr(app, rid)
    assert tr["state"] == "done" and "parts" not in tr
    assert tr["text"] == "Aus 1. #00:00:03-0#\n\nAus 2. #00:10:03-0#\n\n[Teil 3 fehlt]" and tr["covers"] == {"1": 9, "2": 9}


def test_prompt_model_and_duration(app, client, monkeypatch):
    """I-24: the guide's words go along as prompt; a server that refuses the field gets the request without it;
    model and duration are stored."""
    def answer(name, data):
        if "prompt" in data:
            return _R({"detail": "unknown field"}, status=422)
        return _R({"segments": [{"start": 0.0, "end": 2.0, "text": " Ja."}]})

    code, sid, rid, calls = _whisper(app, client, monkeypatch, answer)
    app.config["WHISPER_MODEL"] = "large-v3"
    client.post(f"/guide/{code}/interviews/{rid}/transcribe")
    assert calls[0][1]["prompt"].startswith("Pflegende Angehörige") and "prompt" not in calls[1][1]
    tr = _tr(app, rid)
    assert tr["state"] == "done" and tr["model"] == "large-v3" and isinstance(tr["seconds"], int)


# --- versions and uploads ------------------------------------------------------------------------------------------


def test_version_with_a_file_left_behind(app, client, monkeypatch):
    """L-7: a version whose commit failed leaves its file; the next "Fassung festhalten" takes the next number, and the
    file left goes back into the list an hour later (audit 2026-10-21, F3: before, it was removed)."""
    from app import interview_versions

    code, sid = make(app, client)

    def failing():
        raise RuntimeError("database is locked (simulated)")

    monkeypatch.setattr(db.session, "commit", failing)
    with pytest.raises(RuntimeError):
        client.post(f"/guide/{code}/versions", data={"note": "erste"})
    monkeypatch.undo()
    assert client.post(f"/guide/{code}/versions", data={"note": "erste"}).status_code == 302
    with app.app_context():
        db.session.remove()
        survey = db.session.get(Survey, sid)
        assert [v["n"] for v in survey.interview["versions"]] == [2]
        where = interview_versions.folder(survey)
        old = os.path.join(where, "fassung-1.zip")
        os.utime(old, (time.time() - 7200, time.time() - 7200))
    assert client.post(f"/guide/{code}/versions", data={"note": "ethik"}).status_code == 302
    with app.app_context():
        survey = db.session.get(Survey, sid)
        assert [v["n"] for v in survey.interview["versions"]] == [1, 2, 3] and os.path.exists(old)
        restored = survey.interview["versions"][0]
        assert restored["name"] == "fassung-1.zip" and len(restored["sha256"]) == 64 and restored["by"]


def test_subtitles_start_a_paragraph_per_speaker(app):
    """I-8: each turn its own paragraph and time mark; a WebVTT voice becomes the speaker mark."""
    from app import interview_transcribe

    srt = ("1\n00:00:01,000 --> 00:00:03,000\nI: Wie kam es dazu?\n\n"
           "2\n00:00:03,500 --> 00:00:09,000\nB: Das war so, meine Mutter wurde krank.\n\n"
           "3\n00:00:09,200 --> 00:00:10,000\nund dann kam der Pflegedienst.\n\n"
           "4\n00:00:10,200 --> 00:00:11,000\nI: Und dann?\n").encode("utf-8")
    assert interview_transcribe.read_upload("x.srt", srt) == (
        "I: Wie kam es dazu? #00:00:03-0#\n\n"
        "B: Das war so, meine Mutter wurde krank. und dann kam der Pflegedienst. #00:00:10-0#\n\n"
        "I: Und dann? #00:00:11-0#")
    vtt = "WEBVTT\n\n00:00:01.000 --> 00:00:02.000\n<v I>Wie?</v>\n\n00:00:02.500 --> 00:00:04.000\n<v B>So.</v>\n"
    assert interview_transcribe.read_upload("x.vtt", vtt.encode()) == "I: Wie? #00:00:02-0#\n\nB: So. #00:00:04-0#"


def test_text_files_in_any_encoding_and_word_tables(app):
    """I-9: UTF-16 with and without a byte order mark, UTF-32, Windows-1252; NUL bytes never as Windows-1252; the
    rows of a table in a Word file as "Sprecher: Text #Zeit#"."""
    from docx import Document

    from app import interview_transcribe

    text = "I: Grüß Gott. #00:00:05-0#\n\nB: Ähm, ja."
    for data in (text.encode("utf-16"), text.encode("utf-16-le"), text.encode("utf-16-be"), text.encode("utf-32"),
                 text.encode("utf-8-sig"), text.encode("cp1252")):
        assert interview_transcribe.read_upload("x.txt", data) == text
    with pytest.raises(ValueError):
        interview_transcribe.read_upload("x.txt", b"Gr\xfc\xdf\x00 Gott ohne")  # NUL, no UTF-16 pattern
    doc = Document()
    doc.add_paragraph("Interview I01")
    t = doc.add_table(rows=3, cols=3)
    rows = [("Sprecher", "Zeit", "Text"), ("I", "00:00:01", "Wie kam es dazu?"), ("B", "00:00:04", "Meine Mutter wurde krank.")]
    for r, cells in enumerate(rows):
        for c, value in enumerate(cells):
            t.cell(r, c).text = value
    doc.add_paragraph("Ende")
    buf = io.BytesIO()
    doc.save(buf)
    assert interview_transcribe.read_upload("x.docx", buf.getvalue()) == (
        "Interview I01\n\nI: Wie kam es dazu? #00:00:01-0#\n\nB: Meine Mutter wurde krank. #00:00:04-0#\n\nEnde")


def test_uploaded_transcript_covers_the_recording_as_it_is(app, client):
    """I-3: an uploaded transcript counts for the parts there at the upload; one recorded later is not checked with it."""
    code, sid = make(app, client)
    rid = start(app, client, code)
    _send(client, code, rid, 1, 0, WEBM + b"ONE", rec="a")
    client.post(f"/guide/{code}/interviews/{rid}/transcript/upload",
                data={"file": (io.BytesIO("I: Erzählen Sie.".encode()), "t.txt")}, content_type="multipart/form-data")
    _send(client, code, rid, 2, 0, WEBM + b"TWO", rec="b")
    client.post(f"/guide/{code}/interviews/{rid}/transcript", data={"text": "I: Erzählen Sie.", "checked_shown": "1", "checked": "1"})
    tr = _tr(app, rid)
    assert tr["covers"] == {"1": 7} and tr["checked_parts"] == {"1": 7}
