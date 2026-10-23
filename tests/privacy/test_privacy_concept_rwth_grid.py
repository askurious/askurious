"""The privacy concept in the grid of the RWTH forms (record of processing activities, TOM list after Art. 32, example
table of data flows; comparison 2026-10-20, V1-V6, T1-T6, D1): numbered activities with dates, the installation's
catalogue names and legal bases from its overlay, the new fields, the data flow and the processing agreements."""
import io
import json

from docx import Document


def test_every_activity_has_number_dates_and_a_key(app):
    from app import privacy_concept

    with app.test_request_context():
        rows = privacy_concept.activities()
    assert len(rows) == len(privacy_concept.ACTIVITY_META)  # a new activity needs its line in ACTIVITY_META
    assert [r["no"] for r in rows] == list(range(1, len(rows) + 1))
    assert len({r["key"] for r in rows}) == len(rows)
    assert all(r["since"] <= r["changed"] for r in rows)
    assert not any(r["catalog"] or r["own_basis"] for r in rows)  # the public code names no operator's catalogue
    taking_part = next(r for r in rows if r["key"] == "participation")
    assert "Art. 26" in taking_part["controller"] and "Beschäftigte" in taking_part["subjects"]
    assert "Studierende" in next(r for r in rows if r["key"] == "live")["subjects"]


def test_installation_adds_catalogue_name_and_legal_basis(app, client, tmp_path):
    (tmp_path / "privacy_concept.json").write_text(json.dumps({"activities": {
        "accounts": {"catalog": "Benutzer verwalten", "basis": {"de": "Art. 6 Abs. 1 lit. e DSGVO", "en": "Art. 6 (1) (e)"}},
        "unknown": {"catalog": "x"}}}), encoding="utf-8")
    app.config["INSTANCE_OVERLAY"] = str(tmp_path)
    page = client.get("/admin/privacy").get_data(as_text=True)
    assert "Benutzer verwalten" in page and "Rechtsgrundlage dieser Installation" in page
    assert "Art. 6 Abs. 1 lit. e DSGVO" in page
    text = "\n".join(c.text for t in Document(io.BytesIO(client.get("/admin/privacy.docx").data)).tables
                     for r in t.rows for c in r.cells)
    assert "Benutzer verwalten" in text and "Art. 6 Abs. 1 lit. e DSGVO" in text
    (tmp_path / "privacy_concept.json").write_text("{not json", encoding="utf-8")
    assert client.get("/admin/privacy").status_code == 200  # an unreadable file is left out, not an error


def test_new_fields_flow_and_open_points(app, client):
    page = client.get("/admin/privacy").get_data(as_text=True)
    for words in ("Fachabteilung mit Organisationskennziffer", "Informationssicherheitsbeauftragte Person",
                  "Notfallplan", "4. Datenfluss einer Befragung", "6. Technische und organisatorische Maßnahmen",
                  "Rekrutierung", "Veröffentlichung", "Maßnahmen: siehe Abschnitt 6.", "pg_restore -l",
                  "ob Datenträger und Datenbank verschlüsselt sind", "Auftragsverarbeitung"):
        assert words in page, words
    r = client.post("/admin/privacy", data={"privacy_restore_test": "2026-10-01: vollständig, ok",
                                            "privacy_encryption": "LUKS", "autosave": "1"})
    assert r.get_json()["status"] == "saved"
    page = client.get("/admin/privacy").get_data(as_text=True)
    # the open points go; pg_restore -l stays named in the control types (section 6)
    assert "vollständige Wiederherstellung zur Probe" not in page and "ob Datenträger und Datenbank verschlüsselt sind" not in page


def test_processing_agreements_and_art32_table(app, client):
    app.config.update(OPENAI_API_KEY="sk-test")
    page = client.get("/admin/privacy").get_data(as_text=True)
    assert "OpenAI (Text) ist an: Vertrag zur Auftragsverarbeitung eintragen" in page
    client.post("/admin/privacy", data={"privacy_avv_openai": "seit 2026-10-01, Ordner Verträge", "autosave": "1"})
    page = client.get("/admin/privacy").get_data(as_text=True)
    assert "seit 2026-10-01, Ordner Verträge" in page and "OpenAI (Text) ist an: Vertrag" not in page
    doc = Document(io.BytesIO(client.get("/admin/privacy.docx").data))
    cells = [c.text for t in doc.tables for r in t.rows for c in r.cells]
    assert "Art. 32 Abs. 1" in cells and "lit. b, c" in cells
    assert any("Datenfluss einer Befragung" in p.text for p in doc.paragraphs)


def test_procedure_internal_recipients_third_countries_and_control_types(app, client):
    """The rest of the RWTH templates (2026-10-20, points 1-4): procedure, internal recipients and third countries per
    activity, the control types of the TOM list, the step "Registrierung und Pseudonym" of the data flow."""
    from app import privacy_concept

    app.config.update(OPENAI_API_KEY="", SLACK_WEBHOOK_URL="", ORCID_CLIENT_ID="")
    with app.test_request_context():
        rows = {r["key"]: r for r in privacy_concept.activities()}
        assert set(privacy_concept.INTERNAL) == set(rows)  # every activity names its internal recipients
        assert "Findet nicht statt" in rows["ai"]["third_country"]
        app.config.update(OPENAI_API_KEY="sk-test")
        rows = {r["key"]: r for r in privacy_concept.activities()}
        assert "OpenAI (USA)" in rows["ai"]["third_country"] and "Findet nicht statt" in rows["pool"]["third_country"]
        assert "Admins der Plattform sehen keine Antworten" in rows["participation"]["internal"]
        assert all(r["procedure"] for r in rows.values())
        names = [name for name, _lit, _text in privacy_concept.controls()]
        assert names[:3] == ["Pseudonymisierung", "Verschlüsselung", "Zutrittskontrolle"]
        assert [f["step"] for f in privacy_concept.flow()][2] == "Registrierung und Pseudonym"
    page = client.get("/admin/privacy").get_data(as_text=True)
    for words in ("Übermittlung in ein Drittland", "Empfänger intern", "Verfahren", "Kontrollarten (Art. 32 Abs. 1)",
                  "Weitergabekontrolle", "Registrierung und Pseudonym"):
        assert words in page, words
    cells = [c.text for t in Document(io.BytesIO(client.get("/admin/privacy.docx").data)).tables
             for r in t.rows for c in r.cells]
    assert "Eingabekontrolle" in cells and "Übermittlung in ein Drittland (e)" in cells


def test_deploy_ships_the_overlay_json():
    from pathlib import Path

    script = (Path(__file__).resolve().parents[2] / "deploy/docker/deploy-docker.sh").read_text(encoding="utf-8")
    assert "parts+=(privacy_concept.json)" in script
