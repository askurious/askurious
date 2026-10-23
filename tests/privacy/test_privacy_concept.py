"""Administration > Datenschutz: the Datenschutzkonzept built from code and configuration (app/privacy_concept.py),
the fields admins fill in, and the Word file."""
import io
import re

from docx import Document

from app.extensions import db
from app.models import SiteSetting


def test_page_only_for_admins_and_names_every_section(app, client):
    page = client.get("/admin/privacy").get_data(as_text=True)
    assert re.search(r'aria-current="page" href="/admin/privacy" title="[^"]*">Datenschutz<', page)
    for title in ("Verzeichnis der Verarbeitungstätigkeiten (Art. 30)", "Empfänger und externe Dienste",
                  "Technische und organisatorische Maßnahmen (Art. 32)", "Datenschutzfreundliche Voreinstellungen",
                  "Schwellwertanalyse", "Offene Punkte", "Quellen"):
        assert title in page
    for goal in ("Vertraulichkeit", "Integrität", "Nichtverkettung", "Transparenz", "Intervenierbarkeit",
                 "Datenminimierung"):
        assert f"<h3>{goal}" in page
    other = app.test_client()
    other.get("/auth/dev-login?orcid=0000-0002-1825-0097&name=Other")
    other.post("/terms/", data={"accept": "1"})
    assert other.get("/admin/privacy").status_code == 403
    assert other.get("/admin/privacy.docx").status_code == 403


def test_services_show_what_is_switched_on_here(app, client):
    app.config.update(OPENAI_API_KEY="", SLACK_WEBHOOK_URL="", WHISPER_URL="https://whisper.example.org/v1")
    with app.test_request_context():
        from app import privacy_concept

        rows = {r["switch"]: r for r in privacy_concept.services()}
        assert rows["OPENAI_API_KEY"]["state"][0] is False
        assert rows["SLACK_WEBHOOK_URL"]["state"][0] is False
        assert rows["WHISPER_URL"]["state"] == (True, "an (whisper.example.org)")
        own = [r for r in privacy_concept.services() if r["name"].startswith("Eigener KI-Zugang")]
        assert own and own[0]["state"] == (False, "für alle; 0 Personen mit eigenem Schlüssel")  # for everyone since 2026-10-10
    app.config.update(SLACK_WEBHOOK_URL="https://hooks.slack.com/x")
    page = client.get("/admin/privacy").get_data(as_text=True)
    assert "Slack erhält Auszüge von Beschwerden" in page  # an open point once Slack is on


def test_open_points_name_missing_fields_and_admin_second_factor(app, client):
    app.config.update(ADMIN_TOTP=False)
    page = client.get("/admin/privacy").get_data(as_text=True)
    assert "Nicht ausgefüllt: Verantwortlicher (Name, Anschrift)" in page
    assert "Zweiter Faktor für Admins ist aus" not in page and "nach ORCID" in page


def test_fields_are_saved_trimmed_and_shown(app, client):
    r = client.post("/admin/privacy", data={"privacy_controller": "  RWTH Aachen,\n Lehrstuhl  ",
                                            "privacy_dpo": "x" * 300, "autosave": "1"})
    assert r.get_json()["status"] == "saved"
    with app.app_context():
        assert db.session.get(SiteSetting, "privacy_controller").value == "RWTH Aachen, Lehrstuhl"
        assert len(db.session.get(SiteSetting, "privacy_dpo").value) == 200
        assert db.session.get(SiteSetting, "privacy_authority") is None  # not sent, not stored
    page = client.get("/admin/privacy").get_data(as_text=True)
    assert 'value="RWTH Aachen, Lehrstuhl"' in page
    assert "Nicht ausgefüllt: Verantwortlicher" not in page


def test_word_file_has_the_same_content(app, client):
    client.post("/admin/privacy", data={"privacy_breach": "Lehrstuhlleitung, Tel. 123"})
    r = client.get("/admin/privacy.docx")
    assert r.status_code == 200 and r.headers["Cache-Control"] == "no-store"
    assert re.search(r'filename="datenschutzkonzept-\d{4}-\d\d-\d\d\.docx"', r.headers["Content-Disposition"])
    doc = Document(io.BytesIO(r.data))
    text = "\n".join(p.text for p in doc.paragraphs)
    cells = "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert "Verzeichnis der Verarbeitungstätigkeiten (Art. 30)" in text
    assert "Lehrstuhlleitung, Tel. 123" in text
    assert "KI-Kodierhilfe (nur eigener KI-Zugang)" in cells and "Teilnahme an Befragungen" in text


def test_appendix_impact_assessment_of_the_ai_coding_help(app, client):
    """Appendix 1 (Ralf, 2026-10-10): the DSFA of the AI coding help with the code's own limits, its own field for the
    data protection officer's statement, in the Word file too; section 7 points to it."""
    from app import coding_ai

    page = client.get("/admin/privacy").get_data(as_text=True)
    assert 'id="dsfa-kodierhilfe"' in page and "Anlage 1: Datenschutz-Folgenabschätzung KI-Kodierhilfe" in page
    assert f"höchstens {coding_ai.MAX_TEXTS} je Aufruf, je höchstens {coding_ai.MAX_CHARS} Zeichen" in page
    assert "DSFA in Anlage 1" in page  # the threshold analysis points to it
    assert "Nicht ausgefüllt: Verantwortlicher (Name, Anschrift), Datenschutzbeauftragte Person (Kontakt)" in page
    assert "Stellungnahme der datenschutzbeauftragten Person zur DSFA KI-Kodierhilfe" in page.split("Nicht ausgefüllt:")[1]
    r = client.post("/admin/privacy", data={"privacy_dpia_coding": "2026-11-02: keine Einwände", "autosave": "1"})
    assert r.get_json()["status"] == "saved"
    page = client.get("/admin/privacy").get_data(as_text=True)
    assert 'value="2026-11-02: keine Einwände"' in page
    doc = Document(io.BytesIO(client.get("/admin/privacy.docx").data))
    text = "\n".join(p.text for p in doc.paragraphs)
    cells = "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert "Anlage 1: Datenschutz-Folgenabschätzung KI-Kodierhilfe" in text and "2026-11-02: keine Einwände" in text
    assert "Re-Identifizierung" in cells and "Innovative Nutzung neuer Technologien" in cells


def test_english_page(app, client):
    with app.app_context():
        from app.models import User

        admin = db.session.scalar(db.select(User).where(User.is_admin.is_(True)))
        admin.display = {**(admin.display or {}), "lang": "en"}
        db.session.commit()
    page = client.get("/admin/privacy").get_data(as_text=True)
    assert "Record of processing activities (Art. 30)" in page and "Verzeichnis der" not in page
