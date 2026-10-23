"""UI/UX audit 2026-10-05, rests: "weiß nicht" in the questionnaire documents (V41), the quotas of the methods
section on the pool of the results page, the link of the quota gap message to the quota page (V15); the staged
contact check (V13) is in test_ux_audit5_check.py."""
import io
from datetime import datetime, timedelta

from docx import Document
from werkzeug.datastructures import MultiDict

from app import design, export, methods, questionnaire_docx
from app.extensions import db
from app.models import Response, Survey
from tests.conftest import add
from tests.participation.test_audit_g_concurrency import _coverage
from tests.quotas.test_quotas import build


def _text(doc):
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        parts += [c.text for row in table.rows for c in row.cells]
    return "\n".join(parts)


def test_dont_know_in_word_and_markdown_questionnaire(app, client, survey):
    add(client, "likert")                                                       # 2
    client.post("/surveys/1/editor/items/2", data=MultiDict([
        ("code", "TA"), ("question", "Wie sehen Sie das?"), ("points", "5"), ("labeling", "full"),
        ("answer_mode", "optional"), ("dont_know", "1"), ("stmt_id", "s1"), ("stmt_code", ""),
        ("stmt_text", "Bus und Bahn sind gut ausgebaut")]))
    add(client, "rating")                                                       # 3
    client.post("/surveys/1/editor/items/3", data={"code": "R", "question": "Note?", "answer_mode": "optional",
                                                   "style": "numbers", "min": "0", "max": "10", "dont_know": "1"})
    add(client, "date_time")                                                    # 4
    client.post("/surveys/1/editor/items/4", data={"code": "D", "question": "Wann?", "answer_mode": "optional",
                                                   "parts": "year", "span": "past", "minute_step": "5", "dont_know": "1"})
    with app.app_context():
        s = db.session.get(Survey, 1)
        book = _text(Document(io.BytesIO(questionnaire_docx.build(s, "codebook"))))
        paper = _text(Document(io.BytesIO(questionnaire_docx.build(s, "print"))))
        md = export.questionnaire_md(s)
    assert book.count("„weiß nicht“ = -98, zählt als fehlend.") == 2           # matrix and rating
    assert "D_dk = 1" in book                                                   # the date's own column
    assert paper.count("weiß nicht") == 3 and "-98" not in paper                # a column, two ticks; no codes
    assert md.count('Plus "don\'t know" (exported as -98, counts as missing)') == 2 and "`D_dk` = 1" in md


def test_methods_count_the_quotas_in_the_pool_of_the_page(app, client, survey):
    build(app, client)
    with app.app_context():
        s = db.session.get(Survey, 1)
        quota = s.quotas[0]
        men = quota.cells[0]
        t0 = datetime(2026, 9, 1, 10)
        for _ in range(2):  # dummies keep their decisions in Response.quota, but move no counter
            s.responses.append(Response(status=Response.COMPLETED, source="dummy", is_preview=True, started_at=t0,
                                        completed_at=t0 + timedelta(seconds=300), answers={"2": "1", "3": 30},
                                        quota={str(quota.id): {"cell": men.id, "res": "in", "held": False, "done": True}}))
        db.session.commit()
        real = methods.facts(s)["quotas"][0]
        pool = methods.facts(s, sources=("real", "dummy"))["quotas"][0]
    assert real["done"] == 0 and real["cells"][0] == ("Männer", 2, 0)
    assert pool["done"] == 2 and pool["cells"][0] == ("Männer", 2, 2)


def test_quota_gap_links_to_the_quota_page(app, client, survey):
    _coverage(app, client, [("Jung", "j", 1, [(3, "between", "18..29")]), ("Alt", "a", 1, [(3, "between", "31..45")])])
    with app.app_context():
        gap = [f for f in design.check(db.session.get(Survey, 1)) if "passt in keine Teilquote" in f.message]
    assert [f.link for f in gap] == ["flow_page.quotas"]
    html = client.get("/surveys/1/editor/").data.decode()
    assert '/quotas/' in html.split("passt in keine Teilquote")[0].rsplit("<a ", 1)[1]
