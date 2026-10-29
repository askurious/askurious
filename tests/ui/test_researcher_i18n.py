"""The researchers' and admins' pages in German and English (app/lang.py, concept of 2026-10-06): every text has its
English half, de_en("Konto", "Account") in templates, Python and scripts.

A ratchet while the pages are being translated: tests/i18n_rest.json notes per file how many German texts stand
there without their English half. A file may not get more (new text needs de_en), and once it has fewer, the list
is lowered with

    python tests/ui/test_researcher_i18n.py --update

which only ever lowers counts, never raises one or adds a file. When every count is 0, the list goes.
Counted: text and title/placeholder/aria-label/alt/data-confirm* in templates (also string literals inside {{ }}
and {% %}), messages of flash() in Python, string literals in app/static/*.js. German is recognised by umlauts or
common German words. The participants' pages speak through ui()/tr() (tests/ui/test_ui_i18n.py); what they say inside
ui() is not counted here, and the folder templates/participant is theirs alone.
"""
import ast
import glob
import json
import os
import re
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
REST = os.path.join(ROOT, "tests", "i18n_rest.json")

GERMAN = re.compile(r"[äöüÄÖÜß]|\b(?:der|die|das|dem|den|des|und|oder|nicht|mit|für|ist|sind|ein|eine|einen|einem|einer|Sie|Ihr|Ihre|"
                    r"Ihren|Ihrem|Ihrer|bitte|Bitte|zu|zum|zur|im|auf|von|vom|wird|werden|wurde|noch|alle|kein|keine|nur|hier|"
                    r"wenn|dann|bei|nach|aus|unter|wie|auch|schon|Seite|Befragung|Befragungen|Teilnahme|Teilnahmen|Teilnehmende|"
                    r"Speichern|Anmelden|Abmelden|Einstellungen|Ergebnisse|Fragebogen|Konto|Gruppe|Gruppen|Weiter|Abbrechen|"
                    r"Hinzufügen|Ablauf|Vorschau|Element|Elemente|Frage|Fragen|Antwort|Antworten|Datei|Daten)\b")
SKIP_CALLS = ("de_en", "de_en_n", "de_en_pick", "ui", "url_for", "icon", "csrf", "tr", "uiT")
LITERAL = re.compile(r'"((?:[^"\\]|\\.)*)"|\'((?:[^\'\\]|\\.)*)\'')


REFERENCE = re.compile(r"\(\d{4}[a-z]?\)|, \d{4}[a-z]?\b")  # a source (Kühne & Kroh, 2018): names are no German


def german(text):
    return bool(text and GERMAN.search(text)) and not REFERENCE.search(text)


def _strip_calls(code, names=SKIP_CALLS):
    """`code` without the calls of `names` (their arguments included): strings inside are skipped while counting
    parentheses, so "(höchstens …)" in a text does not end a call early."""
    out, i, n = [], 0, len(code)
    call = re.compile(r"\b(?:" + "|".join(map(re.escape, names)) + r")\s*\(")
    while i < n:
        m = call.search(code, i)
        if not m:
            out.append(code[i:])
            break
        out.append(code[i:m.start()])
        j, depth, quote = m.end(), 1, None
        while j < n and depth:
            c = code[j]
            if quote:
                if c == "\\":
                    j += 1
                elif c == quote:
                    quote = None
            elif c in "\"'`":
                quote = c
            elif c == "(":
                depth += 1
            elif c == ")":
                depth -= 1
            j += 1
        i = j
    return "".join(out)


TAG = re.compile(r"\{\{.*?\}\}|\{%.*?%\}", re.S)
TEXT = re.compile(r">([^<>]+)<")
ATTR = re.compile(r'\b(?:title|placeholder|aria-label|alt|data-confirm(?:-[\w-]+)?)="([^"]*)"')


def template_texts(source):
    """The German texts of a template that are not inside de_en() (or ui() of the participants' pages)."""
    source = re.sub(r"\{#.*?#\}", "", source, flags=re.S)
    found = []
    for tag in TAG.findall(source):
        for m in LITERAL.finditer(_strip_calls(tag)):
            text = m.group(1) if m.group(1) is not None else m.group(2)
            if german(text):
                found.append(text)
    plain = TAG.sub("\x00", source)
    for m in list(TEXT.finditer(plain)) + list(ATTR.finditer(plain)):
        text = " ".join(m.group(1).replace("\x00", " ").split())
        if german(text):
            found.append(text)
    return found


# errors whose text reaches the page (UI/UX audit 2026-10-09, X85): raise …Error("…"), Refused("…"), AIError("…");
# AnswerError speaks to participants (Both, tr) and is not counted here
SHOWN_ERRORS = ("Refused", "AIError")


def _error_message(node):
    """The first argument of a raised error or of Refused(…)/AIError(…) whose text a researcher reads, else None."""
    call = node.exc if isinstance(node, ast.Raise) and isinstance(node.exc, ast.Call) else node
    if not isinstance(call, ast.Call) or not call.args:
        return None
    name = getattr(call.func, "id", "") or getattr(call.func, "attr", "")
    if name == "AnswerError" or not (name.endswith("Error") or name in SHOWN_ERRORS):
        return None
    if call is node and name not in SHOWN_ERRORS:
        return None  # a bare ValueError(...) that is not raised here is no message yet
    return call.args[0]


def _messages(tree):
    """The parts of a module that researchers read: flash(), raised errors and what problems() lists."""
    for node in ast.walk(tree):
        if isinstance(node, ast.Call) and getattr(node.func, "id", "") == "flash" and node.args:
            yield node.args[0]
        elif isinstance(node, (ast.Raise, ast.Call)):
            message = _error_message(node)
            if message is not None:
                yield message
        elif isinstance(node, ast.FunctionDef) and node.name == "problems":
            for inner in ast.walk(node):
                if isinstance(inner, ast.Call) and getattr(inner.func, "attr", "") == "append" and inner.args:
                    yield inner.args[0]


def python_texts(source):
    """German messages of flash() calls, raised errors and problems() that are not de_en()."""
    found, seen = [], set()
    for message in _messages(ast.parse(source)):
        if id(message) in seen:
            continue
        seen.add(id(message))
        parts = [n.value for n in ast.walk(message) if isinstance(n, ast.Constant) and isinstance(n.value, str)
                 and not _inside_de_en(message, n)]
        text = " ".join(parts)
        if german(text):
            found.append(text)
    return found


def _inside_de_en(root, target, names=("de_en", "Both", "tr")):
    for n in ast.walk(root):
        if isinstance(n, ast.Call) and getattr(n.func, "id", "") in names and any(c is target for c in ast.walk(n)):
            return True
    return False


def script_texts(source):
    """German string literals of a script outside de_en() (and uiT() of participant.js)."""
    source = re.sub(r"/\*.*?\*/", "", source, flags=re.S)
    source = re.sub(r"(?m)^\s*//.*$|(?<=[;{}),\s])//[^\n\"']*$", "", source)
    found = []
    for m in re.finditer(r'"((?:[^"\\\n]|\\.)*)"|\'((?:[^\'\\\n]|\\.)*)\'|`([^`]*)`', _strip_calls(source)):
        text = next(g for g in m.groups() if g is not None)
        if german(text):
            found.append(text)
    return found


def counted_files():
    files = {}
    for f in sorted(glob.glob(os.path.join(ROOT, "app", "templates", "**", "*.html"), recursive=True)):
        rel = os.path.relpath(f, ROOT).replace(os.sep, "/")
        if not rel.startswith("app/templates/participant/"):
            files[rel] = template_texts
    for f in sorted(glob.glob(os.path.join(ROOT, "app", "**", "*.py"), recursive=True)):
        files[os.path.relpath(f, ROOT).replace(os.sep, "/")] = python_texts
    for f in sorted(glob.glob(os.path.join(ROOT, "app", "static", "*.js"))):
        files[os.path.relpath(f, ROOT).replace(os.sep, "/")] = script_texts
    return files


def found_now():
    out = {}
    for rel, find in counted_files().items():
        with open(os.path.join(ROOT, rel), encoding="utf-8") as f:
            texts = find(f.read())
        if texts:
            out[rel] = texts
    return out


def rest():
    with open(REST, encoding="utf-8") as f:
        return json.load(f)


def test_no_new_german_text_without_english():
    noted, now = rest(), found_now()
    more = {rel: texts for rel, texts in now.items() if len(texts) > noted.get(rel, 0)}
    assert not more, ("Deutscher Text ohne englische Hälfte, bitte als de_en(\"Deutsch\", \"English\") schreiben "
                      "(app/lang.py): " + "; ".join(f"{rel}: {len(t)} statt {noted.get(rel, 0)}, z. B. „{t[-1][:60]}“"
                                                     for rel, t in sorted(more.items())[:5]))


def test_the_rest_list_only_shrinks():
    noted, now = rest(), found_now()
    fewer = sorted(rel for rel, n in noted.items() if len(now.get(rel, [])) < n)
    assert not fewer, ("Weniger deutsche Texte als in tests/i18n_rest.json notiert, bitte die Liste senken: "
                       "python tests/ui/test_researcher_i18n.py --update (" + ", ".join(fewer[:5]) + ")")


PAIR = re.compile(r'(?:de_en|Both)\(\s*(?:"((?:[^"\\]|\\.)*)"|\'((?:[^\'\\]|\\.)*)\')\s*,\s*(?:"((?:[^"\\]|\\.)*)"|\'((?:[^\'\\]|\\.)*)\')')
FIELDS = re.compile(r"\{(\w*)\}")


def pairs():
    """Every (German, English) literal pair of de_en() and Both() in templates, Python and scripts."""
    out = []
    for pattern in ("app/templates/**/*.html", "app/**/*.py", "app/static/*.js"):
        for f in glob.glob(os.path.join(ROOT, pattern), recursive=True):
            with open(f, encoding="utf-8") as fh:
                source = fh.read()
            for m in PAIR.finditer(source):
                out.append((os.path.relpath(f, ROOT), m.group(1) if m.group(1) is not None else m.group(2),
                            m.group(3) if m.group(3) is not None else m.group(4)))
    return out


def test_every_pair_has_its_english_half():
    found = pairs()
    assert len(found) > 300  # the pattern still finds them
    bad = [(f, de) for f, de, en in found if not en.strip() or (en == de and german(de))]
    assert not bad, "englische Hälfte fehlt oder ist deutsch: " + "; ".join(f"{f}: „{de[:50]}“" for f, de in bad[:5])
    fields = [(f, de) for f, de, en in found if set(FIELDS.findall(de)) != set(FIELDS.findall(en))]
    assert not fields, "Platzhalter verschieden: " + "; ".join(f"{f}: „{de[:50]}“" for f, de in fields[:5])


MACRO = re.compile(r"\{%-?\s*macro\s+\w+\((.*?)\)\s*-?%\}", re.S)


def test_macro_defaults_hold_no_text():
    """A default of a macro parameter cannot be a de_en() pair, and the scan of texts above does not see it
    ("Freigabe beenden" in the English dialog, UI audit W77): a text as default is none, filled inside the macro."""
    bad = []
    for f in glob.glob(os.path.join(ROOT, "app", "templates", "**", "*.html"), recursive=True):
        with open(f, encoding="utf-8") as fh:
            source = fh.read()
        for args in MACRO.findall(source):
            for m in re.finditer(r"=\s*" + LITERAL.pattern, args):
                text = m.group(1) if m.group(1) is not None else m.group(2)
                if re.search(r"[a-zäöüß]\s+\w", text):
                    bad.append(f"{os.path.relpath(f, ROOT)}: „{text}“")
    assert not bad, "Text als Standardwert eines Makros: " + "; ".join(bad)


def _user(app, orcid="0000-0000-0000-001X"):
    from app.extensions import db
    from app.models import User

    with app.app_context():
        return db.session.scalar(db.select(User).filter_by(orcid=orcid)).display


def test_german_stays_the_standard(client):
    page = client.get("/dashboard/").data.decode()
    assert '<html lang="de"' in page and "Meine Befragungen" in page and 'id="sprache"' not in page
    konto = client.get("/account/appearance").data.decode()
    assert 'id="sprache"' in konto and 'value="de" lang="de" data-reload-saved checked' in konto


def test_english_chosen_on_konto(app, client):
    r = client.post("/account/language", data={"lang": "en"}, follow_redirects=True)
    page = r.data.decode()
    assert '<html lang="en"' in page and "Language saved." in page and "<h1>Display and language</h1>" in page
    assert _user(app)["lang"] == "en"
    page = client.get("/dashboard/").data.decode()
    assert "My surveys" in page and "Meine Befragungen" not in page and " Sign out</button>" in page
    # the look and the language live in one column: saving the one keeps the other
    client.post("/account/display", data={"mode": "dark", "look": "askurious", "size": "100"})
    client.post("/account/display/mode", data={"mode": "light"})
    assert _user(app) == {"mode": "light", "look": "askurious", "size": 100, "lang": "en"}
    assert client.post("/account/language", data={"lang": "fr"}).status_code == 400
    client.post("/account/language", data={"lang": "de"})
    assert "Meine Befragungen" in client.get("/dashboard/").data.decode()


def test_messages_and_scripts_follow_the_language(app, client):
    client.post("/account/language", data={"lang": "en"})
    r = client.post("/account/details", data={"name": ""}, follow_redirects=True)  # ORCID accounts: 403
    assert r.status_code == 403
    r = client.post("/account/password", data={"email": "x"}, follow_redirects=True).data.decode()
    assert "Please" in r or "sign-in" in r
    from app.lang import de_en

    with app.test_request_context(headers={"Accept-Language": "en"}):
        assert de_en("Konto", "Account") == "Account" and de_en("{n} Tage", "{n} days", n=3) == "3 days"
    with app.test_request_context(headers={"Accept-Language": "fr, de;q=0.5"}):
        assert de_en("Konto", "Account") == "Konto"
    with app.test_request_context():
        assert de_en("Konto", "Account") == "Konto"


def test_before_signing_in_the_browser_decides(app):
    guest = app.test_client()
    page = guest.get("/", headers={"Accept-Language": "en-GB,en;q=0.9"}).data.decode()
    assert '<html lang="en"' in page and "Sign in with ORCID" in page and "Three ways to ask" in page
    page = guest.get("/", headers={"Accept-Language": "de-DE,de;q=0.9,en;q=0.8"}).data.decode()
    assert '<html lang="de"' in page and "Drei Arten zu fragen" in page
    assert "Sign in" in guest.get("/auth/sign-in", headers={"Accept-Language": "en"}).data.decode()


def test_a_new_account_takes_the_browser_language_once(app):
    new = app.test_client()
    new.get("/auth/dev-login?orcid=0000-0002-1825-0097&name=Neu", headers={"Accept-Language": "en-US,en;q=0.9"})
    assert _user(app, "0000-0002-1825-0097") == {"mode": "light", "lang": "en"}
    other = app.test_client()
    other.get("/auth/dev-login?orcid=0000-0000-0000-0124&name=Andere", headers={"Accept-Language": "de"})
    assert _user(app, "0000-0000-0000-0124") == {"mode": "light"}
    # signing in again later with another browser changes nothing
    again = app.test_client()
    again.get("/auth/dev-login?orcid=0000-0002-1825-0097&name=Neu", headers={"Accept-Language": "de"})
    assert _user(app, "0000-0002-1825-0097")["lang"] == "en"


def test_numbers_and_dates_follow_the_page_language(app, client, survey):
    """E11: researcher pages in English write 6 Oct 2026 and 9.50 / 35%; German stays 06.10.2026, 9,50 and 35 %;
    the participants' templates (here the comments beside the editor's preview) keep their German format."""
    from datetime import datetime

    from werkzeug.datastructures import MultiDict

    from app import units
    from app.extensions import db
    from app.models import PreviewComment, Response, Survey
    from tests.conftest import add

    add(client, "single_choice")                                    # 2
    client.post("/surveys/1/editor/items/2", data=MultiDict([
        ("code", "farbe"), ("question", "Farbe?"), ("option_id", "a"), ("option_code", "1"), ("option_label", "rot"),
        ("option_id", "b"), ("option_code", "2"), ("option_label", "blau")]))
    add(client, "text_input")                                       # 3
    client.post("/surveys/1/editor/items/3", data={"code": "anzahl", "question": "Anzahl?", "validation": "integer"})
    client.post("/surveys/1/editor/preview/comment", data={"block_id": "2", "item_id": "2", "page": "1", "text": "Unklar"})
    stamp = datetime(2026, 10, 6, 12, 5)                            # UTC: 14:05 in Berlin
    with app.app_context():
        s = db.session.get(Survey, 1)
        for i in range(20):
            s.responses.append(Response(status=Response.COMPLETED, started_at=stamp, completed_at=stamp,
                                        answers={"2": "1" if i < 7 else "2", "3": i}))
        s.status, s.field_end = Survey.RUNNING, stamp
        for c in db.session.scalars(db.select(PreviewComment)):
            c.created_at = stamp
        db.session.commit()

    def pages():
        return (client.get("/dashboard/").data.decode(), client.get("/surveys/1/results").data.decode(),
                client.get("/surveys/1/editor/comments").data.decode(),
                client.get("/surveys/1/editor/blocks/2/preview", follow_redirects=True).data.decode())

    dash, res, comments, preview = pages()                          # German: as before
    assert "06.10.2026" in dash and "9,50" in res and ">35,0 % (7)</span>" in res and "06.10.2026 14:05" in comments
    assert "06.10.2026 14:05" in preview
    client.post("/account/language", data={"lang": "en"})
    dash, res, comments, preview = pages()
    assert "until 6 Oct 2026" in dash and "06.10.2026" not in dash
    assert "9.50" in res and "9,50" not in res and ">35.0% (7)</span>" in res
    assert "6 Oct 2026, 14:05" in comments
    assert "06.10.2026 14:05" in preview and "6 Oct 2026" not in preview    # the participants' part keeps its format
    with app.test_request_context("/s/abc"):                         # the participants' blueprint: German numbers
        assert units.page_lang() == "de"
    assert units.page_lang() == "de"                                 # outside a request (exports, jobs)
    assert units.stat(1234.5, 1) == "1.234,5" and units.stat(1234.5, 1, "en") == "1,234.5"


def test_kept_results_are_kept_per_language(app, client, survey):
    """The kept results hold texts made while computing (warnings, notes) in the page's language: an English page
    keeps and reads files of its own, so a German computation never shows German texts to an English reader."""
    from app import results_store
    from app.extensions import db
    from app.models import Response, Survey

    with app.app_context():
        s = db.session.get(Survey, 1)
        s.responses.append(Response(status=Response.COMPLETED, answers={}))
        db.session.commit()
        uid = s.uid
    folder = os.path.join(app.config["UPLOAD_FOLDER"], "hb", uid)
    assert client.get("/surveys/1/results").status_code == 200
    assert "results.bin" in os.listdir(folder) and "results-en.bin" not in os.listdir(folder)
    client.post("/account/language", data={"lang": "en"})
    assert client.get("/surveys/1/results").status_code == 200
    assert "results-en.bin" in os.listdir(folder)
    with app.test_request_context():                                # outside a researcher's page: German
        s = db.session.get(Survey, 1)
        german = results_store._path(s)
    with app.test_request_context(headers={"Accept-Language": "en"}):
        s = db.session.get(Survey, 1)
        assert results_store._path(s) != german and results_store._path(s, {"weighted": True}).endswith("results-weighted-en.bin")
    assert all(name in results_store.KEPT for name in os.listdir(folder) if name.endswith(".bin"))


def _docx_text(body):
    import io

    from docx import Document

    doc = Document(io.BytesIO(body))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        parts += [c.text for row in table.rows for c in row.cells]
    return "\n".join(parts)


def test_questionnaire_word_frame_follows_the_downloader(app, client, survey):
    """The codebook's frame (notes, answer modes, element types) in the downloader's language (E8); questions,
    answers and variable names as written, in the survey's default language."""
    from werkzeug.datastructures import MultiDict

    from tests.conftest import add

    add(client, "likert")                                                       # 2
    client.post("/surveys/1/editor/items/2", data=MultiDict([
        ("code", "TA"), ("question", "Wie sehen Sie das?"), ("points", "5"), ("labeling", "full"),
        ("answer_mode", "required"), ("dont_know", "1"), ("stmt_id", "s1"), ("stmt_code", ""),
        ("stmt_text", "Bus und Bahn sind gut ausgebaut")]))
    german = _docx_text(client.get("/surveys/1/export/questionnaire.docx").data)
    assert "Fragebogen mit Codebuch" in german and "Likert-Skala / Matrix · Pflicht" in german
    assert "„weiß nicht“ = -98, zählt als fehlend." in german
    client.post("/account/language", data={"lang": "en"})
    english = _docx_text(client.get("/surveys/1/export/questionnaire.docx").data)
    assert "Questionnaire with codebook" in english and "Likert scale / matrix · required" in english
    assert "“don't know” = -98, counts as missing." in english and "Fragebogen mit Codebuch" not in english
    assert "Wie sehen Sie das?" in english and "TA_1" in english and "Bus und Bahn sind gut ausgebaut" in english
    paper = _docx_text(client.get("/surveys/1/export/questionnaire-print.docx").data)
    assert "Fragen mit * sind Pflichtangaben." in paper                        # what participants read: the survey's


def test_scale_table_default_language_is_the_downloaders(app, client, survey):
    """The APA tables without an explicit language take the downloader's; the simulator's link passes it."""
    import io
    import zipfile

    from app import manuscript_tables
    from app.extensions import db
    from app.models import Survey

    data = {"n_rows": 10, "overview": {"rows": [{"label": "S", "k": 3, "n": 10, "mean": 3.1, "sd": 0.5, "alpha": 0.8,
                                                  "omega": 0.81}], "corr": [[]]}}
    with app.test_request_context(headers={"Accept-Language": "en"}):
        en = zipfile.ZipFile(io.BytesIO(manuscript_tables.build(db.session.get(Survey, 1), data))).read("word/document.xml").decode()
    with app.test_request_context():
        de = zipfile.ZipFile(io.BytesIO(manuscript_tables.build(db.session.get(Survey, 1), data))).read("word/document.xml").decode()
    assert "Table 1" in en and ">Scale<" in en and ">.80<" in en and ">k<" in en
    assert "Tabelle 1" in de and ">Skala<" in de and ">0,80<" in de


def test_preregistration_texts_follow_the_person_who_starts(app, client, survey):
    """E8: README and research plan of the snapshot in the starter's language; file names stay (app/prereg.py)."""
    import io
    import zipfile

    from app import prereg
    from app.extensions import db
    from app.models import Survey
    from tests.conftest import fill_consent

    with app.app_context():
        s = db.session.get(Survey, 1)
        s.research_question = "Pendeln Teams länger?"
        db.session.commit()
    fill_consent(client)
    client.post("/surveys/1/editor/blocks/2/items", data={"type": "info"})
    client.post("/account/language", data={"lang": "en"})
    client.post("/surveys/1/start", data={})
    with app.app_context():
        [(name, *_rest)] = prereg.listing(db.session.get(Survey, 1))
    with zipfile.ZipFile(io.BytesIO(client.get(f"/surveys/1/export/preregistration/{name}").data)) as z:
        assert {"README.md", "survey.json", "questionnaire.md", "checksums.sha256", "study-plan.md"} <= set(z.namelist())
        readme, plan = z.read("README.md").decode(), z.read("study-plan.md").decode()
    assert readme.startswith("# Pre-registration: ") and "Language: en" in readme and "Präregistrierung" not in readme
    assert "## Research question" in plan and "Pendeln Teams länger?" in plan     # what researchers wrote stays


def test_undo_steps_show_in_the_viewers_language(app, client, survey):
    """E9: a step stores its German name as before and the English one beside it; each person reads their own."""
    from app import undo
    from app.extensions import db
    from app.models import EditStep

    client.post("/surveys/1/editor/blocks/2/items", data={"type": "info"})            # made in German
    with app.app_context():
        step = db.session.scalars(db.select(EditStep).order_by(EditStep.id.desc())).first()
        assert "hinzugefügt" in step.label and "added" in step.data["label_en"]
        old = EditStep(survey_id=step.survey_id, user_id=step.user_id, label="Alter Schritt", data={"parts": {}})
        db.session.add(old)                                                              # from before: German only
        db.session.commit()
    client.post("/account/language", data={"lang": "en"})
    page = client.get("/surveys/1/editor/").data.decode()
    assert "Undo: Alter Schritt" in page and "” added (" in page
    with app.test_request_context(headers={"Accept-Language": "en"}):
        joined = undo.both("Ablauf: {name}", "Flow: {name}", name=undo.both("Ende", "End")) + ", x"
        assert (joined, joined.de) == ("Flow: End, x", "Ablauf: Ende, x")


def test_access_log_actions_in_the_readers_language(app, client, survey):
    """E9: the log keeps the German action as its key and shows it in the reader's language."""
    from app import audit
    from app.extensions import db
    from app.models import AccessLog

    assert set(audit.ACTIONS_EN) == set(audit.ACTIONS)
    assert client.get("/surveys/1/export/data.csv?source=real").status_code == 200
    with app.app_context():
        assert db.session.scalars(db.select(AccessLog.action)).all() == ["Datenexport CSV"]   # stored as before
    assert "Datenexport CSV" in client.get("/admin/log").data.decode()
    client.post("/account/language", data={"lang": "en"})
    page = client.get("/admin/log").data.decode()
    assert "Data export CSV" in page and "Access log" in page and "No. " in page and "kind=data" in page


def test_export_page_and_survey_file_zip_follow_the_downloader(app, client, survey):
    """E8: the export page, and the README.md in the ZIP of "Als Datei exportieren", in the downloader's language;
    file names and the survey file itself stay (app/survey_files.py)."""
    import io
    import json
    import zipfile

    page = client.get("/surveys/1/data").get_data(as_text=True)
    assert "Datenpool und Format" in page and "Als ZIP-Datei herunterladen" in page
    with zipfile.ZipFile(io.BytesIO(client.get("/surveys/1/export/survey.zip").data)) as z:
        german_readme, german_file = z.read("README.md").decode(), z.read("survey.json")
    assert german_readme.startswith("Befragungen aus Askurious")
    client.post("/account/language", data={"lang": "en"})
    page = client.get("/surveys/1/data").get_data(as_text=True)
    assert "Data pool and format" in page and "Download as ZIP file" in page and "Datenpool und Format" not in page
    with zipfile.ZipFile(io.BytesIO(client.get("/surveys/1/export/survey.zip").data)) as z:
        assert {"README.md", "survey.json"} <= set(z.namelist())
        readme, english_file = z.read("README.md").decode(), z.read("survey.json")
    assert readme.startswith("Surveys from Askurious") and "Fragebogen" not in readme
    strip = lambda b: {k: v for k, v in json.loads(b).items() if k != "exported_at"}  # noqa: E731
    assert strip(german_file) == strip(english_file)                                  # the format itself is the same
    assert "Export several surveys" in client.get("/dashboard/export").get_data(as_text=True)


def test_import_messages_follow_the_reader(app):
    """The reasons a file is refused and the notes of a read-in survey in the page language (app/transfer.py)."""
    import pytest
    from flask import g

    from app import transfer

    for lang, text in (("de", "Die Datei ist keine lesbare JSON-Datei."), ("en", "The file is not a readable JSON file.")):
        with app.test_request_context():
            g.ui_lang = lang
            with pytest.raises(transfer.Refused, match=text):
                transfer.loads(b"no json")


def test_coscine_texts_and_metadata_follow_the_uploader(app, client, survey, monkeypatch):
    """E8: the page "In Coscine ablegen" and the metadata Askurious writes (title suffix of checksums.sha256) in the
    uploader's language; file and folder names stay (app/coscine.py)."""
    from app import coscine
    from tests.exports import test_coscine as tc

    fake = tc.FakeCoscine()                                  # as the fixture of tests/exports/test_coscine.py
    app.config["COSCINE_API"] = tc.BASE
    monkeypatch.setattr(coscine.requests, "request", fake)
    tc.on(client)
    tc.connect(client, fake)
    fake.profile = tc.BASIS
    client.post("/account/language", data={"lang": "en"})
    page = client.get("/coscine/survey/1?ziel=p1:r3").get_data(as_text=True)
    assert "WORM resource" in page and "What goes along" in page and "Data package" in page
    r = client.post("/coscine/survey/1", data=tc._basis_form(paket="1"), follow_redirects=True)
    assert "files stored in Coscine" in r.get_data(as_text=True)
    assert any(n.endswith("/checksums.sha256") for n in fake.files)
    meta = next(t for n, t in fake.meta.items() if n.endswith("/checksums.sha256"))
    assert "Checksums" in meta and "Prüfsummen" not in meta


def test_codebook_texts_follow_the_downloader(app, client, survey):
    """E8: the codebook's labels of Askurious's own columns, its types, notes and describing values in the language of
    the person who downloads it; header row, variable names and value labels stay (app/export.py codebook)."""
    import csv
    import io

    def rows():
        body = client.get("/surveys/1/export/codebook.csv").get_data(as_text=True).lstrip("﻿")
        return list(csv.reader(io.StringIO(body)))

    german = rows()
    client.post("/account/language", data={"lang": "en"})
    english = rows()
    assert german[0] == english[0] == ["variable", "label", "type", "values", "notes"]
    assert [r[0] for r in german] == [r[0] for r in english]                    # variable names unchanged
    de, en = {r[0]: r for r in german}, {r[0]: r for r in english}
    assert de["resp_id"][1] == "Laufende Nummer der Teilnahme" and en["resp_id"][1] == "Running number of the participation"
    assert "Schwelle" in de["speeder"][4] or "Regel" in de["speeder"][4]
    assert "threshold" in en["speeder"][4] or "rule" in en["speeder"][4]
    assert en["speeder"][3] == de["speeder"][3] == "0 | 1"
    consent = next(r for r in german if r[2] == "consent")[0]
    # Askurious's own value words follow the reader (audit 2026-10-13, CF-1); answer labels of a question would stay
    assert de[consent][3] == "1 = zugestimmt; 0 = abgelehnt" and en[consent][3] == "1 = agreed; 0 = declined"
    assert de[consent][4].endswith("Pflicht") and en[consent][4].endswith("required")
    from app import export
    from app.extensions import db
    from app.models import Survey

    with app.app_context():  # outside a request German, or the language asked for
        s = db.session.get(Survey, 1)
        assert export.codebook(s)[0][1] == "Laufende Nummer der Teilnahme"
        assert export.codebook(s, lang="en")[0][1] == "Running number of the participation"
        assert export.type_text("Zahl", "en") == "number" and export.type_text("Zahl") == "Zahl"


def test_flow_words_and_codebook_notes_follow_the_reader(app, client, survey, monkeypatch):
    """E8/E9: the flow's words in a codebook (badges, conditions) and the notes of scale scores and vignettes in the
    reader's language, German unchanged by default; translation steps keep both names; erasure ways per page."""
    from app import conditions, erasure, flow, scales, undo
    from app.extensions import db
    from app.models import EditStep, Survey
    from tests.conftest import add

    add(client, "single_choice")                                                # 2
    r = client.post("/surveys/1/translation/fieldwork?lang=en", data={"unit": "item:2", "path": "question", "text": "Hello"})
    assert r.status_code == 200, r.data
    with app.app_context():
        step = db.session.scalars(db.select(EditStep).order_by(EditStep.id.desc())).first()
        assert step.label.startswith("Übersetzung (English) geändert: ")
        assert step.data["label_en"].startswith("Translation (English) changed: ")
        s = db.session.get(Survey, 1)
        s.flow = {"t": "root", "children": [
            {"t": "rand", "id": "r1", "name": "R", "code": "r1", "k": 1, "children": [{"t": "block", "block": 1}]},
            {"t": "branch", "id": "b1", "cond": {"item": "att", "op": "eq", "value": "0"}, "children": [{"t": "block", "block": 2}]},
            {"t": "else", "id": "x1", "of": "b1", "children": []}]}
        db.session.commit()
        assert flow.badges(s)[1] == ["🎲 R: 1 von 1"] and flow.badges(s, "en")[1] == ["🎲 R: 1 of 1"]
        assert flow.badges(s)[2] == ["↳ wenn Aufmerksamkeitsprüfungen: alle bestanden"]
        assert flow.badges(s, "en")[2] == ["↳ if Attention checks: all passed"]
        cond = {"item": "att", "op": "ne", "value": "1"}
        assert conditions.text(s, cond) == "Aufmerksamkeitsprüfungen: nicht eine nicht bestanden"
        assert conditions.text(s, cond, ui="en") == "Attention checks: not one failed"
    monkeypatch.setattr(scales, "partial", lambda item, codes: False)
    assert scales.codebook_note(None, ["a", "b"]).startswith("Mittelwert der Aussagen a, b; umgepolte")
    assert scales.codebook_note(None, ["a", "b"], "en") == ("mean of the statements a, b; reversed statements are recoded; "
                                                            "empty unless all statements are answered (as on the results page)")
    with app.test_request_context(headers={"Accept-Language": "en"}):
        assert erasure.ways()["self"] == "by participants themselves"
        assert undo.both("„{c}“", "“{c}”", c="Q").de == "„Q“"
    assert erasure.WAYS["self"] == "von Teilnehmenden selbst"


def test_learning_aids_follow_the_reader(app, client, survey):
    """E13: reasons, glossary and start checklist in English; the sources stay as they are, German unchanged."""
    from app import learning
    from app.extensions import db
    from app.models import Item
    from tests.conftest import add

    assert set(learning.REASONS_EN) == set(learning.REASONS) and set(learning.TERMS_EN) == set(learning.TERMS)
    assert len(learning.SECTIONS_EN) == len(learning.SECTIONS)
    add(client, "likert")                                                       # 2
    with app.app_context():
        item = db.session.get(Item, 2)
        item.config = {**item.config, "statements": [{"id": f"s{i}"} for i in range(10)]}
        db.session.commit()
    client.post("/account/language", data={"lang": "en"})
    page = client.get("/surveys/1/editor/items/2").data.decode()
    assert "Why? Optional, request or required?" in page and "Warum?" not in page
    assert "Revilla, M. A., Saris, W. E., &amp; Krosnick, J. A. (2014)" in page
    glossary = client.get("/learn/glossary").data.decode()
    assert "Glossary of statistics" in glossary and "Confidence interval" in glossary and "Example:" in glossary
    assert "Choice models" in glossary and "Leiner, D. J. (2019)" in glossary and "Konfidenzintervall" not in glossary
    field = client.get("/surveys/1/field").data.decode()
    assert "Pretest: 0 of 5 test runs, 0 open comments" in field
    assert "Attention check in the Likert questions" in field and "Why? Add an attention check?" in field


def test_constructs_pages_in_english(app, client, survey):
    """E14: names, descriptions, notes and tags of constructs in the reader's language; the item texts of the scales
    stay as published; a linked element's page and the editor's list follow too. German unchanged by default."""
    from tests.conftest import add

    client.post("/account/language", data={"lang": "en"})
    page = client.get("/constructs/").data.decode()
    assert "<h1>Constructs" in page and "General Self-Efficacy Short Scale (ASKU)" in page and ">validated scale <" in page
    assert "Allgemeine Selbstwirksamkeit Kurzskala" not in page and "More filters:" in page
    page = client.get("/constructs/asku?lang=de").data.decode()
    assert "<dt>Scoring</dt><dd>Scale score = mean of the 3 items (1 to 5).</dd>" in page
    assert "free of charge for non-commercial research" in page and "Personality" in page
    assert "Testaussage 1 von ASKU" in page  # the item of the wording file (stand-in in tests/data/instruments.json)
    assert client.get("/constructs/api/search?q=asku").get_json()[0]["name"] == "General Self-Efficacy Short Scale (ASKU)"
    assert "General Self-Efficacy Short Scale (ASKU)" in client.get("/constructs/api/search?q=efficacy").data.decode()
    assert "Big Five Inventory (BFI-10)" in client.get("/surveys/1/editor/").data.decode()
    add(client, "single_choice")
    client.post("/surveys/1/editor/items/2/link", data={"name": "Vorwissen"})
    from app.extensions import db
    from app.models import Item

    with app.app_context():
        key = db.session.get(Item, 2).config["construct"]["key"]
    page = client.get(f"/constructs/linked/{key}").data.decode()
    assert "Shared analysis" in page and "Revision 1" in page and "draft" in page and "Gemeinsame Auswertung" not in page
    client.post("/account/language", data={"lang": "de"})
    assert "Allgemeine Selbstwirksamkeit Kurzskala (ASKU)" in client.get("/constructs/").data.decode()
    assert "Gemeinsame Auswertung" in client.get(f"/constructs/linked/{key}").data.decode()


def test_scripts_speak_the_page_language():
    """de_en() of app.js reads <html lang>: the formatting bar of richtext.js says its buttons in the language of
    the page, German as before (run under Node; skipped without it)."""
    import shutil
    import subprocess

    import pytest

    node = shutil.which("node")
    if node is None:
        pytest.skip("needs Node")
    static = os.path.join(ROOT, "app", "static")
    with open(os.path.join(static, "app.js"), encoding="utf-8") as f:
        de_en_js = re.search(r"^function de_en\(.*$", f.read(), re.M).group(0)
    with open(os.path.join(static, "richtext.js"), encoding="utf-8") as f:
        buttons = re.search(r"var BUTTONS = \[.*?\n  \];", f.read(), re.S).group(0)
    script = (de_en_js + "\nfor (const lang of ['de', 'en']) {\n  globalThis.document = {documentElement: {lang: lang}};\n  "
              + buttons + "\n  console.log(JSON.stringify(BUTTONS.map(function (b) { return [b.text, b.title]; })));\n}\n")
    out = subprocess.run([node], input=script.encode("utf-8"), capture_output=True, check=True).stdout.decode("utf-8")
    german, english = (json.loads(line) for line in out.splitlines())
    assert german[0] == ["F", "Fett (Strg+B)"] and ["Lücke", "Lücke einfügen (ein Klick auf eine Lücke ändert ihre Breite)"] in german
    assert english[0] == ["B", "Bold (Ctrl+B)"] and ["Gap", "Insert a gap (a click on a gap changes its width)"] in english


def test_sample_size_groups_account_deletion_and_feedback_in_english(app, client, survey):
    """Package E3: sample size planner, groups, deleting the account, feedback, pricing and the experimental features
    in English; German unchanged."""
    from app import experimental
    from app.models import WorkGroupMember

    assert set(experimental.FEATURES_EN) == set(experimental.KEYS)
    assert set(WorkGroupMember.ROLES_EN) == set(WorkGroupMember.ROLES)
    client.post("/groups/new", data={"name": "Team A"})
    assert "Gruppe löschen" in client.get("/groups/1").data.decode()
    client.post("/account/language", data={"lang": "en"})
    page = client.get("/sample-size/?survey=1").data.decode()
    assert "What do you want to plan?" in page and "References for this page" in page and "Was wollen Sie" not in page
    assert "Cohen, J. (1988)" in page and "Equivalence test (TOST)" in page
    page = client.get("/groups/1").data.decode()
    assert "Delete group" in page and ">Lead</option>" in page and "Mitglieder" not in page
    page = client.get("/account/delete").data.decode()
    assert "Delete account permanently" in page and "1 group without further members" in page
    assert "Please give the group a name." in client.post("/groups/new", data={"name": " "}, follow_redirects=True).data.decode()
    page = client.post("/feedback/scale", data={"title": ""}, follow_redirects=True).data.decode()
    assert "Please fill in: Title, Description, Justification, Source, DOI." in page and "Suggest a construct" in page
    app.config["DONATION_URL"] = "https://paypal.example/spende"
    with app.test_request_context():
        from flask import g

        g.ui_lang = "en"
        assert experimental.label("coding") == "Categorise open answers"
        assert experimental.features()[0][0] == "ainstein" and experimental.uses("coding")[1] == "surveys"



def test_live_poll_presenter_pages_in_english_audience_in_the_polls_language(app, client, monkeypatch):
    """F: the presenter's pages (wizard, questions, settings, presenting, AI assistant, the controls of the presenter
    view) follow the page language; what the audience sees speaks the live poll's language (here German), and the
    AI's reasons come in the presenter's language while its questions stay in the poll's."""
    import json

    from app import ai
    from app.extensions import db
    from app.models import Survey

    client.post("/account/language", data={"lang": "en"})
    assert "Title and language" in client.get("/presenter/new").data.decode()
    client.post("/presenter/new", data={"title": "Ringvorlesung", "language": "de"})
    with app.app_context():
        survey = db.session.scalar(db.select(Survey).filter_by(title="Ringvorlesung"))
        code, sid, block = survey.code, survey.id, survey.blocks[0].id
    page = client.get(f"/presenter/{code}/setup/question").data.decode()
    assert "Use a template" in page and "Feedback on the session" in page and "Eine Vorlage nehmen" not in page
    seen = {}

    def fake(instructions, prompt, purpose="text", item_type=None):
        seen["instructions"] = instructions
        return json.dumps({"questions": [{"type": "text_input", "question": "Ein Wort zur Vorlesung", "why": "mood"}]})

    monkeypatch.setattr(ai, "generate_text", fake)
    page = client.post(f"/presenter/{code}/ai", data={"description": "Vorlesung"}).data.decode()
    assert "The reasons in English." in seen["instructions"] and "in Deutsch (Sie)" in seen["instructions"]
    assert "Accept selected" in page and "Word cloud" in page and "Ein Wort zur Vorlesung" in page
    client.post(f"/surveys/{sid}/editor/blocks/{block}/items", data={"type": "single_choice"})
    page = client.get(f"/presenter/{code}").data.decode()
    assert "Suggest slides with AI" in page and "Single choice" in page and "Einfachauswahl" not in page
    page = client.get(f"/presenter/{code}/settings").data.decode()
    assert "Allow reactions" in page and "Thumbs up" in page and "Reaktionen erlauben" not in page
    page = client.get(f"/presenter/{code}/present").data.decode()
    assert "Presenter view" in page and "New run" in page and "Neuer Durchgang" not in page
    client.post(f"/presenter/{code}/control", data={"action": "start"})
    page = client.get(f"/presenter/{code}/screen").data.decode()
    assert 'data-ui-lang="en"' in page and "Close answers" in page and '<html lang="de"' in page
    assert "Abstimmung beendet" in page and "Voting closed" not in page                    # data-t: the poll's language
    assert "Gleich geht es los" in app.test_client().get(f"/live/{code}").data.decode()   # the audience: German


def test_admin_system_pages_in_english(app, client):
    """Administration > KI-Verbrauch, IP-Region, Vorlagen, Datenbank and the messages of app/admin.py in English;
    the German default stays as it was."""
    from app import ai_stats, table_notes

    assert set(ai_stats.PURPOSES_EN) == set(ai_stats.PURPOSES) and set(table_notes.NOTES_EN) == set(table_notes.NOTES)
    assert "Kosten (geschätzt)" in client.get("/admin/ai").data.decode()
    client.post("/account/language", data={"lang": "en"})
    page = client.get("/admin/ai").data.decode()
    assert "Cost (estimated)" in page and "By purpose" in page and "Kosten" not in page
    page = client.post("/admin/ai", data={"text_model": ["m"], "text_in": ["1"], "text_out": ["2"]},
                       follow_redirects=True).data.decode()
    assert "Prices saved." in page
    assert "country and federal state from the IP address" in client.get("/admin/database").data.decode()
    page = client.get("/admin/orgs/1").data.decode()
    assert "Institution design" in page and "No template." in page  # W94; the template of the organisation design
    page = client.get("/admin/database").data.decode()
    assert "Database of the app" in page and "Entries per table" in page and "Members of the work groups" in page
    from tests.conftest import PG_URL

    if not PG_URL:  # the form for moving to PostgreSQL is only there while the app runs on SQLite
        page = client.post("/admin/database", data={"action": "check", "host": "", "name": "x", "user": "u"}).data.decode()
        assert "Please enter server, database and user." in page
    page = client.post("/admin/admit", data={"action": "everyone", "value": "0"}, follow_redirects=True).data.decode()
    assert "New accounts now only come with an invitation" in page


def test_admin_people_and_construct_pages_in_english(app, client):
    """Administration > Forschende, Konto, Gruppen, Experimentell and Konstrukte in English; German stays the default."""
    from app.extensions import db
    from app.models import User

    assert ">Konten<" in client.get("/admin/people").data.decode()  # W97: the figure counts accounts
    client.post("/account/language", data={"lang": "en"})
    page = client.get("/admin/people").data.decode()
    assert ">accounts<" in page and "Allow sign-in" in page and "Create account" in page and "Add person …" in page and ">People</a>" in page
    assert ">Konten<" not in page
    with app.app_context():
        other = User(orcid="0000-0001-9999-0002", name="Other Person")
        db.session.add(other)
        db.session.commit()
        uid = other.id
    page = client.get(f"/admin/accounts/{uid}").data.decode()
    assert "Survey types" in page and "Create everywhere" in page and "Experimental features" in page and "Block and delete" in page
    assert "Groups" in client.get("/admin/people").data.decode()
    page = client.get("/admin/experimental").data.decode()
    assert "Choose a person, group or organisation" in page and "Categorise open answers" in page
    page = client.get("/admin/constructs").data.decode()
    assert "Pooled construct data" in page and "Submitted for “Constructs”" in page
    page = client.get("/admin/constructs/compilation").data.decode()
    assert "Compilation for a meta-analysis" in page and "Zusammenstellung für" not in page


def test_ainstein_frame_in_english_survey_content_in_the_survey_language(app, client, monkeypatch):
    """AInstein: the pages and the mentor's fallback language follow the page language (E12); the survey made from
    the project keeps its own language."""
    from app import ai
    from app.extensions import db
    from app.models import Survey
    from tests.conftest import experiments_on

    calls = []

    def fake(instructions, prompt, purpose="text", item_type=None, schema=None):
        calls.append(instructions)
        if purpose == "ainstein_chat":
            return json.dumps({"reply": "Welcome", "canvas": {"frage": "Q"}, "focus": "frage", "name": "Trust"})
        return json.dumps({"goal": "", "design": "cross", "sample": {}, "intervention": "",
                           "hypotheses": [{"text": "A relates to B", "kind": "correlation", "constructs": ["A", "B"]}],
                           "constructs": [{"name": "A", "role": "iv", "measure": "", "n": 4},
                                          {"name": "B", "role": "dv", "measure": "", "n": 4}]})

    monkeypatch.setattr(ai, "generate_text", fake)
    experiments_on(client, "ainstein")
    assert "Mit AInstein beginnen" in client.get("/ainstein/").data.decode()
    client.post("/account/language", data={"lang": "en"})
    assert "Start with AInstein" in client.get("/ainstein/").data.decode()
    where = client.post("/ainstein/new", data={"idea": "one two three four five six seven eight nine ten eleven"}).headers["Location"]
    assert 'English if unsure' in calls[0]
    page = client.get(where).data.decode()
    assert "Conversation and canvas" in page and "Research question &amp; relevance" in page and "Send" in page
    client.post(where + "/plan/derive")
    page = client.get(where + "/design").data.decode()
    assert "Research design and hypotheses" in page and "Relationship" in page and "Independent variable" in page
    page = client.get(where + "/completion").data.decode()
    assert "Create survey" in page and "German and English" in page and "“Work on it”" in page
    page = client.post(where + "/survey", data={"languages": "de"}, follow_redirects=True).data.decode()
    assert "has been created as a draft" in page
    with app.app_context():
        survey = db.session.scalar(db.select(Survey).order_by(Survey.id.desc()))
        assert survey.languages == ["de"]


def test_interview_guide_frame_in_english_guide_content_in_the_survey_language(app, client, monkeypatch):
    """F1: the pages of an interview guide (new, guide, settings, interviews, interview view, export, results) and its
    Word documents follow the page language; the guide's questions stay in the survey's language (here German); the
    AI's reasons come in the page language, its questions in the guide's; undo steps carry both names."""
    import io
    import re
    import zipfile

    from app import ai
    from app.extensions import db
    from app.models import Survey

    client.post("/account/language", data={"lang": "en"})
    page = client.get("/guide/new?start=vorlage").data.decode()
    assert "New interview guide" in page and "Problem-centred interview" in page and "Neuer Interviewleitfaden" not in page
    client.post("/guide/new", data={"title": "Pflege", "language": "de", "template": "witzel", "order": "free",
                                        "target": "erwachsene", "mode": "praesenz"})
    with app.app_context():
        survey = db.session.scalar(db.select(Survey).filter_by(title="Pflege"))
        code, sid = survey.code, survey.id
    page = client.get(f"/guide/{code}").data.decode()
    assert "Topic block" in page and "Check guide with AI" in page and "Einstieg" in page and "Themenblock nach oben" not in page
    page = client.get(f"/guide/{code}/settings").data.decode()
    assert "Target group" in page and "Face-to-face interview" in page and "Zielgruppe" not in page
    page = client.get(f"/guide/{code}/interviews").data.decode()
    assert "Start interview" in page and "in writing (signed)" in page and "Interview beginnen" not in page
    r = client.post(f"/guide/{code}/interviews/new", data={"label": "", "date": "2026-10-12", "mode": "praesenz",
                                                               "consent_form": "schriftlich", "consent": "1", "audio": "1"})
    rid = int(re.search(r"interview=(\d+)", r.headers["Location"]).group(1))
    page = client.get(f"/guide/{code}/conversation?interview={rid}").data.decode()
    assert 'data-ui-lang="en"' in page and '<html lang="de"' in page and "Start recording" in page
    assert "End interview" in page and "Zum Einstieg: Erzählen Sie doch einmal" in page and "Aufnahme starten" not in page
    assert "12 Oct 2026" in client.get(f"/guide/{code}/interviews").data.decode()
    page = client.get(f"/guide/{code}/export").data.decode()
    assert "Interview card" in page and "Record revision" in page and "Fassung festhalten" not in page
    page = client.get(f"/guide/{code}/interviews/analysis").data.decode()
    assert "Case overview" in page and "Coverage" in page and "Abdeckung" not in page
    with zipfile.ZipFile(io.BytesIO(client.get(f"/guide/{code}/export/leitfaden.docx").data)) as z:
        text = re.sub(r"<[^>]+>", "", z.read("word/document.xml").decode())
    assert "Maintenance questions" in text and "Zum Einstieg" in text and "Aufrechterhaltungsfragen" not in text
    with zipfile.ZipFile(io.BytesIO(client.get(f"/guide/{code}/export/einwilligung.docx").data)) as z:
        text = re.sub(r"<[^>]+>", "", z.read("word/document.xml").decode())
    assert "Freiwilligkeit" in text                                                 # what the interviewee signs: German
    seen = {}

    def fake(instructions, prompt, purpose="text", item_type=None):
        seen["instructions"] = instructions
        return json.dumps({"remarks": ["Gut"], "parts": []})

    monkeypatch.setattr(ai, "generate_text", fake)
    page = client.post(f"/guide/{code}/ai/check").data.decode()
    assert "remarks and reasons in English" in seen["instructions"] and "On the guide as a whole" in page
    client.post(f"/guide/{code}/control", data={"action": "add_block"})
    with app.app_context():
        from app.models import EditStep

        step = db.session.scalar(db.select(EditStep).filter_by(survey_id=sid).order_by(EditStep.id.desc()))
        assert step.data["label_en"].startswith("Topic block") and step.label.startswith("Themenblock")
    client.post("/account/language", data={"lang": "de"})
    assert "Neuer Interviewleitfaden" in client.get("/guide/new").data.decode()



def test_legal_texts_read_in_english_but_the_german_version_binds(app, client):

    """E15: terms of use and privacy page have an English reading version with a note that only the German one is

    binding; accepting the terms in English records the version of the German licence.md, and the German pages stay

    as they were."""

    from app import terms

    from app.extensions import db

    from app.models import User



    note = "This translation is provided for convenience. Only the German version is legally binding."

    german = client.get("/terms/").data.decode()

    assert "Nutzungsbedingungen für Forschende" in german and note not in german

    assert "Impressum und Datenschutz" in client.get("/imprint").data.decode()

    client.post("/account/language", data={"lang": "en"})

    with app.app_context():

        user = db.session.scalar(db.select(User).filter_by(orcid="0000-0000-0000-001X"))

        user.terms_version, current = None, terms.version()

        db.session.commit()

    page = client.get("/terms/").data.decode()

    assert page.index(note) < page.index("<h1>Terms of use for researchers</h1>")

    assert "Agree and continue" in page and f"version {current}, legally binding in their German version" in page

    assert "Zustimmen" not in page and "Nutzungsbedingungen für Forschende" not in page

    # the English text translates the German version in force (else the page shows the German text under a note)

    assert f"English translation of the German version {current}" in page

    client.post("/terms/", data={"accept": "1"})

    with app.app_context():

        assert db.session.scalar(db.select(User).filter_by(orcid="0000-0000-0000-001X")).terms_version == current

    page = client.get("/imprint").data.decode()

    assert note in page and "<h1>Legal notice and privacy</h1>" in page and "the operator of this installation" in page

    assert "<h2>Forschende</h2>" not in page and "<h2>Researchers</h2>" in page


def update():
    """Lower the counts in the rest list to what is there now; never raise one or add a file."""
    noted, now = rest(), found_now()
    lowered = {rel: min(n, len(now.get(rel, []))) for rel, n in noted.items()}
    lowered = {rel: n for rel, n in sorted(lowered.items()) if n}
    with open(REST, "w", encoding="utf-8", newline="\n") as f:
        json.dump(lowered, f, ensure_ascii=False, indent=1)
        f.write("\n")
    print(f"{sum(noted.values()) - sum(lowered.values())} weniger, noch {sum(lowered.values())} in {len(lowered)} Dateien")


if __name__ == "__main__":
    if "--update" in sys.argv:
        sys.path.insert(0, ROOT)
        update()
