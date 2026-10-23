"""Exports: the tests of the earlier audit and review files about this area (sorted by area,
architecture concept 2026-10-12, G). Each block names the file it came from; its helpers are in
tests/helpers/audit_<file>.py."""
from tests.helpers.audit_audit_2026_10_06_a import OPEN, OTHER as OTHER__audit_2026_10_06_a, RAW, RAW_ELSEWHERE, RAW_MORE, _member, _person  # noqa: F401
from tests.helpers.audit_audit_2026_10_06_c import _fresh_count, _docx, _file, _pack, _preview, _upload, io, json, pytest  # noqa: F401
from tests.helpers.audit_audit_2026_10_06_d import _no_failures_from_other_tests, OTHER as OTHER__audit_2026_10_06_d, _as  # noqa: F401
from tests.helpers.audit_audit_2026_10_06_e import Survey, _start, db, os  # noqa: F401
from tests.helpers.audit_audit_2026_10_06_f import _likert, _scale_rows, export, quality  # noqa: F401
from tests.helpers.audit_audit_2026_10_12_a import CONSENT_OK, E, Item, User, add, fill_consent, fill_questions  # noqa: F401
from tests.helpers.audit_audit_2026_10_12_c import Response, _states, _with_prompts  # noqa: F401
from tests.helpers.audit_audit_2026_10_12_h import MultiDict, _scale, sem_model  # noqa: F401
from tests.helpers.audit_audit_small import csv  # noqa: F401
from tests.helpers.audit_audit_small2 import _errors, kind_of, postal_code  # noqa: F401
from tests.helpers.audit_cr8_a import app, client, survey  # noqa: F401
from tests.helpers.audit_cr8_b import needs_sqlite, sqlite3  # noqa: F401
from tests.helpers.audit_cr8_e import _every_type, nullcontext  # noqa: F401
from tests.helpers.audit_review_fixes import flow, runtime, test_flow  # noqa: F401


# --- from tests/audits/test_audit13_rest.py: Audit Daten und Berechnungen 2026-10-12, package I (exports, choice models, elements, rest): E-3, E-4, E-5, E-6, ---


def test_e4_codebook_negative_codes_without_apostrophe():
    from app.export import codebook_text, safe_text

    for text in ("-3 = stimme überhaupt nicht zu; 3 = stimme voll zu", "-90 bis 90", "-97 = nicht bearbeitbar"):
        assert codebook_text(text) == text
    assert codebook_text("-2+3+cmd|' /C calc'!A0") == "'-2+3+cmd|' /C calc'!A0"
    assert codebook_text("=HYPERLINK(1)") == safe_text("=HYPERLINK(1)") == "'=HYPERLINK(1)"
    assert codebook_text("-3 = x (HYPERLINK)").startswith("'")


# --- from tests/audits/test_audit_2026_10_06_a.py: Security audit 2026-10-06, package A "Rollen und Objekte": S1, S2, S9, S10, S12, S22, H5 and the side finding on ---


def test_s2_every_export_route_is_classified_and_raw_data_is_refused_to_viewers(app, client, survey):
    from flask import url_for

    found = {r.endpoint for r in app.url_map.iter_rules() if "/export" in r.rule}
    assert found - set(RAW) - RAW_ELSEWHERE - OPEN == set(), "a new export route: put it into RAW (answers) or OPEN"
    _person(app, OTHER__audit_2026_10_06_a, "Viewer")
    _member(app, 1, OTHER__audit_2026_10_06_a, "viewer")
    viewer = _person(app, OTHER__audit_2026_10_06_a, "Viewer")
    with app.test_request_context():
        urls = [url_for(ep, survey_id=1, **args) for ep, args in [*RAW.items(), *RAW_MORE]
                if ep != "surveys.snapshot_file" or args]
        open_csv = url_for("surveys.export_csv", survey_id=1, kind="codebook")
    for url in urls:
        r = viewer.get(url)
        assert r.status_code == 302 and "attachment" not in r.headers.get("Content-Disposition", ""), url
    assert viewer.get(open_csv).status_code == 200


# --- from tests/audits/test_audit_2026_10_06_c.py: Audit of 2026-10-06, package C (uploads and resources): voice messages written only when the page goes on, with a ---


def test_the_parts_of_a_zip_share_one_structure_budget(app, client, survey, monkeypatch):
    from app import survey_files, transfer

    data = _file(client)
    body = json.dumps(data).encode()
    marks = survey_files._structure(body)
    monkeypatch.setattr(transfer, "MAX_STRUCTURE", marks * 2 + 10)                # two such files, not three
    two = _pack({f"s{i}/survey.json": body for i in range(2)})
    three = _pack({f"s{i}/survey.json": body for i in range(3)})
    page, token = _preview(client, raw=two, name="zwei.zip")
    assert token and "Die Datei enthält 2 Befragungen" in page
    page, token = _preview(client, raw=three, name="drei.zip")
    assert token is None and "haben zusammen mehr Einträge" in page


def test_the_parts_of_a_zip_share_one_byte_budget(app, client, survey):
    data = _file(client)
    body = json.dumps(data).encode()
    app.config["IMPORT_MAX_BYTES"] = len(body) * 2 + 10
    try:
        page, token = _preview(client, raw=_pack({f"s{i}/survey.json": body for i in range(3)}), name="drei.zip")
        assert token is None and "zusammen zu groß" in page
    finally:
        app.config["IMPORT_MAX_BYTES"] = 40 * 1024 * 1024


def test_the_import_hands_one_budget_to_all_surveys(app, client, survey, monkeypatch):
    from app import transfer

    body = json.dumps(_file(client)).encode()
    seen, real = [], transfer.create

    def create(data, user, title="", fetch=None, budget=None):
        seen.append(budget)
        return real(data, user, title, fetch=fetch, budget=budget)
    monkeypatch.setattr(transfer, "create", create)
    page = _upload(client, None, raw=_pack({f"s{i}/survey.json": body for i in range(2)}), name="zwei.zip",
                   **{"take-0": "1", "take-1": "1", "title-0": "Neu A", "title-1": "Neu B"})
    assert "2 Befragungen sind als Entwürfe angelegt" in page
    assert len(seen) == 2 and seen[0] is seen[1] and isinstance(seen[0], transfer.Budget)


def test_a_docx_bomb_is_refused_before_it_is_parsed(app, monkeypatch):
    from app import interview_transcribe

    import docx

    bombs, opened = (_docx(60), _docx(3)), []                                    # too much; packed too tightly
    monkeypatch.setattr(docx, "Document", lambda *a: opened.append(1))
    with app.app_context():
        for data in bombs:
            with pytest.raises(ValueError, match="entpackt zu groß"):
                interview_transcribe.read_upload("t.docx", data)
        with pytest.raises(ValueError, match="ließ sich nicht lesen"):
            interview_transcribe.read_upload("t.docx", b"PK\x03\x04 kein zip")
    assert opened == []


def test_an_ordinary_docx_is_read(app):
    from docx import Document

    from app import interview_transcribe

    doc = Document()
    doc.add_paragraph("I: Wie geht es Ihnen?")
    doc.add_paragraph("B: Gut, danke.")
    buf = io.BytesIO()
    doc.save(buf)
    with app.app_context():
        text = interview_transcribe.read_upload("t.docx", buf.getvalue())
    assert "Wie geht es Ihnen?" in text and "Gut, danke." in text


# --- from tests/audits/test_audit_2026_10_06_d.py: Security audit 2026-10-06, package D "Konten": the right to create surveys of invited people (S11), blocking ends ---


def test_the_people_search_takes_no_wildcards_and_three_letters(app, client, survey, monkeypatch):
    from app import orcid_api

    monkeypatch.setattr(orcid_api, "search", lambda q: [])
    _as(app.test_client(), OTHER__audit_2026_10_06_d, "Kollegin Muster")
    assert "Kollegin Muster" not in client.get("/surveys/1?q=%25").data.decode()
    assert "Kollegin Muster" not in client.get("/surveys/1?q=___").data.decode()
    assert "Kollegin Muster" not in client.get("/surveys/1?q=Ko").data.decode()
    assert "Kollegin Muster" in client.get("/surveys/1?q=oll").data.decode()


# --- from tests/audits/test_audit_2026_10_06_e.py: Security audit 2026-10-06, package E: researcher texts of the delay discounting stay inside their data block (S3), ---


def test_downloads_of_signed_in_people_are_not_kept(app, client, survey):
    _start(client)
    from app import prereg

    with app.app_context():
        [(name, *_)] = prereg.listing(db.session.get(Survey, 1))
    r = client.get(f"/surveys/1/export/preregistration/{name}")             # send_file says "no-cache"
    assert r.status_code == 200 and r.headers["Cache-Control"] == "no-store"
    static = client.get("/static/app.css?v=x")
    assert static.status_code == 200 and "no-store" not in (static.headers.get("Cache-Control") or "")
    with app.app_context():
        uid = db.session.get(Survey, 1).uid
    folder = os.path.join(app.config["UPLOAD_FOLDER"], uid)
    os.makedirs(folder, exist_ok=True)
    with open(os.path.join(folder, "bild.png"), "wb") as f:
        f.write(b"\x89PNG\r\n\x1a\n" + b"\0" * 20)
    media = client.get(f"/media/{uid}/bild.png")
    # kept by the browser, never by a shared cache, and not a year long (audit 2026-10-18, INJ-3)
    cache = {part.strip() for part in media.headers["Cache-Control"].split(",")}
    assert media.status_code == 200 and "private" in cache and "max-age=86400" in cache
    assert "public" not in cache and "immutable" not in cache


# --- from tests/audits/test_audit_2026_10_06_f.py: Audit 2026-10-06 (security and performance), package F: performance of the evaluation pages. ---


def test_p2_quality_values_once_per_request(app, client, survey, monkeypatch):
    _likert(app, client, _scale_rows(n=12))
    calls = []
    values = quality.values
    monkeypatch.setattr(quality, "values", lambda *a, **k: calls.append(1) or values(*a, **k))
    with app.test_request_context():
        s = db.session.get(Survey, 1)
        first = export.rows(s, False, ["real"], with_hb=False)
        second = export.rows(s, False, ["real"], with_hb=False, with_quality=False)
        summary = quality.summary(s, ["real"], False)
        export.marked_ids(s)
        assert len(calls) == 12                                             # once per participation
        assert [r["qc_flags"] for r in first] == [r["qc_flags"] for r in second]
        assert all("_notable" not in r for r in first)
        # the same values as without the memo
        ctx = quality.context(s, export._items(s), export._responses(s, False, ["real"]))
        alone = [{k: v for k, v in quality.values(s, r, ctx).items()} for r in export._responses(s, False, ["real"])]
        assert [{k: r[k] for k in a} for r, a in zip(first, alone)] == alone
        assert summary["n"] == 12
        # a participation that changed in between is computed anew
        response = export._responses(s, False, ["real"])[0]
        response.position = 7
        before = len(calls)
        quality.request_values(s, response, quality.context(s, export._items(s)))
        assert len(calls) == before + 1
    with app.test_request_context():  # the next request reads them as kept per process (audit 2026-10-09b, P1)
        before = len(calls)
        export.rows(db.session.get(Survey, 1), False, ["real"], with_hb=False)
        assert len(calls) == before


# --- from tests/audits/test_audit_2026_10_12_a.py: Audit Daten und Berechnungen 2026-10-12, package A (crashes and wrong figures): S-1, S-3, W-1, M-1, SEM-1, SEM-2, ---


def test_voice_text_typed_without_the_setting_is_exported(app, client, survey):
    with app.app_context():
        for u in db.session.scalars(db.select(User)):
            u.experiments = ["voice"]
        db.session.commit()
    fill_consent(client)
    add(client, "voice")  # item 2
    client.post(E + "/items/2", data={"code": "v", "question": "Erzählen Sie", "notice": "Wird gespeichert.",
                                      "answer_mode": "optional", "max_seconds": "60"})  # text_fallback off
    fill_questions(app)
    client.post("/surveys/1/start", data={})
    with app.app_context():
        url = f"/s/{db.session.get(Survey, 1).code}"
        assert db.session.get(Item, 2).config["text_fallback"] is False
    app.config["VOICE_MAX_MB"] = 0  # the page offers only the text field
    p = app.test_client()
    p.get(url)
    p.post(url, data=CONSENT_OK)
    assert p.post(url, data={"nav": "next", "i2_text": "Meine getippte Antwort"}).status_code == 302
    csv = client.get("/surveys/1/export/data.csv").data.decode("utf-8-sig")
    assert "v_text" in csv and "Meine getippte Antwort" in csv
    with app.app_context():
        item = db.session.get(Item, 2)
        _, _, notes = item.kind.codebook_info(item, "v_text", "de")
        assert any("Tippen möglich" in n for n in notes)


# --- from tests/audits/test_audit_2026_10_12_c.py: Audit data and calculations 2026-10-12, package C (diary study): E-1, E-2, D-1 to D-7, D-10, D-12 to D-15 and ---


def test_answered_entry_left_out_by_drop_is_not_missed(app, client):
    from app import export

    _with_prompts(app, client)
    with app.test_request_context():
        survey = db.session.get(Survey, 1)
        full = _states(export.to_csv(survey, include_incomplete=True))
        dropped = _states(export.to_csv(survey, drop_speeders=True))
        assert full.count("beantwortet") > 1
        assert dropped.count("verpasst") == full.count("verpasst")
        assert dropped.count("beantwortet") == full.count("beantwortet") - 1


def test_test_or_dummy_export_gets_no_missed_rows_of_real_persons(app, client):
    from app import export

    _with_prompts(app, client)
    with app.test_request_context():
        survey = db.session.get(Survey, 1)
        assert _states(export.to_csv(survey, sources=("dummy",))) == []
        assert _states(export.to_csv(survey, sources=("test",))) == []


def test_a_broken_off_entry_is_abgebrochen_in_every_pool(app, client):
    from app import export

    prompts = _with_prompts(app, client, first_speeder=False)
    with app.test_request_context():
        survey = db.session.get(Survey, 1)
        broken = db.session.scalar(db.select(Response).filter_by(pid="e-0"))
        broken.status, broken.completed_at = Response.ACTIVE, None
        db.session.commit()
        only_done = export.to_csv(survey)
        with_all = export.to_csv(survey, include_incomplete=True)
        assert _states(only_done).count("abgebrochen") == 1 and _states(with_all).count("abgebrochen") == 1
        assert _states(only_done).count("verpasst") == _states(with_all).count("verpasst")
        assert len(prompts) >= 1


# --- from tests/audits/test_audit_2026_10_12_h.py: Audit Daten und Berechnungen 2026-10-12, package H (SEM rest and R script): SEM-3 to SEM-7, SEM-9 to SEM-11, ---


def test_pls_measurement_model_script(app, client, survey):
    for number, code, k in ((2, "tb", 4), (3, "einst", 3), (4, "absicht", 3)):
        _scale(client, number, code, k)
    client.post("/surveys/1/researchdesign", data=MultiDict({"research_question": "x", "sem_method": "pls"}))
    with app.app_context():
        s = db.session.get(Survey, 1)
        assert sem_model.build(s)["mode"] == "measurement" and sem_model.method_of(s) == "pls"
        text = sem_model.r_script(s, "d.rds")
        assert "seminr schätzt kein Modell ohne Pfade" in text and "relationships(" not in text


# --- from tests/audits/test_audit_small.py: Smaller points of the data audit: codebook values, orders in loops, screen-out times, long JSON in .sav. ---


def test_screen_out_has_end_time_and_duration(app, client, survey):
    from tests.editor.test_flow import run, setup_design

    setup_design(app, client)
    client.post("/surveys/1/start", data={"use_end": "1", "field_end": "2031-01-01"})
    with app.app_context():
        uid = db.session.get(Survey, 1).code
    young = run(app, uid, 16)
    young.post(f"/s/{uid}", data={"nav": "next"})
    row = list(csv.DictReader(io.StringIO(client.get("/surveys/1/export/data.csv?all=1").data.decode("utf-8-sig"))))[-1]
    assert row["status"] == "screenout" and row["finished"] and row["duration_s"] != ""


def test_separated_section_export_writes_decimal_commas(app, client, survey):
    from tests.quotas.test_split import build, participate

    uid = build(app, client)
    with app.app_context():
        note = db.session.get(Item, 4)
        note.config = {**note.config, "validation": "decimal", "multiline": False}
        db.session.commit()
    for n in range(5):
        participate(app, uid, "1", f"p{n}@x.de", "2,5")
    text = client.get("/surveys/1/contact/export.csv?sep=semicolon").data.decode("utf-8-sig")
    assert ";2,5" in text and "2.5" not in text
    assert ",2.5" in client.get("/surveys/1/contact/export.csv").data.decode("utf-8-sig")


# --- from tests/audits/test_audit_small2.py: Small fixes from the audit of 2026-10-01 (S4, S6, S7, S8, S10, S11, D-1, D-3, D-4). ---


def test_postal_codes_with_an_area_are_known():
    assert postal_code.known("52062") and postal_code.known("60312") and not postal_code.known("00000")
    assert "60312" not in postal_code.table()                     # the reason for the second look


def test_association_words_are_free_text_for_publishing(app, client, survey):
    add(client, "association")
    with app.app_context():
        item = db.session.get(Item, 2)
        cols = export.open_columns(db.session.get(Survey, 1))
        assert f"{item.code}_1" in cols and f"{item.code}_1_val" not in cols and f"{item.code}_mean" not in cols


def test_a_question_cannot_take_the_name_of_an_export_column(app, client, survey):
    add(client, "text_input"), add(client, "text_input")
    with app.app_context():
        db.session.get(Item, 2).code, db.session.get(Item, 3).code = "status", "Pid"
        db.session.commit()
    errors = _errors(app)
    assert any("„status“ steht doppelt im Export" in e for e in errors) and any("„Pid“ steht doppelt" in e for e in errors)
    with app.app_context():
        db.session.get(Item, 2).code, db.session.get(Item, 3).code = "beruf", "alter"
        db.session.commit()
    assert not any("doppelt im Export" in e for e in _errors(app))


def test_position_map_exports_cant_say_and_the_shown_order(app, client, survey):
    add(client, "position2d")
    kind = kind_of("position2d")
    with app.app_context():
        item = db.session.get(Item, 2)
        item.code = "tech"
        item.config = {**item.config, "allow_skip": True,
                       "options": [{"id": "a", "code": "ki"}, {"id": "b", "code": "wk"}]}
        db.session.commit()
        cols = [c for c, _ in kind.export_columns(item)]
        assert cols == ["tech_ki_x", "tech_ki_y", "tech_ki_skip", "tech_wk_x", "tech_wk_y", "tech_wk_skip"]
        values = kind.export_values(item, {"ki": [10.0, 20.0], "_skip": ["wk"]})
        assert (values["tech_ki_x"], values["tech_ki_skip"], values["tech_wk_x"], values["tech_wk_skip"]) == (10.0, 0, None, 1)
        assert set(kind.export_values(item, None).values()) == {None}                   # never seen: all empty
        assert kind.order_codes(item) == {"a": "tech_ki", "b": "tech_wk"}
        assert "tech_order" in [c for c, _ in export.columns(db.session.get(Survey, 1))]
        assert kind.codebook_info(item, "tech_wk_skip", "de")[0] == "Auswahl"
        item.config = {**item.config, "options": [{"id": "a", "code": "ki"}, {"id": "b", "code": "_skip"}]}
        assert any("Teil der Variablennamen" in p for p in kind.validate(item))


# --- from tests/audits/test_cr8_a.py: Code review 2026-10-07, package A "Falsche oder verlorene Daten": regression tests for T1, E1, E2, E8, E9, E10, ---


def test_e11_file_is_read_without_python_docx(app, client, survey, monkeypatch):
    import sys

    from tests.conftest import add
    from tests.exports.test_transfer import _file, _upload

    add(client, "text_input")
    data = _file(client)
    monkeypatch.setitem(sys.modules, "docx", None)  # "from docx import ..." raises ImportError
    assert "ist als Entwurf angelegt" in _upload(client, data, title="Ohne Word")


# --- from tests/audits/test_cr8_b.py: Code review 2026-10-07, package B: SQLite write locks and housekeeping (K2/D1, K3/D2, K4, E4, K8, T8/D6, K9). ---


@needs_sqlite
def test_freeze_runs_without_write_lock(app, client, survey, monkeypatch):
    from app import prereg

    fill_consent(client)
    add(client, "text_input")
    fill_questions(app)
    seen = {}
    real = prereg._files

    def spy(s, lang="de"):
        other = sqlite3.connect(db.engine.url.database, timeout=0)
        try:
            other.execute("BEGIN IMMEDIATE")
            other.rollback()
            seen["locked"] = False
        except sqlite3.OperationalError as e:
            seen["locked"] = str(e)
        other.close()
        return real(s, lang)

    monkeypatch.setattr(prereg, "_files", spy)
    client.post("/surveys/1/start", data={"use_end": "1", "field_end": "2031-01-01"})
    assert seen == {"locked": False}
    with app.app_context():
        s = db.session.get(Survey, 1)
        assert s.status == Survey.RUNNING and len(prereg.listing(s)) == 1


# --- from tests/audits/test_cr8_e.py: Code review 2026-10-07, package E: computing time of results and exports. ---


def test_r4_export_is_the_same_with_and_without_the_plan(app, client, survey, monkeypatch):
    from app import export
    from app.items import base

    _every_type(app, client)
    with app.test_request_context():
        s = db.session.get(Survey, 1)
        planned = export.rows(s, with_hb=False)
        csv_planned = export.to_csv(s)
        monkeypatch.setattr(base, "export_plan", nullcontext)  # every participation builds anew, as before
        unplanned = export.rows(s, with_hb=False)
        csv_unplanned = export.to_csv(s)
    assert len(planned) > 10 and list(planned) == list(unplanned)
    assert csv_planned == csv_unplanned


# --- from tests/audits/test_cr8_h.py: Code review 2026-10-07, package H (AI calls and the external signature): I1 no SDK retries and limits below the ---


def test_i9_bibtex_names():
    from app import bibliography

    e = {"text": "x", "kind": "misc", "authors": [("Research and Development Corporation", ""), ("Smith & Co", ""),
                                                   ("Müller", "Anna")], "year": 2020, "title": "Daten"}
    out = bibliography.bibtex([e])
    assert "author = {{Research and Development Corporation} and {Smith \\& Co} and Müller, Anna}" in out


# --- from tests/audits/test_review_fixes.py: What the code review of 2026-10-02 found and what holds since (code-review-2026-10-02.md, S1 to S7). ---


def test_signature_follows_the_structure_not_the_wording(app, client, survey):
    test_flow.setup_design(app, client)
    with app.app_context():
        s = db.session.get(Survey, 1)
        before = runtime._signature(s)
        root = flow.tree(s)
        end, _ = flow.find(root, "e1")
        end.update(text={"de": "Anderer Text"}, name="Zu jung", link="https://example.org/off", link_mode="button")
        assert runtime._signature(s) == before
        end["kind"] = "excluded"  # another status at this end: structure
        assert runtime._signature(s) != before


# --- from tests/audits/test_review_medium.py: The medium findings of the code review of 2026-10-02 (code-review-2026-10-02.md, M5 to M16). ---


def test_downloads_work_with_a_title_outside_latin_1(app, client, survey):
    with app.app_context():
        db.session.get(Survey, 1).title = "Badanie łączności 调查"
        db.session.commit()
    for path in ("export/codebook.csv", "export/questionnaire.docx", "export/survey.json?files=0"):
        answer = client.get(f"/surveys/1/{path}")
        assert answer.status_code == 200, path
        answer.headers["Content-Disposition"].encode("latin-1")


def test_word_questionnaire_with_characters_xml_cannot_hold(app, client, survey):
    from docx import Document

    add(client, "single_choice")
    with app.app_context():
        item = db.session.get(Item, 2)
        item.texts = {"de": {"question": "Vor\x0bNach \x01\x0c￾?", "options": {}}}
        db.session.get(Survey, 1).blocks[1].titles = {"de": "Block\x0bEins"}
        db.session.commit()
    for path in ("export/questionnaire.docx", "export/questionnaire-print.docx"):
        answer = client.get(f"/surveys/1/{path}")
        assert answer.status_code == 200, path
        text = "\n".join(p.text for p in Document(io.BytesIO(answer.data)).paragraphs)
        assert "VorNach ?" in text


def test_import_refuses_a_file_of_nothing_but_structure(app, client, survey):
    from app import transfer

    raw = b'{"format": "queryous-survey", "version": 1, "x": [' + b"[]," * (transfer.MAX_STRUCTURE // 2) + b"[]]}"
    with pytest.raises(transfer.Refused, match="mehr Einträge"):
        transfer.loads(raw)
