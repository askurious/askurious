"""Word exports drop characters XML cannot hold from the start of the app on (security and code audit 2026-10-13,
DA-9): the patch of python-docx came only with app/questionnaire_docx.py, so a methods text or a manuscript table with
a "\\x0b" from old data failed with a 500 until a questionnaire export had run in the same process."""
import shutil
import sys

import pytest


def test_the_patch_is_active_after_create_app(tmp_path, monkeypatch, _db_template):
    import app as package
    from app import constructs, create_app, extensions
    from tests.conftest import PG_URL, _config

    if PG_URL:
        pytest.skip("the app is built from the SQLite template here")
    from docx import Document
    from docx.oxml.text.run import CT_R

    # as in a fresh process: python-docx without the patch, questionnaire_docx not imported yet
    if getattr(CT_R.add_t, "only_xml", False):
        monkeypatch.setattr(CT_R, "add_t", CT_R.add_t.__closure__[0].cell_contents)
    monkeypatch.delitem(sys.modules, "app.questionnaire_docx", raising=False)
    monkeypatch.delattr(package, "questionnaire_docx", raising=False)
    shutil.copyfile(_db_template, tmp_path / "test.db")
    monkeypatch.setattr(extensions, "add_missing_columns", lambda app: None)
    monkeypatch.setattr(constructs, "ensure_seed", lambda: None)
    create_app(_config(tmp_path))
    assert getattr(CT_R.add_t, "only_xml", False)
    doc = Document()
    doc.add_paragraph("Zeile\x0bmit Steuerzeichen\x0c")  # python-docx refused this before
    assert doc.paragraphs[0].text == "Zeilemit Steuerzeichen"
