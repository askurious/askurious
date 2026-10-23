"""Daten › Dokumentation › "Datenfluss (Word)": the data flow of one survey from its settings, in the layout of the
RWTH example table (comparison 2026-10-20, F)."""
import io

from docx import Document

from app.extensions import db
from app.models import Survey
from tests.conftest import fill_consent


def _cells(client):
    r = client.get("/surveys/1/export/datenfluss.docx")
    assert r.status_code == 200 and r.headers["Cache-Control"] == "no-store"
    assert r.headers["Content-Disposition"].endswith('.docx"') and "datenfluss" in r.headers["Content-Disposition"]
    doc = Document(io.BytesIO(r.data))
    return doc, [[c.text for c in row.cells] for row in doc.tables[0].rows]


def test_row_on_the_data_page_and_steps_from_the_settings(app, client, survey):
    assert "Datenfluss (Word)" in client.get("/surveys/1/data").get_data(as_text=True)
    fill_consent(client, open_data="1", repository="OSF", license="CC BY 4.0")
    with app.app_context():
        s = db.session.get(Survey, 1)
        s.paradata, s.erase_code, s.retention = True, True, {"personal": 6, "answers": 120}
        db.session.commit()
    doc, rows = _cells(client)
    assert rows[0] == ["Schritt", "Daten", "Maßnahmen", "Verantwortliche"]
    steps = {r[0]: r for r in rows[1:]}
    assert list(steps) == ["Einwilligung", "Rekrutierung", "Registrierung und Pseudonym", "Erhebung", "Übermittlung", "Speicherung", "Auswertung",
                           "Veröffentlichung", "Löschung"]
    assert "Hash des Löschcodes" in steps["Registrierung und Pseudonym"][1]
    assert "grobe Geräteangaben" in steps["Erhebung"][1] and "Löschcode" in steps["Erhebung"][2]
    assert "OSF" in steps["Veröffentlichung"][1] and "CC BY 4.0" in steps["Veröffentlichung"][1]
    assert "6 Monate" in steps["Löschung"][2] and "10 Jahre" in steps["Löschung"][2]
    assert "Philipp" in steps["Einwilligung"][3]
    assert "Datenfluss der Befragung „Test“" in doc.paragraphs[0].text


def test_without_open_data_and_without_a_period(app, client, survey):
    fill_consent(client)
    with app.app_context():
        db.session.get(Survey, 1).retention = {"personal": None, "answers": None}
        db.session.commit()
    _doc, rows = _cells(client)
    steps = {r[0]: r for r in rows[1:]}
    assert "keine Veröffentlichung der Daten vorgesehen" in steps["Veröffentlichung"][1]
    assert "gute wissenschaftliche Praxis" in steps["Löschung"][2]
    assert "keine Kennungen gespeichert" in steps["Rekrutierung"][2]
