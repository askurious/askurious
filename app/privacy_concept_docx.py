"""The Datenschutzkonzept (app/privacy_concept.py) as a Word document, in the language of the person who downloads it;
the same data as the page Administration > Datenschutz. Built with python-docx."""
import io

from docx import Document
from docx.shared import Cm, Pt, RGBColor

from app import questionnaire_docx  # noqa: F401 - drops characters XML cannot hold from every text (its _only_xml)
from app.lang import de_en

GREY = RGBColor(0x66, 0x66, 0x5F)


def _table(doc, head, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(head))
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, head):
        cell.text = ""
        cell.paragraphs[0].add_run(text).bold = True
    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text or ""
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            if widths:
                cell.width = Cm(widths[i])
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
    doc.add_paragraph()
    return table


def _muted(doc, text):
    run = doc.add_paragraph().add_run(text)
    run.font.color.rgb = GREY
    run.font.size = Pt(9)


def build(c):
    """c: privacy_concept.build(). Returns the .docx as bytes."""
    doc = Document()
    doc.styles["Normal"].font.size = Pt(10)
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2)
    h = c["headings"]
    m = c["meta"]
    doc.add_heading(de_en("Datenschutzkonzept {brand}", "Data protection concept {brand}", brand=m["brand"]), 0)
    _muted(doc, de_en("Stand {date} · {base} · Nutzungsbedingungen Version {terms}",
                      "As of {date} · {base} · terms of use version {terms}", **m))
    for text in c["intro"]:
        doc.add_paragraph(text)

    doc.add_heading(h["scope"][1], 1)
    _table(doc, (de_en("Angabe", "Item"), de_en("Eintrag", "Entry")),
           [(c["labels"][k][0], v["value"] or de_en("— nicht ausgefüllt —", "— not filled in —"))
            for k, v in c["org"].items()], (6, 11))

    doc.add_heading(h["roles"][1], 1)
    _table(doc, (de_en("Personengruppe", "Group of people"), de_en("Wie bestimmt", "How set"),
                 de_en("Sieht und verarbeitet", "Sees and processes")),
           [(r["who"], r["how"], r["sees"]) for r in c["roles"]], (4, 5, 8))

    doc.add_heading(h["record"][1], 1)
    for a in c["activities"]:
        doc.add_heading(f'{a["no"]}. {a["name"]}', 2)
        _table(doc, (de_en("Art. 30 Abs. 1", "Art. 30 (1)"), ""), [
            *([(de_en("Katalogname", "Catalogue name"), a["catalog"])] if a["catalog"] else []),
            (de_en("Eingeführt, zuletzt geändert", "Introduced, last changed"), f'{a["since"]} · {a["changed"]}'),
            (de_en("Verfahren", "Procedure"), a["procedure"]),
            (de_en("Verantwortlich (a)", "Controller (a)"), a["controller"]),
            (de_en("Zweck (b)", "Purpose (b)"), a["purpose"]),
            (de_en("Betroffene (c)", "Data subjects (c)"), a["subjects"]),
            (de_en("Datenkategorien (c)", "Categories of data (c)"), a["data"]),
            (de_en("Empfänger intern (d)", "Internal recipients (d)"), a["internal"]),
            (de_en("Empfänger (d, e)", "Recipients (d, e)"), a["recipients"]),
            (de_en("Übermittlung in ein Drittland (e)", "Transfer to a third country (e)"), a["third_country"]),
            (de_en("Löschfristen (f)", "Deletion periods (f)"), a["deletion"]),
            (de_en("Rechtsgrundlage (Vorschlag)", "Legal basis (proposal)"), a["basis"]),
            *([(de_en("Rechtsgrundlage dieser Installation", "Legal basis of this installation"), a["own_basis"])]
              if a["own_basis"] else []),
        ], (4.5, 12.5))
    _muted(doc, de_en("Maßnahmen (g): siehe Abschnitt 6.", "Measures (g): see section 6."))

    doc.add_heading(h["flow"][1], 1)
    _table(doc, (de_en("Schritt", "Step"), de_en("Daten", "Data"), de_en("Maßnahmen", "Measures"), de_en("Wer", "Who")),
           [(f["step"], f["data"], f["measures"], f["who"]) for f in c["flow"]], (2.6, 5, 6.4, 3))

    doc.add_heading(h["services"][1], 1)
    _table(doc, (de_en("Dienst", "Service"), de_en("Gesendete Daten", "Data sent"), de_en("Ausgelöst von", "Triggered by"),
                 de_en("Schalter", "Switch"), de_en("Hier", "Here"), de_en("Ort, Grundlage", "Location, basis"),
                 de_en("Auftragsverarbeitung", "Processing agreement")),
           [(s["name"], s["data"], s["who"], s["switch"], s["state"][1], s["where"], s["avv"] or "—")
            for s in c["services"]], (2.4, 4, 2, 2, 2.1, 2.6, 1.9))

    doc.add_heading(h["measures"][1], 1)
    doc.add_paragraph(de_en("Die Maßnahmen sind nach den Gewährleistungszielen des Standard-Datenschutzmodells geordnet; "
                            "so entsprechen sie Art. 32 Abs. 1 DSGVO:", "The measures are ordered by the protection goals "
                            "of the Standard Data Protection Model; this is how they answer Art. 32 (1) GDPR:"))
    _table(doc, (de_en("Gewährleistungsziel", "Protection goal"), de_en("Art. 32 Abs. 1", "Art. 32 (1)"),
                 de_en("Gegenstand", "Subject")), c["sdm_art32"], (4.5, 2.5, 10))
    doc.add_paragraph(de_en("Dieselben Maßnahmen nach den Kontrollarten der TOM-Liste:",
                            "The same measures by the control types of the TOM list:"))
    _table(doc, (de_en("Kontrollart", "Control type"), "Art. 32", de_en("Maßnahmen", "Measures")), c["controls"],
           (4, 1.8, 11.2))
    for goal, rows in c["measures"]:
        doc.add_heading(goal, 2)
        _table(doc, (de_en("Maßnahme", "Measure"), de_en("Fundstelle", "Where")), rows, (12, 5))

    doc.add_heading(h["defaults"][1], 1)
    for text in c["defaults"]:
        doc.add_paragraph(text, style="List Bullet")

    doc.add_heading(h["risks"][1], 1)
    doc.add_paragraph(de_en("Kriterien nach WP 248 rev.01; zwei oder mehr sprechen in der Regel für eine DSFA: {k}.",
                            "Criteria of WP 248 rev.01; two or more usually call for an impact assessment: {k}.",
                            k="; ".join(c["criteria"])))
    _table(doc, (de_en("Tätigkeit", "Activity"), de_en("Zutreffende Kriterien", "Criteria met"),
                 de_en("Ergebnis (Vorschlag)", "Result (proposal)")), c["risks"], (4, 6, 7))

    doc.add_heading(h["breach"][1], 1)
    doc.add_paragraph(c["breach"])
    doc.add_paragraph(c["org"]["privacy_breach"]["value"] or de_en("Meldeweg: — nicht ausgefüllt —",
                                                                    "Reporting route: — not filled in —"))

    doc.add_heading(h["open"][1], 1)
    for text in c["open"] or [de_en("Keine.", "None.")]:
        doc.add_paragraph(text, style="List Bullet")

    doc.add_heading(h["sources"][1], 1)
    for text in c["sources"]:
        doc.add_paragraph(text)

    doc.add_page_break()
    doc.add_heading(h["dpia"][1], 1)
    _dpia(doc, c["dpia"], c["labels"]["privacy_dpia_coding"][0])
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _dpia(doc, d, statement_label):
    """The impact assessment of the AI coding help (privacy_concept.dpia_coding()), section by section."""
    doc.add_paragraph().add_run(d["result"]).bold = True
    for text in d["scope"]:
        doc.add_paragraph(text)
    doc.add_paragraph(f'{statement_label}: {d["statement"] or de_en("— nicht ausgefüllt —", "— not filled in —")}')
    doc.add_heading(de_en("Beschreibung der Verarbeitung (Art. 35 Abs. 7 lit. a)",
                          "Description of the processing (Art. 35 (7) (a))"), 2)
    _table(doc, (de_en("Angabe", "Item"), de_en("Eintrag", "Entry")), d["facts"], (4.5, 12.5))
    doc.add_heading(de_en("Schwellwertanalyse", "Threshold analysis"), 2)
    _table(doc, (de_en("Kriterium (WP 248)", "Criterion (WP 248)"), de_en("Erfüllt", "Met"),
                 de_en("Begründung", "Reason")), d["threshold"], (6, 2.5, 8.5))
    doc.add_heading(de_en("Notwendigkeit und Verhältnismäßigkeit (lit. b)", "Necessity and proportionality (b)"), 2)
    for text in d["necessity"]:
        doc.add_paragraph(text, style="List Bullet")
    doc.add_heading(de_en("Risiken (lit. c)", "Risks (c)"), 2)
    _table(doc, (de_en("Nr.", "No."), de_en("Risiko", "Risk"), de_en("Vor den Maßnahmen", "Before the measures"),
                 de_en("Danach, über OpenAI", "After, over OpenAI"),
                 de_en("Danach, Anbieter in Deutschland", "After, provider in Germany")), d["risks"], (1.2, 7.3, 2.8, 2.8, 2.9))
    doc.add_heading(de_en("Maßnahmen (lit. d)", "Measures (d)"), 2)
    _table(doc, (de_en("Vorhanden", "In place"), de_en("Gegen", "Against"), de_en("Fundstelle", "Where")),
           d["measures"], (9, 2.5, 5.5))
    doc.add_paragraph().add_run(de_en("Vorgeschlagen:", "Proposed:")).bold = True
    for text in d["proposed"]:
        doc.add_paragraph(text, style="List Bullet")
    _muted(doc, d["review"])
