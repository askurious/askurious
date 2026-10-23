"""A draft of the methods section and the paragraph on exclusions, from what Askurious knows about a survey (concept
"Askurious für Studierende und Berufseinsteiger", package F): sample, instruments with their sources and
reliability, exclusions, analysis plan and the results per hypothesis. German or English (S5).

It is a draft: the numbers are those of the data pool of the results page (the saved pool, without what quality
control marks), the sentences are plain templates the researcher rewrites. Word when python-docx is there, else
Markdown.
"""
import io

from app import study_plan

TEXT = {
    "de": {
        "title": "Methode (Entwurf aus Askurious)", "sample": "Stichprobe", "instruments": "Instrumente",
        "exclusions": "Datenbereinigung", "analysis": "Auswertung", "results": "Ergebnisse je Hypothese",
        "n": "An der Befragung „{title}“ nahmen {n} Personen vollständig teil{target}.",
        "target": " (angestrebt: {target})", "recruit": "Rekrutiert wurde über: {how}.",
        "period": "Erhoben wurde von {start} bis {end}.",
        "power": "Die Stichprobengröße wurde vorab mit einer Poweranalyse bestimmt (exakte nichtzentrale Verteilungen; "
                 "Faul et al., 2009): {parts}. Maßgeblich war {h} mit N = {n}; bei {e} % erwarteten Ausschlüssen waren dafür {c} "
                 "Abschlüsse nötig.",
        "scale": "{label} wurde mit {k} Aussagen gemessen ({source}); α = {alpha}, ω = {omega}.",
        "scale_nosrc": "{label} wurde mit {k} Aussagen gemessen; α = {alpha}, ω = {omega}.",
        "single": "{label} ({code}) wurde mit einer Frage erhoben.",
        "excl": "Von {n} vollständigen Teilnahmen wurden {marked} ausgeschlossen: {parts}.",
        "excl_none": "Keine der {n} vollständigen Teilnahmen wurde ausgeschlossen.",
        "speeders": "{k} wegen zu kurzer Bearbeitungszeit ({rule})",
        "flagged": "{k} wegen {least} oder mehr Auffälligkeiten im Antwortverhalten (z. B. Aufmerksamkeitsprüfung, Antwortmuster)",
        "rules": "Vorab festgelegte Ausschlussregeln: {rules}",
        "hyp": "H{i}: {text} Geprüft mit: {test}.", "result": "H{i} ({verdict}): {sentence}",
        "explorative": " (explorativ, nach dem Start ergänzt)", "sensitivity": "Sensitivitätsanalyse:", "draft": "Entwurf: bitte prüfen und in eigenen Worten fassen.",
        # what APA JARS asks for and the draft did not say (audit 2026-10-05, I-9)
        "holm": "Für die {k} getesteten Hypothesen wurde nach Holm (1979) korrigiert; das Urteil folgt dem korrigierten p.",
        "tests": "Alle Tests wurden zweiseitig mit α = ,05 gerechnet; gerichtete Hypothesen gelten nur bei einem Effekt in der "
                 "vorhergesagten Richtung als gestützt. Für mehrere Hypothesen wurde nicht korrigiert. Gruppen mit weniger als "
                 "{min_group} Antworten gingen nicht in die Tests ein; bei kleinen erwarteten Häufigkeiten gilt der exakte Test "
                 "nach Fisher (2 × 2) bzw. ein Monte-Carlo-Permutationstest. χ²-Tests wurden ohne Kontinuitätskorrektur gerechnet. "
                 "Der Mann-Whitney-U-Test war ohne Bindungen exakt, solange beide Gruppen kleiner als 50 waren, mit Bindungen "
                 "unter 20 Fällen exakt (bedingte Permutationsverteilung), sonst die Normalnäherung mit Bindungs- und "
                 "Stetigkeitskorrektur. Das p der Rangkorrelation nach Spearman kommt ohne Bindungen aus der exakten "
                 "Verteilung (bis 12 Paare), darüber aus einer an die exakten Momente angeglichenen Beta-Näherung (Pitman, "
                 "1937; David, Kendall & Stuart, 1951), mit Bindungen aus der t-Verteilung.",
        # the structural equation model (app/sem_model.py)
        "sem_cb": "Die mit „Strukturgleichungsmodell“ geplanten Hypothesen wurden gemeinsam in einem kovarianzbasierten "
                  "Strukturgleichungsmodell geprüft (Maximum Likelihood mit robusten Standardfehlern und Satorra-Bentler-"
                  "korrigiertem χ²; Satorra & Bentler, 1994). Skalen gingen als latente Konstrukte mit ihren Aussagen als "
                  "Indikatoren ein. Der Modellfit wurde mit χ², CFI, TLI, RMSEA und SRMR beurteilt (Hu & Bentler, 1999; CFI, "
                  "TLI und RMSEA normaltheoretisch aus dem ML-χ² und robust nach Brosseau-Liard & Savalei, 2014), "
                  "das Messmodell mit AVE, Composite Reliability und HTMT (Fornell & Larcker, 1981; Henseler et al., 2015).",
        "sem_pls": "Die mit „Strukturgleichungsmodell“ geplanten Hypothesen wurden gemeinsam mit PLS-SEM geprüft (Modus A, "
                   "Pfadgewichtung, für Paare ohne Richtung Faktorgewichtung; Hair et al., 2022), Skalen als Konstrukte aus ihren Aussagen. Die Signifikanz der Pfade "
                   "folgt aus einem Bootstrap mit {draws} Ziehungen. Das Messmodell wurde mit AVE, Composite Reliability "
                   "und HTMT beurteilt (Fornell & Larcker, 1981; Henseler et al., 2015).",
        "sem_ind": "Indirekte Effekte (Vermittlung) wurden über die gemeinsame Signifikanz beider Pfade geprüft "
                   "(MacKinnon et al., 2002; als p zählt das größere der beiden Pfad-p) und als Produkt der "
                   "standardisierten Pfade mit einem Perzentil-Bootstrap-Intervall ({draws} Ziehungen) berichtet "
                   "(Preacher & Hayes, 2008).",
        "sem_cov": "Hypothesen über einen Zusammenhang ohne Richtung gingen als Kovarianz in das Modell ein, bei im "
                   "Modell erklärten Größen als Kovarianz der Residuen; berichtet wird die standardisierte Korrelation.",
        "sem_mod": "Moderationen wurden als latente Wechselwirkung mit Produktindikatoren geprüft (doppelt zentrierte, "
                   "nach der Reihenfolge der Aussagen gepaarte Indikatoren; Marsh et al., 2004; Lin et al., 2010), bei "
                   "PLS zweistufig über das Produkt der Konstruktwerte (Hair et al., 2022); berichtet werden der Pfad der "
                   "Wechselwirkung und die einfachen Steigungen bei ±1 SD des Moderators (SD aus dem geschätzten Modell), "
                   "bei einem Moderator aus zwei Gruppen je Gruppe (Hayes, 2015). Ein Moderator mit zwei Gruppen "
                   "wurde im kovarianzbasierten Modell über ein Mehrgruppenmodell mit gleichen Ladungen (metrische "
                   "Invarianz) und einen Wald-Test des Pfadunterschieds geprüft.",
        "sem_fit": "Das Strukturgleichungsmodell (N = {n}) passte so zu den Daten: {fit}.",
        # the models over the waves and of a diary study (longitudinal and multilevel concept 2026-10-17)
        "waves": "Die {k} Wellen wurden über den persönlichen Code verknüpft; jede verknüpfte Person ging mit den Wellen "
                 "ein, an denen sie teilnahm (Full Information Maximum Likelihood mit robusten Standardfehlern, MLR; "
                 "Enders & Bandalos, 2001; Yuan & Bentler, 2000).",
        "waves_models": "Über die Wellen wurden die Messinvarianz über die Zeit (konfigural, metrisch, skalar; Widaman et "
                        "al., 2010) und latente Wachstumskurven geschätzt (Bollen & Curran, 2006).",
        "lagged": "Hypothesen über eine Wirkung über die Zeit wurden im Random-Intercept-Cross-Lagged-Panel-Modell "
                  "geprüft, das stabile Unterschiede zwischen Personen von Wirkungen innerhalb der Personen trennt "
                  "(Hamaker et al., 2015; bei zwei Wellen im Cross-Lagged-Panel-Modell).",
        "mixed": "Zusammenhänge in den Einträgen der Tagebuchstudie wurden mit linearen gemischten Modellen geprüft "
                 "(REML, zufälliger Achsenabschnitt je Person{slope}); die Einflussgröße wurde in die Abweichung vom "
                 "Personenmittel (innerhalb) und das Personenmittel (zwischen) zerlegt (Enders & Tofighi, 2007){lag}, "
                 "p mit Freiheitsgraden nach Satterthwaite (Kuznetsova et al., 2017).",
        "mixed_slope": " und zufälliger Steigung", "mixed_lag": "; das Erklärende stammte aus dem vorigen Eintrag",
        "sem_controls": "Kontrolliert wurde für {c}: mit je einem Pfad auf jede im Modell erklärte Größe, ohne eigene "
                        "Hypothese.",
        "sem_fiml": "Fehlende Werte wurden mit Full Information Maximum Likelihood behandelt (Enders & Bandalos, 2001), "
                    "mit robusten Standardfehlern und dem Test nach Yuan und Bentler (2000, MLR).",
        "sem_second": "Als Faktoren zweiter Ordnung über ihren Subskalen gingen ein: {c}.",
        "sem_compare": "Vorab festgelegt wurde ein Konkurrenzmodell ohne {c}, verglichen mit dem skalierten "
                       "χ²-Differenztest (Satorra & Bentler, 2001), AIC und BIC.",
        "sem_wlsmv": "Die Aussagen gingen als geordnete Kategorien ein: polychorische Korrelationen, geschätzt mit "
                     "diagonal gewichteten kleinsten Quadraten und robusten Korrekturen (WLSMV; Muthén, 1984); χ², CFI, "
                     "TLI und RMSEA sind mittelwert- und varianzkorrigiert (Asparouhov & Muthén, 2010).",
        "sem_invariance": "Die Messinvarianz über {g} wurde in Mehrgruppenmodellen stufenweise geprüft (konfigural, "
                          "metrisch, skalar, strikt; Vandenberg & Lance, 2000): Eine Stufe galt als nicht gehalten, wenn "
                          "der CFI um 0,010 oder mehr sank und zugleich RMSEA oder SRMR stiegen (Chen, 2007); dazu der "
                          "skalierte χ²-Differenztest (Satorra & Bentler, 2001).",
        "sem_pls_more": "Berichtet werden R² und korrigiertes R², die Effektstärke f² je Pfad (Cohen, 1988), die VIF der "
                        "Einflussgrößen und die Vorhersagekraft der erklärten Konstrukte mit PLSpredict (10-fache "
                        "Kreuzvalidierung; Shmueli et al., 2019).",
        "sem_formative": "Formativ gemessen (Modus B) wurden {c}; ihre Gewichte wurden mit dem Bootstrap geprüft, ihre "
                         "Kollinearität mit dem VIF (Hair et al., 2022).",
        "sem_single": "Einzelfragen gingen als Konstrukte mit einer angenommenen Reliabilität ein ({c}); ihre "
                      "Fehlervarianz wurde auf (1 − Reliabilität) ihrer Varianz festgelegt (Petrescu, 2013).",
    },
    "en": {
        "title": "Method (draft from Askurious)", "sample": "Sample", "instruments": "Measures",
        "exclusions": "Data cleaning", "analysis": "Analysis", "results": "Results per hypothesis",
        "n": "{n} people completed the survey “{title}”{target}.",
        "target": " (target: {target})", "recruit": "Participants were recruited via: {how}.",
        "period": "Data were collected from {start} to {end}.",
        "power": "The sample size was set in advance by a power analysis (exact noncentral distributions; Faul et al., "
                 "2009): {parts}. Decisive was {h} with N = {n}; with {e}% expected exclusions, {c} completions were "
                 "needed for it.",
        "scale": "{label} was measured with {k} items ({source}); α = {alpha}, ω = {omega}.",
        "scale_nosrc": "{label} was measured with {k} items; α = {alpha}, ω = {omega}.",
        "single": "{label} ({code}) was measured with a single question.",
        "excl": "Of {n} completed participations, {marked} were excluded: {parts}.",
        "excl_none": "None of the {n} completed participations was excluded.",
        "speeders": "{k} for completing too fast ({rule})",
        "flagged": "{k} for {least} or more signs of careless responding (e.g. attention check, response patterns)",
        "rules": "Exclusion rules set in advance: {rules}",
        "hyp": "H{i}: {text} Tested with: {test}.", "result": "H{i} ({verdict}): {sentence}",
        "explorative": " (exploratory, added after the start)", "sensitivity": "Sensitivity analysis:", "draft": "Draft: please check and rewrite in your own words.",
        "holm": "For the {k} tested hypotheses, p was corrected after Holm (1979); the verdict follows the corrected p.",
        "tests": "All tests were two-sided with α = .05; directional hypotheses count as supported only for an effect in the "
                 "predicted direction. No correction was applied for multiple hypotheses. Groups with fewer than {min_group} "
                 "answers were not tested; with small expected counts, Fisher's exact test (2 × 2) or a Monte Carlo permutation "
                 "test was used. χ² tests were computed without continuity correction. The Mann-Whitney U test was exact "
                 "without ties while both groups were smaller than 50, exact with ties below 20 cases (conditional "
                 "permutation distribution), else the normal approximation with tie and continuity correction. The p of "
                 "Spearman's rank correlation comes from the exact distribution (up to 12 pairs) or above that a beta "
                 "approximation fitted to the exact moments (Pitman, 1937; David, Kendall & Stuart, 1951) without ties, "
                 "from the t distribution with ties.",
        "sem_cb": "The hypotheses planned as “structural equation model” were tested together in a covariance-based "
                  "structural equation model (maximum likelihood with robust standard errors and the Satorra-Bentler "
                  "scaled χ²; Satorra & Bentler, 1994). Scales entered as latent constructs with their items as "
                  "indicators. Model fit was judged by χ², CFI, TLI, RMSEA and SRMR (Hu & Bentler, 1999; CFI, TLI and "
                  "RMSEA as normal-theory values from the ML χ² and robust after Brosseau-Liard & Savalei, 2014), the "
                  "measurement model by AVE, composite reliability and HTMT (Fornell & Larcker, 1981; Henseler et al., 2015).",
        "sem_pls": "The hypotheses planned as “structural equation model” were tested together with PLS-SEM (mode A, path "
                   "weighting, factor weighting for pairs without a direction; Hair et al., 2022), with scales as constructs of their items. The significance of the paths "
                   "comes from a bootstrap with {draws} draws. The measurement model was judged by AVE, composite "
                   "reliability and HTMT (Fornell & Larcker, 1981; Henseler et al., 2015).",
        "sem_ind": "Indirect effects (mediation) were tested by the joint significance of both paths (MacKinnon et al., "
                   "2002; the larger of the two path p counts) and reported as the product of the standardised paths "
                   "with a percentile bootstrap interval ({draws} draws; Preacher & Hayes, 2008).",
        "sem_cov": "Hypotheses on a relationship without a direction entered the model as a covariance, for variables "
                   "explained in the model as a covariance of their residuals; the standardised correlation is reported.",
        "sem_mod": "Moderations were tested as latent interactions with product indicators (double mean centred indicators "
                   "paired in the order of the items; Marsh et al., 2004; Lin et al., 2010), with PLS in two stages through "
                   "the product of the construct scores (Hair et al., 2022); the path of the interaction and the simple "
                   "slopes at ±1 SD of the moderator (SD from the estimated model), for a moderator with two groups in "
                   "each group (Hayes, 2015), are reported. A moderator with two groups was tested in the "
                   "covariance-based model by a multi-group model with equal loadings (metric invariance) and a Wald "
                   "test of the difference of the path.",
        "sem_fit": "The structural equation model (N = {n}) fitted the data as follows: {fit}.",
        "waves": "The {k} waves were linked by the personal code; every linked person entered with the waves they took "
                 "part in (full information maximum likelihood with robust standard errors, MLR; Enders & Bandalos, "
                 "2001; Yuan & Bentler, 2000).",
        "waves_models": "Over the waves, measurement invariance over time (configural, metric, scalar; Widaman et al., "
                        "2010) and latent growth curves (Bollen & Curran, 2006) were estimated.",
        "lagged": "Hypotheses on an effect over time were tested in the random intercept cross-lagged panel model, which "
                  "separates stable differences between people from effects within people (Hamaker et al., 2015; with "
                  "two waves in the cross-lagged panel model).",
        "mixed": "Relationships in the entries of the diary survey were tested with linear mixed models (REML, a random "
                 "intercept per person{slope}); the predictor was split into the deviation from the person's mean "
                 "(within) and the person's mean (between; Enders & Tofighi, 2007){lag}, p with Satterthwaite degrees "
                 "of freedom (Kuznetsova et al., 2017).",
        "mixed_slope": " and a random slope", "mixed_lag": "; the predictor came from the entry before",
        "sem_controls": "The model controlled for {c}, with one path onto every variable it explains, without a "
                        "hypothesis of their own.",
        "sem_fiml": "Missing values were handled by full information maximum likelihood (Enders & Bandalos, 2001), with "
                    "robust standard errors and the test after Yuan and Bentler (2000, MLR).",
        "sem_second": "Second-order factors over their subscales: {c}.",
        "sem_compare": "A competing model without {c} was set in advance and compared by the scaled χ² difference test "
                       "(Satorra & Bentler, 2001), AIC and BIC.",
        "sem_wlsmv": "The statements entered as ordered categories: polychoric correlations, estimated by diagonally "
                     "weighted least squares with robust corrections (WLSMV; Muthén, 1984); χ², CFI, TLI and RMSEA are "
                     "mean- and variance-adjusted (Asparouhov & Muthén, 2010).",
        "sem_invariance": "Measurement invariance across {g} was tested step by step in multi-group models (configural, "
                          "metric, scalar, strict; Vandenberg & Lance, 2000): a level was taken as not holding when the "
                          "CFI dropped by .010 or more while the RMSEA or the SRMR rose (Chen, 2007), together with the "
                          "scaled χ² difference test (Satorra & Bentler, 2001).",
        "sem_pls_more": "R² and adjusted R², the effect size f² per path (Cohen, 1988), the VIF of the predictors and the "
                        "predictive power of the explained constructs with PLSpredict (10-fold cross-validation; Shmueli "
                        "et al., 2019) are reported.",
        "sem_formative": "{c} were measured formatively (mode B); their weights were tested with the bootstrap, their "
                         "collinearity with the VIF (Hair et al., 2022).",
        "sem_single": "Single questions entered as constructs with an assumed reliability ({c}); their error variance "
                      "was fixed at (1 − reliability) of their variance (Petrescu, 2013).",
    },
}
RECRUITMENT_EN = {"convenience": "convenience sample (acquaintances, mailing lists)", "snowball": "snowball sampling",
                  "panel": "online panel", "quota": "quota sample", "random": "random sample", "census": "full census of the group"}
ANALYSES_EN = {"r": "Pearson correlation", "rho": "Spearman rank correlation", "t": "Welch's t-test for two groups",
               "anova": "Welch's ANOVA for several groups", "u": "Mann-Whitney U test", "kw": "Kruskal-Wallis test",
               "chi2": "χ² test (cross table)", "describe": "description only", "": "–",
               "sem": "structural equation model", "panel": "cross-lagged panel model over the waves",
               "mixed": "multilevel model (entries within people)"}
VERDICT_EN = {"unterstützt": "supported", "nicht unterstützt": "not supported", "Gegenteil": "opposite", "offen": "open", "beschrieben": "described"}


def _num(v, lang, d=2):
    """v with d decimals rounded half up (0,125 -> 0,13; Mathe-Audit 2026-10-06, Z-9), a decimal comma in German."""
    from app.units import decimal_text

    return "–" if v is None else decimal_text(v, d, lang=lang)


def _bounded(v, lang):
    """_num of α or ω, which cannot exceed 1: in English without the leading zero (APA, as manuscript_tables.fmt with
    bounded=True); German keeps it (audit 2026-10-10, B-20)."""
    text = _num(v, lang)
    return text.replace("0.", ".", 1) if lang == "en" else text


def exclusions(survey, summary, lang="de"):
    """The paragraph on exclusions from quality.summary (the page "Qualitätskontrolle")."""
    from app import export

    t = TEXT[lang]
    n, marked, speeders = summary["n"], summary["marked"], summary["speeders"]
    if not marked:
        return t["excl_none"].format(n=n)
    rule = (export.rule_text(survey, lang) if export.speeder_rule(survey) == "rsi"
            else (f"unter {survey.speeder_pct} % der Median-Dauer" if lang == "de" else f"below {survey.speeder_pct}% of the median time"))
    parts = []
    if speeders:
        parts.append(t["speeders"].format(k=speeders, rule=rule))
    if marked - speeders > 0:
        parts.append(t["flagged"].format(k=marked - speeders, least=summary["rules"]["min"]))
    return t["excl"].format(n=n, marked=marked, parts="; ".join(parts))


def _sentence(text):
    text = text.strip()
    return text if text.endswith((".", "?", "!")) else text + "."


def sections(survey, data, table, summary, lang="de", n=None, facts=None, sem=None):
    """[(heading, [paragraphs])] of the draft. n: the completed participations of the data pool (default: the real
    ones). facts: the kept facts of the method page (methods.facts) of the same pool: their sentences on choice
    experiments, Standort and tasks stand in for those made from `table`, which is then only needed for the
    hypotheses (audit 2026-10-10, Q-Rest). sem: a function giving the structural equation model of the same pool as
    the results page keeps it (None: estimated here from `table`)."""
    from app import constructs, hypotheses, methods
    from app.items.base import plain
    from app.surveys import pool_counts

    t = TEXT[lang]
    plan = study_plan.plan_of(survey)
    out = []
    sample = []
    target = plan["sample"]["target"]
    sample.append(t["n"].format(title=survey.title, n=pool_counts(survey)["real"] if n is None else n,
                                target=t["target"].format(target=target) if target else ""))
    if plan["sample"]["recruitment"]:
        how = plan["sample"]["recruitment"]
        sample.append(t["recruit"].format(how=RECRUITMENT_EN.get(how, how) if lang == "en" else study_plan.RECRUITMENT[how]))
    log = [e for e in (survey.status_log or []) if isinstance(e, dict)]
    starts = [e.get("at") for e in log if e.get("event") == "start" and e.get("at")]
    ends = [e.get("at") for e in log if e.get("event") == "close" and e.get("at")]
    if starts and ends:
        sample.append(t["period"].format(start=str(starts[0])[:10], end=str(ends[-1])[:10]))
    ps = study_plan.power_summary(survey)
    if ps:
        parts = [f"H{i} {study_plan.power_line(survey, plan['hypotheses'][i - 1], lang).split(': ', 1)[1]}"
                 for i, p in ps["rows"] if p["n"]]
        sample.append(t["power"].format(parts="; ".join(parts), h=", ".join(f"H{i}" for i in ps["by"]), n=ps["n"],
                                        e=ps["exclusions"], c=ps["completions"]))
    out.append((t["sample"], sample))

    measures = []
    rows = {(r["code"], r.get("sub")): r for r in ((data or {}).get("overview") or {}).get("rows", [])}
    for block in survey.blocks:
        for item in block.items:
            if not item.kind.collects_answer or item.type in ("consent", "geolocation") + methods.CHOICE_TYPES + methods.TASK_TYPES:
                continue  # choice experiments, Standort, IAT and delay discounting get their own sentences below
            info = constructs.info(item)
            label = (info or {}).get("label") or plain(item.text("question", survey.default_language))[:80] or item.code
            row = rows.get((item.code, None))
            if row and row.get("k"):
                key = "scale" if info and info.get("citation") else "scale_nosrc"
                measures.append(t[key].format(label=label, k=row["k"], source=(info or {}).get("citation", ""),
                                              alpha=_bounded(row.get("alpha"), lang), omega=_bounded(row.get("omega"), lang)))
            elif info:
                measures.append(t["single"].format(label=label, code=item.code)
                                + (f" ({info['citation']})" if info.get("citation") else ""))
    if table is not None or facts is not None:  # the sentences of the method page (app/methods.py) on choice
        # experiments and Standort: kept with its facts, or made from the table
        for c in facts["choice"] if facts is not None else methods._choice(survey, table):
            measures.append(c["text_de"] if lang == "de" else methods._choice_en(c) + ".")
        for loc in facts["location"] if facts is not None else methods._location(survey, table):
            measures.append(methods.location_de(loc) if lang == "de" else methods.location_en(loc) + ".")
        # voice messages (G7); facts kept before they existed have none
        for v in (facts.get("voice") or []) if facts is not None else methods._voice(survey, table):
            measures.append(methods.voice_de(v) if lang == "de" else methods.voice_en(v) + ".")
        for task in facts["tasks"] if facts is not None else methods._tasks(survey, table):
            measures.append(task["text_de"] if lang == "de" else task["text_en"] + ".")
        from app import coding  # categorised open answers

        for c in coding.methods_sentences(survey):
            measures.append(c["text_de"] if lang == "de" else c["text_en"])
    if measures:
        out.append((t["instruments"], measures))

    cleaning = [exclusions(survey, summary, lang)] if summary else []
    if plan["rules"]:
        cleaning.append(t["rules"].format(rules=plan["rules"]))
    if cleaning:
        out.append((t["exclusions"], cleaning))

    hyps = plan["hypotheses"]
    if hyps:
        analysis = []
        for i, h in enumerate(hyps, 1):
            a = h.get("analysis") or study_plan.suggest(survey, h)
            test = ANALYSES_EN.get(a, a) if lang == "en" else study_plan.ANALYSES.get(a, "–")
            analysis.append(t["hyp"].format(i=i, text=_sentence(h["text"]), test=test)
                            + (t["explorative"] if h.get("after_start") else ""))
        if any((h.get("analysis") or study_plan.suggest(survey, h)) not in ("describe", "") and not study_plan.untested(h)
               for h in hyps):
            from app.compare import MIN_GROUP

            tests = t["tests"].format(min_group=MIN_GROUP)
            k = len(study_plan.tested(survey))
            if plan.get("holm") and k >= 2:  # the sentence on no correction gives way to the one on Holm (PA5)
                tests = tests.replace(" Für mehrere Hypothesen wurde nicht korrigiert.", " " + t["holm"].format(k=k))                              .replace(" No correction was applied for multiple hypotheses.", " " + t["holm"].format(k=k))
            analysis.append(tests)
        from app import sem_model, waves as waves_mod

        linked = len(waves_mod.of(survey))
        if linked >= 2:  # the models over the waves (app/panel_models.py)
            from app import panel_models

            analysis.append(t["waves"].format(k=linked))
            if panel_models.settings(survey)["vars"]:
                analysis.append(t["waves_models"])
            if any(h.get("kind") == "lagged" for h in hyps):
                analysis.append(t["lagged"])
        if survey.sampling is not None and any(
                (h.get("analysis") or study_plan.suggest(survey, h)) in ("r", "rho") or h.get("kind") == "effect" for h in hyps):
            opts = plan["diary"]  # the multilevel model of a diary study (app/diary_model.py)
            analysis.append(t["mixed"].format(slope=t["mixed_slope"] if opts["slope"] else "",
                                              lag=t["mixed_lag"] if opts["lagged"] else ""))
        uses_sem = sem_model.in_use(survey) and survey.sampling is None
        if uses_sem:
            analysis.append(t["sem_pls" if sem_model.method_of(survey) == "pls" else "sem_cb"].format(draws=sem_model.BOOT_DRAWS))
            if any(h.get("kind") == "mediation" for h in hyps):
                analysis.append(t["sem_ind"].format(draws=sem_model.BOOT_DRAWS))
            if any(h.get("kind") == "correlation" and h.get("analysis") == "sem" for h in hyps):
                analysis.append(t["sem_cov"])
            if any(h.get("kind") == "moderation" for h in hyps):
                analysis.append(t["sem_mod"])
            opts = plan["sem"]
            if opts["controls"]:
                analysis.append(t["sem_controls"].format(c=", ".join(opts["controls"])))
            if opts["second"] and sem_model.method_of(survey) != "pls":
                analysis.append(t["sem_second"].format(c=", ".join(opts["second"])))
            if opts["alternative"] and sem_model.method_of(survey) != "pls":
                numbers = {h.get("id"): n for n, h in enumerate(plan["hypotheses"], 1)}
                left = [(f"H{numbers.get(k[2:], '?')}" if k[0] == "h" else
                         (f"den direkten Effekt in H{numbers.get(k[2:], '?')}" if lang == "de" else
                          f"the direct effect in H{numbers.get(k[2:], '?')}")) for k in opts["alternative"]]
                analysis.append(t["sem_compare"].format(c=", ".join(left)))
            if opts["ordinal"] and sem_model.method_of(survey) != "pls":
                analysis.append(t["sem_wlsmv"])
            elif opts["missing"] == "fiml" and sem_model.method_of(survey) != "pls":
                analysis.append(t["sem_fiml"])
            if opts["invariance"]:
                group = (("die Sprache der Befragung" if lang == "de" else "the language of the survey")
                         if opts["invariance"] == study_plan.LANG_GROUPS else opts["invariance"])
                analysis.append(t["sem_invariance"].format(g=group))
            if sem_model.method_of(survey) == "pls":
                analysis.append(t["sem_pls_more"])
                if opts["formative"]:
                    analysis.append(t["sem_formative"].format(c=", ".join(opts["formative"])))
            if opts["single"] and sem_model.method_of(survey) != "pls":
                rel = (lambda r: f"{r:.2f}" if lang == "en" else f"{r:.2f}".replace(".", ","))
                analysis.append(t["sem_single"].format(c=", ".join(f"{c}: {rel(r)}" for c, r in opts["single"].items())))
        out.append((t["analysis"], analysis))
        if table is not None:
            found = hypotheses.evaluate(survey, table, lang, correct=False)
            sem_result = ((sem() if sem else sem_model.estimate(survey, table)) if uses_sem else None)
            if sem_result is not None:  # the results page's model on the same pool, with its bootstrap if there is one
                boot = sem_model.boot_state(survey, sem_result) or {}
                sem_model.apply(found, sem_result, hyps, boot.get("result"), lang)
            hypotheses.with_holm(survey, found, lang)
            lines = []
            if sem_result is not None and sem_model.fit_sentence(sem_result.get("fit"), lang):
                lines.append(t["sem_fit"].format(n=sem_result["n"], fit=sem_model.fit_sentence(sem_result["fit"], lang)))
            for i, r in enumerate(found, 1):
                if r.get("sentence"):
                    verdict = r["verdict"] if lang == "de" else VERDICT_EN.get(r["verdict"], r["verdict"])
                    sentence = r["sentence"].replace("<sub>s</sub>", "s")
                    lines.append(t["result"].format(i=i, verdict=verdict, sentence=sentence)
                                 + (" " + t["sensitivity"] + " " + r["sensitivity"]["text"] if r.get("sensitivity") else ""))
            if lines:
                out.append((t["results"], lines))
    return out


def markdown(survey, parts, lang="de"):
    t = TEXT[lang]
    lines = [f"# {t['title']}", "", f"_{t['draft']}_", ""]
    for heading, paragraphs in parts:
        lines += [f"## {heading}", ""] + [p + "\n" for p in paragraphs]
    return "\n".join(lines)


def docx(survey, parts, lang="de"):
    """The draft as .docx bytes (python-docx; ImportError when it is missing on the server)."""
    from docx import Document
    from docx.shared import Pt

    t = TEXT[lang]
    doc = Document()
    doc.styles["Normal"].font.name, doc.styles["Normal"].font.size = "Times New Roman", Pt(12)
    doc.add_heading(t["title"], level=1)
    doc.add_paragraph().add_run(t["draft"]).italic = True
    for heading, paragraphs in parts:
        doc.add_heading(heading, level=2)
        for p in paragraphs:
            doc.add_paragraph(p)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
