"""Market simulator on the results page (concept wahlmodelle-punkt6-konzept.md, stage C; W8 to W12, W14).

CBC and ACBC: one to six products (two at the start), each a level per attribute (a price marked as numbers: any number within the range of the
price levels of the study, nothing outside it), and "Keine davon" if the study had it. MaxDiff: two to ten statements
("Auswahl aus Aussagen"). The utilities come from the current hierarchical Bayes result (one row per person,
app/conjoint_hb.py, app/maxdiff_hb.py); without one from the aggregate logit, which knows no differences between
persons. Shares (app/choice_predict.py): share of preference (logit per person, then the mean) and, with HB, first
choice over the HB draws (per person the share of its kept draws in which the product has the highest utility, ties
split, then the mean over persons; this is first choice over the posterior draws, not the classic first-choice
rule on one point estimate per person, and the page names it "über die HB-Ziehungen"; audit data storage and calculations
2026-10-21, D8), each with a 95 % percentile
bootstrap interval over persons (only the sampling spread over persons). With HB a person's shares are averaged over
the kept draws stored with the result, not computed from the shrunken posterior means (audit 2026-10-05, W-1). The
notes of the model that bear on the shares stand above them (W-4): with HB its warnings and persons with a positive
price coefficient, with the pooled logit a price that is not clearly negative and levels only the prior bounds. With the weighting of the results page (?gewichtet=1, app/weights.py) the means
over persons are weighted; persons without a weight are left out, as on the rest of the page.

The page sends the scenario with GET (nothing stored): `sim` names the question, then per product j
`sim_<j>_<attribute>` (level index) and `sim_<j>_price`, `sim_n` (products), `sim_none`, `sim_act` (add / drop<j> /
opt / optoff / take<j>); MaxDiff `sim_s` (statement indices). `sim_load` takes a saved scenario. With a script
(static/simulator.js) the fold alone is fetched and replaced (`sim_part`, surveys.results), the page does not load
again (Ralf, 2026-10-05).

Ratings-based conjoint (CVA, app/items/cva.py; concept aca-cva-konzept.md, C6): the same simulator on the part-worths of
the regression per person (results entry "cva", person_pw), basis "persons": one row per person, as the HB draws of one
person. The part-worths are in points of the answer scale; for the logit rule they are converted to logit units with
the range of the scale as 6 units (Festlegung von Askurious: the scale maps onto -3 to 3, a usual span of logit
utilities), times an exponent
the page sets (`sim_exp`, default 1; Orme, 2019: the exponent tunes how sharply the shares follow the utilities).
First choice does not depend on it. No price coefficient and no "Keine davon": the levels of a price are part-worths.

ACBC (Ralf, 2026-10-05): the same simulator on the individual HB utilities; without a current HB estimate none (the
aggregate ACBC logit rests on concepts near each person's own BYO profile, its price coefficient comes out too small),
no "Keine davon" (the threshold "kommt nicht infrage" of the screening is no choice of nothing in a market: it rests on
the screening scale, not on the tournament's), and with a summed price the price is its own row (any number within the
prices shown in the study), set to the list price of the levels (base plus surcharges) when a level changes; "Produkt
optimieren" gives every variant its list price unless the price is held.

"Produkt optimieren" (CBC, ACBC, `sim_opt`): one product of the scenario, some attributes held at its levels (`sim_fix`),
every combination of the levels of the others (a free price: the price levels of the study) computed against the
other products of the scenario and "Keine davon"; the variants with the highest share of preference first. The
single product design problem (Kohli & Krishnamurti, 1987) by complete enumeration instead of their heuristic: the
spaces of a study are small enough (at most MAX_VARIANTS). Alone, without other products and without "none", every
variant has a share of 100 %: then they are ranked by their mean utility.

Saved scenarios live in item.config["sim"]: {"scenarios": [...], "prev": the list before the last change}. They do
not change the questionnaire: the results cache, the HB key and the design lock leave this key out (without_sim),
saving is allowed in every state of the survey, and duplicating and survey files take it along (clean). Undo: the
simulator's own "Rückgängig" (prev), in every state; not the arrows of the editor (app/undo.py): they put the whole
question back, normalised, which could change the key of its HB estimate.
"""
import contextvars
from decimal import Decimal

import numpy as np

from app import choice_predict, units
from app.lang import de_en
from app.units import number

# the language the shown numbers are written in (Mathe-Audit 2026-10-06, Z-1): the page's (1,299.5 and 33.3 in English,
# 1.299,5 and 33,3 in German); the Word tables set German and convert themselves (word_tables, add_to_docx)
_NUMBERS = contextvars.ContextVar("simulator_numbers", default=None)


def _lang():
    return _NUMBERS.get() or units.page_lang()


def _machine(x):
    """A price for a field or the address: a decimal point, no thousands separators, no exponent and no digits lost
    (format "g" gave "12345.7" for 12345.67, taken over at the next "Berechnen", and "1.23457e+06"; Z-11)."""
    text = format(Decimal(repr(float(x))), "f")
    return text.rstrip("0").rstrip(".") if "." in text else text

KEY = "sim"
PRICE_ROW = 99  # "festhalten" of the summed price of ACBC (no attribute index)
MIN_PRODUCTS, MAX_PRODUCTS = 1, 6
START_PRODUCTS = 2
MAX_VARIANTS = 5000  # combinations of the free attributes "Produkt optimieren" computes
TOP_VARIANTS = 10
RIVALS = 200  # "Produkt optimieren": the best variants among which the bootstrap counts how often each is the best
SMALL_SEGMENT = 30  # sum of memberships or effective persons of a segment below which its shares are uncertain
MIN_STATEMENTS, MAX_STATEMENTS = 2, 10
MAX_SCENARIOS = 20
NAME_LEN = 60
LETTERS = "ABCDEF"

ORME = ("Orme, B. K. (2019). Market simulators for conjoint analysis. In Getting started with conjoint analysis: "
        "Strategies for product design and pricing research (4th ed., pp. 89–105). Research Publishers.")
KOHLI = ("Kohli, R., & Krishnamurti, R. (1987). A heuristic approach to product design. Management Science, 33(12), "
         "1523–1533. https://doi.org/10.1287/mnsc.33.12.1523")
SMITH = ("Smith, J. E., & Winkler, R. L. (2006). The optimizer's curse: Skepticism and postdecision surprise in decision "
         "analysis. Management Science, 52(3), 311–322. https://doi.org/10.1287/mnsc.1050.0451")


# --- storage ----------------------------------------------------------------------------------------------------

RESULT_KEYS = (KEY, "lc")  # saved scenarios and the shown segments (app/segments.py): settings of the results page


def without_sim(config):
    """The config without the saved scenarios and the shown segments: what the questionnaire, its design and its
    results rest on."""
    config = config or {}
    return {k: v for k, v in config.items() if k not in RESULT_KEYS} if any(k in config for k in RESULT_KEYS) else config


def saved(item):
    return list(((item.config or {}).get(KEY) or {}).get("scenarios") or [])


def can_undo(item):
    return ((item.config or {}).get(KEY) or {}).get("prev") is not None


def _store(item, scenarios):
    old = saved(item)
    item.config = {**(item.config or {}), KEY: {"scenarios": scenarios, "prev": old}}


def undo_last(item):
    """Back to the scenarios before the last change (and again forward: the change becomes the one to take back)."""
    box = (item.config or {}).get(KEY) or {}
    if box.get("prev") is None:
        return False
    item.config = {**item.config, KEY: {"scenarios": box["prev"], "prev": box.get("scenarios") or []}}
    return True


def clean(raw):
    """The saved scenarios of a survey file: only the known shape, never an error (an unfit one is left out)."""
    if not isinstance(raw, dict) or not isinstance(raw.get("scenarios"), list):
        return {"scenarios": [], "prev": None}
    out = []
    for s in raw["scenarios"][:MAX_SCENARIOS]:
        if not isinstance(s, dict) or not isinstance(s.get("name"), str) or not s["name"].strip():
            continue
        row = {"name": s["name"].strip()[:NAME_LEN]}
        if isinstance(s.get("statements"), list):
            row["statements"] = [x for x in s["statements"] if isinstance(x, str)][:MAX_STATEMENTS]
        elif isinstance(s.get("products"), list):
            products = []
            for p in s["products"][:MAX_PRODUCTS]:
                if not isinstance(p, dict):
                    continue
                levels = p.get("levels") if isinstance(p.get("levels"), dict) else {}
                price = p.get("price")
                products.append({"levels": {k: v for k, v in levels.items() if isinstance(k, str) and isinstance(v, str)},
                                 "price": price if type(price) in (int, float) else None})
            row.update(products=products, none=s.get("none") is True)
        else:
            continue
        if type(s.get("exp")) in (int, float) and MIN_EXP <= s["exp"] <= MAX_EXP:
            row["exp"] = float(s["exp"])  # the exponent of a CVA scenario
        if type(s.get("segment")) is int and 1 <= s["segment"] <= 6:
            row["segment"] = s["segment"]  # the segment of the latent class solution (app/segments.py)
        out.append(row)
    return {"scenarios": out, "prev": None}


# --- the model of a question -------------------------------------------------------------------------------------

def _hb_result(item, entry):
    """The stored HB result if the page shows it as current, else None."""
    if (((entry or {}).get("hb") or {}).get("state")) != "current":
        return None
    from app import conjoint_hb, maxdiff_hb

    result = maxdiff_hb.load(item)
    if not result:
        return None
    if item.type == "maxdiff":
        return result if result.get("format") == maxdiff_hb.FORMAT else None
    return result if result.get("format") == conjoint_hb.FORMAT and result.get("type") == item.type else None


def _price(item):
    """(index of the price attribute, its level numbers) if the price enters linearly, else (None, None)."""
    if item.type in ("cva", "aca"):  # part-worths per level, no price coefficient
        return None, None
    kind = item.kind
    pa = kind.price_attribute(item)
    numbers = kind.level_numbers(item, pa) if pa is not None else None
    if not numbers or len(set(numbers)) < 2:
        return None, None
    return pa, numbers


def _price_levels(numbers, price):
    """The price levels a price stands for when checking prohibitions: the level(s) with that number, a price between
    two levels both neighbours (a prohibition of either holds for it; audit 2026-10-06, L-1), outside the levels the
    nearest level, the highest above them and the lowest below (audit 2026-10-10, W-5)."""
    if price is None:
        return []
    same = [n for n, x in enumerate(numbers) if abs(x - price) < 1e-9]
    if same:
        return same
    below = [x for x in numbers if x < price]
    above = [x for x in numbers if x > price]
    near = ({max(below)} if below else set()) | ({min(above)} if above else set())
    return [n for n, x in enumerate(numbers) if x in near]


def _summed(item):
    """ACBC with a summed price: the price is no attribute but a row of its own (list price: base plus surcharges)."""
    return item.type == "acbc" and item.kind.price_on(item)


def _list_price(item, levels):
    return float(item.kind.byo_price(item, [lv or 0 for lv in levels]))


def _range(item, model, basis):
    """(lowest, highest) price the simulator computes with, or None without a price: CBC the price levels, ACBC the
    prices shown in the study (stored with the HB result)."""
    pa, numbers = _price(item)
    if pa is not None:
        return min(numbers), max(numbers)
    if _summed(item) and basis == "hb" and model.get("price_range") and (model.get("coding") or {}).get("price_scale"):
        lo, hi = model["price_range"]
        return (float(lo), float(hi)) if hi > lo else None
    return None


def _cbc_model(item, entry):
    """(model, pids, "hb" | "pooled") or (None, None, None). hb: the stored result (shares over its kept draws);
    pooled: the model of the aggregate logit as app/choice_predict.py builds it."""
    pa, numbers = _price(item)
    if item.type in ("cva", "aca"):
        return _cva_model(item, entry)
    result = _hb_result(item, entry)
    if result is not None:
        return result, list(result["pids"]), "hb"
    if item.type == "acbc":  # only on the individual utilities, see above
        return None, None, None
    model = (entry or {}).get("model")
    if model and model.get("attributes"):
        codes = [a["code"] for a in item.config.get("attributes", [])]
        center = sum(numbers) / len(numbers) if numbers else 0.0  # Coding.price_center: the mean of the price levels
        return choice_predict.pooled_model(model, codes, center), [None], "pooled"
    return None, None, None


CVA_LOGIT_RANGE = 6.0  # logit units the range of a rating scale stands for (Festlegung von Askurious: -3 to 3)
MIN_EXP, MAX_EXP = 0.1, 10.0


def _cva_model(item, entry):
    """The part-worths per person of a ratings-based conjoint in logit units ({"util": persons x levels, "spans"}),
    their pids and "persons"; (None, None, None) without any. The range of the scale stands for CVA_LOGIT_RANGE: 0 to
    10 or -4 to 4 of a pair alike."""
    if item.type == "aca":  # ACA part-worths: the prior's units (ranges 1 to 4), with calibration logits; taken as they are
        factor = 1.0
    else:
        lo, hi = item.kind.scale_range(item)
        factor = CVA_LOGIT_RANGE / (hi - lo)
    spans, start = [], 0
    for a, n in enumerate(item.kind.level_counts(item)):
        spans.append((a, start, n))
        start += n
    if item.type == "cva" and (((entry or {}).get("hb") or {}).get("state")) == "current":  # the HB regression (app/cva_hb.py), over its draws
        from app import cva_hb

        stored = cva_hb._stored(item)
        if stored:
            return {"util": cva_hb.draw_parameters(stored) * factor, "spans": spans, "hb": True}, list(stored["pids"]),                 "persons"
    people = (entry or {}).get("person_pw") or []
    if not people:
        return None, None, None
    util = np.array([[x * factor for pw in row for x in pw] for _, row in people], dtype=float)
    return {"util": util, "spans": spans, "hb": False}, [pid for pid, _ in people], "persons"


def _maxdiff_model(item, entry):
    """({statement index: column}, utilities draws x persons x columns, pids, "hb" | "pooled") or None; the pooled
    model is one draw of one person."""
    result = _hb_result(item, entry)
    if result is not None:
        from app import maxdiff_hb

        return {int(i): c for c, i in enumerate(result["est"])}, maxdiff_hb.draw_utilities(result), \
            list(result["pids"]), "hb"
    mnl = (entry or {}).get("mnl")
    if mnl and not mnl.get("disconnected") and mnl.get("rows"):
        rows = mnl["rows"]
        return {r["n"] - 1: c for c, r in enumerate(rows)}, np.array([[[float(r["u"]) for r in rows]]]), [None], "pooled"
    return None


def _weights(pids, weighting):
    """Weights per row (persons without one: 0), or None without weighting."""
    if not weighting:
        return None
    by_pid = weighting.get("pids") or {}
    return [float(by_pid.get(pid) or 0.0) for pid in pids]


# --- reading a scenario -----------------------------------------------------------------------------------------

def _int(value):
    try:
        return int(str(value))
    except (TypeError, ValueError):
        return None


def _float(text):
    """A price from the address or a saved scenario: as the number field sends it ("1299.5"); a decimal comma too,
    and with both separators the last one is the decimal mark ("1.299,99", "1,299.99"; N-5: the latter was 1,29999)."""
    if text is None:
        return None
    text = str(text).strip().replace(" ", "")
    if "," in text and "." in text:
        dec = "," if text.rfind(",") > text.rfind(".") else "."
        text = text.replace("." if dec == "," else ",", "").replace(dec, ".")
    elif "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        x = float(text)
    except ValueError:
        return None
    return x if np.isfinite(x) else None


def _cbc_default(item, counts, pa, numbers, j):
    levels = [None if a == pa else j % n if n else None for a, n in enumerate(counts)]
    price = numbers[j % len(numbers)] if pa is not None else _list_price(item, levels) if _summed(item) else None
    return {"levels": levels, "price": price, "price_text": None}


def _cbc_from_saved(item, scenario, counts, pa, numbers):
    attrs = item.config.get("attributes", [])
    products = []
    for j, p in enumerate((scenario.get("products") or [])[:MAX_PRODUCTS]):
        product = _cbc_default(item, counts, pa, numbers, j)
        for a, attr in enumerate(attrs):
            if a == pa:
                continue
            ids = [lv["id"] for lv in attr.get("levels", [])]
            want = (p.get("levels") or {}).get(attr["id"])
            if want in ids:
                product["levels"][a] = ids.index(want)
        if (pa is not None or _summed(item)) and type(p.get("price")) in (int, float):
            product["price"] = float(p["price"])
        products.append(product)
    while len(products) < MIN_PRODUCTS:
        products.append(_cbc_default(item, counts, pa, numbers, len(products)))
    return {"products": products, "none": bool(scenario.get("none")), "opt": None}


def _default_fixed(item, pa):
    """Held at the start: the price of CBC (a free price ends at the lowest); ACBC none (a free price is the list
    price of each variant)."""
    return {pa} if pa is not None else set()


def _cbc_from_args(item, args, counts, pa, numbers):
    n = min(max(_int(args.get("sim_n")) or START_PRODUCTS, MIN_PRODUCTS), MAX_PRODUCTS)
    products = []
    for j in range(n):
        product = _cbc_default(item, counts, pa, numbers, j)
        for a, count in enumerate(counts):
            if a == pa:
                continue
            lv = _int(args.get(f"sim_{j}_{a}"))
            if lv is not None and 0 <= lv < count:
                product["levels"][a] = lv
        if (pa is not None or _summed(item)) and f"sim_{j}_price" in args:
            product["price_text"] = str(args.get(f"sim_{j}_price"))[:40]
            product["price"] = _float(product["price_text"])
            if pa is None and not product["price_text"].strip():  # ACBC, emptied: the list price of the levels
                product["price"], product["price_text"] = _list_price(item, product["levels"]), None
        products.append(product)
    act = str(args.get("sim_act") or "")
    # "Produkt optimieren": which product, which attributes are held (default: the price, a free price would end at
    # the lowest), on or off. The choice is read also while it is off, so ticking before "Varianten durchrechnen"
    # stays ticked (Ralf, 2026-10-05: the tick vanished at once)
    target = _int(args.get("sim_opt_p"))
    values = args.getlist("sim_fix") if hasattr(args, "getlist") else args.get("sim_fix") or []
    fixed = ({a for a in (_int(v) for v in values) if a is not None and (0 <= a < len(counts) or a == PRICE_ROW)}
             if args.get("sim_fixset") == "1" else _default_fixed(item, pa))
    choice = {"product": target if target is not None and 0 <= target < len(products) else 0, "fixed": fixed}
    opt = choice if (args.get("sim_opt") == "1" or act == "opt") and act != "optoff" else None
    if act == "add" and len(products) < MAX_PRODUCTS:
        products.append(_cbc_default(item, counts, pa, numbers, len(products)))
    elif act.startswith("drop") and len(products) > MIN_PRODUCTS:
        k = _int(act[4:])
        if k is not None and 0 <= k < len(products):
            products.pop(k)
            choice["product"] = (choice["product"] - 1 if choice["product"] > k else 0 if choice["product"] == k
                                 else choice["product"])
    elif act.startswith("take") and opt is not None:
        # a variant of "Produkt optimieren" becomes the product: "take<a>:<level>,...,p:<price>"
        for part in act[4:].split(","):
            key, _, value = part.partition(":")
            j = opt["product"]
            if key == "p" and (pa is not None or _summed(item)) and _float(value) is not None:
                products[j]["price"], products[j]["price_text"] = _float(value), None
            elif (_int(key) is not None and 0 <= _int(key) < len(counts) and _int(key) != pa
                  and _int(value) is not None and 0 <= _int(value) < counts[_int(key)]):
                products[j]["levels"][_int(key)] = _int(value)
    return {"products": products, "none": args.get("sim_none") == "1", "opt": opt, "choice": choice}


def _state(item, args, own, read_args, from_saved, default):
    scenarios = saved(item)
    if own:
        k = _int(args.get("sim_load"))
        if k is not None and 0 <= k < len(scenarios):
            return from_saved(scenarios[k]), k
        return read_args(), None
    if scenarios:
        return from_saved(scenarios[0]), 0
    return default(), None


def _exp(value):
    """The exponent of a CVA scenario: a number from MIN_EXP to MAX_EXP, else 1."""
    x = value if type(value) in (int, float) else _float(str(value or ""))
    return float(x) if x is not None and MIN_EXP <= x <= MAX_EXP else 1.0


def cbc_state(item, args, own=True):
    """(scenario, index of the saved scenario shown or None) of a CBC question; a CVA question also with "exp", the
    exponent of the logit rule."""
    counts = item.kind.level_counts(item)
    pa, numbers = _price(item)
    state, k = _state(item, args, own, lambda: _cbc_from_args(item, args, counts, pa, numbers),
                      lambda s: {**_cbc_from_saved(item, s, counts, pa, numbers), "exp": _exp(s.get("exp"))},
                      lambda: {"products": [_cbc_default(item, counts, pa, numbers, j) for j in range(START_PRODUCTS)],
                               "none": False, "opt": None})
    if item.type in ("cva", "aca") and "exp" not in state:
        state["exp"] = _exp(args.get("sim_exp")) if own else 1.0
    return state, k


def maxdiff_state(item, args, own=True):
    n = len(item.config.get("statements", []))
    ids = [s["id"] for s in item.config.get("statements", [])]

    def read_args():
        values = args.getlist("sim_s") if hasattr(args, "getlist") else args.get("sim_s") or []
        return {"statements": sorted({i for i in (_int(v) for v in values) if i is not None and 0 <= i < n})}

    def from_saved(s):
        return {"statements": sorted(ids.index(x) for x in set(s.get("statements") or []) if x in ids)}

    return _state(item, args, own, read_args, from_saved, lambda: {"statements": list(range(min(n, 5)))})


# --- shares -----------------------------------------------------------------------------------------------------

def _pct(x):
    return None if x is None else number(100 * float(x), 1, _lang())


def _per_person(rule, U):
    """A rule of app/choice_predict.py (shares, first_choice) applied to every row of U."""
    return np.array([rule(list(u)) for u in np.asarray(U, dtype=float)])


def _rows(labels, P, weights, F=None, none_last=False):
    """Rows per concept: share of preference (and first choice, F: per person, HB only), sorted by share of
    preference, "none" last."""
    draws = choice_predict.DRAWS if len(P) > 1 else 0
    sop, sop_lo, sop_hi = choice_predict.mean_shares(P, weights, draws)
    if sop is None:
        return None
    fc = fc_lo = fc_hi = None
    if F is not None:
        fc, fc_lo, fc_hi = choice_predict.mean_shares(F, weights, draws)
    rows = []
    for j, label in enumerate(labels):
        rows.append({"label": label, "none": none_last and j == len(labels) - 1, "share": float(sop[j]),
                     "sop": _pct(sop[j]), "sop_lo": _pct(None if sop_lo is None else sop_lo[j]),
                     "sop_hi": _pct(None if sop_hi is None else sop_hi[j]),
                     "fc": _pct(None if fc is None else fc[j]), "fc_lo": _pct(None if fc_lo is None else fc_lo[j]),
                     "fc_hi": _pct(None if fc_hi is None else fc_hi[j]), "width": number(100 * float(sop[j]), 1, "en")})
    rows.sort(key=lambda r: (r["none"], -r["share"]))
    return rows


def _money(x):
    return number(x, 0 if float(x).is_integer() else 2, _lang())


def _points(x):
    """A difference of shares in percentage points, signed, one decimal (no "−0,0")."""
    x = 100 * float(x)
    return number(0.0 if abs(x) < 0.05 else x, 1, _lang(), signed=True)


def _notes(entry, basis):
    """What the model says about the trust in its shares (W-4): with HB its warnings and the share of persons whose
    price coefficient is positive; with the pooled logit a price coefficient that is not clearly negative and the
    attributes whose part-worths only the prior (ridge) bounds."""
    notes = []
    if basis == "persons":
        low = (entry or {}).get("low")
        if low:
            notes.append(de_en("{n} Personen haben ein R² unter {lim}: Ihre Teilnutzen folgen den Bewertungen nur lose, so "
                               "gehen sie in die Anteile ein.",
                               "{n} people have an R² under {lim}: their part-worths follow their ratings only loosely, "
                               "and that is how they enter the shares.", n=low,
                               lim=number((entry or {}).get("r2_low", 0.3), 1, _lang())))
        return notes
    if basis == "hb":
        res = (((entry or {}).get("hb") or {}).get("result")) or {}
        notes += list(res.get("warnings") or [])
        price = res.get("price")
        if price and price.get("negative") is not None and price["negative"] < 1:
            notes.append(de_en("Bei {p} % der Personen ist der Preiskoeffizient positiv: Für sie macht ein höherer "
                               "Preis ein Produkt attraktiver, und so gehen sie in die Anteile ein.",
                               "For {p}% of the people the price coefficient is positive: for them a higher price "
                               "makes a product more attractive, and that is how they enter the shares.",
                               p=number(100 * (1 - price['negative']), 1, _lang())))
        return notes
    model = (entry or {}).get("model") or {}
    beta = (model.get("price") or {}).get("beta")
    if beta is not None and not model.get("wtp_ok"):
        notes.append(de_en("Der Preiskoeffizient des aggregierten Modells ist nicht deutlich negativ: Wie sich die "
                           "Anteile mit dem Preis ändern, ist nicht belastbar.",
                           "The price coefficient of the aggregate model is not clearly negative: how the shares "
                           "change with the price is not robust."))
    loose = [row["label"] for row in model.get("attributes") or [] if row.get("loose")]
    if loose:
        notes.append(de_en("Die Teilnutzen von {which} legen nicht die Daten fest, nur die Ridge-Grenze des Modells "
                           "(nie oder immer gewählt oder nicht bestimmbar): Anteile von Produkten, die sich darin "
                           "unterscheiden, nur als Richtung lesen.",
                           "The part-worths of {which} are not set by the data, only by the ridge bound of the model "
                           "(never or always chosen or not determinable): read shares of products that differ in them "
                           "only as a direction.", which=", ".join(loose)))
    if model.get("converged") is False:
        notes.append(de_en("Das aggregierte Modell ist nicht konvergiert; die Anteile sind nicht verlässlich.",
                           "The aggregate model has not converged; the shares are not reliable."))
    return notes


def _none_offered(item):
    # dual response estimates the constant of "none" from the follow-up question (app/cbc_analysis.py)
    return item.type == "conjoint" and (bool(item.config.get("none")) or item.kind.dual(item))


def _similar(products, pa, span):
    """Pairs of products ("A und B") that are the same in every attribute or differ in one only (or only in the
    price): the logit rule, also per person and draw, gives them together too much share of preference (IIA; audit
    2026-10-05, M-2)."""
    pairs = []
    for i in range(len(products)):
        for j in range(i + 1, len(products)):
            p, q = products[i], products[j]
            differ = sum(1 for a, (x, y) in enumerate(zip(p["levels"], q["levels"])) if a != pa and x != y)
            if span is not None and p["price"] is not None and q["price"] is not None and abs(p["price"] - q["price"]) > 1e-9:
                differ += 1
            if differ <= 1:
                pairs.append(de_en("{a} und {b}", "{a} and {b}", a=LETTERS[i], b=LETTERS[j]))
    return pairs


def cbc_compute(item, entry, state, lang, weighting=None):
    """Shares of a CBC scenario: {"rows", "basis", "persons", "error", "notes"}."""
    model, pids, basis = _cbc_model(item, entry)
    if model is None:
        return {"rows": None, "basis": None, "persons": 0, "error": None}
    attrs = item.kind.attributes(item, lang)
    pa, _ = _price(item)
    span = _range(item, model, basis)
    for j, p in enumerate(state["products"]):
        if span is not None:
            if p["price"] is None:
                return {"rows": None, "basis": basis, "persons": 0,
                        "error": de_en("Produkt {p}: Bitte einen Preis als Zahl eingeben.",
                                       "Product {p}: please enter a price as a number.", p=LETTERS[j])}
            if not span[0] - 1e-9 <= p["price"] <= span[1] + 1e-9:
                return {"rows": None, "basis": basis, "persons": 0,
                        "error": de_en("Produkt {p}: Der Preis {price} liegt außerhalb der Preisspanne der Befragung "
                                       "({lo} bis {hi}). Außerhalb der gezeigten Preise rechnet der Simulator nicht.",
                                       "Product {p}: the price {price} lies outside the price range of the survey "
                                       "({lo} to {hi}). The simulator does not compute outside the prices shown.",
                                       p=LETTERS[j], price=_money(p['price']), lo=_money(span[0]), hi=_money(span[1]))}
    with_none = state["none"] and _none_offered(item) and model.get("asc") is not None
    profiles = [p["levels"] for p in state["products"]]
    prices = [p["price"] if span is not None else None for p in state["products"]]
    if basis in ("hb", "persons"):
        params, spans, center = _params(model, basis, None)
        U = choice_predict.draw_utilities(params, spans, profiles, prices, with_none, center)
        P, F = choice_predict.draw_shares(U)
        if basis == "persons":  # the logit rule with the exponent of the page; first choice does not depend on it
            P = choice_predict.draw_shares(U * state.get("exp", 1.0))[0]
    else:
        P = _per_person(choice_predict.shares, [choice_predict.task_utilities(model, profiles, prices, with_none)])
        F = None
    weights = _weights(pids, weighting) if basis in ("hb", "persons") else None
    if weights is not None and not sum(weights):
        return {"rows": None, "basis": basis, "persons": 0, "error": _no_weighted(basis)}
    labels = []
    for j, p in enumerate(state["products"]):
        parts = [attrs[a]["level_labels"][lv] for a, lv in enumerate(p["levels"]) if lv is not None and a != pa]
        if span is not None:
            parts.append(_money(p["price"]))
        # the letter for the Word table, which names the product in its own language (add_to_docx)
        labels.append({"name": de_en("Produkt {p}", "Product {p}", p=LETTERS[j]), "letter": LETTERS[j],
                       "detail": " · ".join(x for x in parts if x)})
    if with_none:
        none_label = item.t(lang).get("none_label") or "Keine davon"
        labels.append({"name": none_label, "detail": ""})
    rows = _rows(labels, P, weights, F, none_last=with_none)
    return {"rows": rows, "basis": basis, "error": None, "notes": _notes(entry, basis),
            "similar": _similar(state["products"], pa, span), "none_shown": with_none,
            "persons": int(np.count_nonzero(weights)) if weights is not None else len(pids),
            # for the shares per segment (_segment_table): per person and product, the labels in the order of P
            "_P": P, "_pids": pids, "_labels": labels, "_task": (profiles, prices, with_none)}


def _no_weighted(basis):
    if basis == "persons":
        return de_en("Keine gewichtete Person hat eigene Teilnutzen: ohne Gewichtung ansehen.",
                     "No weighted person has part-worths of their own: view without weighting.")
    return de_en("Keine gewichtete Person hat ein HB-Ergebnis: ohne Gewichtung ansehen.",
                 "No weighted person has an HB result: view without weighting.")


def _params(model, basis, counts):
    """(parameters as conjoint_hb.draw_parameters gives them: util draws x persons x levels, price, asc; spans;
    price centre) of either model; the pooled logit is one draw of one person."""
    if basis == "hb":
        from app import conjoint_hb

        return conjoint_hb.draw_parameters(model), model["spans"], (model.get("coding") or {}).get("price_center") or 0.0
    if basis == "persons":  # CVA: the draws of the HB regression, or one draw per person of the regression per person
        util = model["util"] if model["util"].ndim == 3 else model["util"][None, :, :]
        return {"util": util, "price": None, "asc": None}, model["spans"], 0.0
    spans, values, start = [], [], 0
    for a, values_a in sorted(model["pw"].items()):
        spans.append((a, start, len(values_a)))
        values += [float(v) for v in values_a]
        start += len(values_a)
    one = lambda x: None if x is None else np.array([[float(x)]])
    return ({"util": np.array(values, dtype=float).reshape(1, 1, -1), "price": one(model.get("price")),
             "asc": one(model.get("asc"))}, spans, model.get("center", 0.0))


def cbc_optimize(item, entry, state, lang, weighting=None):
    """"Produkt optimieren": the variants of one product (held attributes at its levels, every combination of the
    others) against the rest of the scenario. {"rows" (the best TOP_VARIANTS), "count", "criterion" ("sop": share of
    preference, "utility": alone, the mean utility), "free" (attribute labels), "error", "product"}."""
    import itertools

    opt = state.get("opt")
    model, pids, basis = _cbc_model(item, entry)
    if opt is None or model is None:
        return None
    attrs = item.kind.attributes(item, lang)
    counts = item.kind.level_counts(item)
    pa, numbers = _price(item)
    span = _range(item, model, basis)
    # ACBC: a free summed price follows the levels (list price of the variant), a held one stays as typed
    listed = pa is None and span is not None and PRICE_ROW not in opt["fixed"]
    j = opt["product"]
    product = state["products"][j]
    out = {"product": LETTERS[j], "rows": None, "count": 0, "error": None, "criterion": "sop", "free": [],
           "listed": listed}
    free = [a for a, n in enumerate(counts) if a not in opt["fixed"] and (n or a == pa)]
    if not free:
        out["error"] = de_en("Alle Eigenschaften sind festgehalten: mindestens eine freigeben.",
                             "All properties are fixed: release at least one.")
        return out
    choices = [sorted(set(numbers)) if a == pa else list(range(counts[a])) for a in free]
    total = int(np.prod([len(c) for c in choices]))
    out["count"] = total
    out["free"] = [attrs[a]["label"] for a in free] + ([de_en("Preis", "Price")] if listed else [])
    if total > MAX_VARIANTS:
        out["error"] = de_en("{n} Kombinationen sind zu viele (höchstens {m}): bitte mehr Eigenschaften festhalten.",
                             "{n} combinations are too many (at most {m}): please fix more properties.",
                             n=number(total, 0, _lang()), m=number(MAX_VARIANTS, 0, _lang()))
        return out
    held_price = span is not None and not listed and pa not in free
    if held_price and (product["price"] is None or not span[0] - 1e-9 <= product["price"] <= span[1] + 1e-9):
        out["error"] = de_en("Produkt {p}: Bitte einen Preis innerhalb der Preisspanne der Befragung eingeben.",
                             "Product {p}: please enter a price within the price range of the survey.", p=LETTERS[j])
        return out
    params, spans, center = _params(model, basis, counts)
    if basis == "persons":  # the logit rule with the exponent of the page
        params = {**params, "util": params["util"] * state.get("exp", 1.0)}
    util, price, asc = params["util"], params.get("price"), params.get("asc")
    offset = {a: start for a, start, _ in spans}
    # the utility of the held part of the product, then of the others of the scenario (and "none")
    base = np.zeros(util.shape[:2])
    for a, lv in enumerate(product["levels"]):
        if a not in free and a != pa and lv is not None and a in offset:
            base = base + util[:, :, offset[a] + lv]
    if held_price and price is not None:
        base = base + price * (product["price"] - center)
    others = [p for k, p in enumerate(state["products"]) if k != j]
    if any(span is not None and (p["price"] is None or not span[0] - 1e-9 <= p["price"] <= span[1] + 1e-9)
           for p in others):
        out["error"] = de_en("Bitte für alle Produkte einen Preis innerhalb der Preisspanne der Befragung eingeben.",
                             "Please enter a price within the price range of the survey for all products.")
        return out
    with_none = state["none"] and _none_offered(item) and asc is not None
    rival = choice_predict.draw_utilities(params, spans, [p["levels"] for p in others],
                                          [p["price"] if span is not None else None for p in others], with_none, center) \
        if (others or with_none) else None
    if rival is None:
        out["criterion"] = "utility"
    else:
        top = rival.max(axis=2)
        log_rival = top + np.log(np.exp(rival - top[:, :, None]).sum(axis=2))
        ties = (rival >= top[:, :, None] - 1e-12).sum(axis=2)
    variants = list(itertools.product(*choices))
    # forbidden pairs of levels (prohibitions of the design) are no product
    def profile(v):
        levels = list(product["levels"])
        for a, x in zip(free, v):
            if a != pa:
                levels[a] = x
        return levels

    out["forbidden"] = 0
    if item.kind.prohibitions(item):
        # the price takes part as its level: a prohibition with a price level holds for that price (audit 2026-10-06,
        # L-1; before, the price stood as None in the profile and such a prohibition never held)
        held = _price_levels(numbers, product["price"]) if pa is not None and pa not in free else None
        out["between"] = held is not None and not any(abs(x - product["price"]) < 1e-9 for x in numbers)
        def allowed(v):
            levels = profile(v)
            if pa is None:
                return item.kind.profile_allowed(item, levels)
            at = held if held is not None else _price_levels(numbers, v[free.index(pa)])
            return all(item.kind.profile_allowed(item, levels[:pa] + [lv] + levels[pa + 1:]) for lv in at)

        kept = [v for v in variants if allowed(v)]
        out["forbidden"] = len(variants) - len(kept)
        variants = kept
        if not variants:
            out["error"] = de_en("Jede Kombination enthält ein verbotenes Paar von Stufen.",
                                 "Every combination contains a prohibited pair of levels.")
            return out
    list_prices = np.array([_list_price(item, profile(v)) for v in variants]) if listed else None
    if listed:
        # a list price outside the prices shown is no price the model knows: the variant is left out like a
        # forbidden pair, not computed with the nearest shown price (audit 2026-10-05, M-4)
        inside = [c for c, x in enumerate(list_prices) if span[0] - 1e-9 <= x <= span[1] + 1e-9]
        out["outside"] = len(variants) - len(inside)
        if not inside:
            out["error"] = de_en("Jede Variante liegt mit ihrem Listenpreis außerhalb der gezeigten Preise ({lo} bis "
                                 "{hi}); dort rechnet der Simulator nicht. Den Preis festhalten oder Eigenschaften mit "
                                 "Aufpreisen festhalten.",
                                 "Every variant lies outside the prices shown ({lo} to {hi}) with its list price; the "
                                 "simulator does not compute there. Fix the price or fix properties with surcharges.",
                                 lo=_money(span[0]), hi=_money(span[1]))
            return out
        variants = [variants[c] for c in inside]
        list_prices = list_prices[inside]
    # the variants computed, after the prohibitions and the list prices outside (audit 2026-10-06, L-4)
    out["count"] = len(variants)
    weights = _weights(pids, weighting) if basis in ("hb", "persons") else None
    if weights is not None and not sum(weights):
        out["error"] = _no_weighted(basis)
        return out
    w = np.ones(len(pids)) if weights is None else np.asarray(weights, dtype=float)
    D, N = util.shape[:2]
    chunk = max(1, 2_000_000 // max(1, D * N))
    P_all, F_all, score = [], [], []
    for start in range(0, len(variants), chunk):
        part = variants[start:start + chunk]
        U = np.repeat(base[:, :, None], len(part), axis=2)
        if listed and price is not None:
            x = list_prices[start:start + chunk]
            U = U + price[:, :, None] * (x[None, None, :] - center)
        for f, a in enumerate(free):
            values = np.array([v[f] for v in part])
            if a == pa:
                if price is not None:
                    U = U + price[:, :, None] * (values[None, None, :] - center)
            elif a in offset:
                U = U + util[:, :, offset[a] + values]
        if rival is None:
            P = U.mean(axis=0)                                             # persons x variants: mean utility
            F = None
        else:
            diff = U - log_rival[:, :, None]
            P = (1.0 / (1.0 + np.exp(-np.clip(diff, -700, 700)))).mean(axis=0)
            gap = U - top[:, :, None]
            F = ((gap > 1e-12) + (np.abs(gap) <= 1e-12) / (ties[:, :, None] + 1.0)).mean(axis=0)
        P_all.append(P)
        if F is not None:
            F_all.append(F)
        score.append(w @ P / w.sum())
    P = np.concatenate(P_all, axis=1)
    F = np.concatenate(F_all, axis=1) if F_all else None
    score = np.concatenate(score)
    ranked = sorted(range(len(variants)), key=lambda c: (-score[c], c))
    order = ranked[:TOP_VARIANTS]
    current = tuple(product["price"] if a == pa else product["levels"][a] for a in free)
    draws = choice_predict.DRAWS if len(P) > 1 else 0
    # The best of many variants is chosen on the same sample, so its share is too high (optimizer's curse; Smith &
    # Winkler, 2006) and an interval per variant says nothing about the ranks. From the same bootstrap samples of
    # persons for all variants: how often each is the best, and its distance to rank 1 with a paired interval
    # (audit 2026-10-05, M-1). Among the RIVALS best only: one further down is never the best.
    paired = None
    if out["criterion"] == "sop" and draws:
        rivals = ranked[:RIVALS]
        M = choice_predict.bootstrap_means(P[:, rivals], weights, draws)
        win = M >= M.max(axis=1, keepdims=True) - 1e-12
        paired = {c: ((win[:, i] / win.sum(axis=1)).mean(), M[:, i] - M[:, 0]) for i, c in enumerate(rivals)}
        out["paired"] = True
    rows = []
    for c in order:
        v = variants[c]
        row = {"levels": [_money(x) if a == pa else attrs[a]["level_labels"][x] for a, x in zip(free, v)]
               + ([_money(list_prices[c])] if listed else []),
               "take": ",".join([f"p:{_machine(x)}" if a == pa else f"{a}:{x}" for a, x in zip(free, v)]
                                + ([f"p:{_machine(list_prices[c])}"] if listed else [])),
               "current": all(abs(float(x) - float(y)) < 1e-9 if a == pa else x == y
                              for a, x, y in zip(free, v, current) if y is not None) and None not in current
               and (not listed or (product["price"] is not None and abs(product["price"] - list_prices[c]) < 0.005))}
        if out["criterion"] == "utility":
            row["value"] = number(float(score[c]), 2, _lang())
        else:
            row.update(sop=_pct(score[c]), width=number(100 * float(score[c]), 1, "en"))
            if paired is not None:
                best, gaps = paired[c]
                row["best"] = _pct(best)
                if c != ranked[0]:
                    lo, hi = np.percentile(gaps, [2.5, 97.5])
                    row.update(gap=_points(score[c] - score[ranked[0]]), gap_lo=_points(lo), gap_hi=_points(hi),
                               tie=bool(hi >= 0))
            if F is not None and basis in ("hb", "persons"):  # CVA: the exponent does not change who is first
                row["fc"] = _pct(float(w @ F[:, c] / w.sum()))
        rows.append(row)
    out.update(rows=rows, basis=basis)
    return out


def maxdiff_compute(item, entry, state, lang, weighting=None):
    found = _maxdiff_model(item, entry)
    if found is None:
        return {"rows": None, "basis": None, "persons": 0, "error": None}
    cols, util, pids, basis = found
    chosen = [i for i in state["statements"] if i in cols]
    if not MIN_STATEMENTS <= len(chosen) <= MAX_STATEMENTS:
        return {"rows": None, "basis": basis, "persons": 0,
                "error": de_en("Bitte {lo} bis {hi} Aussagen ankreuzen.", "Please tick {lo} to {hi} statements.",
                               lo=MIN_STATEMENTS, hi=MAX_STATEMENTS)}
    weights = _weights(pids, weighting) if basis == "hb" else None
    if weights is not None and not sum(weights):
        return {"rows": None, "basis": basis, "persons": 0,
                "error": de_en("Keine gewichtete Person hat ein HB-Ergebnis: ohne Gewichtung ansehen.",
                               "No weighted person has an HB result: view without weighting.")}
    texts = item.kind.labels(item, lang)
    P, F = choice_predict.draw_shares(util[:, :, [cols[i] for i in chosen]])
    labels = [{"name": f"{i + 1}. {texts[i]}", "detail": ""} for i in chosen]
    rows = _rows(labels, P, weights, F if basis == "hb" else None)
    extra = {"_P": P, "_pids": pids, "_labels": labels, "_chosen": chosen}
    if basis == "hb":
        notes = list(((((entry or {}).get("hb") or {}).get("result")) or {}).get("warnings") or [])
    else:
        notes = list(((entry or {}).get("mnl") or {}).get("warnings") or [])
    return {"rows": rows, "basis": basis, "error": None, "notes": notes,
            "persons": int(np.count_nonzero(weights)) if weights is not None else util.shape[1], **extra}


# --- the page ---------------------------------------------------------------------------------------------------

def params(item, state):
    """The scenario as address parameters [(name, value)] (hidden fields of the save form)."""
    out = [("sim", str(item.id))]
    if item.type == "maxdiff":
        return out + [("sim_s", str(i)) for i in state["statements"]]
    pa, _ = _price(item)
    out.append(("sim_n", str(len(state["products"]))))
    for j, p in enumerate(state["products"]):
        out += [(f"sim_{j}_{a}", str(lv)) for a, lv in enumerate(p["levels"]) if lv is not None and a != pa]
        if pa is not None or _summed(item):
            out.append((f"sim_{j}_price", p.get("price_text") if p.get("price_text") is not None
                        else ("" if p["price"] is None else _machine(p["price"]))))
    if state["none"]:
        out.append(("sim_none", "1"))
    if item.type in ("cva", "aca"):
        out.append(("sim_exp", _machine(state.get("exp", 1.0))))
    return out


def _keep(args):
    """The parameters of the page that are not the simulator's (pool, weighting, grouping)."""
    items = args.items(multi=True) if hasattr(args, "items") else []
    return [(k, v) for k, v in items if not k.startswith("sim") and k != "go"]


def _segments_of(item, entry):
    """(stored latent class result, shown K, view of the segments) when the question has a current one with at
    least two segments (app/segments.py), else None."""
    view_ = (entry or {}).get("lc") or {}
    res = view_.get("result")
    if not res or res.get("K", 1) < 2 or view_.get("state") != "current":
        return None
    from app import segments

    stored = segments.load(item)
    if not stored:
        return None
    return stored, segments.shown_k(item, stored), res


def _segment_weighting(stored, K, seg, weighting):
    """A weighting (as app/weights.py gives it to the simulator: {"pids": {pid: weight}}) that counts every person
    with the probability to belong to segment seg (1..K), times the weight of the page if there is one (concept
    latente-klassen-konzept.md, LK12: the HB persons stay as they are, the segment only weights them)."""
    from app import segments

    tau = segments.solution(stored, K)["tau"]
    base = (weighting or {}).get("pids")
    return {"pids": {pid: (float(base.get(pid) or 0.0) if base is not None else 1.0) * float(t[seg - 1])
                     for pid, t in zip(stored["pids"], tau)}}


def _class_shares(item, res, out, center):
    """Shares per segment from the latent class parameters themselves (without HB): segments x products; CBC from
    the part-worths, price and constant of each segment, MaxDiff from its utilities."""
    K = res["K"]
    if item.type == "maxdiff":
        util = {st["index"]: [v["util"] for v in st["values"]] for st in res.get("statements") or []}
        U = np.array([[util[i][c] if i in util else 0.0 for i in out.get("_chosen") or []] for c in range(K)])
    else:
        profiles, prices, with_none = out["_task"]
        U = []
        for c in range(K):
            model = {"pw": {r["index"]: [lv["values"][c]["beta"] for lv in r["levels"]] for r in res["attributes"]},
                     "price": res["price"]["beta"][c] if res.get("price") else None, "center": center,
                     "asc": res["asc"][c]["beta"] if res.get("asc") and with_none else None}
            U.append(choice_predict.task_utilities(model, profiles, prices, with_none))
        U = np.array(U, dtype=float)
    P = np.exp(U - U.max(axis=1, keepdims=True))
    return P / P.sum(axis=1, keepdims=True)


def _segment_table(item, out, found, weighting):
    """Share of preference of every product in every segment: with HB the persons weighted with their membership
    (and the weight of the page), without HB from the parameters of the segments."""
    stored, K, res = found
    labels = out.get("_labels") or []
    if not labels or out.get("_P") is None:
        return None
    if out["basis"] in ("hb", "persons"):
        from app import segments

        tau = {pid: t for pid, t in zip(stored["pids"], segments.solution(stored, K)["tau"])}
        base = (weighting or {}).get("pids")
        T = np.array([tau.get(pid, [0.0] * K) for pid in out["_pids"]], dtype=float)
        w0 = np.array([float(base.get(pid) or 0.0) if base is not None else 1.0 for pid in out["_pids"]])
        S = np.array([np.average(out["_P"], axis=0, weights=T[:, c] * w0) if (T[:, c] * w0).sum() > 0
                      else np.full(len(labels), np.nan) for c in range(K)])
    else:
        center = 0.0
        if item.type != "maxdiff" and stored.get("coding"):
            from app import conjoint_hb

            center = conjoint_hb._coding(stored["coding"]).price_center or 0.0
        S = _class_shares(item, res, out, center)
    order = {id(row["label"]): n for n, row in enumerate(out.get("rows") or [])}
    rows = [{"label": label, "shares": [None if np.isnan(S[c, j]) else _pct(S[c, j]) for c in range(K)]}
            for j, label in enumerate(labels)]
    rows.sort(key=lambda r: order.get(id(r["label"]), 0))  # the order of the bars above
    # the size of every segment: the sum of the membership probabilities and the effective number of persons (M-5)
    from app import segments

    T = np.array(segments.solution(stored, K)["tau"], dtype=float)
    if out["basis"] in ("hb", "persons"):
        base = (weighting or {}).get("pids")
        index = {pid: n for n, pid in enumerate(stored["pids"])}
        mine = [index[pid] for pid in out["_pids"] if pid in index]
        w0 = np.array([float(base.get(stored["pids"][n]) or 0.0) if base is not None else 1.0 for n in mine])
        T = T[mine]
    else:
        w0 = np.ones(len(T))
    sizes = [{"sum": number(float(T[:, c].sum()), 1, _lang()), "effective": number(_effective(T[:, c] * w0), 1, _lang()),
              "small": _small(float(T[:, c][w0 > 0].sum()), _effective(T[:, c] * w0))} for c in range(K)]
    out_ = {"names": [s["name"] for s in res["segments"]], "rows": rows, "basis": out["basis"], "sizes": sizes}
    if out["basis"] not in ("hb", "persons"):
        # without HB the bars above come from the pooled logit, which ignores the segments: the mixture of the
        # segments (sum_c pi_c P_cj) differs from it where segments exist (audit 2026-10-05, S-4)
        pi = np.array([s["share"] for s in res["segments"]], dtype=float)
        mix = pi @ S
        by_label = {id(label): j for j, label in enumerate(labels)}
        for r in rows:
            r["mixture"] = _pct(mix[by_label[id(r["label"])]])
    return out_


def _small(total, effective):
    """A segment is small when the sum of its memberships or the effective number of its weights is under
    SMALL_SEGMENT: with memberships under 1 the effective number is never smaller than their sum, so alone it misses a
    blurred segment of few persons (audit 2026-10-06, L-5)."""
    return min(total, effective) < SMALL_SEGMENT


def _effective(w):
    """Effective number of persons of weights w (Kish): (sum w)^2 / sum w^2."""
    w = np.asarray(w, dtype=float)
    return float(w.sum() ** 2 / (w ** 2).sum()) if (w ** 2).sum() > 0 else 0.0


def _segment_size(stored, K, seg, pids, weighting):
    """The size of segment seg among the persons of the shares: the sum of their membership probabilities and the
    effective number of persons of the weights the shares use (audit 2026-10-05, M-5: nearly every person has a
    probability above 0, so the count of persons says little)."""
    from app import segments

    tau = {pid: t[seg - 1] for pid, t in zip(stored["pids"], segments.solution(stored, K)["tau"])}
    w = _weights(pids, weighting)
    eff = _effective(w)
    total = sum(float(tau.get(pid, 0.0)) for pid in pids)
    counted = total if w is None else sum(float(tau.get(pid, 0.0)) for pid, x in zip(pids, w) if x > 0)
    return {"sum": number(total, 1, _lang()), "effective": number(eff, 1, _lang()), "small": _small(counted, eff)}


def _segment_name(found, segment):
    """The name of a saved scenario's segment in the shown solution, None without one."""
    if not found or not segment or segment > found[1]:
        return None
    return found[2]["segments"][segment - 1]["name"]


def _chosen_segment(item, args, own):
    """The segment of the scenario: the list in the address, else the one a saved scenario was stored with (the one
    loaded, or the first when the page opens), else all persons (0)."""
    if own and args.get("sim_seg") is not None:
        return _int(args.get("sim_seg")) or 0
    scenarios = saved(item)
    k = _int(args.get("sim_load")) if own else (0 if scenarios else None)
    if k is not None and 0 <= k < len(scenarios):
        return scenarios[k].get("segment") or 0
    return 0


def view(item, entry, args, lang, weighting=None, editor=False):
    """Everything the fold of one question shows, or None if the question has no simulator. With segments
    (app/segments.py) the list "Segment" weights the HB persons with their membership (sim_seg), and a table shows
    the shares in every segment."""
    own = str(args.get("sim") or "") == str(item.id)
    lc_found = _segments_of(item, entry)
    seg = _chosen_segment(item, args, own) if lc_found else 0
    seg = seg if lc_found and 0 <= seg <= lc_found[1] else 0
    page_weighting = weighting
    if seg:
        basis = _cbc_model(item, entry)[2] if item.type != "maxdiff" else ((_maxdiff_model(item, entry) or [None] * 4)[3])
        if basis in ("hb", "persons"):  # "persons": the part-worths per person of a CVA (or its HB)
            weighting = _segment_weighting(lc_found[0], lc_found[1], seg, weighting)
        else:
            seg = 0  # without HB there are no persons to weight: the table below shows the segments
    if item.type in ("conjoint", "acbc", "cva", "aca"):
        state, loaded = cbc_state(item, args, own)
        out = cbc_compute(item, entry, state, lang, weighting)
        attrs = item.kind.attributes(item, lang)
        pa, numbers = _price(item)
        model, _, basis = _cbc_model(item, entry)
        span = _range(item, model, basis) if model is not None else None
        out["attributes"] = [{"index": a, "label": attr["label"], "levels": attr["level_labels"], "price": a == pa}
                             for a, attr in enumerate(attrs)]
        if pa is None and span is not None:  # ACBC: the summed price as a row of its own
            out["attributes"].append({"index": PRICE_ROW, "label": item.t(lang).get("price_label") or de_en("Preis", "Price"),
                                      "levels": [], "price": True})
        out["price"] = {"min": span[0], "max": span[1], "min_text": _money(span[0]), "max_text": _money(span[1]),
                        "summed": pa is None} if span is not None else None
        out["products"] = [{**p, "letter": LETTERS[j],
                            "price_value": p.get("price_text") if p.get("price_text") is not None
                            else ("" if p["price"] is None else _machine(p["price"]))}
                           for j, p in enumerate(state["products"])]
        # "None of these" only if the model has a constant for it: with dual response and no answered follow-up
        # question it has none, and a constant of 0 would give "none" an arbitrary share (audit 2026-10-06, W-2)
        out["none_offered"] = _none_offered(item) and model is not None and model.get("asc") is not None
        out["none_on"] = bool(state["none"])
        out["max_products"], out["min_products"] = MAX_PRODUCTS, MIN_PRODUCTS
        opt = state.get("opt")
        out["opt_on"] = opt is not None
        choice = state.get("choice") or {"product": 0, "fixed": _default_fixed(item, pa)}
        out["opt_product"] = choice["product"]
        out["opt_fixed"] = sorted(choice["fixed"])
        out["opt"] = cbc_optimize(item, entry, state, lang, weighting) if opt else None
        out["max_variants"] = MAX_VARIANTS
        out["exp"] = _machine(state.get("exp", 1.0)) if item.type in ("cva", "aca") else None
        out["cva_hb"] = bool(item.type == "cva" and model is not None and model.get("hb"))
    elif item.type == "maxdiff":
        state, loaded = maxdiff_state(item, args, own)
        out = maxdiff_compute(item, entry, state, lang, weighting)
        found = _maxdiff_model(item, entry)
        usable = set(found[0]) if found else set()
        out["statements"] = [{"index": i, "label": label, "on": i in state["statements"], "usable": i in usable}
                             for i, label in enumerate(item.kind.labels(item, lang))]
    else:
        return None
    if out["basis"] is None:
        return None
    out["segments"] = None
    if lc_found:
        out["segments"] = {"selected": seg, "options": [s["name"] for s in lc_found[2]["segments"]],
                           "filter": out["basis"] in ("hb", "persons"),
                           "table": _segment_table(item, out, lc_found, page_weighting) if out.get("rows") else None}
    if out["basis"] == "pooled":  # the note under the pooled model names a running estimation (W74)
        from app import maxdiff_hb

        marker = maxdiff_hb.running(item)
        out["hb_running"] = bool(marker and marker.get("state") == "running")
    if seg and out.get("_pids") is not None:
        out["segment_size"] = _segment_size(lc_found[0], lc_found[1], seg, out["_pids"], weighting)
    for key in [k for k in out if k.startswith("_")]:
        del out[key]  # the per-person arrays are not for the page
    out.update(item_id=item.id, kind=item.type, open=own,
               state_params=params(item, state) + ([("sim_seg", str(seg))] if seg else []), keep=_keep(args),
               weighted=bool(page_weighting) and out["basis"] == "hb", loaded=loaded, editor=editor,
               scenarios=[{"index": k, "name": s["name"], "segment": _segment_name(lc_found, s.get("segment"))}
                          for k, s in enumerate(saved(item))],
               undo=can_undo(item), draws=choice_predict.DRAWS)
    return out


TYPES = ("conjoint", "acbc", "maxdiff", "cva", "aca")


def _part(entry):
    """The part of a results entry the simulator rests on (results.compute: "cbc", "acbc", "cva" or "maxdiff")."""
    return entry.get({"cbc": "cbc", "acbc": "acbc", "cva": "cva", "aca": "aca"}.get(entry.get("kind"), "maxdiff"))


def page(data, args, weighted=False, editor=False, kept=None):
    """{item id: view} for the CBC and MaxDiff questions of the results page (results.compute()). kept: {item id:
    (marker, view)} of the scenarios a page opens with (no scenario of its own in the address), kept with the results
    under their fingerprint (surveys.results); a view whose marker still fits is taken from it, the others are computed
    and put into it (performance audit 2026-10-06, P6: the shares over every HB draw were half the time of a page from
    the store). None: computed on every view."""
    weighting = data.get("weighting") if weighted else None
    out = {}
    for entry in data.get("items") or []:
        item = entry.get("item")
        part = _part(entry)
        if item is None or part is None or item.type not in TYPES:
            continue
        lang = item.survey.default_language
        own = str(args.get("sim") or "") == str(item.id)
        marker = _opening_marker(item, part, lang, weighting) if kept is not None and not own else None
        if marker is not None and item.id in kept and kept[item.id][0] == marker:
            shown = kept[item.id][1]
            shown = None if shown is None else {**shown, "editor": editor, "keep": _keep(args)}
        else:
            shown = view(item, part, args, lang, weighting, editor)
            if marker is not None:
                kept[item.id] = (marker, None if shown is None else {k: v for k, v in shown.items() if k not in ("editor", "keep")})
        if shown is not None:
            out[item.id] = shown
    return out


def _opening_marker(item, part, lang, weighting):
    """What the scenario a page opens with rests on besides the kept results: the question with its saved scenarios
    and chosen segments, the HB and segment estimates as the page shows them (their state and key, the time and size
    of their files), the language and the weighting."""
    import hashlib
    import json
    import os

    from app import maxdiff_hb, segments

    def stamp(path):
        try:
            st = os.stat(path)
        except OSError:
            return None
        return [st.st_mtime_ns, st.st_size]

    hb, lc = part.get("hb") or {}, part.get("lc") or {}
    # the language of the page too: the numbers of the view are written in it (Z-1)
    basis = [item.id, item.config, lang, _lang(), bool(weighting), [hb.get("state"), hb.get("input_key")],
             [lc.get("state"), lc.get("input_key")], stamp(maxdiff_hb._paths(item)[0]), stamp(segments._paths(item)[0])]
    return hashlib.sha256(json.dumps(basis, sort_keys=True, default=str).encode()).hexdigest()


def sources(sims):
    """Source groups for the one fold of the page (results.page_sources)."""
    if not sims:
        return []
    from app import hb as hb_module, intervals
    from app.choice_model import TRAIN

    hb = any(s["basis"] == "hb" and s.get("rows") for s in sims.values())
    persons = any(s["basis"] == "persons" and s.get("rows") for s in sims.values())
    groups = [(de_en("Marktsimulator: Präferenzanteile nach dem Logit-Modell je Person (Share of Preference)",
                     "Market simulator: preference shares by the logit model per person (share of preference)")
               + (de_en(" und nach dem höchsten Nutzen (First Choice)", " and by the highest utility (first choice)")
                  if hb or persons else "")
               + de_en(", Train (2009) Kapitel 3, Orme (2019) Kapitel 10", ", Train (2009) chapter 3, Orme (2019) chapter 10"),
               [TRAIN, ORME])]
    if hb:
        groups.append((de_en("Präferenzanteile je Person als Mittel über die behaltenen Ziehungen ihrer HB-Nutzen statt "
                             "aus den geschrumpften Posterior-Mitteln",
                             "Preference shares per person as the mean over the kept draws of their HB utilities "
                             "instead of from the shrunk posterior means"), [hb_module.ROSSI]))
        groups.append((de_en("95-%-Intervalle der Präferenzanteile: Perzentil-Bootstrap über Personen mit {d} Ziehungen "
                             "(nur die Stichprobenstreuung über Personen)",
                             "95% intervals of the preference shares: percentile bootstrap over people with {d} draws "
                             "(only the sampling variation over people)", d=choice_predict.DRAWS), [intervals.EFRON]))
    if persons:
        groups.append((de_en("Marktsimulator der Bewertungs-Conjoint: Teilnutzen je Person aus der Regression, für die "
                             "Logit-Regel in Logit-Einheiten umgerechnet (Spannweite der Skala = 6 Einheiten, "
                             "Festlegung von Askurious), mal einem einstellbaren Exponenten; 95-%-Intervalle per "
                             "Perzentil-Bootstrap über Personen ({d} Ziehungen)",
                             "Market simulator of the ratings-based conjoint: part-worths per person from the regression, "
                             "converted to logit units for the logit rule (range of the scale = 6 units, "
                             "Askurious default), times an adjustable exponent; 95% intervals by percentile bootstrap "
                             "over people ({d} draws)", d=choice_predict.DRAWS), [ORME, intervals.EFRON]))
    if any(s.get("kind") in ("conjoint", "acbc", "cva", "aca") for s in sims.values()):
        groups.append((de_en("Produkt optimieren (Marktsimulator): alle Kombinationen der freien Eigenschaften "
                             "durchgerechnet (vollständige Suche statt einer Heuristik), geordnet nach dem "
                             "Präferenzanteil",
                             "Optimising a product (market simulator): all combinations of the free properties computed "
                             "(exhaustive search instead of a heuristic), sorted by preference share"), [KOHLI, ORME]))
        if hb:
            groups.append((de_en("Produkt optimieren: Anteil der besten Variante auf derselben Stichprobe gesucht und "
                                 "deshalb eher zu hoch; je Variante aus denselben Bootstrap-Stichproben, wie oft sie die "
                                 "beste ist, und ihr Abstand zu Rang 1 mit gepaartem 95-%-Intervall",
                                 "Optimising a product: share of the best variant searched on the same sample and therefore "
                                 "tending to be too high; per variant from the same bootstrap samples how often it is "
                                 "the best, and its gap to rank 1 with a paired 95% interval"), [SMITH, intervals.EFRON]))
    return groups


# --- saving -----------------------------------------------------------------------------------------------------

def scenario_from_form(item, form):
    """The scenario a save form sends, in the stored shape (ids, not positions)."""
    name = (form.get("sim_name") or "").strip()[:NAME_LEN]
    segment = _int(form.get("sim_seg")) or 0
    extra = {"segment": segment} if 1 <= segment <= 6 else {}  # saved with its segment, loaded with it again
    if item.type == "maxdiff":
        state, _ = maxdiff_state(item, form)
        ids = [s["id"] for s in item.config.get("statements", [])]
        return {"name": name, "statements": [ids[i] for i in state["statements"]][:MAX_STATEMENTS], **extra}
    state, _ = cbc_state(item, form)
    pa, _ = _price(item)
    attrs = item.config.get("attributes", [])
    products = []
    for p in state["products"]:
        levels = {attrs[a]["id"]: attrs[a]["levels"][lv]["id"] for a, lv in enumerate(p["levels"])
                  if lv is not None and a != pa}
        products.append({"levels": levels, "price": p["price"] if pa is not None or _summed(item) else None})
    if item.type in ("cva", "aca") and state.get("exp", 1.0) != 1.0:
        extra["exp"] = state["exp"]
    return {"name": name, "products": products, "none": bool(state["none"] and _none_offered(item)), **extra}


def save(item, form):
    """Store the scenario of the form under its name (the same name replaces it). (index, message or None)."""
    scenario = scenario_from_form(item, form)
    scenarios = saved(item)
    if not scenario["name"]:
        scenario["name"] = de_en("Szenario {n}", "Scenario {n}", n=len(scenarios) + 1)
    names = [s["name"] for s in scenarios]
    if scenario["name"] in names:
        k = names.index(scenario["name"])
        scenarios[k] = scenario
    elif len(scenarios) >= MAX_SCENARIOS:
        return None, de_en("Höchstens {n} Szenarien je Frage: bitte erst eines löschen.",
                           "At most {n} scenarios per question: please delete one first.", n=MAX_SCENARIOS)
    else:
        scenarios.append(scenario)
        k = len(scenarios) - 1
    _store(item, scenarios)
    return k, None


def delete(item, index):
    scenarios = saved(item)
    if index is None or not 0 <= index < len(scenarios):
        return None
    gone = scenarios.pop(index)
    _store(item, scenarios)
    return gone["name"]


# --- Word ---------------------------------------------------------------------------------------------------------

DOCX = {
    "de": {"table": "Tabelle {n}", "title": "Marktsimulator: {question}, Szenario „{name}“", "product": "Produkt",
           "statement": "Aussage", "sop": "Share of Preference", "fc": "First Choice", "note": "Anmerkung.",
           "hb": ("Präferenzanteile in Prozent aus den individuellen Nutzen (Hierarchical Bayes, n = {n}), je Person "
                  "gemittelt über die behaltenen Ziehungen: Share of Preference = Logit-Wahrscheinlichkeit; First Choice = "
                  "Anteil der Ziehungen, in denen das Produkt den höchsten Nutzen hat (Gleichstand geteilt); dann über "
                  "die Personen gemittelt. In eckigen Klammern: 95-%-Intervall (Perzentil-Bootstrap über Personen, "
                  "{draws} Ziehungen; nur die Stichprobenstreuung über Personen)."),
           "pooled": ("Präferenzanteile in Prozent aus dem aggregierten Logit-Modell (n = {n}); es kennt keine "
                      "Unterschiede zwischen Personen, deshalb ohne First Choice und ohne Intervall."),
           "persons": ("Präferenzanteile in Prozent aus den Teilnutzen je Person der Bewertungs-Conjoint (Regression je "
                       "Person, n = {n}), in Logit-Einheiten umgerechnet: Share of Preference = Logit-Wahrscheinlichkeit; "
                       "First Choice = Anteil der Personen, für die das Produkt den höchsten Nutzen hat. In eckigen "
                       "Klammern: 95-%-Intervall (Perzentil-Bootstrap über Personen, {draws} Ziehungen)."),
           "market": " Präferenzanteile sind keine Marktanteile (Bekanntheit, Verfügbarkeit und Wettbewerb fehlen).",
           "segment": ", Segment „{segment}“",
           "weighted": (" Nur Segment „{segment}“: jede Person gewichtet mit ihrer Wahrscheinlichkeit, zum Segment zu "
                        "gehören (Latent-Class-Logit); Summe der Zugehörigkeiten = {size}, effektive Fallzahl des gewichteten "
                        "Mittels {effective}."),
           "gone": " Segment {segment} ist nicht mehr vorhanden: alle Personen.",
           "no_hb": " Segment „{segment}“ nur mit HB-Nutzen: ohne sie alle Personen."},
    "en": {"table": "Table {n}", "title": "Market Simulator: {question}, Scenario “{name}”", "product": "Product",
           "statement": "Statement", "sop": "Share of preference", "fc": "First choice", "note": "Note.",
           "hb": ("Shares of preference in percent from individual utilities (hierarchical Bayes, n = {n}), per person "
                  "averaged over the kept posterior draws: share of preference = logit probability; first choice = share "
                  "of draws in which the product has the highest utility (ties split); then averaged over people. In "
                  "square brackets: 95% interval (percentile bootstrap over people, {draws} resamples; sampling "
                  "variation over people only)."),
           "pooled": ("Shares of preference in percent from the aggregate logit model (n = {n}); it assumes no "
                      "differences between people, hence no first choice and no interval."),
           "persons": ("Shares of preference in percent from the part-worths per person of the ratings-based conjoint "
                       "(regression per person, n = {n}), converted to logit units: share of preference = logit "
                       "probability; first choice = share of people for whom the product has the highest utility. In "
                       "square brackets: 95% interval (percentile bootstrap over people, {draws} resamples)."),
           "market": " Shares of preference are not market shares (awareness, distribution, and competition are not modelled).",
           "segment": ", Segment “{segment}”",
           "weighted": (" Segment “{segment}” only: each person weighted by their probability of belonging to the segment "
                        "(latent class logit); sum of memberships = {size}, effective sample size of the weighted mean "
                        "{effective}."),
           "gone": " Segment {segment} no longer exists: all people.",
           "no_hb": " Segment “{segment}” needs the HB utilities: without them, all people."},
}


def word_tables(data, lang="de"):
    """[(question, scenario name, basis, persons, rows, kind, segment)] of every saved scenario, without the weighting
    of the page (as the Word file) but with the segment the scenario was saved with, as the page shows it (audit
    2026-10-05, M-3). segment: None, or {"name", "state": "weighted" (with "size", "effective"), "gone" (no current
    solution has it: all persons) or "no_hb" (without HB nothing to weight: all persons)}. Numbers in German, whatever
    the page: add_to_docx writes them in the language of the file."""
    token = _NUMBERS.set("de")
    try:
        return _word_tables(data)
    finally:
        _NUMBERS.reset(token)


def _word_tables(data):
    out = []
    for entry in data.get("items") or []:
        item = entry.get("item")
        part = _part(entry)
        if item is None or part is None or item.type not in TYPES:
            continue
        lc_found = _segments_of(item, part)
        for k, scenario in enumerate(saved(item)):
            seg = scenario.get("segment") or 0
            weighting, segment = None, None
            if seg:
                name = _segment_name(lc_found, seg)
                if name is None:
                    segment = {"name": str(seg), "state": "gone"}
                else:
                    basis = (_maxdiff_model(item, part) or [None] * 4)[3] if item.type == "maxdiff" \
                        else _cbc_model(item, part)[2]
                    if basis == "hb":
                        weighting = _segment_weighting(lc_found[0], lc_found[1], seg, None)
                        segment = {"name": name, "state": "weighted"}
                    else:
                        segment = {"name": name, "state": "no_hb"}
            if item.type == "maxdiff":
                state, _ = maxdiff_state(item, {"sim_load": str(k)})
                found = maxdiff_compute(item, part, state, item.survey.default_language, weighting)
            else:
                state, _ = cbc_state(item, {"sim_load": str(k)})
                found = cbc_compute(item, part, state, item.survey.default_language, weighting)
            if found.get("rows"):
                persons = found["persons"]
                if weighting is not None:
                    size = _segment_size(lc_found[0], lc_found[1], seg, found["_pids"], weighting)
                    segment.update(size=size["sum"], effective=size["effective"])
                    persons = len(found["_pids"])
                out.append((entry.get("question") or item.code, scenario["name"], found["basis"], persons,
                            found["rows"], item.type, segment))
    return out


def add_to_docx(doc, data, lang, number_from):
    """One APA table per saved scenario after the tables already in `doc` (app/manuscript_tables.py)."""
    from docx.shared import Pt

    from app.manuscript_tables import _border, _write

    t = DOCX.get(lang, DOCX["de"])
    n = number_from
    for question, name, basis, persons, rows, kind, segment in word_tables(data, lang):
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
        p = doc.add_paragraph()
        p.add_run(t["table"].format(n=n)).bold = True
        p.paragraph_format.space_after = Pt(0)
        p = doc.add_paragraph()
        title = t["title"].format(question=question[:80], name=name)
        if segment and segment["state"] == "weighted":
            title += t["segment"].format(segment=segment["name"])
        p.add_run(title).italic = True
        head = [t["statement"] if kind == "maxdiff" else t["product"], t["sop"]] + ([t["fc"]] if basis in ("hb", "persons") else [])
        table = doc.add_table(rows=1 + len(rows), cols=len(head))
        for j, text in enumerate(head):
            cell = table.rows[0].cells[j]
            _write(cell, text, right=j > 0)
            _border(cell, ("top", "bottom"))
        sep = "; " if lang == "de" else ", "

        def value(x, lo, hi):
            text = x if lang == "de" else x.replace(",", ".")
            if lo is not None:
                text += f" [{lo}{sep}{hi}]" if lang == "de" else f" [{lo.replace(',', '.')}{sep}{hi.replace(',', '.')}]"
            return text

        for i, row in enumerate(rows):
            name = f"{t['product']} {row['label']['letter']}" if row["label"].get("letter") else row["label"]["name"]
            label = name + (f" ({row['label']['detail']})" if row["label"]["detail"] else "")
            values = [label, value(row["sop"], row["sop_lo"], row["sop_hi"])]
            if basis in ("hb", "persons"):
                values.append(value(row["fc"], row["fc_lo"], row["fc_hi"]))
            for j, text in enumerate(values):
                cell = table.rows[i + 1].cells[j]
                _write(cell, text, right=j > 0)
                if i == len(rows) - 1:
                    _border(cell, ("bottom",))
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.add_run(t["note"]).italic = True
        text = " " + t[basis].format(n=persons, draws=choice_predict.DRAWS)
        if segment:
            text += t[segment["state"]].format(segment=segment["name"], size=segment.get("size"),
                                               effective=segment.get("effective"))
        p.add_run(text + t["market"])
        n += 1
    return n
