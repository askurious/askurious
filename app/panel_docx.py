"""The models over the waves of a panel as tables for a manuscript (.docx, APA style 7th edition; longitudinal and
multilevel concept 2026-10-17, E14): Table 1 the cross-lagged panel model, Table 2 the growth curves, Table 3 the
measurement invariance over time; the numbers of the kept block (app/panel_models.py). German with the decimal comma,
English without a leading zero where a value cannot exceed 1 (app/manuscript_tables.py)."""
import io

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from app.manuscript_tables import fmt
from app.sem_docx import _note, _p, _table, _title

TEXT = {
    "de": {"t1": "Tabelle {k}", "note": "Anmerkung.", "cross": "Kreuzverzögerte Pfade über die Wellen", "path": "Pfad",
           "growth": "Latente Wachstumskurven", "var": "Größe", "inv": "Messinvarianz über die Zeit", "level": "Stufe",
           "stab": "Stabilität", "riclpm": ("Random-Intercept-Cross-Lagged-Panel-Modell (Hamaker et al., 2015) mit {n} "
                                             "verknüpften Personen; Pfade zwischen den Abweichungen innerhalb der Personen."),
           "clpm": "Cross-Lagged-Panel-Modell mit {n} verknüpften Personen.",
           "equal": " Pfade über die Abstände gleichgesetzt.", "fiml": (" Full Information Maximum Likelihood mit robusten "
                                                                       "Standardfehlern (MLR); β standardisiert, p zweiseitig."),
           "fit": " Modellfit: χ²({df}) = {chi}, CFI = {cfi}, RMSEA = {rmsea}, SRMR = {srmr}.",
           "g_body": ("I = Ausgangsniveau, S = Steigung je {unit}; Mittelwerte mit Standardfehler, Varianzen; r = "
                      "Korrelation von I und S. Full Information Maximum Likelihood (MLR)."),
           "wave": "Welle", "week": "Woche", "i_body": ("Konfigural: gleiche Struktur; metrisch: gleiche Ladungen; skalar: "
                                                         "dazu gleiche Achsenabschnitte. Δχ² skaliert (Satorra & Bentler, 2001)."),
           "levels": {"configural": "konfigural", "metric": "metrisch", "scalar": "skalar"}, "holds": "hält"},
    "en": {"t1": "Table {k}", "note": "Note.", "cross": "Cross-Lagged Paths Over the Waves", "path": "Path",
           "growth": "Latent Growth Curves", "var": "Variable", "inv": "Measurement Invariance Over Time", "level": "Level",
           "stab": "stability", "riclpm": ("Random intercept cross-lagged panel model (Hamaker et al., 2015) with {n} linked "
                                            "people; paths between the deviations within people."),
           "clpm": "Cross-lagged panel model with {n} linked people.",
           "equal": " Paths set equal over the intervals.", "fiml": (" Full information maximum likelihood with robust "
                                                                    "standard errors (MLR); β standardized, p two-sided."),
           "fit": " Model fit: χ²({df}) = {chi}, CFI = {cfi}, RMSEA = {rmsea}, SRMR = {srmr}.",
           "g_body": ("I = initial level, S = slope per {unit}; means with standard errors, variances; r = correlation of "
                      "I and S. Full information maximum likelihood (MLR)."),
           "wave": "wave", "week": "week", "i_body": ("Configural: same structure; metric: equal loadings; scalar: also "
                                                       "equal intercepts. Δχ² scaled (Satorra & Bentler, 2001)."),
           "levels": {"configural": "configural", "metric": "metric", "scalar": "scalar"}, "holds": "holds"},
}


def _chi(fit):
    s = fit.get("scaled")
    return s["chi2"] if s else fit.get("chi2")


def _fit_text(fit, t, lang):
    return t["fit"].format(df=fit.get("df"), chi=fmt(_chi(fit), lang), cfi=fmt(fit.get("cfi"), lang, True, 3),
                           rmsea=fmt(fit.get("rmsea"), lang, True, 3), srmr=fmt(fit.get("srmr"), lang, True, 3))


def build(result, lang="de"):
    """The tables as .docx bytes from the kept block."""
    lang = lang if lang in TEXT else "de"
    t = TEXT[lang]
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Cm(2.5)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Times New Roman", Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    number = 1
    waves = result.get("waves") or []
    c = result.get("cross") or {}
    if c.get("paths"):
        _title(doc, number, t["cross"], t)
        rows = []
        single = c.get("equal") or len(waves) == 2
        for p in c["paths"]:
            name = f"{p['from_label']} → {p['to_label']}" + (f" ({t['stab']})" if p["auto"] else "")
            if single:
                rows.append([name, fmt(p.get("est"), lang), fmt(p.get("se"), lang), fmt(p.get("std"), lang), _p(p.get("p"), lang)])
            else:
                rows.append([name] + [f"{fmt(i['std'], lang)} ({_p(i['p'], lang)})" for i in p["intervals"]])
        head = [t["path"]] + (["B", "SE", "β", "p"] if single else [f"→ {w} β (p)" for w in waves[1:]])
        _table(doc, head, rows, italic=set(range(1, len(head))) if single else set())
        body = (t["riclpm"] if c.get("kind") == "riclpm" else t["clpm"]).format(n=c.get("n"))
        if c.get("equal"):
            body += t["equal"]
        body += t["fiml"] + _fit_text(c.get("fit") or {}, t, lang)
        _note(doc, body, t)
        number += 1
    grow = [v for v in result.get("vars") or [] if (v.get("growth") or {}).get("means")]
    if grow:
        doc.add_paragraph()
        _title(doc, number, t["growth"], t)
        rows = []
        for v in grow:
            g = v["growth"]
            m, var = g["means"], g["variances"]
            rows.append([v["label"], f"{fmt(m['I']['est'], lang)} ({fmt(m['I']['se'], lang)})",
                         f"{fmt(m['S']['est'], lang)} ({fmt(m['S']['se'], lang)})", fmt(var["I"]["est"], lang),
                         fmt(var["S"]["est"], lang), fmt((g.get("corr") or {}).get("std"), lang, True),
                         fmt(g["fit"].get("cfi"), lang, True, 3), fmt(g["fit"].get("rmsea"), lang, True, 3)])
        _table(doc, [t["var"], "M(I) (SE)", "M(S) (SE)", "Var(I)", "Var(S)", "r", "CFI", "RMSEA"], rows,
               italic={1, 2, 3, 4, 5, 6, 7})
        _note(doc, t["g_body"].format(unit=t["week"] if result.get("real_time") else t["wave"]), t)
        number += 1
    inv = [v for v in result.get("vars") or [] if v.get("invariance")]
    if inv:
        doc.add_paragraph()
        _title(doc, number, t["inv"], t)
        rows = []
        for v in inv:
            for r in v["invariance"]["levels"]:
                f = r["fit"]
                if not f.get("converged"):
                    continue
                d = r.get("diff") or {}
                rows.append([f"{v['label']}: {t['levels'][r['level']]}", f"{fmt(_chi(f), lang)} ({f.get('df')})",
                             fmt(f.get("cfi"), lang, True, 3), fmt(f.get("rmsea"), lang, True, 3),
                             fmt(f.get("srmr"), lang, True, 3),
                             f"{fmt(d.get('chi2'), lang)} ({d.get('df')}), {_p(d.get('p'), lang)}" if d else "",
                             fmt(r.get("d_cfi"), lang, False, 3) if r.get("d_cfi") is not None else "",
                             "✓" if r.get("holds") else ""])
        _table(doc, [t["level"], "χ² (df)", "CFI", "RMSEA", "SRMR", "Δχ² (Δdf), p", "ΔCFI", t["holds"]], rows,
               italic={2, 3, 4, 6})
        _note(doc, t["i_body"], t)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
