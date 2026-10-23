"""One interview's transcript as a Word or PDF document with a transcript head and line-numbered body (Ralf
2026-10-18, concept in the thread "Transkript-Export", all points as recommended), and all of them as a ZIP.

The layout follows the institute's template ("Transkriptionsvorlage"): a head of seven lines (Kürzel, date and place,
length, interviewer, interviewee, transcribed on, transcribed by) in Tahoma 12, not numbered; then each turn its own
paragraph "I:" or "B:", a tab and the words, in Courier New 11, single spacing, the interviewer's turns bold, every
line numbered (Word's own line numbering, continuous, 0.5 cm from the text). Page A4, margins 2.5/2.0/2.5/2.5 cm,
the page header "<prefix>_<Kürzel>" in Tahoma 10, the page number bottom right.

The rules are those of the simple transcription after Dresing & Pehl (2018), which is what the Whisper draft and the
checking on the transcript page give (paragraphs per speaker, punctuation, time marks #hh:mm:ss-0#); the head names
them. GAT 2 (Selting et al., 2009) is not claimed: even its minimal transcript wants one line per intonation phrase,
measured pauses and overlaps, which an automatic draft does not hold. Time marks are left out unless asked for (the
template has none). The results of the tasks of the guide stand as a line of their own at their time, in square
brackets, italic and without a speaker, "[Aufgabe „Wichtigkeit“: 1 Preis, 2 Qualität, 3 Marke]" (app/interview_tasks.py,
E5); the head names these insertions; they can be left out like the time marks (`tasks`).

The head fills itself from the interview (identifier, date, length = the last time mark, who created it, when it was
checked, Whisper's model and who checked it) and from the Kurzfragebogen questions picked in the guide's settings
(occupation, age, gender); every value can be changed on the transcript page and is kept per interview:

  Response.interview["transcript_head"] = {key: text}   only the values that differ from the automatic ones

A transcript not yet checked is exported, marked "Entwurf" in the head and the page header. Names in the text are not
replaced: the page says to replace them before exporting.

The PDF is laid out here with reportlab (as app/paper_pdf.py), with the standard fonts Courier and Helvetica (same
widths as Courier New, close to Tahoma; Latin script only, as the paper version); it numbers every printed line, which
matches Word's numbering almost always but not by guarantee (Word wraps by its own measure)."""
import io
import re
import zipfile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from app.items.base import plain

KEYS = ("kuerzel", "datum", "dauer", "interviewer", "befragte", "am", "von")
LABELS = {"kuerzel": "Kürzel", "datum": "Datum und Ort der Aufnahme", "dauer": "Dauer der Aufnahme",
          "interviewer": "Interviewer(in)", "befragte": "Befragte(r)", "am": "Transkribiert am", "von": "Transkribiert von"}
LABELS_EN = {"kuerzel": "Identifier", "datum": "Date and place of recording", "dauer": "Length of recording",
             "interviewer": "Interviewer", "befragte": "Interviewee", "am": "Transcribed on", "von": "Transcribed by"}
RULES = ("Transkriptionsregeln: einfache Transkription nach Dresing & Pehl (2018)",
         "Transcription rules: simple transcription after Dresing & Pehl (2018)")
DRAFT = ("Entwurf, nicht geprüft", "draft, not checked")
MARK = re.compile(r"\s*#\d{1,2}:\d{2}:\d{2}-\d#")
SPEAKER = re.compile(r"^\s*([A-ZÄÖÜ][\w\-]{0,15}):\s*")
PICKS = ("head_job", "head_age", "head_gender")  # the guide's settings (app/interview.py): codes of Kurzfragebogen questions

# the template's page: A4, margins in cm, header and footer 1.25 cm from the edge, Word's default tab stop
MARGINS = {"top": 2.5, "bottom": 2.0, "left": 2.5, "right": 2.5}
EDGE = 1.25
TAB = 1.25
NUMBER_GAP = 0.5
BODY_SIZE, HEAD_SIZE, PAGE_SIZE = 11, 12, 10


def lang_of(survey, ui=None):
    """The document's language: the survey's if German or English, else the downloader's (as the consent form)."""
    from app.interview_docx import _paper_ui
    from app.items.base import reader

    return "en" if _paper_ui(survey, reader(ui)) == "en" else "de"


def _t(lang, de, en):
    return en if lang == "en" else de


def _day(iso, lang):
    iso = (iso or "")[:10]
    try:
        y, m, d = (int(x) for x in iso.split("-"))
    except ValueError:
        return ""
    return f"{y:04d}-{m:02d}-{d:02d}" if lang == "en" else f"{d:02d}.{m:02d}.{y:04d}"


def length(text):
    """The last time mark of the transcript as h:mm:ss (the interview clock), or ""."""
    marks = re.findall(r"#(\d{1,2}):(\d{2}):(\d{2})-\d#", text or "")
    if not marks:
        return ""
    s = max(int(h) * 3600 + int(m) * 60 + int(x) for h, m, x in marks)
    return f"{s // 3600:02d}:{s // 60 % 60:02d}:{s % 60:02d}"


def interviewee(survey, row, lang):
    """Occupation, age and gender from the Kurzfragebogen questions picked in the settings, joined by commas."""
    from app import interview, interview_data
    from app.interview_results import answer_text

    st = interview.state(survey)
    items = {i.code: i for i in interview_data.form_items(survey, "short")}
    out = []
    for key in PICKS:
        item = items.get(st.get(key) or "")
        if item is not None:
            text = answer_text(item, (row.answers or {}).get(str(item.id)), survey.default_language or lang)
            if text:
                out.append(plain(text))
    return ", ".join(out)


def automatic(survey, row, lang):
    """The head's values as the interview gives them."""
    iv = row.interview or {}
    tr = iv.get("transcript") or {}
    # "Philipp, mit Whisper (large-v3)": who checked it, and the machine that drafted it (none for an upload)
    name = tr.get("checked_by") if tr.get("checked") else ""
    machine = f"Whisper ({tr['model']})" if tr.get("model") and tr.get("source") != "upload" else ""
    by = f"{name}, {_t(lang, 'mit', 'with')} {machine}" if name and machine else name or machine
    return {"kuerzel": iv.get("label", ""), "datum": _day(iv.get("date"), lang), "dauer": length(tr.get("text")),
            "interviewer": iv.get("by_name", ""), "befragte": interviewee(survey, row, lang),
            "am": _day(tr.get("checked"), lang) if tr.get("checked") else _t(lang, *DRAFT), "von": by}


def head(survey, row, lang):
    """[(key, label, value, automatic value)] with the changes kept for this interview."""
    auto = automatic(survey, row, lang)
    kept = (row.interview or {}).get("transcript_head") or {}
    labels = LABELS_EN if lang == "en" else LABELS
    return [(k, labels[k], kept[k] if isinstance(kept.get(k), str) else auto[k], auto[k]) for k in KEYS]


def kept_head(form):
    """The changes to keep from the transcript page's fields head_<key>: a value differs from the automatic one the
    page showed beside it (auto_<key>), so that a page saved after the automatic value moved on keeps none by mistake."""
    def norm(value):
        return " ".join((value or "").split())[:200]

    return {k: norm(form.get(f"head_{k}")) for k in KEYS
            if f"head_{k}" in form and norm(form.get(f"head_{k}")) != norm(form.get(f"auto_{k}"))}


def page_title(survey, kuerzel, draft, lang):
    from app import interview

    prefix = interview.state(survey).get("head_title") or plain(survey.title_in(survey.default_language))
    text = "_".join(x for x in ("_".join(prefix.split()), kuerzel) if x)
    return text + (f" ({_t(lang, 'Entwurf', 'draft')})" if draft else "")


def _turn(line, marks):
    """(speaker or "", words, interviewer?) of one line, or None for an empty one."""
    if not marks:
        line = MARK.sub("", line)
    line = line.strip()
    if not line:
        return None
    m = SPEAKER.match(line)
    speaker, words = (m.group(1), line[m.end():]) if m else ("", line)
    return speaker, words, bool(speaker) and speaker[0] == "I"


def turns(text, marks=False):
    """[(speaker or "", words, interviewer?)]: one per non-empty line of the transcript."""
    return [t for t in (_turn(line, marks) for line in (text or "").replace("\r\n", "\n").split("\n")) if t]


def body(survey, row, lang, marks=False, tasks=True):
    """[(speaker or "", words, interviewer?, inserted?)]: the turns, and with `tasks` the line of each task at its
    time (app/interview_tasks.py), inserted = True, shown in italics without a speaker."""
    from app import interview_tasks

    text = ((row.interview or {}).get("transcript") or {}).get("text")
    extra = interview_tasks.inserts(survey, row, lang, marks) if tasks else []
    lines, put = interview_tasks.merge(text, extra) if extra else ((text or "").replace("\r\n", "\n").split("\n"), set())
    out = []
    for n, line in enumerate(lines):
        if n in put:
            out.append(("", line, False, True))
            continue
        t = _turn(line, marks)
        if t:
            out.append((*t, False))
    return out


def _parts(survey, row, lang, marks, tasks=True):
    tr = (row.interview or {}).get("transcript") or {}
    values = head(survey, row, lang)
    kuerzel = next(v for k, _, v, _ in values if k == "kuerzel")
    turns_ = body(survey, row, lang, marks, tasks)
    rules = [(_t(lang, *RULES), None)]
    if any(t[3] for t in turns_):  # the head names the insertions (Dresing & Pehl, E5)
        from app.interview_tasks import HEAD_NOTE

        rules.append((_t(lang, *HEAD_NOTE), None))
    return ([(label, value) for _, label, value, _ in values] + rules, turns_,
            page_title(survey, kuerzel, not tr.get("checked"), lang))


# --- Word -------------------------------------------------------------------------------------------------------


def _font(style_or_run_font, name, size, element):
    style_or_run_font.name, style_or_run_font.size = name, Pt(size)
    rpr = element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)


def _flag(ppr, tag):
    el = OxmlElement(tag)
    ppr.append(el)


def _page_field(paragraph):
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    return run


def docx(survey, row, ui=None, marks=False, tasks=True):
    lang = lang_of(survey, ui)
    values, turns_, title = _parts(survey, row, lang, marks, tasks)
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(MARGINS["top"]), Cm(MARGINS["bottom"])
    section.left_margin, section.right_margin = Cm(MARGINS["left"]), Cm(MARGINS["right"])
    section.header_distance = section.footer_distance = Cm(EDGE)
    numbers = OxmlElement("w:lnNumType")  # every line, continuous, 0.5 cm from the text (284 twips, as the template)
    for attr, value in (("w:countBy", "1"), ("w:distance", "284"), ("w:restart", "continuous")):
        numbers.set(qn(attr), value)
    cols = section._sectPr.find(qn("w:cols"))
    (cols.addprevious if cols is not None else section._sectPr.append)(numbers)

    normal = doc.styles["Normal"]
    _font(normal.font, "Courier New", BODY_SIZE, normal.element)
    fmt = normal.paragraph_format
    fmt.space_before = fmt.space_after = Pt(0)
    fmt.line_spacing = 1.0
    _flag(normal.element.get_or_add_pPr(), "w:suppressAutoHyphens")
    kopf = doc.styles.add_style("Transkriptkopf", WD_STYLE_TYPE.PARAGRAPH)
    kopf.base_style = normal
    _font(kopf.font, "Tahoma", HEAD_SIZE, kopf.element)
    _flag(kopf.element.get_or_add_pPr(), "w:suppressLineNumbers")

    for name in ("Header", "Footer"):  # also the page number Word fills into the field
        _font(doc.styles[name].font, "Tahoma", PAGE_SIZE, doc.styles[name].element)
    page_head = section.header.paragraphs[0]
    _font(page_head.add_run(title).font, "Tahoma", PAGE_SIZE, page_head.runs[0]._r)
    foot = section.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = _page_field(foot)
    _font(run.font, "Tahoma", PAGE_SIZE, run._r)

    for label, value in values:
        p = doc.add_paragraph(style="Transkriptkopf")
        p.add_run(label if value is None else f"{label}: {value}")
    doc.add_paragraph(style="Transkriptkopf")
    for speaker, words, bold, inserted in turns_:
        p = doc.add_paragraph()
        run = p.add_run(f"{speaker}:" if speaker else words)
        if speaker:
            run.add_tab()
            run.add_text(words)
        run.bold = bold
        if inserted:
            run.italic = True
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# --- PDF --------------------------------------------------------------------------------------------------------

CM = 72 / 2.54
BODY_LEADING = BODY_SIZE * 1.133   # Courier New, single spacing
HEAD_LEADING = HEAD_SIZE * 1.207   # Tahoma, single spacing


def _safe(text):
    from app.paper_pdf import _safe as safe

    return safe(text)


def _wrap(text, font, size, first, rest):
    """Lines of `text` in `font`: the first no wider than `first` points, the others than `rest`; an overlong word is
    broken."""
    from reportlab.pdfbase.pdfmetrics import stringWidth

    lines, line, width = [], "", first
    for word in text.split(" "):
        candidate = f"{line} {word}" if line else word
        if stringWidth(candidate, font, size) <= width or not word:
            line = candidate
            continue
        if line:
            lines.append(line)
            width = rest
        while stringWidth(word, font, size) > width:
            cut = max(1, next((n for n in range(len(word), 0, -1) if stringWidth(word[:n], font, size) <= width), 1))
            lines.append(word[:cut])
            word, width = word[cut:], rest
        line = word
    lines.append(line)
    return lines


def pdf(survey, row, ui=None, marks=False, tasks=True):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from reportlab.pdfgen import canvas

    lang = lang_of(survey, ui)
    values, turns_, title = _parts(survey, row, lang, marks, tasks)
    w, h = A4
    left, right = MARGINS["left"] * CM, w - MARGINS["right"] * CM
    top, bottom = h - MARGINS["top"] * CM, MARGINS["bottom"] * CM
    width = right - left
    out = io.BytesIO()
    c = canvas.Canvas(out, pagesize=A4, invariant=1)
    c.setTitle(_safe(title))
    state = {"page": 0, "y": 0.0, "number": 0}

    def new_page():
        if state["page"]:
            c.showPage()
        state["page"] += 1
        c.setFont("Helvetica", PAGE_SIZE)
        c.drawString(left, h - EDGE * CM - PAGE_SIZE, _safe(title))
        c.drawRightString(right, EDGE * CM, str(state["page"]))
        state["y"] = top

    def line_room(leading):
        if state["y"] - leading < bottom:
            new_page()
        state["y"] -= leading
        return state["y"] + leading - leading * 0.8  # the baseline in the line's box

    new_page()
    for label, value in values:
        text = _safe(label if value is None else f"{label}: {value}")
        for line in _wrap(text, "Helvetica", HEAD_SIZE, width, width):
            y = line_room(HEAD_LEADING)
            c.setFont("Helvetica", HEAD_SIZE)
            c.drawString(left, y, line)
    line_room(HEAD_LEADING)
    tab = TAB * CM
    for speaker, words, bold, inserted in turns_:
        font = "Courier-Bold" if bold else "Courier-Oblique" if inserted else "Courier"
        words = _safe(words)
        start = 0.0
        if speaker:
            label = _safe(f"{speaker}:")
            used = stringWidth(label, font, BODY_SIZE)
            start = (int(used // tab) + 1) * tab  # the next tab stop after the label
        lines = _wrap(words, font, BODY_SIZE, width - start, width)
        for n, line in enumerate(lines):
            y = line_room(BODY_LEADING)
            state["number"] += 1
            c.setFont("Courier", BODY_SIZE)
            c.drawRightString(left - NUMBER_GAP * CM, y, str(state["number"]))
            c.setFont(font, BODY_SIZE)
            if n == 0 and speaker:
                c.drawString(left, y, label)
                c.drawString(left + start, y, line)
            else:
                c.drawString(left, y, line)
    c.showPage()
    c.save()
    return out.getvalue()


# --- names and the ZIP ------------------------------------------------------------------------------------------

FORMATS = {"docx": (docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
           "pdf": (pdf, "application/pdf")}


def file_name(survey, row, ext):
    label = "".join(ch for ch in (row.interview or {}).get("label", str(row.id)) if ch.isalnum() or ch in "-_") or str(row.id)
    return f"{survey.code}-{label}-transkript.{ext}"


def has_text(row):
    return bool(((row.interview or {}).get("transcript") or {}).get("text"))


def zip_all(survey, rows, formats, ui=None, marks=False, tasks=True):
    """One file per interview with a transcript and format, in one ZIP."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for row in rows:
            if has_text(row):
                for ext in formats:
                    z.writestr(file_name(survey, row, ext), FORMATS[ext][0](survey, row, ui, marks, tasks))
    return buf.getvalue()
