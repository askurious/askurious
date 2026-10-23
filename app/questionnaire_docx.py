"""The questionnaire as a Word document (.docx), in the survey's default language, in two versions:

- "codebook": for the research documentation or an ethics board. Blocks and pages as headings, each question with
  its variable name, type and answer mode, validated constructs with their source, random orders, flow conditions
  and loop tables.
- "print": a paper questionnaire as participants see it. Only what is on the pages: block titles where they are
  shown, questions, answers to tick, required questions marked with *; no variable names, sources or notes. A block
  in a loop is printed once per table row, with the row's values filled in.

Answers to tick: ○ one, ☐ several; Likert statements and semantic differentials as tables. Item types without a
layout of their own list what questionnaire_md would list for them. Built with python-docx; nothing of the
participants' answers is in it.

The questions, answers and what participants read on paper stay in the survey's default language; the frame of the
codebook (headings, notes, answer modes, the red box of elements not on paper) is in the language of the person who
downloads (`ui`, app.lang.current(); German outside a request; concept "Mehrsprachige Oberfläche", E8)."""
import io
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.items.base import DONT_KNOW_CODE, PLACEHOLDER, dont_know, plain

# Characters XML cannot hold. A manual line break pasted from Word arrives as \x0b, a text from a PDF may bring
# \x0c or others; they are kept in the questionnaire, and python-docx refuses a whole document over one of them
# ("All strings must be XML compatible"). Every text of a document passes CT_R.add_t (docx/oxml/text/run.py,
# python-docx 1.2), whatever it was added with (run, paragraph, heading, table cell): they are dropped there.
_NOT_XML = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")


def _only_xml():
    from docx.oxml.text.run import CT_R

    add_t = CT_R.add_t
    if getattr(add_t, "only_xml", False):
        return

    def clean_add_t(self, text):
        return add_t(self, _NOT_XML.sub("", text))

    clean_add_t.only_xml = True
    CT_R.add_t = clean_add_t


_only_xml()

GREY = RGBColor(0x66, 0x66, 0x5F)
MODES = {"required": "Pflicht", "nudge": "Hinweis, wenn leer", "optional": "freiwillig"}
MODES_EN = {"required": "required", "nudge": "prompt if empty", "optional": "optional"}
VERSIONS = ("codebook", "print")
# live feedback on earlier answers of all participants: nothing to print, left out of both versions
LIVE = ("sample_compare", "sample_map", "sample_crowd", "score_band", "topic_map", "topic_profile", "scale_profile")
# answered by interaction (choice tasks, sorting, placing, maps): a red note instead of a paper version
NOT_ON_PAPER = ("conjoint", "acbc", "maxdiff", "paired", "cva", "aca", "assign", "position2d", "card_sort", "q_sort", "association", "map_marker", "geolocation")
DATE_FORMS = {"year": "____", "month": "__ / ____", "date": "__.__.____", "time": "__:__", "datetime": "__.__.____  __:__"}
WIDTH = 17.0  # cm between the margins


class Ctx:
    def __init__(self, doc, lang, version, variables=None, ui="de"):
        self.doc, self.lang, self.version, self.variables = doc, lang, version, variables or {}
        self.ui = ui  # the language of the frame texts: the downloader's (German or English)

    def say(self, de, en, **values):
        """A frame text in the downloader's language, with its values filled in."""
        text = en if self.ui == "en" else de
        return text.format(**values) if values else text

    def kind_label(self, kind):
        """The element type's name in the downloader's language (ItemType.label is a Both, which reads the page
        language; `ui` may be given explicitly)."""
        from app.lang import Both

        for cls in type(kind).__mro__:
            found = cls.__dict__.get("label")
            if isinstance(found, Both):
                return found.en if self.ui == "en" else found.de
            if found is not None:
                return str(found)
        return str(kind.label)

    @property
    def codebook(self):
        return self.version == "codebook"

    def fill(self, value):
        """The text with the loop row's values put in for {{name}} (print; placeholders stay in the codebook); an
        answer to an earlier question (piping) is its text for an empty answer ({{code|text}}), or a gap."""
        value = str(value or "")
        if self.codebook:
            return value

        def sub(m):
            if m.group(1) in self.variables:
                return str(self.variables[m.group(1)])
            return m.group(2).strip() if m.group(2) is not None else "________"

        return PLACEHOLDER.sub(sub, value)

    def text(self, value):
        return plain(self.fill(value))


# --- layout helpers ---------------------------------------------------------------------------


def _small(paragraph, text, mono=False):
    run = paragraph.add_run(text)
    run.font.size, run.font.color.rgb = Pt(8.5), GREY
    if mono:
        run.font.name = "Courier New"
    return run


def _note(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic, run.font.color.rgb, run.font.size = True, GREY, Pt(9)
    return p


def _grid(table):
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER


def _matrix(table):
    """The cells of a table as rows of cells, read once: python-docx's table.cell(r, c) builds the list of all cells
    on every call, which made a large grid quadratic (performance audit 2026-10-09b, P3). The tables here have no
    merged cells, so tc[r][c] is the same cell as table.cell(r, c)."""
    return [row.cells for row in table.rows]


def _widths(table, widths):
    """Fixed column widths in cm (Word and LibreOffice ignore cell widths alone)."""
    table.autofit = False  # writes <w:tblLayout w:type="fixed"/> in its proper place
    total = table._tbl.tblPr.find(qn("w:tblW"))
    total.set(qn("w:type"), "dxa")
    total.set(qn("w:w"), str(int(Cm(sum(widths)) / 635)))
    for col, width in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        col.set(qn("w:w"), str(int(Cm(width) / 635)))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Cm(width)


def _shade(cell, fill="EEF2F8"):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def _center(cell, text, size=10):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(text).font.size = Pt(size)


def _runs(paragraph, text):
    """One line of text with its **bold** parts bold; other formatting marks dropped."""
    for n, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part:
            paragraph.add_run(plain(part)).bold = n % 2 == 1


def _paragraphs(ctx, text):
    """A longer text (consent, info) paragraph by paragraph, with list items ("- ") as bullets."""
    text = ctx.fill(text)
    for para in (text or "").replace("\r\n", "\n").split("\n\n"):
        lines = [line.strip() for line in para.split("\n") if line.strip()]
        if not lines:
            continue
        p = ctx.doc.add_paragraph()
        for n, line in enumerate(lines):
            if n:
                p.add_run().add_break()
            _runs(p, "• " + line[2:] if line.startswith("- ") else line)


# --- items ------------------------------------------------------------------------------------


def _question(ctx, item):
    from app import constructs

    kind, t = item.kind, item.t(ctx.lang)
    p = ctx.doc.add_paragraph()
    p.paragraph_format.space_before, p.paragraph_format.keep_with_next = Pt(10), True
    text = ctx.text((t.get("question") or "").strip())
    if text:
        p.add_run(text).bold = True
    if ctx.codebook:
        p.add_run("  ")
        _small(p, item.code, mono=True)
        modes = MODES_EN if ctx.ui == "en" else MODES
        _small(p, f" · {ctx.kind_label(kind)} · {modes.get(item.answer_mode, item.answer_mode)}")
        construct = constructs.info(item)
        if construct and ctx.ui == "en":  # the changes in English (constructs.modifications; info() has them in German)
            found = constructs.find(construct["key"])
            construct["modifications"] = constructs.modifications(item, found) if found else construct["modifications"]
        if construct and construct["linked"]:
            _note(ctx.doc, ctx.say("Verknüpftes Element „{label}“, Fassung {version}", "Linked element “{label}”, revision {version}",
                                   label=construct["label"], version=construct["version"] or 1)
                  + (ctx.say(" (in dieser Befragung angepasst)", " (adapted in this survey)")
                     if construct["unlocked"] or construct["modifications"] else "") + ".")
        elif construct:
            _note(ctx.doc, (ctx.say("Verändertes Konstrukt ", "Modified construct ") if construct["modifications"]
                            else ctx.say("Konstrukt ", "Construct "))
                  + ctx.say("„{label}“: {citation}", "“{label}”: {citation}", label=construct["label"],
                            citation=construct["citation"])
                  + (ctx.say(" Änderungen: {mods}.", " Changes: {mods}.", mods="; ".join(construct["modifications"]))
                     if construct["modifications"] else ""))
        from app import item_condition

        if item_condition.of(item):  # "Nur zeigen, wenn …", said like the conditions of the flow at the block
            _note(ctx.doc, ctx.say("Nur gezeigt, wenn {cond}.", "Shown only if {cond}.",
                                   cond=item_condition.text(item, ctx.lang, ctx.ui)))
        if kind.order_codes(item):
            _note(ctx.doc, ctx.say("Reihenfolge der Antworten bzw. Aussagen zufällig je Person.",
                                   "Order of the answers or statements random per person."))
    elif item.answer_mode == "required":
        p.add_run(" *").bold = True
    return t


def _dk_label(lang):
    """The words of the "weiß nicht" answer as participants see them (templates/items/_dont_know.html)."""
    return "weiß nicht" if lang in ("de", "de-du") else "don't know"


def _dk_note(ctx, item):
    """Codebook: how "weiß nicht" of a matrix or a rating is stored (V41); a date names its column <code>_dk in
    questionnaire_lines."""
    if ctx.codebook and dont_know(item):
        _note(ctx.doc, ctx.say("„weiß nicht“ = {code}, zählt als fehlend.", "“don't know” = {code}, counts as missing.",
                               code=DONT_KNOW_CODE))


def _choice(ctx, item):
    kind = item.kind
    mark = "☐" if item.type == "multiple_choice" else "○"
    if kind.dropdown(item) and ctx.codebook:
        _note(ctx.doc, ctx.say("als Auswahlliste", "as a dropdown list"))
    elif kind.search(item) and ctx.codebook:
        _note(ctx.doc, ctx.say("als Suchfeld mit Vorschlägen", "as a search field with suggestions"))
    for o in kind.options(item, ctx.lang):
        p = ctx.doc.add_paragraph()
        p.paragraph_format.left_indent, p.paragraph_format.space_after = Cm(0.6), Pt(0)
        p.add_run(f"{mark}  {ctx.text(o['label'])}")
        if o.get("text"):
            p.add_run("  ____________________")
        if ctx.codebook:
            _small(p, f"   {o['code']}", mono=True)


def _likert(ctx, item):
    kind = item.kind
    labels = [ctx.text(l) for l in kind.labels(item, ctx.lang)]
    statements = kind.statements(item, ctx.lang)
    dk = dont_know(item)  # "weiß nicht": one more column after the scale, not shaded like it (V41)
    table = ctx.doc.add_table(rows=1 + len(statements), cols=1 + len(labels) + dk)
    _grid(table)
    tc = _matrix(table)
    first = 7.0 if len(labels) + dk <= 7 else 5.5
    _widths(table, [first] + [(WIDTH - first) / max(1, len(labels) + dk)] * (len(labels) + dk))
    for n, label in enumerate(labels, start=1):
        _center(tc[0][n], label, size=8)  # empty in the middle of a scale labelled at the ends only
        _shade(tc[0][n])
    if dk:
        _center(tc[0][len(labels) + 1], _dk_label(ctx.lang), size=8)
    _shade(tc[0][0])
    for r, s in enumerate(statements, start=1):
        cell = tc[r][0]
        cell.text = ""
        cell.paragraphs[0].add_run(ctx.text(s["text"]))
        if s.get("when"):
            from app import rows

            _small(cell.paragraphs[0], ("  nur wenn " if ctx.lang in ("de", "de-du") else "  only if ") + rows.text(item, s, ctx.lang, "de" if ctx.lang in ("de", "de-du") else "en"))
        if ctx.codebook:
            _small(cell.paragraphs[0], f"  {s['export_code']}"
                   + (ctx.say(" (umgepolt)", " (reverse-scored)") if s.get("reverse") else ""), mono=True)
        for n in range(1, len(labels) + 1 + dk):
            _center(tc[r][n], "○")
    _dk_note(ctx, item)


def _side_by_side(ctx, item):
    """A table: the rows down the side, one column per question; its answers named in the head (concept side-by-side)."""
    from app.items.side_by_side import column_code

    kind = item.kind
    cols, statements = kind.columns(item), kind.statements(item, ctx.lang)
    table = ctx.doc.add_table(rows=1 + len(statements), cols=1 + len(cols))
    _grid(table)
    tc = _matrix(table)
    first = 4.5
    _widths(table, [first] + [(WIDTH - first) / max(1, len(cols))] * len(cols))
    _shade(tc[0][0])
    for n, col in enumerate(cols, start=1):
        head = tc[0][n]
        head.text = ""
        head.paragraphs[0].add_run(ctx.text(kind.title(item, col, ctx.lang))).bold = True
        if col["kind"] in ("single", "multi", "dropdown"):
            labels = [ctx.text(l) or str(k) for k, l in enumerate(kind.labels(item, col, ctx.lang), start=1)]
            if col.get("dont_know"):
                labels.append(_dk_label(ctx.lang))
            head.add_paragraph()
            _small(head.paragraphs[-1], " / ".join(labels))
        if ctx.codebook:
            _small(head.paragraphs[0], f"  {column_code(n - 1, col)}", mono=True)
        _shade(head)
    for r, s in enumerate(statements, start=1):
        cell = tc[r][0]
        cell.text = ""
        cell.paragraphs[0].add_run(ctx.text(s["text"]))
        if ctx.codebook:
            _small(cell.paragraphs[0], f"  {s['export_code']}", mono=True)
        for n, col in enumerate(cols, start=1):
            points = (col.get("points") or 2) + (1 if col.get("dont_know") else 0)
            mark = {"single": "○ " * points, "multi": "☐ " * points, "dropdown": "▾ ________",
                    "number": "______ " + (kind.unit(item, col, ctx.lang) or ""), "text": "____________"}[col["kind"]]
            _center(tc[r][n], mark.strip())


def _semantic(ctx, item):
    kind = item.kind
    points, pairs = item.config["points"], kind.pairs(item, ctx.lang)
    table = ctx.doc.add_table(rows=len(pairs), cols=points + 2)
    _grid(table)
    tc = _matrix(table)
    _widths(table, [4.0] + [(WIDTH - 9.0) / points] * points + [5.0])
    for r, pair in enumerate(pairs):
        left = tc[r][0]
        left.text = ""
        left.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        left.paragraphs[0].add_run(ctx.text(pair["left"]))
        for n in range(points):
            _center(tc[r][n + 1], "○")
        right = tc[r][points + 1]
        right.text = ""
        right.paragraphs[0].add_run(ctx.text(pair["right"]))
        if pair.get("when"):
            from app import rows

            _small(right.paragraphs[0], ("  nur wenn " if ctx.lang in ("de", "de-du") else "  only if ") + rows.text(item, pair, ctx.lang, "de" if ctx.lang in ("de", "de-du") else "en"))
        if ctx.codebook:
            _small(right.paragraphs[0], f"  {pair['export_code']}"
                   + (ctx.say(" (umgepolt)", " (reverse-scored)") if pair.get("reverse") else ""), mono=True)


def _picture(ctx, item, variables=None):
    """The approved image of an image item (for a loop row: the one for its values) at most 14 cm wide; Word takes
    PNG, JPEG and GIF, anything else (WebP) is converted. Returns False if there is none."""
    import os

    from flask import current_app

    from app import image_rights, media, transfer

    asset = item.kind.asset_for(item, variables)
    path = os.path.join(current_app.config["UPLOAD_FOLDER"], item.block.survey.uid, asset.filename) if asset else ""
    if not asset or not os.path.isfile(path):
        return False
    try:
        seen = getattr(ctx, "picture_paths", set())
        remaining = transfer.export_limit() - (0 if path in seen else getattr(ctx, "picture_bytes", 0))
        if remaining <= 0:
            raise transfer.ExportTooLarge(ctx.say("Die Bilder sind für den Word-Export zu groß. Bitte den Fragebogen als ZIP exportieren.",
                                                 "The images are too large for Word export. Please export the questionnaire as ZIP."))
        with open(path, "rb") as source:
            raw = source.read(remaining + 1)
        if len(raw) > remaining:
            raise transfer.ExportTooLarge(ctx.say("Die Bilder sind für den Word-Export zu groß. Bitte den Fragebogen als ZIP exportieren.",
                                                 "The images are too large for Word export. Please export the questionnaire as ZIP."))
        with media.decoding(), media.open_image(raw, ["PNG", "JPEG", "GIF", "WEBP"]) as img:
            width_px = img.width
            data = io.BytesIO()
            if img.format in ("PNG", "JPEG", "GIF"):
                data.write(raw)
            else:
                rights = image_rights.read(img)
                img.thumbnail((1600, 1600))  # Word displays at most 14 cm; keep raster conversion bounded
                converted = img.convert("RGBA" if img.mode in ("RGBA", "LA", "P") else "RGB")
                converted.info.clear()
                converted.save(data, "PNG", **image_rights.save_options(rights, "PNG"))
        if data.tell() > remaining:
            raise transfer.ExportTooLarge(ctx.say("Die Bilder sind für den Word-Export zu groß. Bitte den Fragebogen als ZIP exportieren.",
                                                 "The images are too large for Word export. Please export the questionnaire as ZIP."))
        size = data.tell()
        data.seek(0)
        p = ctx.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(data, width=Cm(min(14.0, max(4.0, width_px / 96 * 2.54))))
        if path not in seen:
            ctx.picture_bytes = getattr(ctx, "picture_bytes", 0) + size
            ctx.picture_paths = seen | {path}  # python-docx reuses identical pictures across loop rows
        return True
    except transfer.ExportTooLarge:
        raise
    except Exception:  # noqa: BLE001 - a broken file leaves a note, not a failed export
        return False


def _image(ctx, item):
    """Image and image-to-mark items: the picture; in the codebook for a loop with one image per row only the
    first, with a note."""
    from app import loops

    dynamic = item.kind.is_dynamic(item)
    variables = ctx.variables
    if dynamic and not variables and loops.active(item.block):
        rows = loops.all_variables(item.block, ctx.lang)
        variables = rows[0] if rows else {}
        if ctx.codebook:
            _note(ctx.doc, ctx.say("Ein Bild je Loop-Zeile ({n} Bilder); hier das der ersten Zeile.",
                                   "One image per loop row ({n} images); here the one of the first row.",
                                   n=len(item.kind.required_keys(item))))
    if not _picture(ctx, item, variables) and ctx.codebook:
        _note(ctx.doc, ctx.say("Bild fehlt oder ist noch nicht freigegeben.", "Image missing or not approved yet."))


def _warning(ctx, item):
    """A red box: this element is not in the paper version."""
    table = ctx.doc.add_table(rows=1, cols=1)
    _grid(table)
    tc = _matrix(table)
    cell = tc[0][0]
    _shade(cell, "F8D7D5")
    cell.text = ""
    run = cell.paragraphs[0].add_run(ctx.say("„{label}“ lässt sich auf Papier nicht sinnvoll darstellen und fehlt "
                                             "hier; es wird nur online beantwortet.",
                                             "“{label}” cannot be shown sensibly on paper and is left out here; it is "
                                             "answered online only.", label=ctx.kind_label(item.kind)))
    run.font.size, run.font.color.rgb = Pt(9), RGBColor(0x8A, 0x1C, 0x14)


def _slider(ctx, item):
    kind, t = item.kind, item.t(ctx.lang)
    c = item.config
    lo, hi = kind.show(item, c.get("min", 0), ctx.lang), kind.show(item, c.get("max", 100), ctx.lang)
    left, right = ctx.text(t.get("left")), ctx.text(t.get("right"))
    rows = kind.rows(item, ctx.lang) or [{"text": "", "export_code": item.code}]
    table = ctx.doc.add_table(rows=len(rows), cols=3)
    _grid(table)
    tc = _matrix(table)
    _widths(table, [5.5, 7.0, 4.5])
    for r, row in enumerate(rows):
        cell = tc[r][0]
        cell.text = ""
        cell.paragraphs[0].add_run(ctx.text(row["text"]))
        if ctx.codebook:
            _small(cell.paragraphs[0], f"  {row['export_code']}", mono=True)
        line = tc[r][1]
        line.text = ""
        p = line.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(f"{lo} ├────────────────┤ {hi}").font.size = Pt(9)
        if left or right:
            p.add_run().add_break()
            _small(p, f"{left}{' … ' if left and right else ''}{right}")
        _center(tc[r][2], "Wert: ________" if ctx.lang in ("de", "de-du") else "Value: ________", size=9)


def _write_in(ctx, item, total=None):
    """Ranking and constant sum: a field before every option (a rank, or an amount and the total)."""
    for o in item.kind.options(item, ctx.lang):
        p = ctx.doc.add_paragraph()
        p.paragraph_format.left_indent, p.paragraph_format.space_after = Cm(0.6), Pt(2)
        p.add_run("______  " + ctx.text(o["label"]))
        if ctx.codebook:
            _small(p, f"   {o['code']}", mono=True)
    if total is not None:
        p = ctx.doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.add_run(("Summe" if ctx.lang in ("de", "de-du") else "Total") + f": {total}").bold = True


def _lines(ctx, lines):
    for line in lines:
        text = ctx.text(line).strip().lstrip("-").strip()
        if not ctx.codebook:
            text = re.sub(r"`[^`]*`\s*", "", text).strip()  # variable names
        if text:
            p = ctx.doc.add_paragraph(text)
            p.paragraph_format.left_indent, p.paragraph_format.space_after = Cm(0.6), Pt(0)


def _item(ctx, item):
    kind = item.kind
    if item.type == "page_break" or item.type in LIVE:
        return
    if item.type == "info":
        t = item.t(ctx.lang)
        _paragraphs(ctx, t.get("question") or t.get("body") or "")
        return
    t = _question(ctx, item)
    if item.type in NOT_ON_PAPER:
        _warning(ctx, item)
        if ctx.codebook and kind.questionnaire_lines(item, ctx.lang):
            _lines(ctx, kind.questionnaire_lines(item, ctx.lang))
        return
    if item.type in ("image", "image_mark"):
        _image(ctx, item)
        if item.type == "image":
            return
    if item.type in ("single_choice", "multiple_choice", "federal_state"):
        _choice(ctx, item)
    elif item.type == "slider":
        _slider(ctx, item)
    elif item.type == "ranking":
        _note(ctx.doc, "Bitte nummerieren Sie: 1 = an erster Stelle." if ctx.lang == "de"
              else "Bitte nummeriere: 1 = an erster Stelle." if ctx.lang == "de-du" else "Please number them: 1 = first.")
        _write_in(ctx, item)
    elif item.type == "constant_sum":
        from app import units

        c = item.config
        if item.kind.personal(item):
            source = item.kind._source(item)[0]
            name = f"„{source.code}“" if source is not None else "?"
            total = (f"die Antwort auf {name}" if ctx.lang in ("de", "de-du") else f"the answer to {name}")
        else:
            total = units.fmt(c.get("total", 100), c.get("format"), c.get("currency"), item.text("unit", ctx.lang) or "", ctx.lang)
        _write_in(ctx, item, total=total)
    elif item.type == "postal_code":
        ctx.doc.add_paragraph("☐☐☐☐☐")
    elif item.type == "date_time":
        ctx.doc.add_paragraph(DATE_FORMS.get(kind.asked(item), "__________")
                              + (f"     ○  {_dk_label(ctx.lang)}" if dont_know(item) else ""))
        if ctx.codebook:
            _lines(ctx, kind.questionnaire_lines(item, ctx.lang))
    elif item.type == "psm":
        if t.get("product"):
            _paragraphs(ctx, t["product"])
        unit = kind.unit_label(item, ctx.lang) or ""
        for k in kind.sequence(item):
            p = ctx.doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.add_run(ctx.text(t.get("q_" + k, "")) + f"   ________ {unit}")
            if ctx.codebook:
                _small(p, f"   {item.code}_{k}", mono=True)
    elif item.type == "likert":
        _likert(ctx, item)
    elif item.type == "semantic":
        _semantic(ctx, item)
    elif item.type == "side_by_side":
        _side_by_side(ctx, item)
    elif item.type == "consent":
        _paragraphs(ctx, t.get("body") or "")
        if item.kind.pool_notice(item, ctx.lang):
            _paragraphs(ctx, item.kind.pool_notice(item, ctx.lang))
        choices = [ctx.text(item.kind.ui_text(item, ctx.lang, k)) for k in ("agree", "disagree")]
        ctx.doc.add_paragraph("     ".join(f"○  {c}" for c in choices if c))
    elif item.type == "text_input":
        ctx.doc.add_paragraph("_" * 60)
        if ctx.codebook and item.config.get("validation", "none") != "none":
            _note(ctx.doc, ctx.say("Eingabeprüfung: {v}", "Input check: {v}", v=item.config["validation"]))
    elif item.type == "rating":
        lo, hi = kind.bounds(item)
        if ctx.codebook:
            _note(ctx.doc, ctx.say("Bewertung ({style}) von {lo:g} bis {hi:g}", "Rating ({style}) from {lo:g} to {hi:g}",
                                   style=item.config["style"], lo=lo, hi=hi))
        else:
            ctx.doc.add_paragraph(("   ".join(f"○ {v:g}" for v in range(int(lo), int(hi) + 1)) if hi - lo <= 10
                                   else f"{lo:g} ______________________________ {hi:g}")
                                  + (f"      ○ {_dk_label(ctx.lang)}" if dont_know(item) else ""))
        _dk_note(ctx, item)
    elif kind.questionnaire_lines(item, ctx.lang):
        _lines(ctx, kind.questionnaire_lines(item, ctx.lang))


# --- blocks and the document ------------------------------------------------------------------


def _pages(ctx, block):
    pages = block.pages()
    for p_no, page in enumerate(pages, start=1):
        if len(pages) > 1 and ctx.codebook:
            ctx.doc.add_heading(ctx.say("Seite {n}", "Page {n}", n=p_no), level=2)
        elif p_no > 1:
            ctx.doc.add_paragraph().paragraph_format.space_after = Pt(6)
        for item in page:
            _item(ctx, item)


def _codebook_block(ctx, survey, block, n, badges):
    from app import loops

    title = block.title(ctx.lang)
    ctx.doc.add_heading(f"Block {n}: {plain(title) if title else ctx.say('ohne Titel', 'untitled')}", level=1)
    for badge in badges.get(block.id, []):
        _note(ctx.doc, ctx.say("Ablauf: {badge}", "Flow: {badge}", badge=plain(str(badge))))
    if block.randomize_items:
        _note(ctx.doc, ctx.say("Elemente innerhalb jeder Seite in zufälliger Reihenfolge.",
                               "Elements within each page in random order."))
    loop = loops.config(block)
    if loops.active(block) and loops.blocks_in(survey, loops.node_of(block))[0].id == block.id:
        _note(ctx.doc, ctx.say("Schleife {code}: ", "Loop {code}: ", code=loop['code'])
              + loops.selection_text(survey, loop, en=ctx.ui == "en")
              + (ctx.say(", zufällige Reihenfolge", ", random order") if loop["order"] == "random" else "")
              + ctx.say("; Antworten je Zeile in Spalten mit Suffix _<Zeile>",
                        "; answers per row in columns with the suffix _<row>"))
        table = ctx.doc.add_table(rows=1 + len(loop["rows"]), cols=1 + len(loop["columns"]))
        _grid(table)
        tc = _matrix(table)
        tc[0][0].text = "#"
        for c_no, column in enumerate(loop["columns"], start=1):
            tc[0][c_no].text = column["name"]
            _shade(tc[0][c_no])
        for r_no, row in enumerate(loop["rows"], start=1):
            tc[r_no][0].text = str(r_no)
            for c_no, column in enumerate(loop["columns"], start=1):
                tc[r_no][c_no].text = loops.value(loop, row, column, ctx.lang, ctx.lang)
    _pages(ctx, block)


def _print_block(ctx, block):
    """As participants see it: the title only where it is shown; a block in a loop once per table row."""
    from app import loops

    rows = [r["id"] for r in loops.config(block)["rows"]] if loops.active(block) else [None]
    for row in rows:
        ctx.variables = loops.variables(block, row, ctx.lang) if row else {}
        title = block.title(ctx.lang)
        if block.show_title and title:
            ctx.doc.add_heading(ctx.text(title), level=1)
        _pages(ctx, block)
    ctx.variables = {}


def build(survey, version="codebook", ui=None):
    """The questionnaire as .docx bytes; `version` "codebook" or "print"; `ui` the language of the codebook's frame
    texts ("de" or "en"; None: the downloader's, app.lang.current())."""
    from app import flow
    from app import lang as ui_lang
    from app.models import utcnow
    from app.timeutil import to_local

    ui = ui if ui in ui_lang.LANGUAGES else ui_lang.current()
    lang = survey.default_language
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    ctx = Ctx(doc, lang, version if version in VERSIONS else "codebook", ui=ui)
    from app import constructs

    constructs.prefetch(survey)  # the construct of every element without a query each (P12)

    doc.add_heading(plain(survey.title_in(lang)), level=0)
    if ctx.codebook:
        now = to_local(utcnow())
        _note(doc, ctx.say("Fragebogen mit Codebuch, Fassung {version}, Stand {date} · Befragung {code} · "
                           "Variablennamen in Schreibmaschinenschrift; ○ eine Antwort, ☐ mehrere möglich",
                           "Questionnaire with codebook, revision {version}, as of {date} · survey {code} · "
                           "variable names in typewriter font; ○ one answer, ☐ several possible",
                           version=survey.qversion or 1, code=survey.code,
                           date=f"{now:%d.%m.%Y}" if ui == "de" else f"{now:%Y-%m-%d}"))
        badges = flow.badges(survey, ui)
        for n, block in enumerate(survey.blocks, start=survey.block_start):
            _codebook_block(ctx, survey, block, n, badges)
    else:
        from app.ui_i18n import tr

        _note(doc, tr(lang, "Fragen mit * sind Pflichtangaben. ○ eine Antwort, ☐ mehrere Antworten möglich.",
                      "Questions marked * are required. ○ one answer, ☐ several answers possible."))
        for block in survey.blocks:
            _print_block(ctx, block)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
