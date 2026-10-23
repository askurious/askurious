"""The structural equation model of the results page as tables for a manuscript (.docx, APA style 7th edition; SEM
extensions 2026-10-17, point 15): Table 1 the paths with unstandardised and standardised values, Table 2 the indirect
effects where there are any, Table 3 the measurement model; the fit and what the model rests on in the notes. The
numbers are those of the page (app/sem_model.py, its kept estimate and bootstrap). German with the decimal comma,
English without a leading zero where a value cannot exceed 1 (app/manuscript_tables.py)."""
import io

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from app import sem_model
from app.manuscript_tables import _border, _write, fmt, fmt_ci

TEXT = {
    "de": {"t1": "Tabelle {k}", "paths": "Pfadkoeffizienten des Strukturgleichungsmodells", "path": "Pfad",
           "ind": "Indirekte Effekte", "via": "Weg", "measure": "Messmodell", "construct": "Konstrukt / Item",
           "weight": "Gewicht", "note": "Anmerkung.", "direct": "direkter Effekt", "control": "Kontrollvariable",
           "cb": ("N = {n} vollständige Teilnahmen. Maximum Likelihood mit robusten Standardfehlern (Sandwich) und "
                  "Satorra-Bentler-korrigiertem χ²; β = standardisierter Pfad mit 95-%-Konfidenzintervall; p zweiseitig."),
           "pls": ("N = {n} vollständige Teilnahmen. PLS-SEM (Modus A, Pfadgewichtung); β = standardisierter Pfad; "
                   "Intervall und p aus dem Bootstrap mit {d} Ziehungen; f² = Effektstärke (Cohen, 1988)."),
           "noboot": " Der Bootstrap ist noch nicht gerechnet: ohne Intervalle und p.",
           "fit": " Modellfit: {fit}.", "r2": " R²: {r2}.",
           "ind_body": ("Indirekter Effekt = Produkt der standardisierten Pfade; 95-%-Perzentil-Bootstrap-Intervall "
                        "({d} Ziehungen). p nach gemeinsamer Signifikanz beider Pfade (das größere der beiden Pfad-p; "
                        "MacKinnon et al., 2002)."),
           "m_body": ("λ = standardisierte Ladung; AVE = durchschnittlich erfasste Varianz; CR = Composite Reliability "
                      "(Fornell & Larcker, 1981); HTMT = größtes Heterotrait-Monotrait-Verhältnis mit einem anderen "
                      "Konstrukt (Henseler et al., 2015)."),
           "hyp": "H{n}"},
    "en": {"t1": "Table {k}", "paths": "Path Coefficients of the Structural Equation Model", "path": "Path",
           "ind": "Indirect Effects", "via": "Path", "measure": "Measurement Model", "construct": "Construct / item",
           "weight": "Weight", "note": "Note.", "direct": "direct effect", "control": "control variable",
           "cb": ("N = {n} complete cases. Maximum likelihood with robust (sandwich) standard errors and the "
                  "Satorra-Bentler scaled χ²; β = standardized path with its 95% confidence interval; p two-sided."),
           "pls": ("N = {n} complete cases. PLS-SEM (mode A, path weighting); β = standardized path; interval and p from "
                   "the bootstrap with {d} resamples; f² = effect size (Cohen, 1988)."),
           "noboot": " The bootstrap has not been run yet: no intervals and p.",
           "fit": " Model fit: {fit}.", "r2": " R²: {r2}.",
           "ind_body": ("Indirect effect = product of the standardized paths; 95% percentile bootstrap interval ({d} "
                        "resamples). p by the joint significance of both paths (the larger of the two path p; MacKinnon "
                        "et al., 2002)."),
           "m_body": ("λ = standardized loading; AVE = average variance extracted; CR = composite reliability (Fornell & "
                      "Larcker, 1981); HTMT = largest heterotrait-monotrait ratio with another construct (Henseler et al., "
                      "2015)."),
           "hyp": "H{n}"},
}


def _p(p, lang):
    if p is None:
        return "–"
    if p < 0.001:
        return "< .001" if lang == "en" else "< ,001"
    text = fmt(p, lang, False, 3)  # p without its leading zero in both languages, as the results page writes it
    return text[1:] if text.startswith("0") else text


def _title(doc, number, title, t):
    p = doc.add_paragraph()
    p.add_run(t["t1"].format(k=number)).bold = True
    p.paragraph_format.space_after = Pt(0)
    p = doc.add_paragraph()
    p.add_run(title).italic = True


def _table(doc, head, rows, italic=()):
    """An APA table: lines only above and below the head and at the end; the first column wide (the paths and
    statements), the numbers narrow."""
    table = doc.add_table(rows=1 + len(rows), cols=len(head))
    table.autofit = False
    wide = [j for j, name in enumerate(head) if "CI" in name or "KI" in name]  # intervals need room for "[0,48; 0,70]"
    rest = (16.0 - 6.0 - 3.2 * len(wide)) / max(len(head) - 1 - len(wide), 1)
    widths = [Cm(6.0) if j == 0 else Cm(3.2) if j in wide else Cm(rest) for j in range(len(head))]
    for j, column in enumerate(table.columns):  # the grid (LibreOffice) and every cell (Word)
        column.width = widths[j]
    for row in table.rows:
        for j, cell in enumerate(row.cells):
            cell.width = widths[j]
    for j, name in enumerate(head):
        cell = table.rows[0].cells[j]
        _write(cell, name, italic=j in italic, right=j > 0)
        _border(cell, ("top", "bottom"))
    for i, values in enumerate(rows):
        for j, text in enumerate(values):
            cell = table.rows[i + 1].cells[j]
            _write(cell, text, right=j > 0)
            if i == len(rows) - 1:
                _border(cell, ("bottom",))


def _note(doc, text, t):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.add_run(t["note"]).italic = True
    p.add_run(" " + text)


def build(survey, result, boot=None, lang="de"):
    """The tables as .docx bytes; result: sem_model.estimate of the page, boot: its bootstrap (boot_state result)."""
    lang = lang if lang in TEXT else "de"
    t = TEXT[lang]
    cb = result.get("method") != "pls"
    fit = result.get("fit") or {}
    names = result.get("nodes") or {}
    label = (lambda v: (names.get(v) or (v, v))[1])
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Cm(2.5)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Times New Roman", Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    number = 1
    # Table 1: the paths, the hypotheses that named them
    by_path = {}
    for n, entry in sorted((result.get("hyps") or {}).items()):
        for pair in entry.get("paths") or []:
            by_path.setdefault(tuple(pair), []).append(n)
    controls = result.get("controls") or {}
    rows = []
    for a, b in (result.get("spec") or {}).get("paths") or []:
        row = sem_model.path_row(result, a, b, boot)
        if row is None:
            continue
        hyps = by_path.get((a, b))
        tag = (", ".join(t["hyp"].format(n=n) for n in hyps) if hyps else
               t["control"] if a in controls else t["direct"])
        values = [f"{label(a)} → {label(b)} ({tag})"]
        if cb:
            values += [fmt(row.get("est"), lang), fmt(row.get("se"), lang)]
        values += [fmt(row["std"], lang), fmt_ci(row.get("ci"), lang, bounded=False) or "–", _p(row.get("p"), lang)]
        if not cb:
            values.append(fmt((fit.get("f2") or {}).get((a, b)), lang))
        rows.append(values)
    if rows:
        _title(doc, number, t["paths"], t)
        head = [t["path"]] + (["b", "SE"] if cb else []) + ["β", "95% CI" if lang == "en" else "95-%-KI", "p"] + ([] if cb else ["f²"])
        _table(doc, head, rows, italic=set(range(1, len(head))))
        body = (t["cb"] if cb else t["pls"]).format(n=result.get("n"), d=sem_model.BOOT_DRAWS)
        if not cb and not boot:
            body += t["noboot"]
        sentence = sem_model.fit_sentence(fit, lang) if cb else ""
        if sentence:
            body += t["fit"].format(fit=sentence)
        r2 = fit.get("r2") or {}
        if r2:
            body += t["r2"].format(r2="; ".join(f"{label(v)} {fmt(x, lang, True)}" for v, x in r2.items() if v in names))
        _note(doc, body, t)
        number += 1
    # Table 2: the indirect effects
    ind = sem_model.indirect_table(result, boot)
    if ind:
        doc.add_paragraph()
        _title(doc, number, t["ind"], t)
        rows = []
        for r in ind:
            legs = []
            for e in (result.get("hyps") or {}).get(r["hyps"][0], {}).get("indirect") or []:
                if " → ".join(label(v) for v in e) == r["path"]:
                    x, m, y = e
                    legs = [sem_model.path_row(result, x, m, boot), sem_model.path_row(result, m, y, boot)]
            ps = [leg.get("p") for leg in legs if leg]
            p = max(ps) if len(ps) == 2 and all(v is not None for v in ps) else None
            rows.append([f"{r['path']} ({t['hyp'].format(n=r['hyps'][0])})", fmt(r["std"], lang),
                         fmt_ci(r.get("boot_ci"), lang, bounded=False) or "–", _p(p, lang)])
        _table(doc, [t["via"], "β", "95% CI" if lang == "en" else "95-%-KI", "p"], rows, italic={1, 2, 3})
        _note(doc, t["ind_body"].format(d=sem_model.BOOT_DRAWS), t)
        number += 1
    # Table 3: the measurement model
    mtable, _ = sem_model.measurement_table(result)
    if mtable:
        doc.add_paragraph()
        _title(doc, number, t["measure"], t)
        rows = []
        for c in mtable:
            rows.append([c["label"], "", fmt(c["ave"], lang, True), fmt(c["cr"], lang, True), fmt(c["htmt_max"], lang, True)])
            for i in c["items"]:
                rows.append([f" {i['code']} {i['text']}".rstrip(), fmt(i["std"], lang, True), "", "", ""])
        _table(doc, [t["construct"], "λ", "AVE", "CR", "HTMT"], rows, italic={1, 2, 3, 4})
        _note(doc, t["m_body"], t)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
