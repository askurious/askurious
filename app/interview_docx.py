"""The Word documents of an interview guide (app/interview.py, concept interviewleitfaden-konzept.md 6.1):

- "leitfaden": head with title, research question and date; per Themenblock the table of the guide scheme of
  Helfferich (2011) with four columns (narrative prompt, content aspects as boxes to tick, maintaining questions,
  concrete questions), the hint and the planned duration below the prompt, "K" for a Kernfrage; hints and stimuli
  of a block as text and picture; Kurzfragebogen and Postskript at the end as forms; the sources.
- "interviewkarte": only the prompts and the aspects, large, one page per Themenblock, for the interview without a device.
- "einwilligung": information and consent for the interviewee to sign, built from the settings (recording, how
  long recording and transcripts are kept), the guide (topic, planned duration, mode) and the members (who is
  responsible). A starting point the researchers adapt in Word, not a legal text; no AI writes it (decision I11).
- "kurzfragebogen", "postskript": the form to fill in on paper (as the print version of a questionnaire,
  app/questionnaire_docx.py), with a field for the interview's identifier.

Built with python-docx and the helpers of questionnaire_docx; nothing of an interview is in it (there is nothing).

Languages (concept "Mehrsprachige Oberfläche", E8, as app/questionnaire_docx.py): the guide's questions, aspects and
forms stay in the survey's default language; the frame of the documents the researchers use (headings, notes, column
titles of the guide, the postscript's head) is in the language of the person who downloads (`ui`, app.lang.current()).
What the interviewee reads and signs or fills in (the consent and the short questionnaire's frame) is in the survey's
language where it is German or English, else in the downloader's."""
import io

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from app import interview
from app.items.base import plain
from app.questionnaire_docx import Ctx, _grid, _item, _note, _paragraphs, _shade, _widths

COLUMNS = ("Leitfrage / Erzählaufforderung", "Inhaltliche Aspekte (Check)", "Aufrechterhaltungsfragen", "Konkrete Fragen")
COLUMNS_EN = ("Guide question / narrative prompt", "Content aspects (check)", "Maintenance questions", "Specific questions")
WIDTHS = (6.2, 4.0, 3.4, 3.4)  # cm, 17 between the margins


def _document():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    return doc


def _paper_ui(survey, ui):
    """The language of what the interviewee reads: the survey's if it is German or English, else the downloader's."""
    lang = survey.default_language or ""
    return "de" if lang in ("de", "de-du") else "en" if lang.startswith("en") else ui


def _head(doc, survey, what, ui="de", guide=False):
    """Title and the line under it; `guide`: the Word guide, which also names the research question."""
    from app.models import utcnow
    from app.timeutil import to_local

    now = to_local(utcnow())
    doc.add_heading(plain(survey.title_in(survey.default_language)), level=0)
    _note(doc, f"{what} · Stand {now:%d.%m.%Y} · Interviewleitfaden {survey.code}" if ui != "en"
          else f"{what} · as of {now:%Y-%m-%d} · interview guide {survey.code}")
    if survey.research_question and guide:
        p = doc.add_paragraph()
        p.add_run("Research question: " if ui == "en" else "Forschungsfrage: ").bold = True
        p.add_run(plain(survey.research_question))


def _cell_lines(cell, lines, box=False):
    cell.text = ""
    for n, line in enumerate(lines):
        p = cell.paragraphs[0] if n == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.add_run(("☐ " if box else "") + plain(line)).font.size = Pt(9.5)


def _probe_cell(ctx, cell, item):
    """The concrete questions: a condition in small italics above its question, a Nachfrage indented with ↳ and its
    "Wann stellen?" (concept bedingungen-leitfaden, Ralf 2026-10-19)."""
    cell.text = ""
    first = True

    def line(text, note, sub):
        nonlocal first
        if note:
            p = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            p.paragraph_format.space_after = Pt(0)
            if sub:
                p.paragraph_format.left_indent = Cm(0.4)
            run = p.add_run(note)
            run.italic, run.font.size = True, Pt(8.5)
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(2)
        if sub:
            p.paragraph_format.left_indent = Cm(0.4)
        p.add_run(("↳ " if sub else "") + plain(text)).font.size = Pt(9.5)

    for row in item.kind.probe_tree(item, ctx.lang, ctx.ui):
        line(row["text"], row["cond_text"], False)
        for sub in row["subs"]:
            line(sub["text"], " · ".join(n for n in (sub["when"], sub["cond_text"]) if n), True)


def _guide_table(ctx, items, subs=frozenset()):
    """The Leitfragen of a run as the four columns of Helfferich; an Unterfrage (ids `subs`) indented with ↳ and its
    "Wann stellen?" under it, without a duration (concept unterfragen, Ralf 2026-10-18)."""
    rows = [i for i in items if i.type == "guide_question"]
    table = ctx.doc.add_table(rows=1 + len(rows), cols=4)
    _grid(table)
    for c, title in enumerate(COLUMNS_EN if ctx.ui == "en" else COLUMNS):
        cell = table.cell(0, c)
        cell.text = ""
        cell.paragraphs[0].add_run(title).bold = True
        _shade(cell, "EFEEE9")
    for r, item in enumerate(rows, start=1):
        kind = item.kind
        first = table.cell(r, 0)
        first.text = ""
        p = first.paragraphs[0]
        sub = item.id in subs
        if sub:
            p.paragraph_format.left_indent = Cm(0.5)
            p.add_run("↳ ")
        if (item.config or {}).get("core") and not sub:
            p.add_run(ctx.say("K  ", "C  ")).bold = True
        p.add_run(plain(item.text("question", ctx.lang)) or ctx.say("(ohne Erzählaufforderung)", "(without narrative prompt)"))
        hint = item.text("hint", ctx.lang)
        meta = first.add_paragraph()
        if sub:
            meta.paragraph_format.left_indent = Cm(0.5)
            parts = [plain(item.text("when", ctx.lang)), kind.case_text(item, ctx.lang, ctx.ui), plain(hint)]
            text = " · ".join(x for x in parts if x)
        else:
            text = (f"{plain(hint)} · " if hint else "") + ctx.say("ca. {n} Min.", "approx. {n} min", n=kind.minutes(item))
        if text:
            run = meta.add_run(text)
            run.italic, run.font.size = True, Pt(8.5)
        _cell_lines(table.cell(r, 1), [t for _, t in kind.rows(item, "aspects", ctx.lang)], box=True)
        _cell_lines(table.cell(r, 2), [t for _, t in kind.rows(item, "keep", ctx.lang)])
        _probe_cell(ctx, table.cell(r, 3), item)
    _widths(table, WIDTHS)


def _guide_block(ctx, block):
    """A Themenblock: the hints and stimuli where they stand, the Leitfragen together in tables between them."""
    subs = {i.id for i, main in interview.outline(block) if main is not None}
    run = []
    for item in block.items:
        if item.type == "guide_question":
            run.append(item)
            continue
        if run:
            _guide_table(ctx, run, subs)
            run = []
        if item.type in interview.TASK_ITEM_TYPES:  # a task (app/interview_tasks.py): named as one
            from app.interview_tasks import kind_name

            label = ctx.doc.add_paragraph().add_run(ctx.say("Aufgabe", "Task") + " · " + kind_name(item, ctx.ui))
            label.italic, label.font.size = True, Pt(8.5)
        if item.type != "page_break":
            _item(ctx, item)
    if run:
        _guide_table(ctx, run, subs)


def _form(ctx, block, title, ui=None):
    """A form on paper; `ui`: the language of its frame if not the downloader's (the short questionnaire)."""
    ctx.doc.add_heading(title, level=1)
    ctx.doc.add_paragraph("Interview identifier: ______________     Date: ______________" if (ui or ctx.ui) == "en"
                          else "Interviewkennung: ______________     Datum: ______________")
    for item in block.items:
        _item(ctx, item)


def _leitfaden(survey, doc, ctx):
    from app.interview import ORDERS, ORDERS_EN, state

    _head(doc, survey, ctx.say("Interviewleitfaden", "Interview guide"), ctx.ui, guide=True)
    st = state(survey)
    total = sum(interview.planned_minutes(b) for b in interview.guide_blocks(survey))
    _note(doc, ctx.say("Geplante Dauer ca. {total} Min. · Reihenfolge: {order} · K = Kernfrage, soll in jedem "
                       "Interview vorkommen · ☐ inhaltliche Aspekte nur abhaken, wenn sie von selbst zur Sprache kommen",
                       "Planned duration approx. {total} min · order: {order} · C = core question, should come up in every "
                       "interview · ☐ tick content aspects only when they come up unprompted",
                       total=total, order=(ORDERS_EN if ctx.ui == "en" else ORDERS)[st["order"]][0]))
    for which, block in interview.ordered_blocks(survey):
        if which == "guide":
            minutes = interview.planned_minutes(block)
            doc.add_heading(plain(block.title(ctx.lang)) + (ctx.say(" (ca. {n} Min.)", " (approx. {n} min)", n=minutes) if minutes else ""), level=1)
            _guide_block(ctx, block)
        else:
            _form(ctx, block, _part(ctx.ui, which) + (ctx.say(" (nach dem Gespräch)", " (after the conversation)") if which == "post" else ""))
    doc.add_heading(ctx.say("Quellen", "Sources"), level=2)
    for source in interview.guide_sources(survey):
        _note(doc, source)


def _part(ui, which):
    """The name of a part (Kurzfragebogen, Postskript) in the language `ui`."""
    return (interview.PARTS_EN if ui == "en" else interview.PARTS)[which]


def _karte(survey, doc, ctx):
    _head(doc, survey, ctx.say("Interviewkarte", "Interview card"), ctx.ui)
    first = True
    for which, block in interview.ordered_blocks(survey):
        if which != "guide":
            continue
        if not first:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        first = False
        doc.add_heading(plain(block.title(ctx.lang)), level=1)
        subs = {i.id for i, main in interview.outline(block) if main is not None}
        for item in block.items:
            if item.type == "guide_question":
                sub = item.id in subs  # an Unterfrage: indented under its Leitfrage, smaller (concept unterfragen)
                indent = Cm(0.9) if sub else Cm(0)
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(4 if sub else 10)
                p.paragraph_format.left_indent = indent
                when = plain(item.text("when", ctx.lang)) if sub else ""
                if sub:
                    p.add_run("↳ ").font.size = Pt(13)
                if when:
                    w = p.add_run(when + ": ")
                    w.italic, w.font.size = True, Pt(12)
                core = ctx.say("K  ", "C  ") if (item.config or {}).get("core") and not sub else ""
                run = p.add_run(core + plain(item.text("question", ctx.lang)))
                run.font.size = Pt(13 if sub else 15)
                for _, text in item.kind.rows(item, "aspects", ctx.lang):
                    line = doc.add_paragraph()
                    line.paragraph_format.left_indent, line.paragraph_format.space_after = indent + Cm(0.6), Pt(0)
                    line.add_run("☐  " + plain(text)).font.size = Pt(13)
            elif item.type == "info":
                _paragraphs(ctx, item.text("question", ctx.lang))
            elif item.type in interview.TASK_ITEM_TYPES:  # a task: its prompt, the material is on the device or on paper
                p = doc.add_paragraph()
                p.paragraph_format.left_indent, p.paragraph_format.space_before = Cm(0.9), Pt(6)
                p.add_run(ctx.say("Aufgabe: ", "Task: ")).italic = True
                p.add_run(plain(item.text("question", ctx.lang))).font.size = Pt(13)


def _offline(survey):
    """Whether the template names the offline mode: the guide allows it (app/interview_offline.py)."""
    from app import interview_offline

    return interview_offline.switched_on(survey)


def _screen(survey):
    """Whether the template names the interviewee's own screen: the guide has tasks, and the screen is offered to the
    person who downloads it or the guide used it (app/interview_screen.py)."""
    from app import interview_screen, interview_tasks

    return (bool(interview_tasks.tasks(survey)) or bool(interview_screen.short_items(survey))) and interview_screen.offered(survey)


def _contact(survey):
    """The contact e-mail and the data protection officer the researchers entered in the settings (Einwilligung und
    Datenschutz), each "" if none; the template then keeps its place to fill in […]."""
    consent = getattr(survey, "consent_item", None)
    config = (consent.config or {}) if consent is not None else {}
    return (config.get("contact_email") or "").strip(), (config.get("dpo") or "").strip()


def _consent(survey, doc, ctx):
    from app.interview import AUDIO_KEEP, DATA_KEEP, guide_blocks, planned_minutes, state

    st = state(survey)
    minutes = sum(planned_minutes(b) for b in guide_blocks(survey))
    people = [m for m in survey.members if not m.is_former]
    doc.add_heading("Information und Einwilligung zur Teilnahme an einem Interview", level=0)
    _note(doc, "Vorlage aus Askurious: bitte an Ihre Studie, Ihre Einrichtung und die Vorgaben Ihrer Ethikkommission "
               "anpassen und die markierten Stellen […] ausfüllen.")

    def part(title, *paragraphs):
        doc.add_heading(title, level=2)
        for text in paragraphs:
            doc.add_paragraph(text)

    topic = plain(survey.title_in(survey.default_language))
    part("Worum es geht",
         f"Wir laden Sie zu einem Interview im Rahmen der Studie „{topic}“ ein."
         + (f" Die Studie geht der Frage nach: {plain(survey.research_question)}" if survey.research_question else "")
         + " Uns interessieren Ihre Erfahrungen und Sichtweisen; es gibt keine richtigen oder falschen Antworten.")
    where = {"praesenz": "persönlich vor Ort", "video": "als Videogespräch", "telefon": "am Telefon",
             "gemischt": "je nach Absprache persönlich, als Videogespräch oder am Telefon"}[st["mode"]]
    part("Ablauf",
         f"Das Interview dauert etwa {minutes or '[…]'} Minuten und findet {where} statt. Sie können jederzeit eine "
         "Pause machen, Fragen nicht beantworten oder das Interview beenden, ohne dass Ihnen daraus ein Nachteil entsteht.")
    if st["audio"]:
        kept = ("Die Tonaufnahme wird eine Woche, nachdem das Transkript geprüft ist, gelöscht." if st["audio_keep"] == "transcript"
                else f"Die Tonaufnahme wird {AUDIO_KEEP[st['audio_keep']]} geschützt aufbewahrt und danach gelöscht.")
        part("Tonaufnahme", "Das Interview wird als Ton aufgenommen und anschließend wörtlich verschriftlicht (Transkript). " + kept)
    if _screen(survey):  # the interviewee's own screen for tasks (app/interview_screen.py)
        part("Aufgaben auf Ihrem eigenen Gerät",
             "Für einzelne Aufgaben im Interview (etwa Begriffe in eine Rangfolge bringen oder Karten sortieren) und für "
             "den Kurzfragebogen mit Angaben zu Ihrer Person können Sie eine Seite auf Ihrem eigenen Handy oder Computer öffnen, über einen Link oder QR-Code, den Sie im Interview "
             "bekommen. Die Seite zeigt nur die jeweilige Aufgabe oder den Kurzfragebogen. Gespeichert wird nur Ihr Ergebnis als Teil des "
             "Interviews, mit der Zeit, zu der Sie es gegeben haben, und wie das übrige Interview gelöscht. Angaben zu "
             "Ihrem Gerät und Ihre IP-Adresse werden nicht gespeichert; Ihr Browser erhält nur ein technisch notwendiges "
             "Sitzungs-Cookie für diese Verbindung. Sie können die Aufgaben auch ohne eigenes Gerät lösen.")
    if _offline(survey):  # interviews without a network (app/interview_offline.py)
        part("Interview ohne Internetverbindung",
             "Wird das Interview ohne Internetverbindung geführt, bleiben Notizen, Antworten und die Tonaufnahme bis zur "
             "Übertragung verschlüsselt auf dem Gerät der Interviewenden, höchstens sieben Tage, und werden dort nach "
             "der Übertragung gelöscht. Ihr Löschcode gilt schon in dieser Zeit.")
    part("Pseudonymisierung und Aufbewahrung",
         "Im Transkript werden Ihr Name und alle Angaben, die auf Sie schließen lassen (Orte, Namen anderer Personen, "
         "Einrichtungen), durch Platzhalter ersetzt. Die Liste, die Ihren Namen mit Ihrer Interviewkennung verbindet, "
         "wird getrennt von den Daten aufbewahrt und gelöscht, sobald sie nicht mehr gebraucht wird.",
         f"Die pseudonymisierten Transkripte werden {DATA_KEEP[st['data_keep']]} aufbewahrt und dann gelöscht."
         " Veröffentlicht werden nur Ergebnisse und kurze Zitate, die keinen Rückschluss auf Sie zulassen.")
    part("Freiwilligkeit und Widerruf",
         "Die Teilnahme ist freiwillig. Sie sind nicht verpflichtet, an dem Interview teilzunehmen oder diese "
         "Einwilligung zu erteilen, und können sie jederzeit ohne Angabe von Gründen widerrufen; dann werden Ihre "
         "Aufnahme und Ihr Transkript gelöscht, soweit sie noch nicht in anonymisierte Ergebnisse eingegangen sind. "
         "Die Rechtmäßigkeit der bis zum Widerruf erfolgten Verarbeitung bleibt unberührt.",
         "Grundlage der Verarbeitung ist Ihre Einwilligung (Art. 6 Abs. 1 lit. a DSGVO). Sie haben das Recht auf "
         "Auskunft, Berichtigung, Löschung, Einschränkung der Verarbeitung, Datenübertragbarkeit und Widerspruch "
         "(Art. 15 bis 21 DSGVO) sowie auf Beschwerde bei einer Datenschutz-Aufsichtsbehörde (Art. 77 DSGVO).")
    names = "; ".join(f"{m.display_name}" + (f", {m.affiliation}" if m.affiliation else "") for m in people)
    email, dpo = _contact(survey)
    part("Verantwortlich und Kontakt",
         f"{names or '[Name, Einrichtung]'}. E-Mail: {email or '[…]'}. "
         f"Datenschutzbeauftragte Ihrer Einrichtung: {(dpo or '[…]').rstrip('.')}.")
    doc.add_heading("Einwilligung", level=2)
    statements = ["Ich habe die Informationen gelesen und hatte Gelegenheit, Fragen zu stellen.",
                  "Ich nehme freiwillig an dem Interview teil."]
    if st["audio"]:
        statements.append("Ich bin mit der Tonaufnahme und ihrer Verschriftlichung einverstanden.")
    statements.append("Ich bin damit einverstanden, dass kurze, pseudonymisierte Zitate in Veröffentlichungen verwendet werden.")
    for text in statements:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.add_run("☐  " + text)
    doc.add_paragraph()
    doc.add_paragraph("Interviewkennung: ______________")
    doc.add_paragraph("Ort, Datum: ________________________     Unterschrift: ________________________________")
    _note(doc, "Ein Exemplar für die interviewte Person, eines für die Forschenden.")


def _consent_en(survey, doc, ctx):
    """_consent in English (the same structure and statements)."""
    from app.interview import AUDIO_KEEP_EN, DATA_KEEP_EN, guide_blocks, planned_minutes, state

    st = state(survey)
    minutes = sum(planned_minutes(b) for b in guide_blocks(survey))
    people = [m for m in survey.members if not m.is_former]
    doc.add_heading("Information and consent to take part in an interview", level=0)
    _note(doc, "Template from Askurious: please adapt it to your study, your institution and the requirements of your "
               "ethics committee and fill in the marked places […].")

    def part(title, *paragraphs):
        doc.add_heading(title, level=2)
        for text in paragraphs:
            doc.add_paragraph(text)

    topic = plain(survey.title_in(survey.default_language))
    part("What it is about",
         f"We invite you to an interview as part of the study “{topic}”."
         + (f" The study addresses the question: {plain(survey.research_question)}" if survey.research_question else "")
         + " We are interested in your experiences and views; there are no right or wrong answers.")
    where = {"praesenz": "in person", "video": "as a video call", "telefon": "by telephone",
             "gemischt": "in person, as a video call or by telephone, as agreed"}[st["mode"]]
    part("Procedure",
         f"The interview takes about {minutes or '[…]'} minutes and takes place {where}. You can take a break at any "
         "time, leave questions unanswered or end the interview without any disadvantage to you.")
    if st["audio"]:
        kept = ("The audio recording is deleted one week after the transcript has been checked." if st["audio_keep"] == "transcript"
                else f"The audio recording is kept securely for {AUDIO_KEEP_EN[st['audio_keep']]} and then deleted.")
        part("Audio recording", "The interview is audio-recorded and then transcribed word for word (transcript). " + kept)
    if _screen(survey):
        part("Tasks on your own device",
             "For some tasks in the interview (such as putting terms in a ranking or sorting cards) and for the short "
             "questionnaire with details about you, you can open a page on your own phone or computer, through a link or QR code you are given in the interview. The page shows only "
             "the task or the short questionnaire at hand. Only your result is stored, as part of the interview, with the time you gave it, and it "
             "is deleted like the rest of the interview. No details of your device and not your IP address are stored; "
             "your browser receives only a technically necessary session cookie for this connection. You can "
             "also do the tasks without a device of your own.")
    if _offline(survey):
        part("Interview without an internet connection",
             "If the interview is conducted without an internet connection, notes, answers and the audio recording stay "
             "encrypted on the interviewer's device until they are transferred, at most seven days, and are deleted there "
             "after the transfer. Your deletion code already works during this time.")
    data_keep = (f"for {DATA_KEEP_EN[st['data_keep']]}" if st["data_keep"] != "end" else DATA_KEEP_EN["end"])
    part("Pseudonymisation and retention",
         "In the transcript, your name and all details that could identify you (places, names of other people, "
         "institutions) are replaced by placeholders. The list that links your name to your interview identifier "
         "is kept separately from the data and deleted as soon as it is no longer needed.",
         f"The pseudonymised transcripts are kept {data_keep} and then deleted."
         " Only results and short quotations that do not allow conclusions about you are published.")
    part("Voluntary participation and withdrawal",
         "Taking part is voluntary. You are under no obligation to take part in the interview or to give this consent, "
         "and you can withdraw it at any time without giving reasons; your recording and your transcript are then "
         "deleted, as far as they have not yet gone into anonymised results. The lawfulness of the processing carried "
         "out until the withdrawal remains unaffected.",
         "The legal basis of the processing is your consent (Art. 6(1)(a) GDPR). You have the right of access, "
         "rectification, erasure, restriction of processing, data portability and to object (Art. 15 to 21 GDPR) and "
         "the right to lodge a complaint with a data protection supervisory authority (Art. 77 GDPR).")
    names = "; ".join(f"{m.display_name}" + (f", {m.affiliation}" if m.affiliation else "") for m in people)
    email, dpo = _contact(survey)
    part("Responsible and contact",
         f"{names or '[Name, institution]'}. Email: {email or '[…]'}. "
         f"Data protection officer of your institution: {(dpo or '[…]').rstrip('.')}.")
    doc.add_heading("Consent", level=2)
    statements = ["I have read the information and had the opportunity to ask questions.",
                  "I take part in the interview voluntarily."]
    if st["audio"]:
        statements.append("I agree to the audio recording and its transcription.")
    statements.append("I agree that short, pseudonymised quotations are used in publications.")
    for text in statements:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.add_run("☐  " + text)
    doc.add_paragraph()
    doc.add_paragraph("Interview identifier: ______________")
    doc.add_paragraph("Place, date: ________________________     Signature: ________________________________")
    _note(doc, "One copy for the interviewee, one for the researchers.")


def build(survey, kind, ui=None, material_license=None):
    """The document `kind` (interview.DOCUMENTS) as .docx bytes; `ui`: the downloader's language (reader())."""
    from app.items.base import reader

    ui = reader(ui)
    doc = _document()
    ctx = Ctx(doc, survey.default_language, "print", ui=ui)
    if kind == "leitfaden":
        _leitfaden(survey, doc, ctx)
    elif kind == "interviewkarte":
        _karte(survey, doc, ctx)
    elif kind == "einwilligung":
        if _paper_ui(survey, ui) == "en":
            _consent_en(survey, doc, ctx)
        else:
            _consent(survey, doc, ctx)
    else:
        which = "short" if kind == "kurzfragebogen" else "post"
        paper = _paper_ui(survey, ui) if which == "short" else ui  # the short questionnaire is filled in by the interviewee
        block = interview.shown_form(survey, which)
        _head(doc, survey, _part(paper, which), paper)
        if block is not None:
            _form(ctx, block, _part(paper, which), ui=paper)
    if material_license:
        from app import interview_versions

        doc.add_heading("Material licence" if ui == "en" else "Materiallizenz", level=1)
        for paragraph in interview_versions.license_text(material_license, ui).split("\n\n"):
            doc.add_paragraph(paragraph)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
