"""The data flow of one survey as a Word file (Daten › Dokumentation › "Datenfluss (Word)"), in the layout of the
RWTH example table of data flows (comparison 2026-10-20, F): one row per step from consent to deletion with the
data, the measures and who is responsible, filled from the settings of this survey. The steps are those of section 4
of the privacy concept (privacy_concept.flow); this file says what holds for the survey, the concept what the
platform does in general. For the researchers' records of processing and for ethics or data protection reviews."""
import io

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Cm, Pt, RGBColor

from app import questionnaire_docx  # noqa: F401 - drops characters XML cannot hold from every text (its _only_xml)
from app.lang import de_en

GREY = RGBColor(0x66, 0x66, 0x5F)


def _join(*parts):
    return "; ".join(p for p in parts if p)


def _months(n):
    if not n:
        return ""
    return de_en("{n} Jahre", "{n} years", n=n // 12) if n % 12 == 0 and n >= 24 else (
        de_en("1 Jahr", "1 year") if n == 12 else de_en("{n} Monate", "{n} months", n=n))


def _location(level):
    """The location element (items/geolocation.py) at the finest level the survey stores."""
    if not level:
        return ""
    words = {"exact": de_en("genau", "exact"), "100m": de_en("auf etwa 100 m", "to about 100 m"),
             "1km": de_en("auf etwa 1 km", "to about 1 km"), "plz": de_en("Postleitzahl", "postal code")}
    return de_en("Standort ({level})", "location ({level})", level=words.get(level, de_en("Bundesland", "federal state")))


def steps(survey):
    """[{"step", "data", "measures", "who"}] for this survey."""
    from app import coding, external, flow, panel, retention, webservice
    from app.items import geolocation, voice

    consent = survey.consent_item
    c = (consent.config or {}) if consent is not None else {}
    people = consent.kind.researchers(consent) if consent is not None else list(survey.members)
    institutions = consent.kind.institutions(consent) if consent is not None and hasattr(consent.kind, "institutions") else []
    responsible = ", ".join(m.display_name + (f" ({m.affiliation})" if m.affiliation else "") for m in people) or \
        de_en("die Forschenden der Befragung", "the researchers of the survey")
    if len(institutions) > 1:
        responsible += " · " + de_en("gemeinsam Verantwortliche (Art. 26 DSGVO)", "joint controllers (Art. 26 GDPR)")
    platform = de_en("Betreiber der Plattform im Auftrag (Art. 28 DSGVO)", "Provider of the platform on behalf (Art. 28 GDPR)")
    stored_ids = panel.stored_names(survey)
    separated = flow.has_separated(survey) or bool((survey.contact or {}).get("enabled"))
    periods = retention.of(survey)
    keys = consent.kind.confirmation_keys(consent) if consent is not None and hasattr(consent.kind, "confirmation_keys") else []
    location = geolocation.finest(survey)
    ai = coding.ai_informed(survey)
    hosts = external.hosts(survey) + webservice.privacy_hosts(survey)
    code = any(i.type == "personal_code" for b in survey.blocks for i in b.items)
    chat = _ai_chat_where(survey)

    return [
        {"step": de_en("Einwilligung", "Consent"),
         "data": _join(de_en("Zustimmung oder Ablehnung", "Agreement or refusal"),
                       de_en("ausdrückliche Einwilligung in besondere Kategorien (Art. 9)", "explicit consent to special "
                             "categories (Art. 9)") if c.get("special_categories") else "",
                       de_en("Einwilligung in die Tonaufnahme", "consent to the audio recording") if "audio" in keys else "",
                       de_en("Einwilligung in den KI-Chat", "consent to the AI chat") if "ai_chat" in keys else ""),
         "measures": _join(de_en("Einwilligungstext mit Rechtsgrundlage, Widerruf, Empfängern und Speicherdauer; "
                                 "Datenschutzseite der Befragung", "Consent text with legal basis, withdrawal, recipients "
                                 "and retention; privacy page of the survey"),
                           de_en("Mindestalter {n} bestätigt", "minimum age {n} confirmed", n=c["min_age"])
                           if c.get("min_age") else ""),
         "who": responsible},
        {"step": de_en("Rekrutierung", "Recruitment"),
         "data": _join(de_en("Teilnahmelink", "Participation link"),
                       de_en("Kennung aus dem Aufruf-Link ({ids})", "identifier from the link ({ids})",
                             ids=", ".join(stored_ids)) if stored_ids else "",
                       de_en("Weiterleitung am Ende an {hosts}", "forwarding at the end to {hosts}",
                             hosts=", ".join(panel.link_hosts(survey))) if panel.link_hosts(survey) else ""),
         "measures": _join(de_en("Kennung pseudonym, gelöscht {span} nach dem Ende",
                                 "identifier pseudonymous, deleted {span} after the end", span=_months(periods["personal"])) if stored_ids and periods["personal"]
                           else (de_en("keine Kennungen gespeichert", "no identifiers stored") if not stored_ids else ""),
                           de_en("an das Panel nur Kennung und Stand, keine Antworten", "only identifier and state go to "
                                 "the panel, no answers") if panel.link_hosts(survey) else ""),
         "who": responsible},
        {"step": de_en("Registrierung und Pseudonym", "Registration and pseudonym"),
         "data": _join(de_en("zufällige Teilnahme-Kennung der Plattform", "random participation ID of the platform"),
                       de_en("Kennung aus dem Aufruf-Link ({ids})", "identifier from the link ({ids})",
                             ids=", ".join(stored_ids)) if stored_ids else "",
                       de_en("persönlicher Code für mehrere Wellen", "personal code for several waves") if code else "",
                       de_en("Hash des Löschcodes", "hash of the deletion code") if getattr(survey, "erase_code", False)
                       else ""),
         "measures": _join(de_en("keine Namen in den Antworten", "no names in the answers"),
                           de_en("die Zuordnung Kennung → Person führt nur, wer rekrutiert (Panelanbieter oder "
                                 "Forschende), nicht die Plattform", "the mapping identifier → person is kept only by "
                                 "whoever recruits (panel provider or researchers), not by the platform")
                           if stored_ids else ""),
         "who": responsible},
        {"step": de_en("Erhebung", "Collection"),
         "data": _join(de_en("Antworten, Zeitpunkte, Sprache", "Answers, times, language"),
                       de_en("grobe Geräteangaben", "coarse device information") if survey.paradata else "",
                       de_en("Zeit je Seite", "time per page") if getattr(survey, "page_timing", False) else "",
                       de_en("Fensterwechsel und Einfügen (Anzahl)", "leaving the window and pasting (counts)")
                       if getattr(survey, "focus_paste", False) else "",
                       de_en("Land und Bundesland aus der IP-Adresse", "country and federal state from the IP address")
                       if getattr(survey, "geo_ip", False) else "",
                       _location(location),
                       de_en("Sprachaufnahmen", "voice recordings") if voice.used(survey) else "",
                       de_en("Nachrichten des KI-Chats und die Antworten der KI, vollständig", "messages of the AI chat "
                             "and the answers of the AI, in full") if chat else "",
                       de_en("getrennt gespeicherte Angaben (z. B. E-Mail für eine Verlosung)", "details stored apart "
                             "(e.g. an email address for a prize draw)") if separated else ""),
         "measures": _join(de_en("keine IP-Adressen, keine Tracking-Cookies", "no IP addresses, no tracking cookies"),
                           de_en("Löschen im Browser bis 15 Minuten nach dem Ende", "deleting in the browser up to 15 "
                                 "minutes after the end"),
                           de_en("Löschcode für jede Person", "deletion code for every person")
                           if getattr(survey, "erase_code", False) else "",
                           de_en("getrennte Angaben ohne Verbindung zu den Antworten", "separate details without a link "
                                 "to the answers") if separated else ""),
         "who": de_en("Teilnehmende; ", "Participants; ") + platform},
        {"step": de_en("Übermittlung", "Transfer"),
         "data": _join(de_en("alle Eingaben zwischen Browser und Server", "all input between browser and server"),
                       de_en("zufällige Teilnahme-Kennung und Angaben an {hosts}",
                             "random participation ID and details to {hosts}", hosts=", ".join(hosts)) if hosts else "",
                       de_en("Nachrichten des KI-Chats an den Modellzugang der Befragung ({where})",
                             "messages of the AI chat to the survey's model access ({where})", where=chat) if chat else ""),
         "measures": _join(de_en("TLS mit HSTS; sichere Cookies", "TLS with HSTS; secure cookies"),
                           de_en("E-Mail-, Web-Adressen und Telefonnummern vor dem Senden maskiert; keine Kennung der "
                                 "Teilnahme an den KI-Dienst", "email and web addresses and phone numbers masked before "
                                 "sending; no identifier of the participation to the AI service") if chat else ""),
         "who": platform},
        {"step": de_en("Speicherung", "Storage"),
         "data": de_en("Datenbank und Dateien auf dem Server der Plattform; Sicherungen", "Database and files on the "
                       "server of the platform; backups"),
         "measures": _join(de_en("Rechte je Rolle", "rights per role"),
                           de_en("Siegelkette gegen stille Änderungen", "seal chain against silent changes")
                           if getattr(survey, "seal", False) else "",
                           de_en("Sicherungen alle vier Stunden, 14 Tage", "backups every four hours, 14 days")),
         "who": platform},
        {"step": de_en("Auswertung", "Analysis"),
         "data": _join(de_en("Ergebnisse und Exporte", "Results and exports"),
                       de_en("Freitexte an den KI-Dienst für Kategorienvorschläge", "free texts to the AI service for "
                             "suggested categories") if ai else ""),
         "measures": _join(de_en("Rohdaten nur mit Bearbeiten oder Verwalten; Zugriffsprotokoll", "raw data only with "
                                 "Edit or Manage; access log"),
                           de_en("KI ohne Speicherung beim Anbieter für späteren Abruf; Teilnehmende informiert",
                                 "AI without storage at the provider for later retrieval; participants informed") if ai else ""),
         "who": responsible},
        {"step": de_en("Veröffentlichung", "Publication"),
         "data": (de_en("anonyme Daten (Open Data){where}{licence}", "anonymous data (open data){where}{licence}",
                        where=(" · " + c["repository"]) if c.get("repository") else "",
                        licence=(" · " + c["license"]) if c.get("license") else "") if c.get("open_data")
                  else de_en("keine Veröffentlichung der Daten vorgesehen", "no publication of the data planned")),
         "measures": de_en("offene Antworten nur mit eigenem Häkchen; keine Tonaufnahmen", "open answers only with a tick "
                           "of their own; no voice recordings") if c.get("open_data") else "—",
         "who": responsible},
        {"step": de_en("Löschung", "Deletion"),
         "data": de_en("Kennungen und getrennte Angaben; alle Antworten", "Identifiers and separate details; all answers"),
         "measures": _join(de_en("Kennungen und getrennte Angaben {span} nach dem Ende",
                                 "identifiers and separate details {span} after the end", span=_months(periods["personal"])) if periods["personal"] else "",
                           de_en("alle Antworten {span} nach dem Ende", "all answers {span} after the end",
                                 span=_months(periods["answers"])) if periods["answers"] else
                           de_en("keine Frist für die Antworten gesetzt: Kriterium gute wissenschaftliche Praxis",
                                 "no period set for the answers: criterion good scientific practice"),
                           de_en("Löschen durch Teilnehmende und Forschende jederzeit", "deletion by participants and "
                                 "researchers at any time")),
         "who": responsible + "; " + platform},
    ]


def build(survey):
    """The .docx as bytes, in the language of the person who downloads it."""
    from app.timeutil import local_today

    doc = Document()
    doc.styles["Normal"].font.size = Pt(10)
    for section in doc.sections:  # landscape, like the example table: four wide columns
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
        section.left_margin = section.right_margin = Cm(1.8)
    doc.add_heading(de_en("Datenfluss der Befragung „{title}“", "Data flow of the survey “{title}”",
                          title=survey.title_in(survey.default_language)), 1)
    run = doc.add_paragraph().add_run(de_en("Stand {date} · aus den Einstellungen der Befragung · ID {code}",
                                            "As of {date} · from the settings of the survey · ID {code}",
                                            date=local_today().isoformat(), code=survey.code))
    run.font.size, run.font.color.rgb = Pt(9), GREY
    rows = steps(survey)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, (de_en("Schritt", "Step"), de_en("Daten", "Data"),
                                                de_en("Maßnahmen", "Measures"), de_en("Verantwortliche", "Responsible"))):
        cell.text = ""
        cell.paragraphs[0].add_run(text).bold = True
    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, (row["step"], row["data"], row["measures"], row["who"])):
            cell.text = text or "—"
    for r in table.rows:
        for i, cell in enumerate(r.cells):
            cell.width = Cm((3.2, 7.5, 9.5, 5.8)[i])
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
    note = doc.add_paragraph().add_run(de_en(
        "Was die Plattform allgemein tut, steht im Datenschutzkonzept des Betreibers; diese Tabelle nennt, was die "
        "Einstellungen dieser Befragung daraus machen. Bitte vor der Weitergabe prüfen und ergänzen.",
        "What the platform does in general is set out in the provider's privacy concept; this table names what the "
        "settings of this survey make of it. Please check and complete it before passing it on."))
    note.font.size, note.font.color.rgb = Pt(9), GREY
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _ai_chat_where(survey):
    """Where the messages of a KI-Chat go (provider and model of the study's access, app/ai_chat.py), "" without one."""
    from app import ai_access, ai_chat

    if not ai_chat.chat_items(survey):
        return ""
    row = ai_chat.access(survey)
    return f"{ai_access.provider_label(row.provider)}, {row.model}" if row else de_en("noch kein Zugang", "no access yet")
